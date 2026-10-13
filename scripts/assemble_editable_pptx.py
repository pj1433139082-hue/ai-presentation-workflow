#!/usr/bin/env python3
"""Build an object-level PPTX from approved assets and the current scene graph."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import re
import sys
import tempfile
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from native_objects import NativeObjectError, add_native_object
from project_snapshot import assert_current, document_snapshot, publish_new_directory
from validate_project import ProjectValidator, speaker_notes_text


VERSION = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,63}\Z")


class AssemblyError(ValueError):
    pass


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(root: Path, name: str) -> dict[str, Any]:
    doc = json.loads((root / name).read_text(encoding="utf-8"))
    if not isinstance(doc, dict):
        raise AssemblyError(f"{name} must be an object")
    return doc


def color(raw: Any, fallback: str) -> RGBColor:
    value = raw if isinstance(raw, str) and re.fullmatch(r"#[0-9A-Fa-f]{6}", raw) else fallback
    return RGBColor.from_string(value[1:])


def box(bbox: list[float], width: int, height: int) -> tuple[int, int, int, int]:
    if (
        not isinstance(bbox, list) or len(bbox) != 4
        or any(type(value) not in {float, int} for value in bbox)
        or bbox[0] < 0 or bbox[1] < 0 or bbox[2] <= 0 or bbox[3] <= 0
        or bbox[0] + bbox[2] > 1 or bbox[1] + bbox[3] > 1
    ):
        raise AssemblyError(f"invalid normalized bbox: {bbox!r}")
    return tuple(round(value * dimension) for value, dimension in zip(bbox, (width, height, width, height)))  # type: ignore[return-value]


def apply_text_role(paragraph: Any, region: dict[str, Any], text_roles: dict[str, Any], foreground: RGBColor) -> str | None:
    font_role = region.get("font_role")
    token = text_roles.get(font_role) if isinstance(font_role, str) else None
    paragraph.font.name = token["font_family"] if isinstance(token, dict) else "Microsoft YaHei"
    paragraph.font.size = Pt(region.get("planned_font_pt", 20))
    paragraph.font.bold = token["bold"] if isinstance(token, dict) else None
    paragraph.font.color.rgb = color(token.get("color"), "#F8FAFC") if isinstance(token, dict) else foreground
    return font_role if isinstance(token, dict) else None


def render_powerpoint(path: Path, render_dir: Path, slide_count: int) -> list[dict[str, Any]]:
    if sys.platform != "win32":
        raise AssemblyError("PowerPoint rendering requires Windows; select another verified renderer")
    import pythoncom  # type: ignore[import-untyped]
    import win32com.client  # type: ignore[import-untyped]

    render_dir.mkdir(parents=True, exist_ok=True)
    pythoncom.CoInitialize()
    app = None
    opened = None
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        opened = app.Presentations.Open(str(path.resolve()), ReadOnly=True, WithWindow=False)
        if opened.Slides.Count != slide_count:
            raise AssemblyError("PowerPoint rendered slide count differs from scene graph")
        records = []
        for index in range(1, slide_count + 1):
            output = render_dir / f"slide-{index:02d}.png"
            if output.exists():
                raise AssemblyError(f"render path already exists: {output}")
            opened.Slides(index).Export(str(output.resolve()), "PNG", 1920, 1080)
            records.append({"path": output.as_posix(), "sha256": digest(output)})
        return records
    except Exception as exc:
        raise AssemblyError(f"PowerPoint could not render the assembled PPTX: {exc}") from exc
    finally:
        if opened is not None:
            opened.Close()
        if app is not None:
            app.Quit()
        pythoncom.CoUninitialize()


def _showcase_docs(root: Path) -> dict[str, dict[str, Any]]:
    names = (
        "project-kind.json", "brief.json", "deck-spec.json", "reference-study.json",
        "style-contract.json", "scene-graph.json", "asset-manifest.json",
        "concept-generation-log.json", "run-state.json",
    )
    return {name: read_json(root, name) for name in names}


def _showcase_preflight(root: Path) -> dict[str, Any]:
    kind = read_json(root, "project-kind.json")
    if kind != {"contract_version": "showcase-v1", "purpose": "showcase"}:
        raise AssemblyError("project-kind.json must exactly select showcase-v1/showcase")
    try:
        from delivery_modes import read_delivery_mode
    except ImportError:
        from .delivery_modes import read_delivery_mode
    if read_delivery_mode(root) != "assembly":
        raise AssemblyError("the selected delivery mode stops before assembly; select assembly explicitly")
    try:
        from showcase_contract import validate_showcase
    except ImportError as exc:
        raise AssemblyError("showcase contract validator is unavailable; production validator is not a fallback") from exc
    report = validate_showcase(root, "concept")
    if not isinstance(report, dict) or report.get("ok") is not True:
        errors = report.get("errors", []) if isinstance(report, dict) else ["validator returned no report"]
        raise AssemblyError("showcase assembly preflight has not passed: " + "; ".join(str(item) for item in errors[:8]))
    return report


def _showcase_snapshot(root: Path, docs: dict[str, dict[str, Any]]) -> dict[Path, str]:
    try:
        snapshot = document_snapshot(root, docs)
        # The concept log calls its source output_path/output_sha256. Bind those
        # generated files explicitly because the shared helper reads output/sha256.
        log = docs["concept-generation-log.json"]
        for row in log.get("slides", []):
            if not isinstance(row, dict):
                continue
            relative = row.get("output_path")
            expected = row.get("output_sha256")
            if not isinstance(relative, str) or not isinstance(expected, str):
                raise AssemblyError("concept log output path/hash is missing")
            path = (root / relative).resolve()
            if not path.is_relative_to(root) or not path.is_file():
                raise AssemblyError("concept log points outside the project or to a missing file")
            current = digest(path)
            if current != expected:
                raise AssemblyError("concept output changed after showcase preflight")
            snapshot[path] = current
        return snapshot
    except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
        raise AssemblyError(f"could not freeze showcase input files: {exc}") from exc


def _verify_showcase_pptx(path: Path, slides: list[dict[str, Any]]) -> None:
    """Reopen the file and compare editable object content with the assembly map."""
    try:
        reopened = Presentation(path)
        if len(reopened.slides) != len(slides):
            raise AssemblyError("reopened PPTX slide count differs from the scene graph")
        for index, slide_record in enumerate(slides):
            slide = reopened.slides[index]
            shapes = {shape.name: shape for shape in slide.shapes}
            for row in slide_record["objects"]:
                shape = shapes.get(row["object_name"])
                if shape is None:
                    raise AssemblyError(f"{row['instance_id']}: named native object is missing after save")
                object_type = row["type"]
                if object_type == "text":
                    if not shape.has_text_frame or shape.text_frame.text != row["text"]:
                        raise AssemblyError(f"{row['instance_id']}: native text changed during PPTX serialization")
                elif object_type == "table":
                    if not shape.has_table:
                        raise AssemblyError(f"{row['instance_id']}: native table did not reopen as a table")
                    table = shape.table
                    if len(table.columns) != row["table_shape"]["column_count"]:
                        raise AssemblyError(f"{row['instance_id']}: table column count changed after save")
                    offset = int(row["table_shape"]["has_visible_label_row"])
                    if offset and table.cell(0, 0).text != row["visible_label"]:
                        raise AssemblyError(f"{row['instance_id']}: visible table label changed after save")
                    for col_index, column in enumerate(row["columns"]):
                        header = column["label"]
                        if table.cell(offset, col_index).text != header:
                            raise AssemblyError(f"{row['instance_id']}: table header changed after save")
                    for row_index, data_row in enumerate(row["rows"], start=1 + offset):
                        for col_index, value in enumerate(data_row["cells"]):
                            if table.cell(row_index, col_index).text != value:
                                raise AssemblyError(f"{row['instance_id']}: table cell changed during PPTX serialization")
                elif object_type == "chart":
                    if not shape.has_chart:
                        raise AssemblyError(f"{row['instance_id']}: native chart did not reopen as a chart")
                    chart = shape.chart
                    if (not chart.has_title or chart.chart_title.text_frame.text != row["title"]
                            or len(chart.series) != len(row["series"])):
                        raise AssemblyError(f"{row['instance_id']}: chart title/series differs after save")
                    categories = [category.label for category in chart.plots[0].categories]
                    if categories != row["categories"]:
                        raise AssemblyError(f"{row['instance_id']}: chart categories changed after save")
                    for series, expected in zip(chart.series, row["series"]):
                        if series.name != expected["name"] or list(series.values) != expected["values"]:
                            raise AssemblyError(f"{row['instance_id']}: chart series values changed after save")
                elif object_type in {"picture", "source-evidence"}:
                    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                        raise AssemblyError(f"{row['instance_id']}: picture did not reopen as an independent picture")
                elif object_type == "connector":
                    if shape.shape_type != MSO_SHAPE_TYPE.LINE:
                        raise AssemblyError(f"{row['instance_id']}: connector did not reopen as a native connector")
                    if row["arrowhead"] == "triangle" and "tailEnd" not in shape._element.xml:
                        raise AssemblyError(f"{row['instance_id']}: native connector arrowhead did not survive serialization")
                elif object_type == "shape" and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    raise AssemblyError(f"{row['instance_id']}: shape flattened into a picture")
    except AssemblyError:
        raise
    except Exception as exc:
        raise AssemblyError(f"could not reopen and inspect showcase PPTX: {exc}") from exc


def assemble_showcase(root: Path, version: str, render: bool = True) -> dict[str, Any]:
    """Assemble an explicitly selected local showcase project from native objects."""
    root = root.resolve()
    if not VERSION.fullmatch(version):
        raise AssemblyError("version must be a safe short identifier")
    if not root.is_dir():
        raise AssemblyError("showcase project directory does not exist")

    _showcase_preflight(root)
    docs = _showcase_docs(root)
    snapshot = _showcase_snapshot(root, docs)
    deck = docs["deck-spec.json"]
    scene_doc = docs["scene-graph.json"]
    manifest = docs["asset-manifest.json"]
    style = docs["style-contract.json"]
    deck_slides = deck.get("slides")
    scene_slides = scene_doc.get("slides")
    if not isinstance(deck_slides, list) or not deck_slides or not isinstance(scene_slides, list):
        raise AssemblyError("showcase deck and scene graph need non-empty slide lists")
    scene_by_id = {row.get("slide_id"): row for row in scene_slides if isinstance(row, dict)}
    deck_ids = [row.get("id") for row in deck_slides if isinstance(row, dict)]
    if (len(deck_ids) != len(deck_slides) or len(set(deck_ids)) != len(deck_ids)
            or len(scene_by_id) != len(scene_slides) or set(scene_by_id) != set(deck_ids)):
        raise AssemblyError("showcase scene graph must cover every deck slide exactly once")
    assets = manifest.get("assets")
    if not isinstance(assets, list):
        raise AssemblyError("showcase asset manifest needs an assets list")
    asset_by_id: dict[str, dict[str, Any]] = {}
    for asset in assets:
        if not isinstance(asset, dict) or not isinstance(asset.get("id"), str) or asset["id"] in asset_by_id:
            raise AssemblyError("showcase asset manifest contains an invalid or duplicate ID")
        asset_by_id[asset["id"]] = asset

    profile = style.get("readability_profile")
    floors = profile.get("font_floors_pt") if isinstance(profile, dict) else None
    if not isinstance(floors, dict) or not floors:
        raise AssemblyError("style contract must declare readability_profile.font_floors_pt")
    numeric_floors = [value for value in floors.values() if type(value) in (int, float) and value > 0]
    if not numeric_floors:
        raise AssemblyError("style contract font floors must include a positive numeric value")
    body_floor = floors.get("body")
    if type(body_floor) not in (int, float) or body_floor <= 0:
        body_floor = max(numeric_floors)

    assembly_root = root / "assembly"
    showcase_root = assembly_root / "showcase"
    output_dir = showcase_root / version
    output_relative = (Path("assembly") / "showcase" / version / f"deck-{version}.pptx").as_posix()
    if output_dir.exists() or (root / output_relative).exists():
        raise AssemblyError(f"showcase version already exists; use a new version: {version}")
    assembly_root.mkdir(parents=True, exist_ok=True)
    showcase_root.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".staging-{version}-", dir=showcase_root))
    staging_pptx = staging / f"deck-{version}.pptx"
    assembled_slides: list[dict[str, Any]] = []
    published = False
    map_published = False
    map_path = root / "assembly-map.json"
    previous_map = map_path.read_bytes() if map_path.is_file() else None
    try:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = str(docs["brief.json"].get("project_id", "Showcase presentation"))
        presentation.core_properties.subject = "Standalone showcase; editable native PowerPoint objects"
        palette = style.get("palette")
        if isinstance(palette, list) and palette:
            background_value = palette[0]
        elif isinstance(palette, dict):
            background_value = palette.get("background", "#FFFFFF")
        else:
            background_value = "#FFFFFF"
        if not isinstance(background_value, str) or not re.fullmatch(r"#[0-9A-Fa-f]{6}", background_value):
            raise AssemblyError("style palette needs a valid background color")
        background = color(background_value, "#FFFFFF")

        copy_lookup_by_slide: dict[str, dict[str, dict[str, Any]]] = {}
        for deck_slide in deck_slides:
            slide_id = deck_slide["id"]
            copy_rows = deck_slide.get("copy")
            if not isinstance(copy_rows, list):
                raise AssemblyError(f"{slide_id}: deck slide must provide its complete copy table")
            copy_by_id: dict[str, dict[str, Any]] = {}
            for copy_row in copy_rows:
                if not isinstance(copy_row, dict) or not isinstance(copy_row.get("id"), str) or copy_row["id"] in copy_by_id:
                    raise AssemblyError(f"{slide_id}: copy IDs must be unique")
                copy_by_id[copy_row["id"]] = copy_row
            titles = [row["text"] for row in copy_rows if row.get("role") == "title"]
            takeaways = [row["text"] for row in copy_rows if row.get("role") == "key-takeaway"]
            if titles != [deck_slide.get("title")] or takeaways != [deck_slide.get("message")]:
                raise AssemblyError(f"{slide_id}: title/message must match the complete copy table")
            copy_lookup_by_slide[slide_id] = copy_by_id

        slide_width, slide_height = presentation.slide_width, presentation.slide_height
        for index, deck_slide in enumerate(deck_slides, start=1):
            slide_id = deck_slide["id"]
            scene_slide = scene_by_id[slide_id]
            if scene_slide.get("canvas") != {"width": 1920, "height": 1080}:
                raise AssemblyError(f"{slide_id}: showcase canvas must be 1920x1080")
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid()
            page_background = scene_slide.get("background_color", background_value)
            if not isinstance(page_background, str) or not re.fullmatch(r"#[0-9A-Fa-f]{6}", page_background):
                raise AssemblyError(f"{slide_id}: background_color must be #RRGGBB")
            slide.background.fill.fore_color.rgb = color(page_background, "#FFFFFF")
            try:
                notes = speaker_notes_text(deck_slide)
            except ValueError as exc:
                raise AssemblyError(f"{slide_id}: {exc}") from exc
            elements = scene_slide.get("elements")
            if not isinstance(elements, list) or not elements or any(not isinstance(row, dict) for row in elements):
                raise AssemblyError(f"{slide_id}: scene elements must be a non-empty list of objects")
            instance_ids = [element.get("instance_id") for element in elements]
            if any(not isinstance(value, str) or not value.strip() for value in instance_ids) or len(set(instance_ids)) != len(instance_ids):
                raise AssemblyError(f"{slide_id}: scene instance IDs must be unique non-empty strings")
            rows = []
            provenance_notes = []
            elements = sorted(
                elements,
                key=lambda row: (row.get("z") if type(row.get("z")) is int else -1, row.get("instance_id", "")),
            )
            for element in elements:
                if element.get("type") == "text":
                    copy_row = copy_lookup_by_slide[slide_id].get(element.get("copy_id"))
                    if (copy_row is None or element.get("text") != copy_row.get("text")
                            or element.get("role") != copy_row.get("role")):
                        raise AssemblyError(f"{slide_id}/{element.get('instance_id')}: native text must exactly match its deck copy row")
                try:
                    row, provenance = add_native_object(
                        slide,
                        element,
                        slide_id=slide_id,
                        slide_width=slide_width,
                        slide_height=slide_height,
                        root=root,
                        assets=asset_by_id,
                        body_font_pt=body_floor,
                    )
                except (NativeObjectError, KeyError, TypeError) as exc:
                    raise AssemblyError(f"{slide_id}/{element.get('instance_id')}: {exc}") from exc
                rows.append(row)
                if provenance:
                    provenance_notes.append(provenance)
            full_notes = "\n".join([notes, *provenance_notes]).strip()
            notes_sha = None
            if full_notes:
                frame = slide.notes_slide.notes_text_frame
                if frame is None:
                    raise AssemblyError(f"{slide_id}: PowerPoint notes body is unavailable")
                frame.text = full_notes
                notes_sha = hashlib.sha256(full_notes.encode("utf-8")).hexdigest()
            assembled_slides.append({
                "slide_id": slide_id,
                "output_index": index,
                "flattened_concept_used": False,
                "objects": rows,
                **({"speaker_notes_sha256": notes_sha} if notes_sha else {}),
            })

        presentation.save(staging_pptx)
        _verify_showcase_pptx(staging_pptx, assembled_slides)
        render_records = []
        if render:
            raw_records = render_powerpoint(staging_pptx, staging / "renders", len(assembled_slides))
            for slide_id, record in zip((row["slide_id"] for row in assembled_slides), raw_records):
                render_path = Path("assembly") / "showcase" / version / "renders" / Path(record["path"]).name
                render_records.append({
                    "slide_id": slide_id,
                    "path": render_path.as_posix(),
                    "sha256": record["sha256"],
                    "pptx_sha256": digest(staging_pptx),
                    "width": 1920,
                    "height": 1080,
                })

        assembly = {
            "contract_version": "showcase-v1",
            "purpose": "showcase",
            "backend": "python-pptx",
            "status": "assembled",
            "review_boundary": "standalone showcase; independent and human review remain pending; formal production G4/G5 is out of scope",
            "output": output_relative,
            "sha256": digest(staging_pptx),
            "input_sha256": {
                "deck_spec": digest(root / "deck-spec.json"),
                "scene_graph": digest(root / "scene-graph.json"),
                "asset_manifest": digest(root / "asset-manifest.json"),
                "style_contract": digest(root / "style-contract.json"),
                "concept_log": digest(root / "concept-generation-log.json"),
            },
            "render_backend": "PowerPoint COM (win32com)" if render else None,
            "renders": render_records,
            "slides": assembled_slides,
        }
        assert_current(snapshot)
        publish_new_directory(staging, output_dir)
        published = True
        assert_current(snapshot)
        map_path = root / "assembly-map.json"
        temp_map_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                newline="\n",
                prefix=".assembly-map-",
                suffix=".tmp",
                dir=root,
                delete=False,
            ) as handle:
                json.dump(assembly, handle, ensure_ascii=False, indent=2)
                handle.write("\n")
                temp_map_path = Path(handle.name)
            temp_map_path.replace(map_path)
            temp_map_path = None
            map_published = True
        finally:
            if temp_map_path is not None and temp_map_path.exists():
                temp_map_path.unlink()
        if render:
            try:
                from showcase_contract import validate_showcase
                final_report = validate_showcase(root, "assembly")
            except ImportError as exc:
                raise AssemblyError("showcase contract validator disappeared after assembly") from exc
            if not final_report.get("ok"):
                details = "; ".join(str(item) for item in final_report.get("errors", [])[:8])
                raise AssemblyError("assembled showcase failed current assembly validation: " + details)
        return assembly
    except Exception:
        if staging.exists():
            shutil.rmtree(staging)
        if map_published:
            if previous_map is None:
                map_path.unlink(missing_ok=True)
            else:
                map_path.write_bytes(previous_map)
        if published and output_dir.exists() and output_dir.resolve().is_relative_to(showcase_root.resolve()):
            shutil.rmtree(output_dir)
        raise


def assemble(root: Path, version: str, *, render: bool = True, pilot_slide: str | None = None,
             preview_slide: str | None = None) -> Path:
    root = root.resolve()
    kind_path = root / "project-kind.json"
    if kind_path.exists() or kind_path.is_symlink():
        if pilot_slide is not None or preview_slide is not None:
            raise AssemblyError("showcase assembly does not use pilot/preview options")
        result = assemble_showcase(root, version, render=render)
        return root / result["output"]
    if not VERSION.fullmatch(version):
        raise AssemblyError("version must be a safe short identifier")
    if pilot_slide and preview_slide:
        raise AssemblyError("pilot and preview slide options are mutually exclusive")
    stage = "pilot" if pilot_slide or preview_slide else "assets"
    # A trial is allowed to prepare its own review evidence; batch production is not.
    validator = ProjectValidator(root, stage, require_pilot_review=pilot_slide is None)
    issues = validator.run()
    if issues:
        raise AssemblyError(f"{stage} preflight has not passed: " + "; ".join(issues[:5]))
    pilot_bindings = validator.pilot_bindings() if pilot_slide else None
    if pilot_slide and (pilot_bindings is None or pilot_bindings["slide_id"] != pilot_slide):
        raise AssemblyError("pilot slide must match the current pilot plan")
    brief = read_json(root, "brief.json")
    if "pptx" not in brief.get("delivery", {}).get("formats", []):
        raise AssemblyError("brief must explicitly request PPTX delivery")
    deck = read_json(root, "deck-spec.json")
    scene = read_json(root, "scene-graph.json")
    manifest = read_json(root, "asset-manifest.json")
    style = read_json(root, "style-contract.json")
    override_path = root / "assembly-style-overrides.json"
    style_overrides = read_json(root, "assembly-style-overrides.json") if override_path.is_file() else {}
    batch_override_path = root / "assembly-style-overrides-batch.json"
    if batch_override_path.is_file():
        batch_overrides = read_json(root, "assembly-style-overrides-batch.json")
        style_overrides.setdefault("slides", {}).update(batch_overrides.get("slides", {}))
    deck_slides = deck.get("slides", [])
    selected_slide = pilot_slide or preview_slide
    if selected_slide:
        deck_slides = [slide for slide in deck_slides if slide.get("id") == selected_slide]
    scene_by_id = {slide["slide_id"]: slide for slide in scene.get("slides", [])}
    assets = {asset["id"]: asset for asset in manifest.get("assets", [])}
    if not deck_slides or (not selected_slide and len(scene_by_id) != len(deck_slides)):
        raise AssemblyError("scene graph must cover every deck slide")
    if len(assets) != len(manifest.get("assets", [])):
        raise AssemblyError("duplicate asset IDs are prohibited")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = brief.get("project_id", "Editable presentation")
    presentation.core_properties.subject = "Object-level assembly from approved components"
    raw_palette = style.get("palette")
    palette: list[Any] = raw_palette if isinstance(raw_palette, list) else []
    background = color(palette[0] if len(palette) > 0 else None, "#0F172A")
    foreground = color(palette[1] if len(palette) > 1 else None, "#F8FAFC")
    accent = color(palette[2] if len(palette) > 2 else None, "#38BDF8")
    semantic_system = style.get("semantic_system") if isinstance(style.get("semantic_system"), dict) else {}
    text_roles = semantic_system.get("text_roles", {})
    frame_roles = semantic_system.get("frame_roles", {})
    assembled_slides = []
    for index, deck_slide in enumerate(deck_slides, start=1):
        slide_id = deck_slide["id"]
        scene_slide = scene_by_id.get(slide_id)
        if not isinstance(scene_slide, dict):
            raise AssemblyError(f"scene graph omits {slide_id}")
        canvas = scene_slide.get("canvas")
        if canvas != {"width": 1920, "height": 1080}:
            raise AssemblyError(f"{slide_id} must use the approved 1920x1080 canvas")
        slide_overrides = style_overrides.get("slides", {}).get(slide_id, {})
        if slide_overrides:
            if slide_overrides.get("concept_sha256") != digest(root / deck_slide["concept_path"]):
                raise AssemblyError(f"{slide_id} text-format overrides are stale against the approved concept")
            if not isinstance(slide_overrides.get("slots"), dict):
                raise AssemblyError(f"{slide_id} text-format overrides need a slots map")
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        try:
            notes = speaker_notes_text(deck_slide)
        except ValueError as exc:
            raise AssemblyError(f"{slide_id}: {exc}") from exc
        if notes:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is None:
                raise AssemblyError(f"{slide_id}: PowerPoint notes body is unavailable")
            notes_frame.text = notes
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = background
        elements = sorted(scene_slide.get("elements", []), key=lambda item: (item.get("z"), item.get("instance_id")))
        used_instances: set[str] = set()
        object_rows = []
        for element in elements:
            instance_id, component_id = element.get("instance_id"), element.get("component_id")
            asset = assets.get(component_id)
            if not isinstance(instance_id, str) or instance_id in used_instances or not isinstance(asset, dict):
                raise AssemblyError(f"{slide_id} contains a missing asset or duplicate scene instance")
            used_instances.add(instance_id)
            if slide_id not in asset.get("slide_ids", [asset.get("slide_id")]) or asset.get("status") != "approved":
                raise AssemblyError(f"{component_id} is not approved for {slide_id}")
            bbox = element.get("bbox")
            geometry = box(bbox, presentation.slide_width, presentation.slide_height)
            if bbox[0] <= .05 and bbox[1] <= .05 and bbox[2] >= .9 and bbox[3] >= .9:
                raise AssemblyError(f"{instance_id} is a prohibited near-full-slide raster")
            if element.get("opacity", 1) != 1:
                raise AssemblyError(f"{instance_id} has unsupported opacity; do not silently flatten it")
            rotation = element.get("rotation", 0)
            if type(rotation) not in {float, int} or not 0 <= rotation < 360:
                raise AssemblyError(f"{instance_id} has invalid rotation")
            asset_path = (root / asset["output"]).resolve()
            if not asset_path.is_relative_to(root) or digest(asset_path) != asset.get("sha256"):
                raise AssemblyError(f"{component_id} source file or hash is stale")
            if asset.get("role") == "ordinary-photo":
                with Image.open(asset_path) as photo_image:
                    source_ratio = photo_image.width / photo_image.height
                placed_ratio = geometry[2] / geometry[3]
                if abs(source_ratio / placed_ratio - 1) > 0.02:
                    raise AssemblyError(f"{component_id} photo placement would distort its aspect ratio; approve a crop or revise the bbox")
            shape = slide.shapes.add_picture(str(asset_path), *geometry)
            shape.name = f"APW:visual:{instance_id}:{component_id}"
            shape.rotation = rotation
            placement_crop = asset.get("placement_crop")
            if isinstance(placement_crop, dict):
                for side in ("left", "top", "right", "bottom"):
                    setattr(shape, f"crop_{side}", placement_crop.get(side, 0))
            object_rows.append({"instance_id": instance_id, "component_id": component_id,
                                "name": shape.name, "bbox": bbox, "z": element["z"],
                                "rotation": rotation, "asset_sha256": asset["sha256"],
                                **({"placement_crop": placement_crop} if isinstance(placement_crop, dict) else {})})
        if {row["component_id"] for row in object_rows} != set(deck_slide.get("components", [])):
            raise AssemblyError(f"{slide_id} scene instances differ from deck components")

        regions = {region["fill_slot_id"]: region for region in deck_slide.get("readability", {}).get("text_regions", [])}
        slot_rows = []
        for slot in deck_slide.get("fill_slots", []):
            slot_id = slot.get("id")
            kind = slot.get("kind")
            if kind not in {"ordinary-text", "ordinary-data", "fixed-evidence"}:
                raise AssemblyError(f"{slide_id}/{slot_id} has unsupported fill-slot kind {kind!r}")
            geometry = box(slot.get("bbox"), presentation.slide_width, presentation.slide_height)
            content = slot.get("content", "")
            if kind in {"ordinary-data", "fixed-evidence"} and content not in {"", None}:
                raise AssemblyError(f"{slide_id}/{slot_id} data or evidence slot must remain empty for user entry")
            if not isinstance(content, str):
                raise AssemblyError(f"{slide_id}/{slot_id} content must be text")
            if not content and regions.get(slot_id, {}).get("role") == "action-title":
                content = deck_slide.get("title", "")
            if not isinstance(content, str):
                raise AssemblyError(f"{slide_id}/{slot_id} title must be text")
            if kind == "ordinary-text" and not content.strip() and slot.get("user_fill_required") is not True:
                raise AssemblyError(f"{slide_id}/{slot_id} ordinary text needs approved content or explicit user_fill_required")
            textbox = slide.shapes.add_textbox(*geometry)
            textbox.name = f"APW:slot:{slot_id}"
            if kind == "ordinary-data":
                # Keep user-entered values blank while making the editable field discoverable.
                frame_token = frame_roles.get(slot.get("frame_role")) if isinstance(slot.get("frame_role"), str) else None
                textbox.line.color.rgb = color(frame_token.get("color"), "#94A3B8") if isinstance(frame_token, dict) else accent
                textbox.line.width = Pt(1.5)
            frame = textbox.text_frame
            frame.clear()
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = frame.paragraphs[0]
            region = regions.get(slot_id, {})
            slot_format = slide_overrides.get("slots", {}).get(slot_id, slot.get("format"))
            display_content = slot_format.get("display_prefix", "") + content if isinstance(slot_format, dict) else content
            rich_runs = slot_format.get("runs") if isinstance(slot_format, dict) else None
            if rich_runs is not None:
                if not isinstance(rich_runs, list) or not rich_runs or any(
                    not isinstance(part, dict) or not isinstance(part.get("text"), str) for part in rich_runs
                ) or "".join(part["text"] for part in rich_runs) != display_content:
                    raise AssemblyError(f"{slide_id}/{slot_id} rich runs must exactly match approved text")
                paragraph.text = ""
            else:
                paragraph.text = display_content
            font_role = apply_text_role(paragraph, region, text_roles, foreground)
            if isinstance(slot_format, dict):
                if "font_pt" in slot_format:
                    paragraph.font.size = Pt(slot_format["font_pt"])
                if "bold" in slot_format:
                    paragraph.font.bold = slot_format["bold"]
                if "color" in slot_format:
                    paragraph.font.color.rgb = color(slot_format["color"], "#F8FAFC")
                if "align" in slot_format:
                    paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                                           "right": PP_ALIGN.RIGHT}[slot_format["align"]]
                if "vertical_align" in slot_format:
                    frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                                             "bottom": MSO_ANCHOR.BOTTOM}[slot_format["vertical_align"]]
            if rich_runs is not None:
                for part in rich_runs:
                    run = paragraph.add_run()
                    run.text = part["text"]
                    run.font.name = paragraph.font.name
                    run.font.size = Pt(part.get("font_pt", paragraph.font.size.pt))
                    run.font.bold = part.get("bold", paragraph.font.bold)
                    run.font.color.rgb = color(part.get("color"), "#" + str(paragraph.font.color.rgb))
            slot_rows.append({"id": slot_id, "kind": kind, "name": textbox.name,
                              "bbox": slot["bbox"], "content": display_content,
                              "font_pt": slot_format.get("font_pt", region.get("planned_font_pt", 20)) if isinstance(slot_format, dict) else region.get("planned_font_pt", 20),
                              **({"font_role": font_role} if font_role is not None else {}),
                              **({"format": slot_format} if isinstance(slot_format, dict) else {})})
        assembled_slides.append({"slide_id": slide_id, "component_ids": list(deck_slide.get("components", [])),
                                 "fill_slot_ids": [slot["id"] for slot in deck_slide.get("fill_slots", [])],
                                 "flattened_concept_used": False, "output_index": index,
                                 "pptx_part": f"ppt/slides/slide{index}.xml", "objects": object_rows,
                                 "slots": slot_rows,
                                 **({"speaker_notes_sha256": hashlib.sha256(notes.encode("utf-8")).hexdigest()} if notes else {})})

    relative = (f"pilot/{version}/deck-{version}.pptx" if pilot_slide
                else f"previews/{version}/deck-{version}.pptx" if preview_slide
                else f"assembly/deck-{version}.pptx")
    output = root / relative
    if output.exists():
        raise AssemblyError(f"output already exists; use a new version: {relative}")
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    render_records = []
    if render:
        render_dir = root / ("reviews/pilot" if pilot_slide else "reviews/previews" if preview_slide
                             else "reviews/assembly") / version
        raw_records = render_powerpoint(output, render_dir, len(assembled_slides))
        output_sha = digest(output)
        render_records = [{"path": Path(record["path"]).relative_to(root).as_posix(),
                           "sha256": record["sha256"], "pptx_sha256": output_sha} for record in raw_records]
    assembly = {
        "backend": "python-pptx", "output": relative, "sha256": digest(output),
        "input_sha256": {name: digest(root / name) for name in ("deck-spec.json", "scene-graph.json", "asset-manifest.json", "style-contract.json")},
        "status": "structure-assembled-visual-review-pending", "render_backend": "powerpoint" if render else None,
        "renders": render_records, "slides": assembled_slides,
    }
    if batch_override_path.is_file():
        assembly["input_sha256"]["assembly-style-overrides-batch.json"] = digest(batch_override_path)
    if pilot_slide:
        assembly["pilot_bindings"] = pilot_bindings
        map_path = root / "pilot" / version / "pilot-assembly-map.json"
    elif preview_slide:
        assembly["preview_slide"] = preview_slide
        map_path = root / "previews" / version / "preview-assembly-map.json"
    else:
        map_path = root / "assembly-map.json"
    map_path.write_text(json.dumps(assembly, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project_dir", type=Path)
    parser.add_argument("--version", required=True)
    parser.add_argument("--no-render", action="store_true", help="Only for diagnostic environments without a renderer")
    parser.add_argument("--pilot-slide", help="Assemble the planned one-slide trial before batch production")
    parser.add_argument("--preview-slide", help="Render an already isolated slide for local batch QA")
    args = parser.parse_args()
    try:
        output = assemble(args.project_dir, args.version, render=not args.no_render,
                          pilot_slide=args.pilot_slide, preview_slide=args.preview_slide)
    except (AssemblyError, OSError, ValueError, KeyError, TypeError, json.JSONDecodeError) as exc:
        parser.exit(1, f"FAIL: {exc}\n")
    print(f"PPTX assembled: {output}")
    kind_path = args.project_dir.resolve() / "project-kind.json"
    if kind_path.exists() or kind_path.is_symlink():
        print("Standalone showcase draft assembled; independent and human review remain pending.")
    else:
        print("Structure assembled; visual consistency is still pending human review.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
