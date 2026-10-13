#!/usr/bin/env python3
"""Validate the public project contract for ai-presentation-workflow."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
import posixpath
import re
import sys
import zlib
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree


VISUAL_ADAPTER_PATH = Path(__file__).resolve().with_name("visual_adapters.py")
VISUAL_ADAPTER_SPEC = importlib.util.spec_from_file_location(
    "ai_presentation_visual_adapters", VISUAL_ADAPTER_PATH
)
if VISUAL_ADAPTER_SPEC is None or VISUAL_ADAPTER_SPEC.loader is None:
    raise RuntimeError("cannot load visual adapter contract")
VISUAL_ADAPTERS = importlib.util.module_from_spec(VISUAL_ADAPTER_SPEC)
VISUAL_ADAPTER_SPEC.loader.exec_module(VISUAL_ADAPTERS)

DECK_QA_PATH = Path(__file__).resolve().with_name("deck_qa_evidence.py")
DECK_QA_SPEC = importlib.util.spec_from_file_location("ai_presentation_deck_qa", DECK_QA_PATH)
if DECK_QA_SPEC is None or DECK_QA_SPEC.loader is None:
    raise RuntimeError("cannot load full-deck QA evidence contract")
DECK_QA = importlib.util.module_from_spec(DECK_QA_SPEC)
DECK_QA_SPEC.loader.exec_module(DECK_QA)

PIXELS_SPEC = importlib.util.spec_from_file_location(
    "ai_presentation_image_pixels", Path(__file__).resolve().with_name("image_pixels.py")
)
if PIXELS_SPEC is None or PIXELS_SPEC.loader is None:
    raise RuntimeError("cannot load decoded pixel cache")
PIXELS = importlib.util.module_from_spec(PIXELS_SPEC)
PIXELS_SPEC.loader.exec_module(PIXELS)

POLICY_SPEC = importlib.util.spec_from_file_location(
    "ai_presentation_review_policy", Path(__file__).resolve().with_name("review_policy.py")
)
if POLICY_SPEC is None or POLICY_SPEC.loader is None:
    raise RuntimeError("cannot load review policy")
REVIEW_POLICY = importlib.util.module_from_spec(POLICY_SPEC)
POLICY_SPEC.loader.exec_module(REVIEW_POLICY)


STAGES = [
    "intake",
    "research",
    "story",
    "style",
    "spec",
    "concept",
    "inventory",
    "pilot",
    "assets",
    "assembly",
    "qa",
    "release",
]

REQUIRED_BY_STAGE = {
    "intake": ["brief.json", "run-state.json"],
    "research": ["research-bundle.json"],
    "story": ["alignment.json", "storyboard.json"],
    "style": [
        "reference-study.json",
        "concept-calibration.json",
        "style-contract.json",
        "capability-report.json",
        "visual-provider-registry.json",
    ],
    "spec": ["deck-spec.json"],
    "concept": ["concept-generation-log.json"],
    "inventory": [
        "scene-graph.json",
        "asset-manifest.json",
        "component-master-manifest.json",
        "decomposition-map.json",
    ],
    "pilot": ["pilot-plan.json"],
    "assets": [],
    "assembly": ["assembly-map.json"],
    "qa": ["qa-report.json"],
    "release": [],
}

VISUAL_ROLES = {
    "illustration",
    "ordinary-photo",
    "photo-cutout",
    "icon",
    "arrow",
    "basic-shape",
    "decorative-shape",
    "panel",
    "table-shell",
    "chart-shell",
    "diagram-node",
    "special-typography",
}

PHOTO_ROLES = {"ordinary-photo", "photo-cutout"}

TRUTH_SENSITIVE_ROLES = {
    "source-screenshot",
    "source-figure",
    "equation",
    "brand-mark",
}

GATES = [
    "brief_authority",
    "story",
    "style_anchors",
    "concept_contact_sheet",
    "decomposition",
    "release",
]

CRITICAL_G1_DECISIONS = {
    "audience",
    "decision_goal",
    "source_authority",
    "truth_boundaries",
    "research_permission",
    "delivery_form",
    "editability_commitment",
}

GATE_REQUIRED_DECISIONS = {
    "story": {
        "slide_sequence",
        "action_titles",
        "slide_messages",
        "evidence_mapping",
        "appendix_boundary",
    },
    "style_anchors": {
        "reference_selection",
        "concept_calibration",
        "composition_profile",
        "readability_profile",
        "visual_direction",
        "style_rules",
        "anchor_opening",
        "anchor_typical",
        "anchor_highest_risk",
    },
    "concept_contact_sheet": {
        "full_deck_contact_sheet",
        "contentful_copy_source_review",
        "readability_review",
        "layout_richness_review",
        "blank_table_chart_regions",
        "special_typography_exceptions",
        "slide_specific_deviations",
    },
    "decomposition": {
        "object_inventory",
        "coverage_review",
        "review_evidence",
    },
    "release": {
        "final_render",
        "qa_summary",
        "concept_differences",
        "qa_exceptions",
        "component_usability",
        "deliverable_index",
    },
}

DEFECT_STAGE_OWNERS = {
    "research": "research",
    "story": "story",
    "style": "style",
    "concept": "concept",
    "inventory": "inventory",
    "asset": "assets",
    "assembly": "assembly",
    "qa": "qa",
}

GATE_ROLLBACK_OWNERS = {
    "brief_authority": "intake",
    "story": "story",
    "style_anchors": "style",
    "concept_contact_sheet": "concept",
    "decomposition": "inventory",
    "release": "qa",
}

QUALITY_PROFILES = {"production", "draft", "validation-fixture"}
READABILITY_FONT_FLOORS = {
    "action-title": ("action_title_min_pt", 32),
    "body": ("body_min_pt", 20),
    "key-caption": ("key_caption_min_pt", 18),
    "source-qualifier": ("source_qualifier_min_pt", 16),
}
READABILITY_VISUAL_FLOORS = {
    "evidence_min_width": 0.42,
    "evidence_min_height": 0.43,
    "fine_ui_preferred_width": 0.50,
    "fine_ui_preferred_height": 0.46,
    "illustration_min_width": 0.30,
    "illustration_min_height": 0.25,
}
REFERENCE_OBSERVATION_FIELDS = {
    "composition",
    "hierarchy",
    "density",
    "image_treatment",
}
SPARSE_DEFAULT_ROLES = {"opening", "transition", "closing"}
COPY_SOURCE_MODES = {"storyboard-exact", "source-exact", "approved-paraphrase"}
CONCEPT_IMAGE_CLASSIFICATIONS = {
    "illustrative",
    "verified-source",
    "pending-placeholder",
}
NARRATIVE_ROLES = {
    "opening",
    "context",
    "tension",
    "evidence",
    "mechanism",
    "proposal",
    "decision",
    "next-step",
    "transition",
    "closing",
    "appendix",
}


def value_sha256(value: Any) -> str:
    canonical = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(canonical).hexdigest()


def png_dimensions(path: Path) -> tuple[int, int] | None:
    data = path.read_bytes()
    if len(data) < 24 or not data.startswith(b"\x89PNG\r\n\x1a\n") or data[12:16] != b"IHDR":
        return None
    width = int.from_bytes(data[16:20], "big")
    height = int.from_bytes(data[20:24], "big")
    return (width, height) if width > 0 and height > 0 else None


def has_meaningful_value(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (list, dict, tuple, set)):
        return bool(value)
    return True


def positive_int_or(value: Any, fallback: int) -> int:
    return value if type(value) is int and value > 0 else fallback


def speaker_notes_text(slide: dict[str, Any]) -> str:
    """Return canonical native-note text without accepting XML-invalid content."""
    value = slide.get("speaker_notes", "")
    if not isinstance(value, str):
        raise ValueError("speaker_notes must be plain text")
    text = value.replace("\r\n", "\n").replace("\r", "\n")
    if any(
        (ord(character) < 32 and character not in "\n\t")
        or 0xD800 <= ord(character) <= 0xDFFF
        or ord(character) in {0xFFFE, 0xFFFF}
        for character in text
    ):
        raise ValueError("speaker_notes contains XML-invalid control characters")
    return text


class ProjectValidator:
    def __init__(self, root: Path, stage: str, *, require_pilot_review: bool = True) -> None:
        self.root = root
        self.stage = stage
        self.stage_index = STAGES.index(stage)
        self.errors: list[str] = []
        self.docs: dict[str, Any] = {}
        self.visual_registry: dict[str, dict[str, Any]] = {}
        self.require_pilot_review = require_pilot_review
        self.pixel_cache = PIXELS.ImagePixels()

    def pilot_asset_ids(self) -> set[str]:
        plan = self.docs.get("pilot-plan.json", {})
        ids = plan.get("component_ids") if isinstance(plan, dict) else None
        return set(ids) if isinstance(ids, list) and all(isinstance(item, str) for item in ids) else set()

    def asset_ready(self, asset: dict[str, Any]) -> bool:
        return self.reaches("assets") or self.stage == "pilot" and asset.get("id") in self.pilot_asset_ids()

    def reaches(self, stage: str) -> bool:
        return self.stage_index >= STAGES.index(stage)

    def quality_profile(self) -> str | None:
        brief = self.docs.get("brief.json")
        if isinstance(brief, dict):
            value = brief.get("quality_profile")
            if isinstance(value, str):
                return value
        return None

    def semantic_system_active(self) -> bool:
        style = self.docs.get("style-contract.json")
        return isinstance(style, dict) and (
            style.get("contract_schema_version") == 2 or "semantic_system" in style
        )

    def required_gate_decisions(self, gate: str) -> set[str]:
        required = set(GATE_REQUIRED_DECISIONS[gate])
        if self.semantic_system_active():
            if gate == "style_anchors":
                required.add("semantic_system")
            elif gate == "concept_contact_sheet":
                required.add("cross_slide_semantics_review")
        return required

    def is_fixture_profile(self) -> bool:
        return self.quality_profile() == "validation-fixture"

    def pilot_required(self) -> bool:
        brief = self.docs.get("brief.json", {})
        formats = brief.get("delivery", {}).get("formats", []) if isinstance(brief, dict) else []
        return self.quality_profile() == "production" and "pptx" in formats

    def error(self, location: str, message: str) -> None:
        self.errors.append(f"{location}: {message}")

    def load_required_documents(self, overrides: dict | None = None) -> None:
        if (self.root / ".inventory-merge").exists():
            self.error("project", "inventory merge is in progress or interrupted; recover it before validation")
            return
        names: list[str] = []
        for stage in STAGES[: self.stage_index + 1]:
            names.extend(REQUIRED_BY_STAGE[stage])
        for name in dict.fromkeys(names):
            if overrides is not None and name in overrides:
                self.docs[name] = overrides[name]
                continue
            if name == "component-master-manifest.json" and self.is_fixture_profile():
                continue
            if name == "pilot-plan.json" and self.stage != "pilot" and not self.pilot_required():
                continue
            path = self.root / name
            if not path.is_file():
                self.error(name, "required artifact is missing")
                continue
            try:
                self.docs[name] = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, UnicodeError, json.JSONDecodeError) as exc:
                self.error(name, f"cannot read valid UTF-8 JSON ({exc})")

    def require_fields(self, doc_name: str, fields: list[str]) -> None:
        doc = self.docs.get(doc_name)
        if not isinstance(doc, dict):
            return
        for field in fields:
            if field not in doc or doc[field] in (None, "", [], {}):
                self.error(doc_name, f"missing non-empty field '{field}'")

    def validate_brief(self) -> None:
        self.require_fields(
            "brief.json",
            [
                "project_id",
                "audience",
                "goal",
                "quality_profile",
                "delivery",
                "editability",
                "capacity_budget",
            ],
        )
        doc = self.docs.get("brief.json")
        if not isinstance(doc, dict):
            return
        template_placeholders = {
            "project_id": "replace-with-project-id",
            "audience": "replace with the primary audience",
            "goal": "replace with the decision or change this deck should produce",
        }
        for field, placeholder in template_placeholders.items():
            if doc.get(field) == placeholder:
                self.error("brief.json", f"{field} still contains a template placeholder")
        quality_profile = doc.get("quality_profile")
        if quality_profile not in QUALITY_PROFILES:
            self.error(
                "brief.json",
                f"quality_profile must be one of {sorted(QUALITY_PROFILES)}",
            )
        state = self.docs.get("run-state.json", {})
        fixture_only = state.get("fixture_only") if isinstance(state, dict) else None
        if quality_profile == "validation-fixture" and fixture_only is not True:
            self.error(
                "run-state.json",
                "validation-fixture quality_profile requires fixture_only=true",
            )
        if fixture_only is True and quality_profile != "validation-fixture":
            self.error(
                "brief.json",
                "fixture_only=true requires quality_profile='validation-fixture'",
            )
        editability = doc.get("editability")
        delivery = doc.get("delivery", {})
        if isinstance(delivery, dict) and delivery.get("speaker_notes", "not-requested") not in ("not-requested", "required"):
            self.error("brief.json", "delivery.speaker_notes must be not-requested or required")
        if isinstance(editability, dict):
            expected = {
                "ordinary_text": "native-editable",
                "data": "native-editable",
                "visual_components": "component-editable",
            }
            for field, value in expected.items():
                if editability.get(field) != value:
                    self.error("brief.json", f"editability.{field} must be '{value}' for this workflow")
            if not isinstance(editability.get("internal_vector_paths"), bool):
                self.error("brief.json", "editability.internal_vector_paths must be boolean")
            elif editability.get("internal_vector_paths"):
                self.error(
                    "brief.json",
                    "editability.internal_vector_paths must be false for transparent PNG components",
                )
        budget = doc.get("capacity_budget")
        if isinstance(budget, dict):
            for field in ("component_limit", "generation_call_limit", "retry_limit_per_asset"):
                value = budget.get(field)
                if not isinstance(value, int) or value <= 0:
                    self.error("brief.json", f"capacity_budget.{field} must be a positive integer")
            if budget.get("retry_limit_per_asset", 99) > 3:
                self.error("brief.json", "capacity_budget.retry_limit_per_asset cannot exceed 3")
            if "deadline" not in budget:
                self.error("brief.json", "capacity_budget.deadline must be present, using null when open")
            deadline = budget.get("deadline")
            if deadline is not None:
                try:
                    datetime.fromisoformat(str(deadline))
                except ValueError:
                    self.error("brief.json", "capacity_budget.deadline must be ISO-8601 or null")
            if budget.get("fallback_delivery") not in {
                "component-library",
                "pptx",
                "pdf",
                "contact-sheet",
            }:
                self.error("brief.json", "capacity_budget.fallback_delivery is required")

    def validate_research(self) -> None:
        name = "research-bundle.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        mode = doc.get("mode")
        if mode not in {"external", "source-only", "not-required"}:
            self.error(name, "mode must be 'external', 'source-only', or 'not-required'")
        if mode == "external":
            self.require_fields(name, ["query_log", "sources", "claims"])
        elif mode == "source-only":
            self.require_fields(name, ["sources", "claims"])
        elif mode == "not-required" and not doc.get("reason"):
            self.error(name, "not-required mode requires a reason")
        source_ids = {
            source.get("id")
            for source in doc.get("sources", [])
            if isinstance(source, dict) and source.get("id")
        }
        self.check_unique(doc.get("sources", []), "id", f"{name}.sources")
        self.check_unique(doc.get("claims", []), "id", f"{name}.claims")
        for index, claim in enumerate(doc.get("claims", [])):
            if not isinstance(claim, dict):
                self.error(f"{name}.claims[{index}]", "must be an object")
                continue
            refs = claim.get("source_ids", [])
            if not refs:
                self.error(f"{name}.claims[{index}]", "claim requires source_ids")
            for source_id in refs:
                if source_id not in source_ids:
                    self.error(f"{name}.claims[{index}]", f"unknown source_id '{source_id}'")

    def validate_alignment(self) -> None:
        name = "alignment.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        material_review = doc.get("material_review")
        if not isinstance(material_review, dict):
            self.error(name, "material_review must be an object")
        else:
            if material_review.get("status") != "completed":
                self.error(name, "material_review.status must be 'completed'")
            source_material_provided = material_review.get("source_material_provided")
            if not isinstance(source_material_provided, bool):
                self.error(name, "material_review.source_material_provided must be boolean")
            inspected_sources = material_review.get("inspected_sources")
            if not isinstance(inspected_sources, list):
                self.error(name, "material_review.inspected_sources must be a list")
            elif source_material_provided and not inspected_sources:
                self.error(name, "material_review.inspected_sources must list supplied material")
            else:
                for index, source in enumerate(inspected_sources):
                    location = f"{name}.material_review.inspected_sources[{index}]"
                    if not isinstance(source, dict):
                        self.error(location, "must be an object")
                        continue
                    path = self.validate_relative_file(source.get("path"), f"{location}.path")
                    if path and source.get("sha256") != hashlib.sha256(path.read_bytes()).hexdigest():
                        self.error(location, "sha256 is stale or missing")
            if source_material_provided is False and not material_review.get("inspection_notes"):
                self.error(name, "material_review.inspection_notes must explain that no material was supplied")
        decisions = doc.get("decisions")
        if not isinstance(decisions, list):
            self.error(name, "decisions must be a list")
            return
        self.check_unique(decisions, "id", f"{name}.decisions")
        decision_map = {
            decision.get("id"): decision
            for decision in decisions
            if isinstance(decision, dict)
            and isinstance(decision.get("id"), str)
            and decision.get("id")
        }
        discovered_facts = doc.get("discovered_facts")
        discovered_decision_ids: set[str] = set()
        if not isinstance(discovered_facts, list):
            self.error(name, "discovered_facts must be a list")
        else:
            self.check_unique(discovered_facts, "id", f"{name}.discovered_facts")
            for index, fact in enumerate(discovered_facts):
                location = f"{name}.discovered_facts[{index}]"
                if not isinstance(fact, dict):
                    self.error(location, "must be an object")
                    continue
                decision_id = fact.get("decision_id")
                if not isinstance(decision_id, str) or not decision_id:
                    self.error(location, "decision_id must be a non-empty string")
                else:
                    discovered_decision_ids.add(decision_id)
                if "value" not in fact:
                    self.error(location, "value is required")
                if not isinstance(fact.get("source_refs"), list) or not fact.get("source_refs"):
                    self.error(location, "source_refs must be a non-empty list")
        frontier = doc.get("question_frontier")
        frontier_ids: list[str] = []
        if not isinstance(frontier, dict):
            self.error(name, "question_frontier must be an object")
        else:
            raw_frontier_ids = frontier.get("decision_ids")
            if not isinstance(raw_frontier_ids, list):
                self.error(name, "question_frontier.decision_ids must be a list")
            else:
                frontier_ids = [
                    decision_id
                    for decision_id in raw_frontier_ids
                    if isinstance(decision_id, str) and decision_id
                ]
                if len(frontier_ids) != len(raw_frontier_ids):
                    self.error(name, "question_frontier decision IDs must be non-empty strings")
                if len(frontier_ids) != len(set(frontier_ids)):
                    self.error(name, "question_frontier decision IDs must be unique")
                if len(frontier_ids) > 5:
                    self.error(name, "question_frontier cannot contain more than 5 decisions")
                available_frontier = [
                    decision.get("id")
                    for decision in decisions
                    if isinstance(decision, dict)
                    and isinstance(decision.get("id"), str)
                    and decision.get("id") not in discovered_decision_ids
                    and decision.get("status") in {"open", "ambiguous"}
                    and decision.get("independent") is True
                ]
                required_count = min(3, len(available_frontier))
                if len(frontier_ids) < required_count:
                    self.error(
                        name,
                        f"question_frontier must include at least {required_count} available decisions",
                    )
                expected_frontier = available_frontier[:5]
                if frontier_ids != expected_frontier:
                    self.error(
                        name,
                        "question_frontier must match the first up-to-five available decisions in ledger order",
                    )
                for decision_id in frontier_ids:
                    if decision_id in discovered_decision_ids:
                        self.error(
                            name,
                            f"discovered fact '{decision_id}' cannot remain in question_frontier",
                        )
                    decision = decision_map.get(decision_id)
                    if not isinstance(decision, dict):
                        self.error(name, f"question_frontier references unknown decision '{decision_id}'")
                        continue
                    if decision.get("status") not in {"open", "ambiguous"}:
                        self.error(name, f"question_frontier decision '{decision_id}' is not open")
                    if decision.get("independent") is not True:
                        self.error(name, f"question_frontier decision '{decision_id}' must be independent")
                    recommendation = decision.get("recommendation")
                    if not isinstance(recommendation, dict):
                        self.error(name, f"question_frontier decision '{decision_id}' needs a recommendation")
                    else:
                        if not has_meaningful_value(recommendation.get("value")):
                            self.error(name, f"question_frontier decision '{decision_id}' needs a recommended value")
                        for field in ("reason", "impact"):
                            if not recommendation.get(field):
                                self.error(
                                    name,
                                    f"question_frontier decision '{decision_id}' needs recommendation {field}",
                                )
        for decision_id in sorted(CRITICAL_G1_DECISIONS):
            decision = decision_map.get(decision_id)
            location = f"{name}.decisions.{decision_id}"
            if not isinstance(decision, dict):
                self.error(location, f"missing critical decision '{decision_id}'")
                continue
            value = decision.get("value")
            if decision.get("status") != "resolved" or not has_meaningful_value(value):
                self.error(location, f"critical decision '{decision_id}' is not resolved")

    def validate_capability_report(self) -> None:
        name = "capability-report.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        if not doc.get("probed_at"):
            self.error(name, "missing probed_at")
        if self.reaches("concept"):
            image_generation = doc.get("image_generation", {})
            if not isinstance(image_generation, dict) or image_generation.get("status") != "verified":
                self.error(name, "image_generation.status must be 'verified' before concept generation")
        if self.reaches("pilot"):
            transparent_output = doc.get("transparent_output", {})
            if not isinstance(transparent_output, dict) or transparent_output.get("status") != "verified":
                self.error(name, "transparent_output.status must be 'verified' before asset generation")
            else:
                sample = self.validate_relative_file(
                    transparent_output.get("sample_path"),
                    f"{name}.transparent_output.sample_path",
                )
                if sample:
                    png_ok, ratio, clear_border, visible = self.inspect_png(sample)
                    if not png_ok or ratio < 0.05 or not clear_border or not visible:
                        self.error(name, "transparent output sample must have real transparent padding and visible content")
        if self.reaches("assembly"):
            brief = self.docs.get("brief.json", {})
            formats = brief.get("delivery", {}).get("formats", []) if isinstance(brief, dict) else []
            assembly = doc.get("assembly", {})
            if "pptx" in formats and (
                not isinstance(assembly, dict) or assembly.get("status") != "verified"
            ):
                self.error(name, "assembly.status must be 'verified' when PPTX is requested")

    def validate_storyboard(self) -> None:
        name = "storyboard.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        slides = doc.get("slides")
        if not isinstance(slides, list) or not slides:
            self.error(name, "slides must be a non-empty list")
            return
        self.check_unique(slides, "id", f"{name}.slides")
        research = self.docs.get("research-bundle.json", {})
        claims = research.get("claims", []) if isinstance(research, dict) else []
        self.check_unique(claims, "id", "research-bundle.json.claims")
        claim_ids = {
            claim.get("id") for claim in claims if isinstance(claim, dict) and claim.get("id")
        }
        for index, slide in enumerate(slides):
            if not isinstance(slide, dict):
                self.error(f"{name}.slides[{index}]", "must be an object")
                continue
            for field in ("id", "title", "message", "narrative_role", "visual_intent"):
                if not slide.get(field):
                    self.error(f"{name}.slides[{index}]", f"missing '{field}'")
            if slide.get("narrative_role") not in NARRATIVE_ROLES:
                self.error(
                    f"{name}.slides[{index}]",
                    f"narrative_role must be one of {sorted(NARRATIVE_ROLES)}",
                )
            policy = slide.get("evidence_policy")
            evidence_ids = slide.get("evidence_ids")
            if policy == "required":
                if not isinstance(evidence_ids, list) or not evidence_ids:
                    self.error(f"{name}.slides[{index}]", "evidence_policy required needs evidence_ids")
            elif policy == "not-applicable":
                if not slide.get("evidence_not_required_reason"):
                    self.error(f"{name}.slides[{index}]", "not-applicable evidence needs a reason")
            else:
                self.error(f"{name}.slides[{index}]", "evidence_policy must be required or not-applicable")
            if isinstance(evidence_ids, list):
                for evidence_id in evidence_ids:
                    if evidence_id not in claim_ids:
                        self.error(f"{name}.slides[{index}]", f"unknown evidence_id '{evidence_id}'")

    def validate_reference_study(self) -> None:
        name = "reference-study.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        profile = self.quality_profile()
        status = doc.get("status")
        if profile == "validation-fixture" and status == "not-required":
            if not doc.get("reason"):
                self.error(name, "not-required fixture study needs a reason")
            return
        if status != "analyzed":
            self.error(name, "status must be 'analyzed' before concept generation")

        contact_sheet = doc.get("reference_contact_sheet")
        if not isinstance(contact_sheet, dict):
            self.error(name, "reference_contact_sheet is required")
        else:
            sheet_path = self.validate_relative_file(
                contact_sheet.get("path"), f"{name}.reference_contact_sheet.path"
            )
            declared_sha = contact_sheet.get("sha256")
            if not re.fullmatch(r"[0-9a-fA-F]{64}", str(declared_sha)):
                self.error(
                    f"{name}.reference_contact_sheet.sha256",
                    "must be a SHA-256 hex digest",
                )
            elif sheet_path and hashlib.sha256(sheet_path.read_bytes()).hexdigest() != str(
                declared_sha
            ).lower():
                self.error(
                    f"{name}.reference_contact_sheet.sha256",
                    "does not match the contact sheet",
                )

        references = doc.get("references")
        minimum = 3 if profile == "production" else 2
        if not isinstance(references, list) or not (minimum <= len(references) <= 6):
            self.error(
                name,
                f"{profile or 'unknown'} reference study requires {minimum}-6 references",
            )
            references = []
        self.check_unique(references, "id", f"{name}.references")
        reference_ids: set[str] = set()
        for index, reference in enumerate(references):
            location = f"{name}.references[{index}]"
            if not isinstance(reference, dict):
                self.error(location, "must be an object")
                continue
            reference_id = reference.get("id")
            if isinstance(reference_id, str) and reference_id:
                reference_ids.add(reference_id)
            for field in ("source_kind", "source_locator", "usage_note"):
                if not has_meaningful_value(reference.get(field)):
                    self.error(location, f"missing non-empty '{field}'")
            if reference.get("status") != "analyzed":
                self.error(location, "status must be 'analyzed'")
            image_path = self.validate_relative_file(
                reference.get("image_path"), f"{location}.image_path"
            )
            declared_sha = reference.get("sha256")
            if not re.fullmatch(r"[0-9a-fA-F]{64}", str(declared_sha)):
                self.error(f"{location}.sha256", "must be a SHA-256 hex digest")
            elif image_path and hashlib.sha256(image_path.read_bytes()).hexdigest() != str(
                declared_sha
            ).lower():
                self.error(f"{location}.sha256", "does not match the reference image")
            observations = reference.get("observations")
            if not isinstance(observations, dict):
                self.error(location, "observations must be an object")
            else:
                for field in REFERENCE_OBSERVATION_FIELDS:
                    if not has_meaningful_value(observations.get(field)):
                        self.error(location, f"observations.{field} must not be empty")
            for field in (
                "transferable_rules",
                "project_translation",
                "do_not_copy",
                "target_slide_roles",
            ):
                value = reference.get(field)
                if not isinstance(value, list) or not value:
                    self.error(location, f"{field} must be a non-empty list")

        selected_ids = doc.get("selected_reference_ids")
        if not isinstance(selected_ids, list) or not selected_ids:
            self.error(name, "selected_reference_ids must be a non-empty list")
            selected_ids = []
        elif any(
            not isinstance(reference_id, str) or reference_id not in reference_ids
            for reference_id in selected_ids
        ):
            self.error(name, "selected_reference_ids contains an unknown reference")

        directions = doc.get("direction_options")
        if not isinstance(directions, list) or not (2 <= len(directions) <= 3):
            self.error(name, "direction_options must contain 2-3 directions")
            directions = []
        self.check_unique(directions, "id", f"{name}.direction_options")
        direction_ids: set[str] = set()
        for index, direction in enumerate(directions):
            location = f"{name}.direction_options[{index}]"
            if not isinstance(direction, dict):
                self.error(location, "must be an object")
                continue
            direction_id = direction.get("id")
            if isinstance(direction_id, str) and direction_id:
                direction_ids.add(direction_id)
            for field in ("name", "structure", "palette", "density", "imagery"):
                if not has_meaningful_value(direction.get(field)):
                    self.error(location, f"missing non-empty '{field}'")
            direction_refs = direction.get("reference_ids")
            if not isinstance(direction_refs, list) or not direction_refs:
                self.error(location, "reference_ids must be a non-empty list")
            elif any(
                not isinstance(reference_id, str) or reference_id not in reference_ids
                for reference_id in direction_refs
            ):
                self.error(location, "reference_ids contains an unknown reference")
            tradeoffs = direction.get("tradeoffs")
            if not isinstance(tradeoffs, list) or not tradeoffs:
                self.error(location, "tradeoffs must be a non-empty list")
        if not isinstance(doc.get("recommended_direction_id"), str) or doc.get(
            "recommended_direction_id"
        ) not in direction_ids:
            self.error(name, "recommended_direction_id must identify a direction option")

    def validate_content_source_lineage(
        self, records: Any, location: str
    ) -> dict[str, dict[str, Any]]:
        if not isinstance(records, list) or not records:
            self.error(location, "source_lineage must contain approved or verified sources")
            return {}
        self.check_unique(records, "id", location)
        sources: dict[str, dict[str, Any]] = {}
        for index, source in enumerate(records):
            source_location = f"{location}[{index}]"
            if not isinstance(source, dict):
                self.error(source_location, "must be an object")
                continue
            source_id = source.get("id")
            kind = source.get("kind")
            if not isinstance(source_id, str) or not source_id.strip():
                self.error(source_location, "id must be a non-empty string")
                continue
            if not isinstance(kind, str) or not kind.strip():
                self.error(source_location, "kind must identify the source authority")
            if source.get("status") not in {"approved", "verified"}:
                self.error(source_location, "status must be approved or verified")
            if not isinstance(source.get("locator"), str) or not source["locator"].strip():
                self.error(source_location, "locator must identify the relevant source passage or region")
            path_value = source.get("path")
            path = self.validate_relative_file(path_value, f"{source_location}.path")
            digest = source.get("sha256")
            if not re.fullmatch(r"[0-9a-fA-F]{64}", str(digest)):
                self.error(source_location, "sha256 must be a SHA-256 hex digest")
            elif path and hashlib.sha256(path.read_bytes()).hexdigest() != str(digest).lower():
                self.error(source_location, "source_lineage sha256 does not match file")
            sources[source_id] = source
        return sources

    def validate_content_copy_map(
        self,
        copy_map: Any,
        sources: dict[str, dict[str, Any]],
        story_slide: Any,
        location: str,
        message_mapping: Any = None,
    ) -> set[str]:
        if not isinstance(copy_map, list) or not copy_map:
            self.error(location, "copy_map must contain every visible text item")
            return set()
        self.check_unique(copy_map, "id", location)
        used_sources: set[str] = set()
        copy_by_story_field: dict[str, list[dict[str, Any]]] = {"title": [], "message": []}
        for index, item in enumerate(copy_map):
            item_location = f"{location}[{index}]"
            if not isinstance(item, dict):
                self.error(item_location, "must be an object")
                continue
            if not isinstance(item.get("id"), str) or not item["id"].strip():
                self.error(item_location, "id must be a non-empty string")
            if not isinstance(item.get("role"), str) or not item["role"].strip():
                self.error(item_location, "role must be a non-empty string")
            text_value = item.get("text")
            if not isinstance(text_value, str) or not text_value.strip():
                self.error(item_location, "text must be the exact visible wording")
                text_value = ""
            truth_status = item.get("truth_status")
            if truth_status not in {
                "verified",
                "planned",
                "pending",
                "unverified",
                "not-applicable",
            }:
                self.error(item_location, "truth_status must distinguish verified, planned, pending, or not-applicable copy")
            elif truth_status != "not-applicable" and (
                not isinstance(item.get("status_label"), str)
                or not item["status_label"].strip()
            ):
                self.error(item_location, "factual copy needs a visible status_label")
            story_field = item.get("storyboard_field")
            if story_field is not None:
                if story_field not in {"title", "message"}:
                    self.error(item_location, "storyboard_field must be title or message")
                elif isinstance(story_slide, dict):
                    copy_by_story_field[story_field].append(item)
                    if story_field == "title" and text_value != story_slide.get("title"):
                        self.error(
                            item_location,
                            "storyboard title must match the approved storyboard exactly",
                        )
            evidence = item.get("sources")
            if not isinstance(evidence, list) or not evidence:
                self.error(item_location, "every copy item needs source IDs, locators, and excerpts")
                continue
            for source_index, reference in enumerate(evidence):
                ref_location = f"{item_location}.sources[{source_index}]"
                if not isinstance(reference, dict):
                    self.error(ref_location, "must be an object")
                    continue
                source_id = reference.get("source_id")
                source = sources.get(source_id) if isinstance(source_id, str) else None
                if not isinstance(source, dict):
                    self.error(ref_location, "source_id must identify a source_lineage record")
                    continue
                used_sources.add(source_id)
                if not isinstance(reference.get("locator"), str) or not reference["locator"].strip():
                    self.error(ref_location, "locator must identify the quoted passage or visual region")
                quote = reference.get("quote")
                if not isinstance(quote, str) or not quote.strip():
                    self.error(ref_location, "quote must preserve the supporting source text")
                    continue
                mode = reference.get("mode")
                if mode not in COPY_SOURCE_MODES:
                    self.error(ref_location, "mode must be storyboard-exact, source-exact, or approved-paraphrase")
                    continue
                if mode == "storyboard-exact":
                    field = item.get("storyboard_field")
                    if (
                        source.get("kind") != "approved-storyboard"
                        or field not in {"title", "message"}
                        or source.get("path") != "storyboard.json"
                        or quote != text_value
                        or not isinstance(story_slide, dict)
                        or quote != story_slide.get(field)
                    ):
                        self.error(ref_location, "storyboard-exact evidence must match the approved storyboard field")
                else:
                    source_path = self.root / str(source.get("path", ""))
                    if mode == "source-exact" and quote != text_value:
                        self.error(ref_location, "source-exact quote must match visible copy")
                    if source_path.is_file():
                        try:
                            source_text = source_path.read_text(encoding="utf-8")
                        except (UnicodeError, OSError):
                            source_text = None
                        if source_text is not None and quote not in source_text:
                            self.error(ref_location, "quoted source text is not present at the recorded source path")
        if isinstance(story_slide, dict):
            titles = copy_by_story_field["title"]
            messages = copy_by_story_field["message"]
            if len(titles) != 1:
                self.error(location, "copy_map must contain exactly one approved storyboard title")
            if not messages:
                self.error(location, "copy_map must map the approved storyboard message to visible copy")
            elif len(messages) == 1 and messages[0].get("text") == story_slide.get("message"):
                if message_mapping is not None:
                    self.error(location, "message_mapping is unnecessary when the exact storyboard message is visible")
            else:
                mapping_location = f"{location}.message_mapping"
                approved_message = story_slide.get("message")
                if not isinstance(message_mapping, dict):
                    self.error(mapping_location, "non-verbatim message copy needs an explicit mapping for human review")
                else:
                    if message_mapping.get("storyboard_message") != approved_message:
                        self.error(mapping_location, "storyboard_message must match the approved narrative claim exactly")
                    mapped_ids = message_mapping.get("copy_ids")
                    expected_ids = {item.get("id") for item in messages}
                    if (
                        not isinstance(mapped_ids, list)
                        or not mapped_ids
                        or not all(isinstance(value, str) and value for value in mapped_ids)
                        or len(mapped_ids) != len(set(mapped_ids))
                        or set(mapped_ids) != expected_ids
                    ):
                        self.error(mapping_location, "copy_ids must identify every visible message realization exactly once")
                    if not isinstance(message_mapping.get("review_note"), str) or not message_mapping["review_note"].strip():
                        self.error(mapping_location, "review_note must explain the semantic coverage and truth boundary")
                    if message_mapping.get("human_gate") != "G3/G4-required":
                        self.error(mapping_location, "non-verbatim message realization requires G3/G4 human review")
        return used_sources

    def validate_concept_image_inventory(
        self,
        records: Any,
        sources: dict[str, dict[str, Any]],
        location: str,
    ) -> set[str]:
        if not isinstance(records, list) or not records:
            self.error(location, "image_inventory must identify representative imagery")
            return set()
        self.check_unique(records, "id", location)
        used_sources: set[str] = set()
        has_representative = False
        for index, visual in enumerate(records):
            visual_location = f"{location}[{index}]"
            if not isinstance(visual, dict):
                self.error(visual_location, "must be an object")
                continue
            if not isinstance(visual.get("id"), str) or not visual["id"].strip():
                self.error(visual_location, "id must be a non-empty string")
            if not isinstance(visual.get("visual_role"), str) or not visual["visual_role"].strip():
                self.error(visual_location, "visual_role must describe the image's role")
            has_representative = has_representative or visual.get("visual_role") == "representative"
            classification = visual.get("classification")
            if classification not in CONCEPT_IMAGE_CLASSIFICATIONS:
                self.error(visual_location, "classification must be illustrative, verified-source, or pending-placeholder")
                continue
            if classification == "illustrative":
                if not isinstance(visual.get("label"), str) or not visual["label"].strip():
                    self.error(visual_location, "illustrative images need an explicit illustrative label")
                if visual.get("evidence_role") != "illustration-only":
                    self.error(visual_location, "illustrative images cannot be presented as evidence")
            elif classification == "pending-placeholder":
                if not isinstance(visual.get("label"), str) or not visual["label"].strip():
                    self.error(visual_location, "pending imagery needs an explicit placeholder label")
                if visual.get("evidence_role") != "placeholder-only":
                    self.error(visual_location, "pending imagery must remain a placeholder")
            else:
                source_id = visual.get("source_id")
                source = sources.get(source_id) if isinstance(source_id, str) else None
                if not isinstance(source, dict) or source.get("kind") not in {
                    "verified-image",
                    "authoritative-image",
                }:
                    self.error(visual_location, "verified-source imagery must link a verified local image source")
                else:
                    used_sources.add(source_id)
                    if visual.get("source_sha256") != source.get("sha256"):
                        self.error(visual_location, "verified-source image hash must match source_lineage")
                    if visual.get("preservation") != "unchanged":
                        self.error(visual_location, "verified-source imagery must be preserved unchanged")
        if not has_representative:
            self.error(location, "image_inventory needs a representative visual")
        return used_sources

    def validate_content_text_review(
        self,
        review: Any,
        copy_map: Any,
        output_path: str | None,
        output_sha256: str | None,
        location: str,
    ) -> set[str]:
        if not isinstance(review, dict):
            self.error(location, "text_review must link exact text to a reviewed transcript")
            return set()
        if review.get("method") not in {"human-transcription", "ocr-transcript-reviewed"}:
            self.error(location, "method must identify a human-checked transcript")
        expected_copy_hash = value_sha256(copy_map)
        if review.get("copy_map_sha256") != expected_copy_hash:
            self.error(location, "copy_map_sha256 does not match the exact visible copy")
        transcript_path_value = review.get("transcript_path")
        transcript_path = self.validate_relative_file(
            transcript_path_value, f"{location}.transcript_path"
        )
        transcript_hash = review.get("transcript_sha256")
        transcript_text: str | None = None
        if not re.fullmatch(r"[0-9a-fA-F]{64}", str(transcript_hash)):
            self.error(location, "transcript_sha256 must be a SHA-256 hex digest")
        elif transcript_path:
            actual_hash = hashlib.sha256(transcript_path.read_bytes()).hexdigest()
            if actual_hash != str(transcript_hash).lower():
                self.error(location, "transcript_sha256 does not match transcript evidence")
            try:
                transcript_text = transcript_path.read_text(encoding="utf-8")
            except (UnicodeError, OSError):
                self.error(location, "transcript evidence must be readable UTF-8 text")
        copy_records = {
            item.get("id"): item
            for item in copy_map
            if isinstance(item, dict) and isinstance(item.get("id"), str)
        } if isinstance(copy_map, list) else {}
        observations = review.get("observations")
        if not isinstance(observations, list):
            self.error(location, "observations must transcribe every copy_map item")
            observations = []
        self.check_unique(observations, "copy_id", f"{location}.observations")
        observed_ids: set[str] = set()
        for index, observation in enumerate(observations):
            observation_location = f"{location}.observations[{index}]"
            if not isinstance(observation, dict):
                self.error(observation_location, "must be an object")
                continue
            copy_id = observation.get("copy_id")
            expected = copy_records.get(copy_id)
            observed_text = observation.get("observed_text")
            if not isinstance(expected, dict):
                self.error(observation_location, "copy_id must identify a copy_map item")
                continue
            if observed_text != expected.get("text"):
                self.error(observation_location, "observed_text must exactly match copy_map")
            if isinstance(transcript_text, str) and isinstance(observed_text, str) and observed_text not in transcript_text:
                self.error(observation_location, "observed_text must appear in transcript evidence")
            status_label = expected.get("status_label")
            if status_label is not None:
                if observation.get("observed_status_label") != status_label:
                    self.error(observation_location, "observed_status_label must exactly match copy_map")
                if isinstance(transcript_text, str) and status_label not in transcript_text:
                    self.error(observation_location, "observed_status_label must appear in transcript evidence")
            if isinstance(copy_id, str):
                observed_ids.add(copy_id)
        if observed_ids != set(copy_records):
            self.error(location, "text_review observations must cover every copy_map item")
        evidence_paths = {transcript_path_value} if isinstance(transcript_path_value, str) else set()
        repair = review.get("repair")
        if repair is not None:
            repair_location = f"{location}.repair"
            if not isinstance(repair, dict) or repair.get("kind") != "deterministic-overlay":
                self.error(repair_location, "repair must record deterministic-overlay evidence")
                return evidence_paths
            if not isinstance(repair.get("tool"), str) or not repair["tool"].strip():
                self.error(repair_location, "tool must identify the deterministic overlay method")
            input_value = repair.get("input_path")
            input_path = self.validate_relative_file(input_value, f"{repair_location}.input_path")
            if input_value == output_path:
                self.error(repair_location, "input_path must preserve a prior image version")
            if input_path and hashlib.sha256(input_path.read_bytes()).hexdigest() != repair.get("input_sha256"):
                self.error(repair_location, "input_sha256 does not match preserved prior image")
            if repair.get("output_path") != output_path or repair.get("output_sha256") != output_sha256:
                self.error(repair_location, "overlay output must match the contentful concept image")
            evidence_value = repair.get("evidence_path")
            evidence_path = self.validate_relative_file(evidence_value, f"{repair_location}.evidence_path")
            if isinstance(evidence_value, str):
                evidence_paths.add(evidence_value)
            if evidence_path and hashlib.sha256(evidence_path.read_bytes()).hexdigest() != repair.get("evidence_sha256"):
                self.error(repair_location, "evidence_sha256 does not match overlay evidence")
        return evidence_paths

    def validate_concept_calibration(self) -> None:
        name = "concept-calibration.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        if self.is_fixture_profile() and doc.get("status") == "not-required":
            if not doc.get("reason"):
                self.error(name, "not-required fixture calibration needs a reason")
            return
        if doc.get("status") != "generated":
            self.error(name, "status must be 'generated'")
        storyboard = self.docs.get("storyboard.json", {})
        slides = storyboard.get("slides", []) if isinstance(storyboard, dict) else []
        slide = next(
            (item for item in slides if isinstance(item, dict) and item.get("id") == doc.get("slide_id")),
            None,
        )
        fixed_content = doc.get("fixed_content")
        if not isinstance(fixed_content, dict):
            self.error(name, "fixed_content must be an object")
            fixed_content = {}
        fixed_fill_slots = doc.get("fixed_fill_slots")
        if not isinstance(fixed_fill_slots, list) or not fixed_fill_slots:
            self.error(name, "fixed_fill_slots must list the shared editable blank regions")
        if slide is None:
            self.error(name, "slide_id must identify a storyboard slide")
        else:
            for field in ("title", "message", "evidence_ids", "visual_intent"):
                if fixed_content.get(field) != slide.get(field):
                    self.error(name, f"fixed_content.{field} must match the storyboard")
        study = self.docs.get("reference-study.json", {})
        direction_ids = {
            item.get("id") for item in study.get("direction_options", [])
            if isinstance(item, dict)
        } if isinstance(study, dict) else set()
        references = {
            item.get("id"): item for item in study.get("references", [])
            if isinstance(item, dict)
        } if isinstance(study, dict) else {}
        board = doc.get("contact_sheet")
        if not isinstance(board, dict):
            self.error(name, "contact_sheet is required")
            board = {}
        tracked_paths = [name]
        for location, item in [(f"{name}.contact_sheet", board)]:
            path = self.validate_relative_file(item.get("path"), f"{location}.path")
            if path:
                tracked_paths.append(item["path"])
                if item.get("sha256") != hashlib.sha256(path.read_bytes()).hexdigest():
                    self.error(location, "sha256 does not match file")
        source_records: dict[str, dict[str, Any]] = {}
        copy_map = doc.get("copy_map")
        if self.quality_profile() in {"production", "draft"}:
            source_records = self.validate_content_source_lineage(
                doc.get("source_lineage"), f"{name}.source_lineage"
            )
            self.validate_content_copy_map(
                copy_map, source_records, slide, f"{name}.copy_map", doc.get("message_mapping")
            )
            for source in source_records.values():
                if isinstance(source.get("path"), str):
                    tracked_paths.append(source["path"])
        variants = doc.get("variants")
        if not isinstance(variants, list) or not 2 <= len(variants) <= 3:
            self.error(name, "variants must contain 2-3 real concept images")
            variants = []
        self.check_unique(variants, "id", f"{name}.variants")
        ids = set()
        comparisons = set()
        for index, variant in enumerate(variants):
            location = f"{name}.variants[{index}]"
            if not isinstance(variant, dict):
                self.error(location, "must be an object")
                continue
            variant_id = variant.get("id")
            if not isinstance(variant_id, str) or not variant_id:
                self.error(location, "id must be a non-empty string")
            else:
                ids.add(variant_id)
            if variant.get("direction_id") not in direction_ids:
                self.error(location, "direction_id must identify a reference-study direction")
            if variant.get("style_strength") not in {"restrained", "balanced", "expressive"}:
                self.error(location, "style_strength is invalid")
            if variant.get("layout_density") not in {"airy", "standard", "dense"}:
                self.error(location, "layout_density is invalid")
            comparisons.add((variant.get("direction_id"), variant.get("style_strength"), variant.get("layout_density")))
            if variant.get("tool") not in {"codex-image_gen", "image_gen"}:
                self.error(location, "tool must identify Codex image_gen")
            prompt = variant.get("prompt")
            if not isinstance(prompt, str) or len(prompt.strip()) < 80:
                self.error(location, "prompt must preserve the exact structured prompt")
            if not isinstance(variant.get("generated_at"), str) or not variant.get("generated_at"):
                self.error(location, "generated_at is required")
            attached = variant.get("reference_images")
            if not isinstance(attached, list) or not attached:
                self.error(location, "reference_images must be a non-empty list")
                attached = []
            for ref_index, ref in enumerate(attached):
                ref_location = f"{location}.reference_images[{ref_index}]"
                if not isinstance(ref, dict) or ref.get("reference_id") not in references:
                    self.error(ref_location, "unknown reference_id")
                    continue
                source = references[ref["reference_id"]]
                if ref.get("path") != source.get("image_path") or ref.get("sha256") != source.get("sha256"):
                    self.error(ref_location, "reference path/hash must match reference-study")
            path = self.validate_relative_file(variant.get("output_path"), f"{location}.output_path")
            if path:
                tracked_paths.append(variant["output_path"])
                if variant.get("output_sha256") != hashlib.sha256(path.read_bytes()).hexdigest():
                    self.error(location, "output_sha256 does not match concept image")
            if self.quality_profile() in {"production", "draft"}:
                self.validate_concept_image_inventory(
                    variant.get("image_inventory"),
                    source_records,
                    f"{location}.image_inventory",
                )
                tracked_paths.extend(
                    self.validate_content_text_review(
                        variant.get("text_review"),
                        copy_map,
                        variant.get("output_path"),
                        variant.get("output_sha256"),
                        f"{location}.text_review",
                    )
                )
        if len(comparisons) < 2:
            self.error(name, "variants must visibly vary direction, strength, or density")
        if doc.get("recommended_variant_id") not in ids:
            self.error(name, "recommended_variant_id must identify a variant")
        recommendation = doc.get("recommendation")
        if not isinstance(recommendation, dict) or not recommendation.get("reason") or not recommendation.get("impact"):
            self.error(name, "recommendation needs reason and impact")
        state = self.docs.get("run-state.json", {})
        gate = state.get("gates", {}).get("style_anchors", {}) if isinstance(state, dict) else {}
        approval = gate.get("calibration_selection") if isinstance(gate, dict) else None
        location = "run-state.json.gates.style_anchors.calibration_selection"
        if not isinstance(approval, dict) or approval.get("status") != "approved":
            self.error(location, "human calibration selection must be approved before anchors")
            return
        actor_error = REVIEW_POLICY.review_actor_error(self.docs.get("run-state.json"), approval, human_only=True)
        if actor_error:
            self.error(location, actor_error)
        if approval.get("selected_variant_id") not in ids:
            self.error(location, "selected_variant_id must identify a generated variant")
        style = self.docs.get("style-contract.json", {})
        if isinstance(style, dict) and approval.get("selected_variant_id") != style.get("selected_calibration_variant_id"):
            self.error(location, "selected variant differs from style-contract")
        for field in ("approver", "approved_at", "artifact_version", "approval_phrase"):
            if not isinstance(approval.get(field), str) or not approval.get(field).strip():
                self.error(location, f"{field} is required")
        try:
            approved_at = approval.get("approved_at")
            parsed_time = datetime.fromisoformat(approved_at) if isinstance(approved_at, str) else None
            if parsed_time is None or parsed_time.utcoffset() is None:
                self.error(location, "approved_at must include timezone offset")
        except ValueError:
            self.error(location, "approved_at must be ISO-8601")
        if approval.get("artifact_version") != doc.get("version"):
            self.error(location, "artifact_version differs from concept-calibration")
        mode = approval.get("approval_mode")
        phrase = approval.get("approval_phrase", "")
        if mode in {"explicit", "local-affirmative"}:
            check_reply = REVIEW_POLICY.local_reply_error if mode == "local-affirmative" else REVIEW_POLICY.explicit_reply_error
            reply_error = check_reply(phrase, self.docs.get("run-state.json"))
            if reply_error:
                self.error(location, reply_error)
        elif mode == "batch-recommendations":
            if (phrase != "全部按推荐" or approval.get("selected_variant_id") != doc.get("recommended_variant_id") or approval.get("recommendations_visible") is not True):
                self.error(location, "batch approval must select the visible recommended variant")
        else:
            self.error(location, "approval_mode must be explicit, local-affirmative, or batch-recommendations")
        hashes = approval.get("approved_artifacts")
        if not isinstance(hashes, dict) or set(tracked_paths) - set(hashes):
            self.error(location, "approved_artifacts must include calibration, board, and every variant")
        else:
            for path_value in tracked_paths:
                path = self.validate_relative_file(path_value, f"{location}.approved_artifacts")
                if path and hashes.get(path_value) != hashlib.sha256(path.read_bytes()).hexdigest():
                    self.error(location, f"approved artifact hash is stale: {path_value}")

    def validate_readability_profile(self, profile: Any) -> None:
        location = "style-contract.json.readability_profile"
        if not isinstance(profile, dict):
            self.error(location, "readability_profile is required")
            return
        for field, floor in READABILITY_FONT_FLOORS.values():
            value = profile.get(field)
            if isinstance(value, bool) or not isinstance(value, (int, float)) or value < floor:
                self.error(location, f"{field} must be at least {floor} pt")
        for field, floor in READABILITY_VISUAL_FLOORS.items():
            value = profile.get(field)
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not floor <= value < 1:
                self.error(location, f"{field} must be at least {floor} and less than 1")

    def validate_semantic_system(self, system: Any) -> None:
        location = "style-contract.json.semantic_system"
        if not isinstance(system, dict):
            self.error(location, "semantic_system is required for contract_schema_version 2")
            return
        if not isinstance(system.get("version"), str) or not system["version"].strip():
            self.error(location, "version must be a non-empty string")
        text_roles = system.get("text_roles")
        if not isinstance(text_roles, dict) or not text_roles:
            self.error(location, "text_roles must be a non-empty role map")
            text_roles = {}
        for role_id, token in text_roles.items():
            role_location = f"{location}.text_roles[{role_id}]"
            if not isinstance(role_id, str) or not role_id or not isinstance(token, dict):
                self.error(role_location, "role ID and token must be valid")
                continue
            if not isinstance(token.get("meaning"), str) or not token["meaning"].strip():
                self.error(role_location, "meaning is required")
            readability_role = token.get("readability_role")
            if not isinstance(readability_role, str) or readability_role not in READABILITY_FONT_FLOORS:
                self.error(role_location, "readability_role is invalid")
            if not isinstance(token.get("font_family"), str) or not token["font_family"].strip():
                self.error(role_location, "font_family is required")
            font_size = token.get("font_size_pt")
            if isinstance(font_size, bool) or not isinstance(font_size, (int, float)) or font_size <= 0:
                self.error(role_location, "font_size_pt must be positive")
            elif readability_role in READABILITY_FONT_FLOORS:
                floor_key, baseline = READABILITY_FONT_FLOORS[readability_role]
                profile = self.docs.get("style-contract.json", {}).get("readability_profile", {})
                floor = profile.get(floor_key, baseline) if isinstance(profile, dict) else baseline
                if isinstance(floor, (int, float)) and font_size < floor:
                    self.error(role_location, "font_size_pt is below the approved readability floor")
            if type(token.get("bold")) is not bool:
                self.error(role_location, "bold must be boolean")
            if not isinstance(token.get("color"), str) or not re.fullmatch(r"#[0-9A-Fa-f]{6}", token["color"]):
                self.error(role_location, "color must be #RRGGBB")
        frame_roles = system.get("frame_roles")
        if not isinstance(frame_roles, dict) or not frame_roles:
            self.error(location, "frame_roles must be a non-empty role map")
            frame_roles = {}
        for role_id, token in frame_roles.items():
            role_location = f"{location}.frame_roles[{role_id}]"
            if not isinstance(role_id, str) or not role_id or not isinstance(token, dict):
                self.error(role_location, "role ID and token must be valid")
                continue
            if not isinstance(token.get("meaning"), str) or not token["meaning"].strip():
                self.error(role_location, "meaning is required")
            line_style = token.get("line_style")
            if not isinstance(line_style, str) or line_style not in {"solid", "dashed", "none"}:
                self.error(role_location, "line_style must be solid, dashed, or none")
            color = token.get("color")
            if line_style == "none" and color is not None:
                self.error(role_location, "a no-line frame must not declare a line color")
            elif line_style in {"solid", "dashed"} and (
                not isinstance(color, str) or not re.fullmatch(r"#[0-9A-Fa-f]{6}", color)
            ):
                self.error(role_location, "visible frame color must be #RRGGBB")
            if type(token.get("visible_label_required")) is not bool:
                self.error(role_location, "visible_label_required must be boolean")
        verified, pending = frame_roles.get("verified"), frame_roles.get("pending")
        if isinstance(verified, dict) and isinstance(pending, dict) and (
            verified.get("line_style"), verified.get("color")
        ) == (pending.get("line_style"), pending.get("color")):
            self.error(location, "verified and pending frame treatments must be distinguishable")
        repeated = system.get("repeated_regions")
        if not isinstance(repeated, dict) or not repeated:
            self.error(location, "repeated_regions must be a non-empty region map")
            repeated = {}
        for region_id, region in repeated.items():
            region_location = f"{location}.repeated_regions[{region_id}]"
            if not isinstance(region_id, str) or not region_id or not isinstance(region, dict):
                self.error(region_location, "region ID and definition must be valid")
                continue
            text_role = region.get("text_role")
            if not isinstance(text_role, str) or text_role not in text_roles:
                self.error(region_location, "text_role must name an approved text role")
            if type(region.get("required")) is not bool:
                self.error(region_location, "required must be boolean")
            anchor = region.get("anchor")
            if not isinstance(anchor, list) or len(anchor) != 2 or any(
                isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= value <= 1
                for value in anchor
            ):
                self.error(region_location, "anchor must be normalized [x,y]")
            tolerance = region.get("tolerance")
            if isinstance(tolerance, bool) or not isinstance(tolerance, (int, float)) or not 0 <= tolerance <= 0.1:
                self.error(region_location, "tolerance must be between 0 and 0.1")

    def validate_slide_semantics(self, slide: dict[str, Any], fill_slots: list[Any], location: str) -> None:
        style = self.docs.get("style-contract.json", {})
        system = style.get("semantic_system") if isinstance(style, dict) else None
        if not isinstance(system, dict):
            return
        text_roles = system.get("text_roles", {})
        frame_roles = system.get("frame_roles", {})
        repeated = system.get("repeated_regions", {})
        if not all(isinstance(value, dict) for value in (text_roles, frame_roles, repeated)):
            return
        semantic = slide.get("semantic_design")
        semantic_location = f"{location}.semantic_design"
        if not isinstance(semantic, dict):
            self.error(semantic_location, "semantic_design is required")
            return
        slots = {slot.get("id"): slot for slot in fill_slots if isinstance(slot, dict) and isinstance(slot.get("id"), str)}
        for slot_id, slot in slots.items():
            if slot.get("kind") == "ordinary-data":
                frame_token = frame_roles.get(slot.get("frame_role")) if isinstance(slot.get("frame_role"), str) else None
                if not isinstance(frame_token, dict) or frame_token.get("line_style") == "none" or frame_token.get("visible_label_required") is True:
                    self.error(f"{semantic_location}.data[{slot_id}]", "blank data slot needs a declared non-status visible frame_role")
        plan = slide.get("readability", {})
        regions = plan.get("text_regions", []) if isinstance(plan, dict) else []
        text_regions = {region.get("fill_slot_id"): region for region in regions if isinstance(region, dict)} if isinstance(regions, list) else {}
        for slot_id, region in text_regions.items():
            region_location = f"{semantic_location}.text[{slot_id}]"
            token_id = region.get("font_role")
            token = text_roles.get(token_id) if isinstance(token_id, str) else None
            if not isinstance(token, dict):
                self.error(region_location, "font_role must name an approved text role")
                continue
            if region.get("role") != token.get("readability_role"):
                self.error(region_location, "readability role differs from semantic text role")
            if region.get("planned_font_pt") != token.get("font_size_pt"):
                self.error(region_location, "planned_font_pt differs from fixed semantic text role")
            if region.get("role") == "action-title":
                content = slots.get(slot_id, {}).get("content")
                if isinstance(content, str) and content.strip() and content != slide.get("title"):
                    self.error(region_location, "action-title copy must match the approved slide title")
        raw_bindings = semantic.get("region_slots")
        if not isinstance(raw_bindings, dict):
            self.error(semantic_location, "region_slots must be an object")
            raw_bindings = {}
        exceptions = semantic.get("exceptions")
        if not isinstance(exceptions, list):
            self.error(semantic_location, "exceptions must be a list")
            exceptions = []
        exception_ids: set[str] = set()
        for index, exception in enumerate(exceptions):
            exception_location = f"{semantic_location}.exceptions[{index}]"
            region_id = exception.get("region_id") if isinstance(exception, dict) else None
            if not isinstance(region_id, str) or region_id not in repeated:
                self.error(exception_location, "region_id must name a repeated region")
                continue
            if not isinstance(exception.get("reason"), str) or not exception["reason"].strip():
                self.error(exception_location, "a concrete anchor-exception reason is required")
            if exception["region_id"] in exception_ids:
                self.error(exception_location, "duplicate region exception")
            exception_ids.add(exception["region_id"])
        for region_id, definition in repeated.items():
            if not isinstance(definition, dict):
                continue
            slot_id = raw_bindings.get(region_id)
            region_location = f"{semantic_location}.region_slots[{region_id}]"
            if slot_id is None:
                if definition.get("required") is True:
                    self.error(region_location, "required repeated region needs a fill slot")
                continue
            slot = slots.get(slot_id)
            if not isinstance(slot, dict) or slot.get("kind") != "ordinary-text":
                self.error(region_location, "repeated region must bind an ordinary-text fill slot")
                continue
            region = text_regions.get(slot_id)
            if not isinstance(region, dict) or region.get("font_role") != definition.get("text_role"):
                self.error(region_location, "repeated region uses the wrong semantic text role")
            if region_id != "page-title" and definition.get("required") is True and not str(slot.get("content", "")).strip():
                self.error(region_location, "required repeated region needs visible copy")
            anchor = definition.get("anchor")
            tolerance = definition.get("tolerance")
            bbox = slot.get("bbox")
            if (
                region_id not in exception_ids and isinstance(anchor, list) and len(anchor) == 2
                and isinstance(tolerance, (int, float)) and not isinstance(tolerance, bool)
                and isinstance(bbox, list) and len(bbox) == 4
                and all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in bbox[:2])
                and any(abs(bbox[index] - anchor[index]) > tolerance for index in (0, 1))
            ):
                self.error(region_location, "repeated region anchor differs from approved semantic system")
        for region_id in raw_bindings:
            if region_id not in repeated:
                self.error(semantic_location, f"unknown repeated region '{region_id}'")
        frames = semantic.get("frame_instances")
        if not isinstance(frames, list):
            self.error(semantic_location, "frame_instances must be a list")
            return
        self.check_unique(frames, "id", f"{semantic_location}.frame_instances")
        for index, frame in enumerate(frames):
            frame_location = f"{semantic_location}.frame_instances[{index}]"
            if not isinstance(frame, dict):
                self.error(frame_location, "must be an object")
                continue
            frame_role = frame.get("role")
            token = frame_roles.get(frame_role) if isinstance(frame_role, str) else None
            if not isinstance(token, dict):
                self.error(frame_location, "role must name an approved frame role")
                continue
            self.validate_bbox(frame.get("bbox"), f"{frame_location}.bbox")
            if token.get("visible_label_required") is True:
                label = slots.get(frame.get("label_slot_id"))
                if not isinstance(label, dict) or label.get("kind") != "ordinary-text" or not str(label.get("content", "")).strip():
                    self.error(frame_location, "status/evidence frame needs a visible ordinary-text label slot")

    def validate_slide_readability(
        self, slide: dict[str, Any], fill_slots: list[Any], profile: Any, location: str
    ) -> None:
        plan = slide.get("readability")
        review_location = f"{location}.readability"
        if not isinstance(plan, dict):
            self.error(review_location, "readability plan is required")
            return
        if not isinstance(profile, dict):
            return
        slots = {
            slot.get("id"): slot
            for slot in fill_slots
            if isinstance(slot, dict) and isinstance(slot.get("id"), str)
        }
        ordinary_ids = {
            slot_id for slot_id, slot in slots.items()
            if slot.get("kind") in {"ordinary-text", "ordinary-data"}
        }
        regions = plan.get("text_regions")
        if not isinstance(regions, list) or not regions:
            self.error(review_location, "text_regions must cover visible ordinary copy")
            regions = []
        self.check_unique(regions, "fill_slot_id", f"{review_location}.text_regions")
        represented: set[str] = set()
        violations: set[str] = set()
        for index, region in enumerate(regions):
            region_location = f"{review_location}.text_regions[{index}]"
            if not isinstance(region, dict):
                self.error(region_location, "must be an object")
                continue
            slot_id = region.get("fill_slot_id")
            if not isinstance(slot_id, str) or slot_id not in ordinary_ids:
                self.error(region_location, "fill_slot_id must identify an ordinary text/data fill slot")
                continue
            represented.add(slot_id)
            role = region.get("role")
            if role not in READABILITY_FONT_FLOORS:
                self.error(region_location, "role must be action-title, body, key-caption, or source-qualifier")
                continue
            font = region.get("planned_font_pt")
            floor_key, baseline = READABILITY_FONT_FLOORS[role]
            floor = profile.get(floor_key, baseline)
            if isinstance(font, bool) or not isinstance(font, (int, float)) or font <= 0:
                self.error(region_location, "planned_font_pt must be a positive number")
            elif isinstance(floor, (int, float)) and font < floor:
                violations.add(f"font:{slot_id}")
            line_count = region.get("line_count")
            if isinstance(line_count, bool) or not isinstance(line_count, int) or line_count < 1:
                self.error(region_location, "line_count must be a positive integer")
        if represented != ordinary_ids:
            self.error(review_location, f"text_regions must cover ordinary fill slots: {sorted(ordinary_ids - represented)}")
        visual = plan.get("main_visual")
        if not isinstance(visual, dict):
            self.error(review_location, "main_visual is required")
        else:
            kind = visual.get("kind")
            if kind not in {"evidence", "fine-ui", "illustration", "not-applicable"}:
                self.error(review_location, "main_visual.kind is invalid")
            elif kind == "not-applicable":
                if not isinstance(visual.get("reason"), str) or not visual["reason"].strip():
                    self.error(review_location, "not-applicable main_visual needs a reason")
            else:
                bbox = visual.get("bbox")
                self.validate_bbox(bbox, f"{review_location}.main_visual.bbox")
                if isinstance(bbox, list) and len(bbox) == 4 and all(
                    isinstance(value, (int, float)) and not isinstance(value, bool) for value in bbox
                ):
                    if kind in {"evidence", "fine-ui"} and not any(
                        isinstance(slot, dict)
                        and slot.get("kind") == "fixed-evidence"
                        and slot.get("bbox") == bbox
                        for slot in fill_slots
                    ):
                        self.error(review_location, "evidence main_visual bbox must match a fixed-evidence fill slot")
                    if kind == "illustration":
                        width_key, height_key = "illustration_min_width", "illustration_min_height"
                    elif kind == "fine-ui" and visual.get("detail_mode") == "full-interface":
                        width_key, height_key = "fine_ui_preferred_width", "fine_ui_preferred_height"
                    else:
                        width_key, height_key = "evidence_min_width", "evidence_min_height"
                    if kind == "fine-ui" and visual.get("detail_mode") not in {"full-interface", "focus-crop"}:
                        self.error(review_location, "fine-ui detail_mode must be full-interface or focus-crop")
                    width_floor = profile.get(width_key)
                    height_floor = profile.get(height_key)
                    if isinstance(width_floor, (int, float)) and isinstance(height_floor, (int, float)) and (
                        bbox[2] < width_floor or bbox[3] < height_floor
                    ):
                        violations.add("visual:main")
        exception = plan.get("exception")
        if exception is not None:
            if not isinstance(exception, dict):
                self.error(review_location, "exception must be an object or null")
            else:
                declared = exception.get("violations")
                if not isinstance(declared, list) or any(not isinstance(value, str) for value in declared):
                    self.error(review_location, "exception.violations must list exact floor IDs")
                    declared = []
                for field in ("reason", "alternative_viewing_route"):
                    if not isinstance(exception.get(field), str) or not exception[field].strip():
                        self.error(review_location, f"exception.{field} is required")
                if violations - set(declared):
                    self.error(review_location, f"underfloor regions need an explicit exception: {sorted(violations - set(declared))}")
        elif violations:
            self.error(review_location, f"underfloor regions need an explicit exception: {sorted(violations)}")
        if self.reaches("concept"):
            full_review = plan.get("full_size_review")
            full_location = f"{review_location}.full_size_review"
            if not isinstance(full_review, dict):
                self.error(full_location, "full_size_review is required for the current contentful concept")
                return
            if full_review.get("status") != "passed" or full_review.get("display_scale_percent") != 100:
                self.error(full_location, "full_size_review must pass at 100% display scale")
            if not isinstance(full_review.get("reviewer"), str) or not full_review["reviewer"].strip():
                self.error(full_location, "reviewer is required")
            evidence_path = self.validate_relative_file(full_review.get("path"), f"{full_location}.path")
            if evidence_path and full_review.get("sha256") != hashlib.sha256(evidence_path.read_bytes()).hexdigest():
                self.error(full_location, "full_size_review sha256 does not match evidence")
            concept_path_value = slide.get("concept_path")
            concept_path = self.root / concept_path_value if isinstance(concept_path_value, str) else None
            if concept_path and concept_path.is_file() and full_review.get("reviewed_concept_sha256") != hashlib.sha256(concept_path.read_bytes()).hexdigest():
                self.error(full_location, "full_size_review must bind the current contentful concept hash")

    def validate_style(self) -> None:
        self.require_fields(
            "style-contract.json",
            [
                "version",
                "direction",
                "primary_visual_family",
                "palette",
                "rules",
            ],
        )
        doc = self.docs.get("style-contract.json")
        if not isinstance(doc, dict):
            return
        reference_study = self.docs.get("reference-study.json", {})
        if isinstance(reference_study, dict) and reference_study.get("status") == "analyzed":
            if doc.get("selected_reference_ids") != reference_study.get(
                "selected_reference_ids"
            ):
                self.error(
                    "style-contract.json",
                    "selected_reference_ids must match reference-study.json",
                )
        calibration = self.docs.get("concept-calibration.json", {})
        if isinstance(calibration, dict) and calibration.get("status") == "generated":
            variants = calibration.get("variants", [])
            selected_id = doc.get("selected_calibration_variant_id")
            selected = next(
                (item for item in variants if isinstance(item, dict) and item.get("id") == selected_id),
                None,
            ) if isinstance(variants, list) else None
            if not isinstance(selected, dict):
                self.error("style-contract.json", "selected_calibration_variant_id must identify a generated variant")
            else:
                for field in ("direction_id", "style_strength", "layout_density"):
                    style_field = "selected_direction_id" if field == "direction_id" else field
                    if doc.get(style_field) != selected.get(field):
                        self.error("style-contract.json", f"{style_field} must match selected calibration variant")
        profile = self.quality_profile()
        self.validate_readability_profile(doc.get("readability_profile"))
        schema_version = doc.get("contract_schema_version", 1)
        if type(schema_version) is not int or schema_version not in {1, 2}:
            self.error("style-contract.json", "contract_schema_version must be 1 or 2")
        if self.semantic_system_active():
            self.validate_semantic_system(doc.get("semantic_system"))
        composition = doc.get("composition_profile")
        if not isinstance(composition, dict):
            self.error("style-contract.json", "composition_profile is required")
            composition = {}
        if composition.get("quality_profile") != profile:
            self.error(
                "style-contract.json",
                "composition_profile.quality_profile must match brief.json",
            )
        if composition.get("visual_richness") not in {
            "restrained",
            "balanced-rich",
            "editorial-rich",
        }:
            self.error(
                "style-contract.json",
                "composition_profile.visual_richness is invalid",
            )
        integer_fields = {
            "layout_archetype_minimum": 1,
            "max_consecutive_same_archetype": 1,
            "min_information_units_per_content_slide": 1,
            "min_visual_layers_per_content_slide": 1,
        }
        for field, minimum in integer_fields.items():
            value = composition.get(field)
            if not isinstance(value, int) or value < minimum:
                self.error(
                    "style-contract.json",
                    f"composition_profile.{field} must be an integer >= {minimum}",
                )
        sparse_roles = composition.get("sparse_allowed_roles")
        if not isinstance(sparse_roles, list) or not sparse_roles:
            self.error(
                "style-contract.json",
                "composition_profile.sparse_allowed_roles must be a non-empty list",
            )
        if profile == "production":
            if doc.get("concept_reference_mode") != "image-and-analysis":
                self.error(
                    "style-contract.json",
                    "production concepts require concept_reference_mode='image-and-analysis'",
                )
            if doc.get("pass_selected_images_to_image_gen") is not True:
                self.error(
                    "style-contract.json",
                    "production concepts must pass selected images to image_gen",
                )
            production_minima = {
                "layout_archetype_minimum": 3,
                "min_information_units_per_content_slide": 3,
                "min_visual_layers_per_content_slide": 2,
            }
            for field, minimum in production_minima.items():
                value = composition.get(field)
                if not isinstance(value, int) or value < minimum:
                    self.error(
                        "style-contract.json",
                        f"production composition_profile.{field} must be >= {minimum}",
                    )
            repeat_limit = composition.get("max_consecutive_same_archetype")
            if not isinstance(repeat_limit, int) or repeat_limit > 2:
                self.error(
                    "style-contract.json",
                    "production max_consecutive_same_archetype must be <= 2",
                )
        try:
            self.visual_registry = VISUAL_ADAPTERS.load_registry(
                self.root / "visual-provider-registry.json"
            )
            VISUAL_ADAPTERS.validate_style_policy(self.visual_registry, doc)
        except (VISUAL_ADAPTERS.AdapterError, OSError, json.JSONDecodeError) as error:
            self.error("visual-provider-registry.json", str(error))
        anchors = doc.get("anchor_concepts")
        required_roles = {"opening", "typical", "highest-risk"}
        if not isinstance(anchors, list) or len(anchors) != 3:
            self.error("style-contract.json", "anchor_concepts must contain exactly three records")
            return
        roles = {anchor.get("role") for anchor in anchors if isinstance(anchor, dict)}
        if roles != required_roles:
            self.error("style-contract.json", f"anchor roles must be {sorted(required_roles)}")
        for index, anchor in enumerate(anchors):
            if not isinstance(anchor, dict) or not anchor.get("slide_id"):
                self.error(f"style-contract.json.anchor_concepts[{index}]", "missing slide_id")
                continue
            if self.reaches("concept"):
                self.validate_relative_file(
                    anchor.get("path"), f"style-contract.json.anchor_concepts[{index}].path"
                )

    def validate_deck(self) -> None:
        name = "deck-spec.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        estimate = doc.get("production_estimate")
        if not isinstance(estimate, dict):
            self.error(name, "production_estimate is required")
        else:
            for field in ("unique_components", "planned_generation_calls", "reserved_retries"):
                value = estimate.get(field)
                if not isinstance(value, int) or value < 0:
                    self.error(name, f"production_estimate.{field} must be a non-negative integer")
            brief = self.docs.get("brief.json", {})
            budget = brief.get("capacity_budget", {}) if isinstance(brief, dict) else {}
            limit = budget.get("component_limit") if isinstance(budget, dict) else None
            if isinstance(limit, int) and estimate.get("unique_components", 0) > limit:
                self.error(name, "production_estimate.unique_components exceeds approved component_limit")
            call_limit = budget.get("generation_call_limit") if isinstance(budget, dict) else None
            planned_calls = estimate.get("planned_generation_calls", 0)
            reserved_retries = estimate.get("reserved_retries", 0)
            if isinstance(call_limit, int) and planned_calls + reserved_retries > call_limit:
                self.error(name, "production estimate exceeds approved generation_call_limit")
            retry_limit = budget.get("retry_limit_per_asset") if isinstance(budget, dict) else None
            unique_components = estimate.get("unique_components", 0)
            if (
                isinstance(retry_limit, int)
                and isinstance(unique_components, int)
                and reserved_retries > unique_components * retry_limit
            ):
                self.error(name, "reserved_retries exceeds retry_limit_per_asset budget")
            deadline = budget.get("deadline") if isinstance(budget, dict) else None
            if deadline is not None:
                estimated_completion = estimate.get("estimated_completion")
                if not estimated_completion:
                    self.error(name, "production_estimate.estimated_completion is required when deadline is set")
                else:
                    try:
                        completion_time = datetime.fromisoformat(str(estimated_completion))
                        deadline_time = datetime.fromisoformat(str(deadline))
                        completion_is_aware = completion_time.utcoffset() is not None
                        deadline_is_aware = deadline_time.utcoffset() is not None
                        if completion_is_aware != deadline_is_aware:
                            self.error(
                                name,
                                "estimated_completion and deadline must use the same timezone-awareness policy",
                            )
                        elif completion_time > deadline_time:
                            self.error(name, "estimated_completion exceeds approved deadline")
                    except (TypeError, ValueError):
                        self.error(name, "estimated_completion must be ISO-8601")
        slides = doc.get("slides")
        if not isinstance(slides, list) or not slides:
            self.error(name, "slides must be a non-empty list")
            return
        self.check_unique(slides, "id", f"{name}.slides")
        style = self.docs.get("style-contract.json", {})
        readability_profile = (
            style.get("readability_profile") if isinstance(style, dict) else None
        )
        composition_profile = (
            style.get("composition_profile", {}) if isinstance(style, dict) else {}
        )
        if not isinstance(composition_profile, dict):
            composition_profile = {}
        selected_reference_values = (
            style.get("selected_reference_ids", []) if isinstance(style, dict) else []
        )
        selected_reference_ids = {
            value for value in selected_reference_values if isinstance(value, str)
        } if isinstance(selected_reference_values, list) else set()
        storyboard = self.docs.get("storyboard.json", {})
        storyboard_slides = storyboard.get("slides", []) if isinstance(storyboard, dict) else []
        if isinstance(storyboard_slides, list) and storyboard_slides:
            story_order = [item.get("id") for item in storyboard_slides if isinstance(item, dict)]
            deck_order = [item.get("id") for item in slides if isinstance(item, dict)]
            if deck_order != story_order:
                self.error(name, "slide IDs and order must match storyboard.json")
        narrative_roles = {
            slide.get("id"): slide.get("narrative_role")
            for slide in storyboard_slides
            if isinstance(slide, dict)
        }
        sparse_values = composition_profile.get("sparse_allowed_roles")
        sparse_roles = {
            value for value in sparse_values if isinstance(value, str)
        } if isinstance(sparse_values, list) else SPARSE_DEFAULT_ROLES
        archetypes: list[str] = []
        rich_content_slide_count = 0
        for index, slide in enumerate(slides):
            location = f"{name}.slides[{index}]"
            if not isinstance(slide, dict):
                self.error(location, "must be an object")
                continue
            components = slide.get("components", [])
            try:
                notes = speaker_notes_text(slide)
            except ValueError as exc:
                self.error(location, str(exc))
            else:
                brief_doc = self.docs.get("brief.json", {})
                brief_delivery = brief_doc.get("delivery", {}) if isinstance(brief_doc, dict) else {}
                if isinstance(brief_delivery, dict) and brief_delivery.get("speaker_notes") == "required" and not notes.strip():
                    self.error(location, "speaker_notes is required for every slide")
            if not isinstance(components, list):
                self.error(location, "components must be a list")
            elif any(not isinstance(component, str) or not component for component in components):
                self.error(location, "component IDs must be non-empty strings")
            elif len(components) != len(set(components)):
                self.error(location, "components must contain unique component IDs")
            composition = slide.get("composition")
            if not isinstance(composition, dict):
                self.error(location, "composition is required")
                composition = {}
            archetype = composition.get("layout_archetype")
            if not isinstance(archetype, str) or not archetype:
                self.error(location, "composition.layout_archetype is required")
            else:
                archetypes.append(archetype)
            complexity = composition.get("complexity")
            if complexity not in {"sparse", "standard", "rich"}:
                self.error(
                    location,
                    "composition.complexity must be sparse, standard, or rich",
                )
            information_units = composition.get("information_units")
            if not isinstance(information_units, list):
                self.error(location, "composition.information_units must be a list")
                information_units = []
            self.check_unique(
                information_units, "id", f"{location}.composition.information_units"
            )
            visual_layers = composition.get("visual_layers")
            if not isinstance(visual_layers, list):
                self.error(location, "composition.visual_layers must be a list")
                visual_layers = []
            self.check_unique(
                visual_layers, "id", f"{location}.composition.visual_layers"
            )
            for unit_kind, records in (
                ("information_units", information_units),
                ("visual_layers", visual_layers),
            ):
                for item_index, item in enumerate(records):
                    if not isinstance(item, dict) or not item.get("role"):
                        self.error(
                            f"{location}.composition.{unit_kind}[{item_index}]",
                            "must contain a non-empty role",
                        )
            slide_reference_ids = composition.get("reference_ids")
            if not isinstance(slide_reference_ids, list):
                self.error(location, "composition.reference_ids must be a list")
                slide_reference_ids = []
            if self.quality_profile() in {"production", "draft"} and not slide_reference_ids:
                self.error(location, "production/draft slides require reference_ids")
            if any(
                not isinstance(reference_id, str)
                or reference_id not in selected_reference_ids
                for reference_id in slide_reference_ids
            ):
                self.error(location, "composition.reference_ids contains an unselected reference")
            if index > 0 and self.quality_profile() in {"production", "draft"}:
                if not has_meaningful_value(composition.get("variation_from_previous")):
                    self.error(location, "composition.variation_from_previous is required")
            narrative_role = narrative_roles.get(slide.get("id"))
            if complexity == "sparse" and narrative_role not in sparse_roles:
                self.error(
                    location,
                    "sparse composition is not allowed for this narrative role",
                )
            is_content_slide = narrative_role not in sparse_roles
            if self.quality_profile() == "production" and is_content_slide:
                min_units = positive_int_or(
                    composition_profile.get("min_information_units_per_content_slide"), 3
                )
                min_layers = positive_int_or(
                    composition_profile.get("min_visual_layers_per_content_slide"), 2
                )
                if len(information_units) < min_units:
                    self.error(
                        location,
                        f"content slide requires at least {min_units} information units",
                    )
                if len(visual_layers) < min_layers:
                    self.error(
                        location,
                        f"content slide requires at least {min_layers} visual layers",
                    )
                if complexity == "rich":
                    rich_content_slide_count += 1
            fill_slots = slide.get("fill_slots", [])
            if not isinstance(fill_slots, list):
                self.error(location, "fill_slots must be a list")
                fill_slots = []
            self.check_unique(fill_slots, "id", f"{location}.fill_slots")
            for slot_index, slot in enumerate(fill_slots):
                if not isinstance(slot, dict):
                    self.error(f"{location}.fill_slots[{slot_index}]", "must be an object")
                    continue
                self.validate_bbox(slot.get("bbox"), f"{location}.fill_slots[{slot_index}].bbox")
            self.validate_slide_readability(slide, fill_slots, readability_profile, location)
            if self.semantic_system_active():
                self.validate_slide_semantics(slide, fill_slots, location)
            if self.reaches("concept"):
                concept_path = slide.get("concept_path")
                if not concept_path:
                    self.error(location, "concept_path is required from concept stage")
                else:
                    self.validate_relative_file(concept_path, f"{location}.concept_path")
        if self.quality_profile() in {"production", "draft"}:
            archetype_minimum = positive_int_or(
                composition_profile.get("layout_archetype_minimum"), 1
            )
            required_archetypes = min(archetype_minimum, len(slides))
            if len(set(archetypes)) < required_archetypes:
                self.error(
                    name,
                    f"deck requires at least {required_archetypes} layout archetypes",
                )
            repeat_limit = positive_int_or(
                composition_profile.get("max_consecutive_same_archetype"), len(slides)
            )
            run_length = 0
            previous: str | None = None
            for archetype in archetypes:
                run_length = run_length + 1 if archetype == previous else 1
                previous = archetype
                if run_length > repeat_limit:
                    self.error(
                        name,
                        "layout archetype repeats more often than composition_profile allows",
                    )
                    break
        if (
            self.quality_profile() == "production"
            and len(slides) >= 5
            and composition_profile.get("visual_richness") != "restrained"
            and rich_content_slide_count == 0
        ):
            self.error(name, "production deck needs at least one rich content slide")

    def validate_concept_generation_log(self) -> None:
        name = "concept-generation-log.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        if self.is_fixture_profile() and doc.get("status") == "not-required":
            if not doc.get("reason"):
                self.error(name, "not-required fixture log needs a reason")
            return
        if doc.get("status") != "generated":
            self.error(name, "status must be 'generated'")
        if self.quality_profile() in {"production", "draft"}:
            for artifact_name, artifact in self.docs.items():
                if not isinstance(artifact, dict):
                    continue
                baseline = artifact.get("current_style_rebaseline")
                if isinstance(baseline, dict) and baseline.get("status") == "pending_user_selection":
                    self.error(
                        f"{artifact_name}.current_style_rebaseline",
                        "style rebaseline is pending user selection; concepts cannot proceed on the old G3 approval",
                    )
            if not isinstance(doc.get("version"), str) or not doc["version"].strip():
                self.error(name, "version must identify the contentful concept set")
            contact_sheet = doc.get("review_contact_sheet")
            if not isinstance(contact_sheet, dict):
                self.error(name, "review_contact_sheet must bind the contentful contact sheet")
            else:
                contact_path = self.validate_relative_file(
                    contact_sheet.get("path"), f"{name}.review_contact_sheet.path"
                )
                if not re.fullmatch(r"[0-9a-fA-F]{64}", str(contact_sheet.get("sha256"))):
                    self.error(name, "review_contact_sheet.sha256 must be a SHA-256 hex digest")
                elif contact_path and hashlib.sha256(contact_path.read_bytes()).hexdigest() != str(
                    contact_sheet["sha256"]
                ).lower():
                    self.error(name, "review_contact_sheet hash does not match file")
        records = doc.get("slides")
        if not isinstance(records, list) or not records:
            self.error(name, "slides must be a non-empty list")
            return
        self.check_unique(records, "slide_id", f"{name}.slides")
        records_by_slide = {
            record.get("slide_id"): record
            for record in records
            if isinstance(record, dict) and record.get("slide_id")
        }
        deck = self.docs.get("deck-spec.json", {})
        deck_slides = deck.get("slides", []) if isinstance(deck, dict) else []
        expected_slide_ids = {
            slide.get("id") for slide in deck_slides if isinstance(slide, dict)
        }
        if set(records_by_slide) != expected_slide_ids:
            self.error(name, "slide coverage must match deck-spec.json")
        reference_study = self.docs.get("reference-study.json", {})
        reference_records = (
            reference_study.get("references", [])
            if isinstance(reference_study, dict)
            else []
        )
        reference_by_id = {
            reference.get("id"): reference
            for reference in reference_records
            if isinstance(reference, dict) and reference.get("id")
        }
        style = self.docs.get("style-contract.json", {})
        storyboard = self.docs.get("storyboard.json", {})
        storyboard_slides = {
            item.get("id"): item
            for item in storyboard.get("slides", [])
            if isinstance(item, dict) and item.get("id")
        } if isinstance(storyboard, dict) else {}
        output_paths: set[str] = set()
        for index, slide in enumerate(deck_slides):
            if not isinstance(slide, dict):
                continue
            slide_id = slide.get("id")
            record = records_by_slide.get(slide_id)
            location = f"{name}.slides[{index}]"
            if not isinstance(record, dict):
                continue
            if record.get("tool") not in {"codex-image_gen", "image_gen"}:
                self.error(location, "tool must identify Codex image_gen")
            prompt = record.get("prompt")
            if not isinstance(prompt, str) or len(prompt.strip()) < 80:
                self.error(location, "prompt must preserve the exact structured prompt")
            if record.get("status") not in {"generated", "approved"}:
                self.error(location, "status must be generated or approved")
            if isinstance(style, dict):
                for field in ("selected_calibration_variant_id", "style_strength", "layout_density"):
                    if record.get(field) != style.get(field):
                        self.error(location, f"{field} must match approved style-contract")
            expected_output = slide.get("concept_path")
            if record.get("output_path") != expected_output:
                self.error(location, "output_path must match deck-spec concept_path")
            if self.quality_profile() in {"production", "draft"}:
                version = record.get("version")
                if not isinstance(version, str) or not version.strip():
                    self.error(location, "version must identify this immutable contentful concept revision")
                if record.get("output_path") in output_paths:
                    self.error(location, "contentful concept output paths must be unique by slide revision")
                if isinstance(record.get("output_path"), str):
                    output_paths.add(record["output_path"])
            output_path = self.validate_relative_file(
                record.get("output_path"), f"{location}.output_path"
            )
            output_sha = record.get("output_sha256")
            if not re.fullmatch(r"[0-9a-fA-F]{64}", str(output_sha)):
                self.error(location, "output_sha256 must be a SHA-256 hex digest")
            elif output_path and hashlib.sha256(output_path.read_bytes()).hexdigest() != str(
                output_sha
            ).lower():
                self.error(location, "output_sha256 does not match the concept file")
            if self.quality_profile() in {"production", "draft"}:
                source_records = self.validate_content_source_lineage(
                    record.get("source_lineage"), f"{location}.source_lineage"
                )
                copy_map = record.get("copy_map")
                self.validate_content_copy_map(
                    copy_map,
                    source_records,
                    storyboard_slides.get(slide_id),
                    f"{location}.copy_map",
                    record.get("message_mapping"),
                )
                self.validate_concept_image_inventory(
                    record.get("image_inventory"),
                    source_records,
                    f"{location}.image_inventory",
                )
                self.validate_content_text_review(
                    record.get("text_review"),
                    copy_map,
                    record.get("output_path"),
                    record.get("output_sha256"),
                    f"{location}.text_review",
                )
            expected_reference_ids = set(
                slide.get("composition", {}).get("reference_ids", [])
                if isinstance(slide.get("composition"), dict)
                else []
            )
            attached = record.get("reference_images")
            if not isinstance(attached, list) or not attached:
                self.error(location, "reference_images must be a non-empty list")
                attached = []
            attached_ids = {
                item.get("reference_id")
                for item in attached
                if isinstance(item, dict) and item.get("reference_id")
            }
            if attached_ids != expected_reference_ids:
                self.error(
                    location,
                    "attached reference IDs must match deck-spec composition.reference_ids",
                )
            for item_index, item in enumerate(attached):
                item_location = f"{location}.reference_images[{item_index}]"
                if not isinstance(item, dict):
                    self.error(item_location, "must be an object")
                    continue
                source = reference_by_id.get(item.get("reference_id"))
                if not isinstance(source, dict):
                    self.error(item_location, "unknown reference_id")
                    continue
                if item.get("path") != source.get("image_path"):
                    self.error(item_location, "path must match reference-study image_path")
                if str(item.get("sha256", "")).lower() != str(
                    source.get("sha256", "")
                ).lower():
                    self.error(item_location, "sha256 must match reference-study")

    def validate_component_master_manifest(self) -> None:
        name = "component-master-manifest.json"
        if self.is_fixture_profile():
            return
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            self.error(name, "is required before inventory can begin")
            return
        if doc.get("status") != "generated":
            self.error(name, "status must be generated after G4 approval")
        state = self.docs.get("run-state.json", {})
        gates = state.get("gates", {}) if isinstance(state, dict) else {}
        g4 = gates.get("concept_contact_sheet", {}) if isinstance(gates, dict) else {}
        if not isinstance(g4, dict) or g4.get("status") != "approved":
            self.error(name, "G4 contentful concept approval is required before component masters")
        approved_artifacts = g4.get("approved_artifacts", {}) if isinstance(g4, dict) else {}
        concept_log = self.docs.get("concept-generation-log.json", {})
        if not isinstance(concept_log, dict):
            self.error(name, "concept-generation-log.json must be present before component masters")
            concept_log = {}
        log_path = self.root / "concept-generation-log.json"
        log_sha256 = hashlib.sha256(log_path.read_bytes()).hexdigest() if log_path.is_file() else None
        if not isinstance(approved_artifacts, dict) or approved_artifacts.get(
            "concept-generation-log.json"
        ) != log_sha256:
            self.error(name, "component masters must link the current G4-approved concept-generation-log.json")
        deck = self.docs.get("deck-spec.json", {})
        slides = deck.get("slides", []) if isinstance(deck, dict) else []
        master_records = doc.get("slides")
        if not isinstance(master_records, list) or not master_records:
            self.error(name, "slides must contain one textless master per deck slide")
            return
        self.check_unique(master_records, "slide_id", f"{name}.slides")
        concepts_by_slide = {
            record.get("slide_id"): record
            for record in concept_log.get("slides", [])
            if isinstance(record, dict) and record.get("slide_id")
        } if isinstance(concept_log.get("slides"), list) else {}
        masters_by_slide = {
            record.get("slide_id"): record
            for record in master_records
            if isinstance(record, dict) and record.get("slide_id")
        }
        expected_ids = {slide.get("id") for slide in slides if isinstance(slide, dict)}
        if set(masters_by_slide) != expected_ids:
            self.error(name, "slide coverage must match deck-spec.json exactly")
        for index, slide in enumerate(slides):
            if not isinstance(slide, dict):
                continue
            slide_id = slide.get("id")
            record = masters_by_slide.get(slide_id)
            location = f"{name}.slides[{index}]"
            if not isinstance(record, dict):
                continue
            concept = concepts_by_slide.get(slide_id)
            if not isinstance(concept, dict):
                self.error(location, "must link an existing contentful concept record")
                continue
            concept_path = slide.get("concept_path")
            concept_sha256 = concept.get("output_sha256")
            approved_concept_sha256 = (
                approved_artifacts.get(concept_path)
                if isinstance(approved_artifacts, dict) and isinstance(concept_path, str)
                else None
            )
            if record.get("approved_contentful_path") != concept_path:
                self.error(location, "approved_contentful_path must match deck-spec concept_path")
            if record.get("approved_contentful_sha256") != concept_sha256:
                self.error(location, "approved_contentful_sha256 must match the contentful concept log")
            if record.get("approved_contentful_sha256") != approved_concept_sha256:
                self.error(location, "contentful concept hash is not approved by G4")
            if record.get("concept_generation_log_sha256") != log_sha256:
                self.error(location, "concept_generation_log_sha256 must match the current G4-approved log")
            if record.get("version") != concept.get("version"):
                self.error(location, "version must match the approved contentful concept")
            master_path_value = record.get("component_master_path")
            if master_path_value == concept_path:
                self.error(location, "component master cannot reuse the contentful full-slide image")
            master_path = self.validate_relative_file(
                master_path_value, f"{location}.component_master_path"
            )
            master_hash = record.get("component_master_sha256")
            if not re.fullmatch(r"[0-9a-fA-F]{64}", str(master_hash)):
                self.error(location, "component_master_sha256 must be a SHA-256 hex digest")
            elif master_path and hashlib.sha256(master_path.read_bytes()).hexdigest() != str(
                master_hash
            ).lower():
                self.error(location, "component_master_sha256 does not match file")
            if record.get("text_policy") != "textless":
                self.error(location, "text_policy must be textless")
            if record.get("ordinary_text_removed") is not True:
                self.error(location, "ordinary_text_removed must be true")
            if record.get("contains_ordinary_text") is not False:
                self.error(location, "contains_ordinary_text must be false")
            expected_slots = [
                {"id": slot.get("id"), "kind": slot.get("kind"), "bbox": slot.get("bbox")}
                for slot in slide.get("fill_slots", [])
                if isinstance(slot, dict)
            ]
            preserved_slots = record.get("preserved_fill_slots")
            if preserved_slots != expected_slots:
                self.error(location, "preserved_fill_slots must preserve deck-spec geometry exactly")
            geometry_review = record.get("geometry_review")
            review_location = f"{location}.geometry_review"
            if not isinstance(geometry_review, dict):
                self.error(review_location, "geometry_review is required before component inventory")
                continue
            if geometry_review.get("status") != "passed" or geometry_review.get("display_scale_percent") != 100:
                self.error(review_location, "geometry_review must pass at 100% display scale")
            if not isinstance(geometry_review.get("reviewer"), str) or not geometry_review["reviewer"].strip():
                self.error(review_location, "geometry_review reviewer is required")
            actor_error = REVIEW_POLICY.review_actor_error(self.docs.get("run-state.json"), geometry_review)
            if actor_error:
                self.error(review_location, actor_error)
            comparison_value = geometry_review.get("comparison_path")
            if comparison_value in (concept_path, master_path_value):
                self.error(review_location, "geometry_review comparison_path must be a separate comparison image")
            comparison_path = self.validate_relative_file(comparison_value, f"{review_location}.comparison_path")
            if comparison_path and geometry_review.get("comparison_sha256") != hashlib.sha256(comparison_path.read_bytes()).hexdigest():
                self.error(review_location, "geometry_review comparison_sha256 does not match evidence")
            if (
                geometry_review.get("reviewed_contentful_sha256") != concept_sha256
                or geometry_review.get("reviewed_master_sha256") != master_hash
            ):
                self.error(review_location, "geometry_review must bind the current contentful and master images")
            checks = geometry_review.get("checks")
            if not isinstance(checks, dict) or any(
                checks.get(field) is not True
                for field in ("composition_preserved", "fill_slots_preserved", "ordinary_text_removed")
            ):
                self.error(review_location, "geometry_review checks must confirm composition, fill slots and text removal")

    def validate_decomposition_map(self, *, require_review: bool = True) -> None:
        name = "decomposition-map.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        if not isinstance(doc.get("version"), str) or not doc["version"].strip():
            self.error(name, "version is required")
        slides = doc.get("slides")
        if not isinstance(slides, list) or not slides:
            self.error(name, "slides must contain every deck slide")
            return
        self.check_unique(slides, "slide_id", f"{name}.slides")
        deck = self.docs.get("deck-spec.json", {})
        scene = self.docs.get("scene-graph.json", {})
        manifest = self.docs.get("asset-manifest.json", {})
        masters = self.docs.get("component-master-manifest.json", {})
        state = self.docs.get("run-state.json", {})
        deck_slides = {
            slide.get("id"): slide for slide in deck.get("slides", [])
            if isinstance(slide, dict) and isinstance(slide.get("id"), str)
        } if isinstance(deck, dict) and isinstance(deck.get("slides"), list) else {}
        scene_slides = {
            slide.get("slide_id"): slide for slide in scene.get("slides", [])
            if isinstance(slide, dict) and isinstance(slide.get("slide_id"), str)
        } if isinstance(scene, dict) and isinstance(scene.get("slides"), list) else {}
        master_slides = {
            slide.get("slide_id"): slide for slide in masters.get("slides", [])
            if isinstance(slide, dict) and isinstance(slide.get("slide_id"), str)
        } if isinstance(masters, dict) and isinstance(masters.get("slides"), list) else {}
        assets = {
            asset.get("id"): asset for asset in manifest.get("assets", [])
            if isinstance(asset, dict) and isinstance(asset.get("id"), str)
        } if isinstance(manifest, dict) and isinstance(manifest.get("assets"), list) else {}
        g4 = state.get("gates", {}).get("concept_contact_sheet", {}) if isinstance(state, dict) and isinstance(state.get("gates"), dict) else {}
        approved = g4.get("approved_artifacts", {}) if isinstance(g4, dict) else {}
        mapped_ids = {
            slide.get("slide_id") for slide in slides
            if isinstance(slide, dict) and isinstance(slide.get("slide_id"), str)
        }
        if mapped_ids != set(deck_slides):
            self.error(name, "slide coverage must match deck-spec.json exactly")

        for slide_index, record in enumerate(slides):
            location = f"{name}.slides[{slide_index}]"
            if not isinstance(record, dict):
                self.error(location, "must be an object")
                continue
            slide_id = record.get("slide_id")
            if not isinstance(slide_id, str):
                self.error(location, "slide_id must be a string")
                continue
            deck_slide = deck_slides.get(slide_id)
            scene_slide = scene_slides.get(slide_id)
            if not isinstance(deck_slide, dict) or not isinstance(scene_slide, dict):
                self.error(location, "must link current deck and scene slide")
                continue
            concept_path = deck_slide.get("concept_path")
            concept_sha = record.get("concept_sha256")
            if record.get("concept_path") != concept_path:
                self.error(location, "concept_path must match current deck-spec")
            concept_file = self.validate_relative_file(concept_path, f"{location}.concept_path")
            if concept_file and concept_sha != hashlib.sha256(concept_file.read_bytes()).hexdigest():
                self.error(location, "concept_sha256 is stale")
            if not isinstance(approved, dict) or approved.get(concept_path) != concept_sha:
                self.error(location, "concept is not current G4-approved content")
            if not self.is_fixture_profile():
                master = master_slides.get(slide_id)
                if not isinstance(master, dict):
                    self.error(location, "textless master is required")
                else:
                    if record.get("master_path") != master.get("component_master_path") or record.get("master_sha256") != master.get("component_master_sha256"):
                        self.error(location, "master binding must match current textless master")
                    master_file = self.validate_relative_file(record.get("master_path"), f"{location}.master_path")
                    if master_file and record.get("master_sha256") != hashlib.sha256(master_file.read_bytes()).hexdigest():
                        self.error(location, "master_sha256 is stale")

            items = record.get("items")
            if not isinstance(items, list) or not items:
                self.error(location, "items must be a non-empty numbered inventory")
                continue
            items_sha256 = value_sha256(items)
            self.check_unique(items, "id", f"{location}.items")
            scene_instances = {
                element.get("instance_id"): element for element in scene_slide.get("elements", [])
                if isinstance(element, dict) and isinstance(element.get("instance_id"), str)
            } if isinstance(scene_slide.get("elements"), list) else {}
            fill_slots = {
                slot.get("id"): slot for slot in deck_slide.get("fill_slots", [])
                if isinstance(slot, dict) and isinstance(slot.get("id"), str)
            } if isinstance(deck_slide.get("fill_slots"), list) else {}
            mapped_instances: set[str] = set()
            mapped_slots: set[str] = set()
            for item_index, item in enumerate(items):
                item_location = f"{location}.items[{item_index}]"
                if not isinstance(item, dict):
                    self.error(item_location, "must be an object")
                    continue
                if not isinstance(item.get("id"), str) or not item["id"].strip():
                    self.error(item_location, "stable item id is required")
                if not isinstance(item.get("label"), str) or not item["label"].strip():
                    self.error(item_location, "human-readable label is required")
                if item.get("inventory_pass") not in {"semantic", "residual"}:
                    self.error(item_location, "inventory_pass must be semantic or residual")
                self.validate_bbox(item.get("bbox"), f"{item_location}.bbox")
                kind = item.get("kind")
                if kind == "visual":
                    instance_id = item.get("instance_id")
                    if not isinstance(instance_id, str) or not instance_id:
                        self.error(item_location, "instance_id must be a non-empty string")
                        element = None
                    else:
                        if instance_id in mapped_instances:
                            self.error(item_location, "two independent objects cannot share one scene instance")
                        mapped_instances.add(instance_id)
                        element = scene_instances.get(instance_id)
                    if not isinstance(element, dict):
                        self.error(item_location, "visual item must link a scene instance")
                    elif any(item.get(field) != element.get(field) for field in ("component_id", "bbox", "z")):
                        self.error(item_location, "visual item must match scene component, bbox and z")
                    component_id = item.get("component_id")
                    asset = assets.get(component_id) if isinstance(component_id, str) else None
                    if not isinstance(asset, dict):
                        self.error(item_location, "component_id must link an asset")
                    elif item.get("role") != asset.get("role"):
                        self.error(item_location, "role must match asset manifest")
                    source_class = item.get("source_class")
                    if (item.get("role") in PHOTO_ROLES and source_class != "ordinary-photo") or (source_class == "ordinary-photo" and item.get("role") not in PHOTO_ROLES):
                        self.error(item_location, "ordinary photo source class must use a photo component role")
                    if (item.get("role") in TRUTH_SENSITIVE_ROLES and source_class != "truth-sensitive") or (source_class == "truth-sensitive" and item.get("role") not in TRUTH_SENSITIVE_ROLES):
                        self.error(item_location, "truth-sensitive source class must use a preserved evidence role")
                    allowed_treatments = {
                        "design-object": {"preserve-pixels", "reconstruct-single"},
                        "ordinary-photo": {"preserve-source", "photo-adjust", "reconstruct-single"},
                        "truth-sensitive": {"preserve-source", "blank-content"},
                    }
                    treatment = item.get("treatment")
                    if not isinstance(source_class, str) or not isinstance(treatment, str) or treatment not in allowed_treatments.get(source_class, set()):
                        self.error(item_location, "source_class and treatment require a valid single-object handling decision")
                    occlusion = item.get("occlusion")
                    if not isinstance(occlusion, str) or occlusion not in {"none", "partially-occluded", "fused"}:
                        self.error(item_location, "occlusion decision is required")
                    if occlusion in {"partially-occluded", "fused"} and treatment == "preserve-pixels":
                        self.error(item_location, "occluded or fused object cannot claim complete original-pixel extraction")
                    if not isinstance(item.get("overlaps"), list):
                        self.error(item_location, "overlaps must list related scene instances, or be empty")
                    elif any(
                        not isinstance(related, str) or related not in scene_instances or related == instance_id
                        for related in item["overlaps"]
                    ):
                        self.error(item_location, "overlaps must reference other current scene instances")
                    padding = item.get("crop_padding")
                    if (
                        not isinstance(padding, list) or len(padding) != 4
                        or any(type(value) not in {int, float} or value < 0 or value > 0.2 for value in padding)
                    ):
                        self.error(item_location, "crop_padding must be four normalized non-negative margins at most 0.2")
                    contains_text_data = item.get("contains_ordinary_text_or_data")
                    decision = item.get("text_data_decision")
                    if type(contains_text_data) is not bool:
                        self.error(item_location, "contains_ordinary_text_or_data must be explicit true or false")
                    if not isinstance(decision, str) or decision not in {"none", "blank-for-fill-slots", "preserve-factual-source"}:
                        self.error(item_location, "text_data_decision must specify removal or factual preservation")
                    elif contains_text_data is False and decision != "none":
                        self.error(item_location, "text_data_decision must be none when no text or data is present")
                    elif contains_text_data is True:
                        if source_class != "truth-sensitive" and decision != "blank-for-fill-slots":
                            self.error(item_location, "ordinary text/data outside factual evidence must be blanked")
                        if source_class == "design-object" and (decision != "blank-for-fill-slots" or treatment != "reconstruct-single"):
                            self.error(item_location, "design object with ordinary text/data must be rebuilt as a blank single object")
                        if source_class == "ordinary-photo" and treatment != "reconstruct-single":
                            self.error(item_location, "photo with removable text/data must use a blank single-object route")
                        if source_class == "truth-sensitive" and not (
                            (decision == "preserve-factual-source" and treatment == "preserve-source")
                            or (decision == "blank-for-fill-slots" and treatment == "blank-content")
                        ):
                            self.error(item_location, "truth-sensitive text/data decision must match preservation or blanking treatment")
                elif kind == "fill-slot":
                    slot_id = item.get("fill_slot_id")
                    if not isinstance(slot_id, str) or not slot_id:
                        self.error(item_location, "fill_slot_id must be a non-empty string")
                        slot = None
                    else:
                        if slot_id in mapped_slots:
                            self.error(item_location, "fill slot is mapped more than once")
                        mapped_slots.add(slot_id)
                        slot = fill_slots.get(slot_id)
                    if not isinstance(slot, dict):
                        self.error(item_location, "fill_slot_id must link a deck fill slot")
                    elif item.get("bbox") != slot.get("bbox") or item.get("role") != slot.get("kind"):
                        self.error(item_location, "fill slot role and bbox must match deck-spec")
                    if item.get("treatment") != "editable":
                        self.error(item_location, "fill slot treatment must be editable")
                elif kind == "omitted":
                    if not isinstance(item.get("reason"), str) or not item["reason"].strip():
                        self.error(item_location, "omitted visible item requires a human-reviewable reason")
                else:
                    self.error(item_location, "kind must be visual, fill-slot, or omitted")
            for instance_id in set(scene_instances) - mapped_instances:
                self.error(location, f"unmapped scene instance '{instance_id}'")
            for slot_id in set(fill_slots) - mapped_slots:
                self.error(location, f"unmapped fill slot '{slot_id}'")
            if not require_review:
                continue
            for field in ("review_image", "review_packet"):
                evidence = record.get(field)
                evidence_location = f"{location}.{field}"
                if not isinstance(evidence, dict):
                    self.error(evidence_location, "review evidence is required")
                    continue
                evidence_file = self.validate_relative_file(evidence.get("path"), f"{evidence_location}.path")
                if evidence_file and evidence.get("sha256") != hashlib.sha256(evidence_file.read_bytes()).hexdigest():
                    self.error(evidence_location, "review evidence hash is stale")
                if evidence.get("reviewed_items_sha256") != items_sha256:
                    self.error(evidence_location, "review evidence must bind current item inventory")
                if field == "review_image" and (evidence.get("reviewed_concept_sha256") != concept_sha or evidence.get("display_scale_percent") != 100):
                    self.error(evidence_location, "numbered image must review the current concept at 100% scale")
                if evidence_file and field == "review_image":
                    if evidence_file in {concept_file, self.root / str(record.get("master_path"))}:
                        self.error(evidence_location, "numbered image must be separate from concept and master")
                    review_dimensions = png_dimensions(evidence_file)
                    if evidence_file.suffix.lower() != ".png" or review_dimensions is None:
                        self.error(evidence_location, "numbered review must be a PNG")
                    concept_dimensions = png_dimensions(concept_file) if concept_file else None
                    if concept_dimensions is None:
                        self.error(evidence_location, "current concept must be a readable PNG")
                    elif review_dimensions != concept_dimensions:
                        self.error(evidence_location, "review image dimensions must match current concept")
                    elif concept_file and isinstance(items, list):
                        try:
                            review_script = Path(__file__).resolve().with_name("render_decomposition_review.py")
                            review_spec = importlib.util.spec_from_file_location("ai_presentation_decomposition_review", review_script)
                            if review_spec is None or review_spec.loader is None:
                                raise RuntimeError("cannot load numbered review renderer")
                            review_module = importlib.util.module_from_spec(review_spec)
                            review_spec.loader.exec_module(review_module)
                            expected = review_module.build_review_image(concept_file, items)
                            with review_module.Image.open(evidence_file) as actual:
                                actual_rgba = actual.convert("RGBA")
                            if expected.size != actual_rgba.size or expected.tobytes() != actual_rgba.tobytes():
                                self.error(evidence_location, "numbered review pixels do not match declared item overlay")
                        except (ImportError, OSError, ValueError, RuntimeError) as exc:
                            self.error(evidence_location, f"cannot verify numbered review overlay: {exc}")
                if evidence_file and field == "review_packet":
                    try:
                        readable = evidence_file.read_text(encoding="utf-8").strip()
                    except UnicodeError:
                        readable = ""
                    if evidence_file.suffix.lower() != ".md" or not readable:
                        self.error(evidence_location, "readable review list must be non-empty Markdown")
                    else:
                        if items_sha256 not in readable:
                            self.error(evidence_location, "readable review list must include the current item hash and numbered rows")
                        else:
                            try:
                                review_script = Path(__file__).resolve().with_name("render_decomposition_review.py")
                                review_spec = importlib.util.spec_from_file_location("ai_presentation_decomposition_review", review_script)
                                if review_spec is None or review_spec.loader is None:
                                    raise RuntimeError("cannot load numbered review renderer")
                                review_module = importlib.util.module_from_spec(review_spec)
                                review_spec.loader.exec_module(review_module)
                                actual_rows = set(readable.splitlines())
                                expected_rows = {
                                    review_module.review_row(item) for item in items
                                    if isinstance(item, dict) and isinstance(item.get("id"), str)
                                }
                                if not expected_rows or not expected_rows.issubset(actual_rows):
                                    self.error(evidence_location, "readable review list must include full current item rows")
                            except (ImportError, OSError, ValueError, RuntimeError) as exc:
                                self.error(evidence_location, f"cannot verify readable review rows: {exc}")
            review = record.get("coverage_review")
            if not isinstance(review, dict) or not isinstance(review.get("reviewer"), str) or not review["reviewer"].strip():
                self.error(location, "human coverage reviewer is required")
            elif any(review.get(field) is not True for field in ("semantic_passed", "residual_passed", "independence_passed")) or review.get("missing_items") != []:
                self.error(location, "human coverage and independence review must pass with no missing items")
            if isinstance(review, dict):
                actor_error = REVIEW_POLICY.review_actor_error(self.docs.get("run-state.json"), review, human_only=True)
                if actor_error:
                    self.error(location, actor_error)

    def validate_scene_graph(self) -> None:
        name = "scene-graph.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        slides = doc.get("slides")
        if not isinstance(slides, list) or not slides:
            self.error(name, "slides must be a non-empty list")
            return
        self.check_unique(slides, "slide_id", f"{name}.slides")
        seen_instances: set[str] = set()
        for slide_index, slide in enumerate(slides):
            if not isinstance(slide, dict):
                self.error(f"{name}.slides[{slide_index}]", "must be an object")
                continue
            for element_index, element in enumerate(slide.get("elements", [])):
                location = f"{name}.slides[{slide_index}].elements[{element_index}]"
                if not isinstance(element, dict):
                    self.error(location, "must be an object")
                    continue
                component_id = element.get("component_id")
                if not isinstance(component_id, str) or not component_id:
                    self.error(location, "component_id must be a non-empty string")
                instance_id = element.get("instance_id")
                if not isinstance(instance_id, str) or not instance_id:
                    self.error(location, "instance_id must be a non-empty string")
                elif instance_id in seen_instances:
                    self.error(location, f"duplicate instance_id '{instance_id}'")
                else:
                    seen_instances.add(instance_id)
                self.validate_bbox(element.get("bbox"), f"{location}.bbox")
                z = element.get("z")
                if type(z) is not int or z < 0:
                    self.error(location, "z must be a non-negative integer")
                rotation = element.get("rotation", 0)
                if type(rotation) not in {int, float} or not 0 <= rotation < 360:
                    self.error(location, "rotation must be a number from 0 to below 360")
                opacity = element.get("opacity", 1)
                if type(opacity) not in {int, float} or not 0 <= opacity <= 1:
                    self.error(location, "opacity must be a number from 0 to 1")

    def validate_visual_provider_provenance(
        self,
        asset: dict[str, Any],
        location: str,
        style: Any,
        style_gate: Any,
        concept_gate: Any,
        style_sha256: str | None,
    ) -> None:
        provider_id = asset.get("provider_id")
        provider = (
            self.visual_registry.get(provider_id)
            if isinstance(provider_id, str)
            else None
        )
        if provider is None:
            self.error(location, f"unknown visual provider: {provider_id!r}")
        else:
            role = asset.get("role")
            if role not in provider.get("supported_component_roles", []):
                self.error(
                    location,
                    f"provider '{provider_id}' does not support component role '{role}'",
                )
            if provider_id == "photo-postprocess-coach":
                photo_classes = {
                    "rectangle": "opaque-photo-rectangle",
                    "cutout": "transparent-photo-cutout",
                }
                if asset.get("photo_output_mode") not in photo_classes or asset.get("output_class") != photo_classes.get(asset.get("photo_output_mode")):
                    self.error(location, "photo output class must distinguish opaque rectangle from transparent cutout")
                if (asset.get("photo_output_mode") == "rectangle" and role != "ordinary-photo") or (asset.get("photo_output_mode") == "cutout" and role != "photo-cutout"):
                    self.error(location, "photo role must match its rectangle or cutout output mode")
            elif asset.get("output_class") != provider.get("output_class"):
                self.error(location, "output_class does not match registered provider")
            if asset.get("provider_transparency_behavior") != provider.get(
                "transparency_behavior"
            ):
                self.error(
                    location,
                    "provider_transparency_behavior does not match registered provider",
                )
            if asset.get("provider_aspect_ratio_behavior") != provider.get(
                "aspect_ratio_behavior"
            ):
                self.error(
                    location,
                    "provider_aspect_ratio_behavior does not match registered provider",
                )
            if asset.get("adapter_id") != provider.get("adapter"):
                self.error(location, "adapter_id does not match registered provider")
            try:
                VISUAL_ADAPTERS.route_provider(
                    self.visual_registry,
                    {
                        "provider_id": provider_id,
                        "page_role": asset.get("page_role"),
                        "source_kind": asset.get("source_kind"),
                        "truth_sensitive": asset.get("truth_sensitive"),
                        "deliverable_role": asset.get("deliverable_role"),
                        "deck_released": self.stage == "release",
                        "slide_id": asset.get("slide_id"),
                        "source_rights_confirmation": asset.get(
                            "source_rights_confirmation"
                        ),
                        "source_photo": asset.get("source_photo"),
                        "special_page_approval": asset.get(
                            "special_page_approval"
                        ),
                    },
                    style,
                    project_root=self.root,
                )
            except VISUAL_ADAPTERS.AdapterError as error:
                self.error(location, str(error))
            if provider.get("adapter") == "material-illustration-v1":
                try:
                    reference_bbox = VISUAL_ADAPTERS.validate_reference_bbox(
                        asset.get("reference_bbox"), f"{location}.reference_bbox"
                    )
                    expected_prompt = VISUAL_ADAPTERS.build_material_prompt(
                        str(role), reference_bbox, asset.get("prompt_profile", "v1")
                    )
                    if asset.get("prompt") != expected_prompt:
                        self.error(
                            location,
                            "material prompt does not match the adapter-derived prompt",
                        )
                except VISUAL_ADAPTERS.AdapterError as error:
                    self.error(location, str(error))

        if isinstance(style, dict):
            if asset.get("style_contract_version") != style.get("version"):
                self.error(location, "style_contract_version does not match style contract")
            approved_style_artifacts = (
                style_gate.get("approved_artifacts", {})
                if isinstance(style_gate, dict)
                else {}
            )
            if (
                not isinstance(style_sha256, str)
                or asset.get("style_contract_sha256") != style_sha256
                or not isinstance(approved_style_artifacts, dict)
                or approved_style_artifacts.get("style-contract.json") != style_sha256
            ):
                self.error(location, "style contract SHA-256 is not the approved artifact")

        references = asset.get("references")
        if not isinstance(references, list) or not references:
            self.error(location, "generated visual requires approved references")
            return
        approved_concepts = (
            concept_gate.get("approved_artifacts", {})
            if isinstance(concept_gate, dict)
            else {}
        )
        for reference_index, reference in enumerate(references):
            reference_location = f"{location}.references[{reference_index}]"
            if not isinstance(reference, dict):
                self.error(reference_location, "must be an object")
                continue
            reference_path = reference.get("path")
            resolved_reference = self.validate_relative_file(
                reference_path, f"{reference_location}.path"
            )
            declared_reference_sha = reference.get("sha256")
            if reference.get("approval_gate") != "concept_contact_sheet":
                self.error(
                    reference_location,
                    "approval_gate must be concept_contact_sheet",
                )
            if not re.fullmatch(r"[0-9a-fA-F]{64}", str(declared_reference_sha)):
                self.error(
                    reference_location,
                    "sha256 must be 64 hexadecimal characters",
                )
            elif resolved_reference and resolved_reference.is_file():
                actual_reference_sha = hashlib.sha256(
                    resolved_reference.read_bytes()
                ).hexdigest()
                if declared_reference_sha.lower() != actual_reference_sha:
                    self.error(
                        reference_location,
                        "approved reference SHA-256 does not match file",
                    )
                if (
                    not isinstance(approved_concepts, dict)
                    or approved_concepts.get(reference_path) != actual_reference_sha
                ):
                    self.error(
                        reference_location,
                        "reference is not approved by concept_contact_sheet",
                    )

    def validate_asset_isolation(self, asset: dict[str, Any], location: str, output_path: Path | None) -> None:
        if not self.asset_ready(asset):
            return
        module_path = Path(__file__).resolve().with_name("render_asset_isolation_review.py")
        spec = importlib.util.spec_from_file_location("ai_presentation_isolation_contract", module_path)
        if spec is None or spec.loader is None:
            self.error(location, "cannot load isolation review contract")
            return
        try:
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            decomposition = self.docs.get("decomposition-map.json", {})
            source_items = module.source_items_for_asset(decomposition, asset.get("id"))
        except (ImportError, OSError, ValueError, TypeError, KeyError) as exc:
            self.error(location, f"cannot resolve current decomposition items: {exc}")
            return
        if not source_items:
            self.error(location, "asset must map to at least one current independent decomposition item")
            return
        routes = {str(item.get("treatment")) for item in source_items}
        extraction = asset.get("extraction")
        if not isinstance(extraction, dict) or not isinstance(extraction.get("route"), str) or extraction.get("route") not in routes or len(routes) != 1:
            self.error(location, "asset extraction route must match every mapped decomposition item")
        elif extraction["route"] == "preserve-pixels" and output_path:
            if extraction.get("output_path") != asset.get("output") or extraction.get("output_sha256") != asset.get("sha256"):
                self.error(location, "pixel-preserved extraction must bind the current output path and hash")
            source_item = extraction.get("source_item")
            match = next(
                (item for item in source_items if source_item == {"slide_id": item["slide_id"], "item_id": item["item_id"]}),
                None,
            )
            if match is None or asset.get("strategy") != "preserved-transparent-png":
                self.error(location, "pixel-preserved asset must identify one approved source item and preservation strategy")
            else:
                slide_records = decomposition.get("slides", []) if isinstance(decomposition, dict) else []
                slide = next(
                    (record for record in slide_records if isinstance(record, dict) and record.get("slide_id") == match["slide_id"]),
                    None,
                )
                source_relative = slide.get("master_path") if isinstance(slide, dict) else None
                source = self.validate_relative_file(extraction.get("source_path"), f"{location}.extraction.source_path")
                mask = self.validate_relative_file(extraction.get("mask_path"), f"{location}.extraction.mask_path")
                if extraction.get("source_path") != source_relative:
                    self.error(location, "pixel-preserved source must be the current textless master")
                if source and extraction.get("source_sha256") != hashlib.sha256(source.read_bytes()).hexdigest():
                    self.error(location, "pixel-preserved source hash is stale")
                if mask and extraction.get("mask_sha256") != hashlib.sha256(mask.read_bytes()).hexdigest():
                    self.error(location, "pixel-preserved mask hash is stale")
                if source and mask:
                    try:
                        extractor_path = Path(__file__).resolve().with_name("extract_preserved_component.py")
                        extractor_spec = importlib.util.spec_from_file_location("ai_presentation_pixel_extractor", extractor_path)
                        if extractor_spec is None or extractor_spec.loader is None:
                            raise RuntimeError("cannot load pixel extractor")
                        extractor = importlib.util.module_from_spec(extractor_spec)
                        extractor_spec.loader.exec_module(extractor)
                        expected, box = extractor.extract_pixels(source, mask, match["bbox"], match["crop_padding"])
                        with extractor.Image.open(output_path) as actual:
                            actual_rgba = actual.convert("RGBA")
                        if extraction.get("source_crop_px") != list(box) or expected.size != actual_rgba.size or expected.tobytes() != actual_rgba.tobytes():
                            self.error(location, "pixel-preserved output differs from current source pixels and mask")
                    except (ImportError, OSError, ValueError, TypeError, RuntimeError) as exc:
                        self.error(location, f"cannot verify pixel preservation: {exc}")
        elif isinstance(extraction, dict) and extraction.get("route") == "reconstruct-single":
            if asset.get("strategy") != "generated-transparent-png":
                self.error(location, "reconstructed single object must use generated-transparent-png strategy")

        history = asset.get("extraction_history", [])
        if not isinstance(history, list):
            self.error(location, "extraction_history must be a list")
        else:
            for index, prior in enumerate(history):
                if not isinstance(prior, dict):
                    self.error(f"{location}.extraction_history[{index}]", "historical extraction must be an object")
                    continue
                if prior.get("route") != "preserve-pixels":
                    continue
                prior_location = f"{location}.extraction_history[{index}]"
                prior_output = self.validate_relative_file(prior.get("output_path"), f"{prior_location}.output_path")
                if prior_output:
                    if not prior_output.is_relative_to((self.root / "assets").resolve()) or prior_output.suffix.lower() != ".png":
                        self.error(prior_location, "historical output must be a PNG under assets/")
                    if prior.get("output_path") == asset.get("output") or prior.get("output_sha256") != hashlib.sha256(prior_output.read_bytes()).hexdigest():
                        self.error(prior_location, "historical output is reused or its hash is stale")
                for field in ("source", "mask"):
                    prior_file = self.validate_relative_file(prior.get(f"{field}_path"), f"{prior_location}.{field}_path")
                    if prior_file and prior.get(f"{field}_sha256") != hashlib.sha256(prior_file.read_bytes()).hexdigest():
                        self.error(prior_location, f"historical {field} hash is stale")

        review = asset.get("isolation_review")
        if not isinstance(review, dict) or review.get("status") != "passed" or not isinstance(review.get("reviewer"), str) or not review["reviewer"].strip():
            self.error(location, "passed human isolation review is required")
            return
        actor_error = REVIEW_POLICY.review_actor_error(self.docs.get("run-state.json"), review)
        if actor_error:
            self.error(location, actor_error)
        checks = review.get("checks")
        required_isolation_checks = ["single_object", "no_neighbor_residue", "edge_clean", "concept_match"]
        if asset.get("role") not in PHOTO_ROLES:
            required_isolation_checks.append("no_ordinary_text_or_data")
        for check in required_isolation_checks:
            if not isinstance(checks, dict) or checks.get(check) is not True:
                self.error(location, f"isolation review check must pass: {check}")
        asset_sha = hashlib.sha256(output_path.read_bytes()).hexdigest() if output_path else None
        if review.get("asset_sha256") != asset_sha or review.get("reviewed_items_sha256") != module.value_sha256(source_items):
            self.error(location, "isolation review must bind the current asset and decomposition items")
        image_path = self.validate_relative_file(review.get("image_path"), f"{location}.isolation_review.image_path")
        packet_path = self.validate_relative_file(review.get("packet_path"), f"{location}.isolation_review.packet_path")
        review_dir = (self.root / "reviews" / "assets").resolve()
        if image_path:
            if not image_path.is_relative_to(review_dir) or image_path.suffix.lower() != ".png" or review.get("image_sha256") != hashlib.sha256(image_path.read_bytes()).hexdigest():
                self.error(location, "isolation image hash or format is invalid")
            elif output_path:
                try:
                    expected = module.build_isolation_image(self.root, output_path, source_items, pixels=self.pixel_cache)
                    try:
                        actual_bytes = image_path.read_bytes()
                        if hashlib.sha256(actual_bytes).hexdigest() != review.get("image_sha256"):
                            raise ValueError("isolation image changed during validation")
                        with module.Image.open(io.BytesIO(actual_bytes)) as actual:
                            actual_rgb = actual.convert("RGB")
                        try:
                            if expected.size != actual_rgb.size or expected.tobytes() != actual_rgb.tobytes():
                                self.error(location, "isolation image pixels do not match current component and concept crops")
                        finally:
                            actual_rgb.close()
                    finally:
                        expected.close()
                except (OSError, ValueError, TypeError, KeyError, module.Image.DecompressionBombError, module.Image.DecompressionBombWarning) as exc:
                    self.error(location, f"cannot verify isolation image: {exc}")
        if packet_path:
            if not packet_path.is_relative_to(review_dir) or packet_path.suffix.lower() != ".md" or review.get("packet_sha256") != hashlib.sha256(packet_path.read_bytes()).hexdigest():
                self.error(location, "isolation review list hash or format is invalid")
            else:
                try:
                    readable = packet_path.read_text(encoding="utf-8")
                except UnicodeError:
                    readable = ""
                if module.value_sha256(source_items) not in readable or any(
                    module.review_row(item) not in readable.splitlines() for item in source_items
                ):
                    self.error(location, "isolation review list must include all current item rows and hash")

    def validate_photo_asset(self, asset: dict[str, Any], location: str, output_path: Path | None) -> None:
        if not self.asset_ready(asset) or asset.get("role") not in PHOTO_ROLES:
            return
        source_record = asset.get("source_photo")
        if not isinstance(source_record, dict):
            self.error(location, "photo asset requires a classified source_photo record")
            return
        source_path = self.validate_relative_file(source_record.get("path"), f"{location}.source_photo.path")
        source_sha = source_record.get("sha256")
        if source_path and source_sha != hashlib.sha256(source_path.read_bytes()).hexdigest():
            self.error(location, "photo source hash is stale")
        mode = asset.get("photo_output_mode")
        strategy = asset.get("strategy")
        if asset.get("content_policy") != "preserve-source-photo":
            self.error(location, "ordinary photo must preserve source content, including incidental scene text")
        if mode == "rectangle":
            if asset.get("role") != "ordinary-photo" or asset.get("background_requirement") != "opaque-rectangle" or strategy not in {"preserved-photo-rectangle", "postprocessed-photo-rectangle"}:
                self.error(location, "ordinary photo must be an opaque, independent rectangle")
        elif mode == "cutout":
            if asset.get("role") != "photo-cutout" or asset.get("background_requirement") != "transparent" or strategy not in {"preserved-transparent-png", "postprocessed-transparent-png"}:
                self.error(location, "photo cutout must be an independently transparent PNG")
        else:
            self.error(location, "photo_output_mode must be rectangle or cutout")
        if output_path:
            try:
                from PIL import Image
                with Image.open(output_path) as image:
                    image.verify()
                with Image.open(output_path) as image:
                    if image.format not in {"PNG", "JPEG"} or image.width < 3 or image.height < 3:
                        self.error(location, "photo output must be a readable PNG or JPEG")
                    elif mode == "rectangle" and image.convert("RGBA").getchannel("A").getextrema() != (255, 255):
                        self.error(location, "opaque photo rectangle contains transparent pixels")
                    elif mode == "cutout" and image.format != "PNG":
                        self.error(location, "transparent photo cutout must be PNG")
            except (ImportError, OSError, ValueError) as exc:
                self.error(location, f"cannot inspect photo output: {exc}")
        if isinstance(strategy, str) and strategy.startswith("preserved-"):
            if source_path and output_path and source_path.read_bytes() != output_path.read_bytes():
                self.error(location, "preserved photo output must match the approved source byte-for-byte")
            return
        if not isinstance(strategy, str) or not strategy.startswith("postprocessed-"):
            return
        if source_path and output_path and source_path == output_path:
            self.error(location, "postprocessed photo must use a new output path")
        processing = asset.get("photo_processing")
        if not isinstance(processing, dict):
            self.error(location, "postprocessed photo requires photo_processing provenance")
            return
        confirmation = processing.get("confirmation")
        diagnosis = processing.get("diagnosis")
        recipe = processing.get("preservation_prompt")
        recipe_sha = hashlib.sha256(recipe.encode("utf-8")).hexdigest() if isinstance(recipe, str) and recipe.strip() else None
        approval_time = None
        if isinstance(confirmation, dict) and isinstance(confirmation.get("approved_at"), str):
            try:
                approval_time = datetime.fromisoformat(confirmation["approved_at"])
            except ValueError:
                pass
        if (
            not isinstance(confirmation, dict)
            or confirmation.get("approval_phrase") != "确认照片后期方案"
            or not isinstance(confirmation.get("approved_by"), str)
            or not confirmation["approved_by"].strip()
            or confirmation.get("source_sha256") != source_sha
            or confirmation.get("direction") != processing.get("direction")
            or not isinstance(diagnosis, dict)
            or confirmation.get("diagnosis_sha256") != diagnosis.get("sha256")
            or recipe_sha is None
            or confirmation.get("preservation_prompt_sha256") != recipe_sha
            or approval_time is None
            or approval_time.utcoffset() is None
        ):
            self.error(location, "photo-processing diagnosis, direction, and recipe lack explicit source-bound human confirmation")
        if not isinstance(diagnosis, dict):
            self.error(location, "photo-processing diagnosis evidence is missing")
        else:
            diagnosis_path = self.validate_relative_file(diagnosis.get("path"), f"{location}.photo_processing.diagnosis.path")
            if diagnosis_path and diagnosis.get("sha256") != hashlib.sha256(diagnosis_path.read_bytes()).hexdigest():
                self.error(location, "photo-processing diagnosis hash is stale")
            elif diagnosis_path:
                try:
                    diagnosis_doc = json.loads(diagnosis_path.read_text(encoding="utf-8"))
                except (OSError, UnicodeError, json.JSONDecodeError):
                    diagnosis_doc = None
                measurements = diagnosis_doc.get("measurements") if isinstance(diagnosis_doc, dict) else None
                if not isinstance(diagnosis_doc, dict) or diagnosis_doc.get("source_sha256") != source_sha or not isinstance(measurements, dict) or not {"luminance", "clipping", "noise", "sharpness", "photo_type"}.issubset(measurements):
                    self.error(location, "photo diagnosis must measure and bind the current source")
        if not isinstance(processing.get("direction"), str) or not processing["direction"].strip() or not isinstance(processing.get("preservation_prompt"), str) or not processing["preservation_prompt"].strip():
            self.error(location, "photo-processing direction and preservation prompt are required")
        attempts = processing.get("attempts")
        if not isinstance(attempts, list) or not attempts:
            self.error(location, "photo-processing attempts must retain each result")
        else:
            brief = self.docs.get("brief.json", {})
            budget = brief.get("capacity_budget", {}) if isinstance(brief, dict) else {}
            retry_limit = budget.get("retry_limit_per_asset") if isinstance(budget, dict) else None
            if isinstance(retry_limit, int) and len(attempts) - 1 > retry_limit:
                self.error(location, "photo-processing attempts exceed retry_limit_per_asset")
            accepted = []
            seen_ids: set[str] = set()
            seen_paths: set[str] = set()
            for index, attempt in enumerate(attempts):
                attempt_location = f"{location}.photo_processing.attempts[{index}]"
                if not isinstance(attempt, dict):
                    self.error(attempt_location, "photo-processing attempt must be an object")
                    continue
                call_id = attempt.get("call_id")
                if not isinstance(call_id, str) or not call_id or call_id in seen_ids:
                    self.error(attempt_location, "photo-processing call_id must be unique")
                else:
                    seen_ids.add(call_id)
                if attempt.get("source_sha256") != source_sha or attempt.get("component_id") != asset.get("id"):
                    self.error(attempt_location, "photo-processing attempt must bind this source and component")
                if not isinstance(attempt.get("tool"), str) or not attempt["tool"].strip() or attempt.get("prompt_or_recipe") != processing.get("preservation_prompt"):
                    self.error(attempt_location, "photo-processing attempt must retain its tool and approved preservation recipe")
                relative = attempt.get("output_path")
                if not isinstance(relative, str) or relative in seen_paths:
                    self.error(attempt_location, "photo-processing output path must be unique")
                else:
                    seen_paths.add(relative)
                    attempt_path = self.validate_relative_file(relative, f"{attempt_location}.output_path")
                    if attempt_path and attempt.get("sha256") != hashlib.sha256(attempt_path.read_bytes()).hexdigest():
                        self.error(attempt_location, "photo-processing attempt output hash is stale")
                if attempt.get("status") == "accepted":
                    accepted.append(attempt)
                elif attempt.get("status") != "rejected":
                    self.error(attempt_location, "photo-processing attempt status must be accepted or rejected")
            if len(accepted) != 1 or accepted[0].get("output_path") != asset.get("output") or accepted[0].get("sha256") != asset.get("sha256"):
                self.error(location, "one accepted photo-processing result must match the current output")
        review = asset.get("photo_review")
        if not isinstance(review, dict) or review.get("status") != "passed" or not isinstance(review.get("reviewer"), str) or not review["reviewer"].strip():
            self.error(location, "postprocessed photo requires passed human source-vs-output review")
            return
        checks = review.get("checks")
        for check in ("subject_identity", "architecture", "perspective", "key_scene_content", "no_factual_additions", "source_text_preserved"):
            if not isinstance(checks, dict) or checks.get(check) is not True:
                self.error(location, f"photo preservation check must pass: {check}")
        if review.get("source_sha256") != source_sha or review.get("output_sha256") != asset.get("sha256"):
            self.error(location, "photo review must bind the current source and output")
        image_path = self.validate_relative_file(review.get("image_path"), f"{location}.photo_review.image_path")
        if image_path and source_path and output_path:
            if not image_path.is_relative_to((self.root / "reviews" / "photos").resolve()) or review.get("image_sha256") != hashlib.sha256(image_path.read_bytes()).hexdigest():
                self.error(location, "photo comparison evidence path or hash is invalid")
            else:
                try:
                    review_script = Path(__file__).resolve().with_name("render_photo_review.py")
                    spec = importlib.util.spec_from_file_location("ai_presentation_photo_review", review_script)
                    if spec is None or spec.loader is None:
                        raise RuntimeError("cannot load photo comparison renderer")
                    module = importlib.util.module_from_spec(spec)
                    spec.loader.exec_module(module)
                    expected = module.comparison_image(source_path, output_path)
                    with module.Image.open(image_path) as actual:
                        actual_rgb = actual.convert("RGB")
                    if expected.size != actual_rgb.size or expected.tobytes() != actual_rgb.tobytes():
                        self.error(location, "photo comparison image does not match current source and output")
                except (ImportError, OSError, ValueError, RuntimeError) as exc:
                    self.error(location, f"cannot verify photo comparison: {exc}")

    def validate_manifest(self) -> None:
        name = "asset-manifest.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        assets = doc.get("assets")
        if not isinstance(assets, list):
            self.error(name, "assets must be a list")
            return
        self.check_unique(assets, "id", f"{name}.assets")
        brief = self.docs.get("brief.json", {})
        budget = brief.get("capacity_budget", {}) if isinstance(brief, dict) else {}
        limit = budget.get("component_limit") if isinstance(budget, dict) else None
        if isinstance(limit, int) and len(assets) > limit:
            self.error(name, "actual asset count exceeds approved component_limit")
        deck = self.docs.get("deck-spec.json", {})
        estimate = deck.get("production_estimate", {}) if isinstance(deck, dict) else {}
        if isinstance(estimate, dict) and estimate.get("unique_components") != len(assets):
            self.error(name, "production_estimate.unique_components must equal manifest asset count")
        generated_assets = [
            asset
            for asset in assets
            if isinstance(asset, dict) and str(asset.get("strategy", "")).startswith("generated-")
        ]
        processed_photo_assets = [
            asset for asset in assets
            if isinstance(asset, dict) and asset.get("role") in PHOTO_ROLES
            and str(asset.get("strategy", "")).startswith("postprocessed-")
        ]
        if isinstance(estimate, dict) and estimate.get("planned_generation_calls", 0) < len(generated_assets) + len(processed_photo_assets):
            self.error(name, "planned_generation_calls is lower than generated and postprocessed asset count")
        actual_calls = 0
        seen_generation_call_ids: set[str] = set()
        seen_generation_output_paths: set[str] = set()
        seen_asset_outputs: set[str] = set()
        style = self.docs.get("style-contract.json", {})
        run_state = self.docs.get("run-state.json", {})
        gates = run_state.get("gates", {}) if isinstance(run_state, dict) else {}
        style_gate = gates.get("style_anchors", {}) if isinstance(gates, dict) else {}
        concept_gate = (
            gates.get("concept_contact_sheet", {}) if isinstance(gates, dict) else {}
        )
        style_path = self.root / "style-contract.json"
        style_sha256 = (
            hashlib.sha256(style_path.read_bytes()).hexdigest()
            if style_path.is_file()
            else None
        )
        for index, asset in enumerate(assets):
            location = f"{name}.assets[{index}]"
            if not isinstance(asset, dict):
                self.error(location, "must be an object")
                continue
            role = asset.get("role")
            if role in PHOTO_ROLES or (role in VISUAL_ROLES and str(asset.get("strategy", "")).startswith("generated-")):
                self.validate_visual_provider_provenance(
                    asset,
                    location,
                    style,
                    style_gate,
                    concept_gate,
                    style_sha256,
                )
            if role in PHOTO_ROLES and str(asset.get("strategy", "")).startswith("postprocessed-"):
                processing = asset.get("photo_processing")
                attempts = processing.get("attempts") if isinstance(processing, dict) else None
                if isinstance(attempts, list):
                    actual_calls += len(attempts)
            slide_ids = asset.get("slide_ids")
            if slide_ids is None:
                slide_ids = [asset.get("slide_id")] if asset.get("slide_id") else []
            if not isinstance(slide_ids, list) or not slide_ids or any(not value for value in slide_ids):
                self.error(location, "asset requires slide_id or non-empty slide_ids")
            elif len(slide_ids) != len(set(slide_ids)):
                self.error(location, "asset slide_ids must be unique")
            if role in {"ordinary-text", "data-value", "body-copy"}:
                self.error(location, f"'{role}' must be a fill slot, not a generated asset")
            if role not in VISUAL_ROLES | TRUTH_SENSITIVE_ROLES:
                self.error(location, f"unknown or missing visual role '{role}'")
            if role in VISUAL_ROLES:
                if role != "special-typography":
                    if role not in PHOTO_ROLES and asset.get("contains_text") is not False:
                        self.error(location, "visual asset requires explicit contains_text=false")
                    if role in PHOTO_ROLES and type(asset.get("contains_text")) is not bool:
                        self.error(location, "photo contains_text must explicitly describe incidental source text")
                    if asset.get("contains_data") is not False:
                        self.error(location, "visual asset requires explicit contains_data=false")
                    if asset.get("contains_data_encoding_marks") is not False:
                        self.error(
                            location,
                            "visual asset requires explicit contains_data_encoding_marks=false",
                        )
                if asset.get("editability") != "component-editable":
                    self.error(location, f"{role} requires editability 'component-editable'")
                strategy = asset.get("strategy", "")
                if role not in PHOTO_ROLES:
                    if asset.get("background_requirement") != "transparent":
                        self.error(location, f"{role} must use a transparent background")
                    if not isinstance(strategy, str) or not strategy.endswith("transparent-png"):
                        self.error(location, f"{role} strategy must end with 'transparent-png'")
                if isinstance(strategy, str) and strategy.startswith("generated-") and self.asset_ready(asset):
                    attempts = asset.get("generation_attempts")
                    if not isinstance(attempts, int) or attempts < 1:
                        self.error(location, "generated asset requires positive generation_attempts")
                    else:
                        actual_calls += attempts
                        retry_limit = budget.get("retry_limit_per_asset") if isinstance(budget, dict) else None
                        if isinstance(retry_limit, int) and attempts - 1 > retry_limit:
                            self.error(location, "generation_attempts exceeds retry_limit_per_asset")
                    history = asset.get("generation_history")
                    if not isinstance(history, list) or not history or len(history) != attempts:
                        self.error(location, "generation_history must record every one-component call")
                    else:
                        accepted = []
                        for call_index, call in enumerate(history):
                            call_location = f"{location}.generation_history[{call_index}]"
                            if not isinstance(call, dict):
                                self.error(call_location, "generation call must be an object")
                                continue
                            call_id = call.get("call_id")
                            if not isinstance(call_id, str) or not call_id or call_id in seen_generation_call_ids:
                                self.error(call_location, "generation call_id must be unique and non-empty")
                            else:
                                seen_generation_call_ids.add(call_id)
                            if call.get("component_id") != asset.get("id"):
                                self.error(call_location, "generation call must request exactly this one component")
                            if call.get("status") == "accepted":
                                accepted.append(call)
                            elif call.get("status") != "rejected":
                                self.error(call_location, "generation call status must be accepted or rejected")
                            if not isinstance(call.get("output_path"), str) or not isinstance(call.get("sha256"), str):
                                self.error(call_location, "generation call must retain its output path and hash")
                            else:
                                if call["output_path"] in seen_generation_output_paths:
                                    self.error(call_location, "generation attempt output path must be immutable and unique")
                                seen_generation_output_paths.add(call["output_path"])
                                call_output = self.validate_relative_file(call["output_path"], f"{call_location}.output_path")
                                if call_output and hashlib.sha256(call_output.read_bytes()).hexdigest() != call["sha256"]:
                                    self.error(call_location, "generation call output hash is stale")
                        if len(accepted) != 1 or accepted[0].get("output_path") != asset.get("output") or accepted[0].get("sha256") != asset.get("sha256"):
                            self.error(location, "one accepted generation call must match the current asset output")
            if role in {"table-shell", "chart-shell"}:
                if asset.get("content_policy") != "blank-fill-slots":
                    self.error(location, f"{role} requires content_policy 'blank-fill-slots'")
                if asset.get("contains_text") is not False:
                    self.error(location, f"{role} requires explicit contains_text=false")
                if asset.get("contains_data") is not False:
                    self.error(location, f"{role} requires explicit contains_data=false")
                if asset.get("contains_data_encoding_marks") is not False:
                    self.error(
                        location,
                        f"{role} requires explicit contains_data_encoding_marks=false",
                    )
                fill_slot_ids = asset.get("fill_slot_ids")
                if not isinstance(fill_slot_ids, list) or not fill_slot_ids:
                    self.error(location, f"{role} requires non-empty fill_slot_ids")
                elif any(
                    not isinstance(fill_slot_id, str)
                    or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", fill_slot_id)
                    for fill_slot_id in fill_slot_ids
                ):
                    self.error(location, f"{role} fill_slot_ids must be safe string ids")
                elif len(set(fill_slot_ids)) != len(fill_slot_ids):
                    self.error(location, f"{role} fill_slot_ids must be unique")
            elif role not in PHOTO_ROLES | {"special-typography"} and asset.get("contains_text"):
                self.error(location, "ordinary text must remain a fill slot")
            if role == "special-typography":
                if not asset.get("exact_text"):
                    self.error(location, "special-typography requires exact_text")
                if not asset.get("language"):
                    self.error(location, "special-typography requires language")
                ocr_result = asset.get("ocr_result")
                if not isinstance(ocr_result, dict) or ocr_result.get("status") != "passed":
                    self.error(location, "special-typography requires ocr_result.status='passed'")
                elif ocr_result.get("text") != asset.get("exact_text"):
                    self.error(location, "ocr_result.text must exactly match exact_text")
                if self.reaches("release") and asset.get("human_approved") is not True:
                    self.error(location, "special-typography requires human_approved=true before release")
            if role in TRUTH_SENSITIVE_ROLES:
                if asset.get("editability") != "fixed-evidence":
                    self.error(location, f"{role} requires editability 'fixed-evidence'")
                if asset.get("strategy") not in {
                    "preserved-raster",
                    "preserved-vector",
                    "preserved-truth-sensitive",
                }:
                    self.error(location, f"{role} requires a non-generative preservation strategy")
                source_path = self.validate_relative_file(asset.get("source"), f"{location}.source")
                source_sha = asset.get("source_sha256", "")
                if not re.fullmatch(r"[0-9a-fA-F]{64}", str(source_sha)):
                    self.error(location, "source_sha256 must be 64 hexadecimal characters")
                elif source_path and source_path.is_file():
                    actual_source_sha = hashlib.sha256(source_path.read_bytes()).hexdigest()
                    if actual_source_sha.lower() != str(source_sha).lower():
                        self.error(location, "declared source_sha256 does not match source file")
            else:
                source_path = None
            if self.asset_ready(asset):
                required_fields = ["output", "sha256", "status"]
                if role in VISUAL_ROLES:
                    if role != "ordinary-photo":
                        required_fields.append("alpha")
                    if str(asset.get("strategy", "")).startswith("generated-"):
                        required_fields.append("prompt")
                for field in required_fields:
                    if not asset.get(field):
                        self.error(location, f"missing '{field}' from assets stage")
                output = asset.get("output")
                if isinstance(output, str):
                    if output in seen_asset_outputs:
                        self.error(location, "asset output path must be unique; reuse one component ID instead")
                    seen_asset_outputs.add(output)
                if output:
                    output_path = self.validate_relative_file(output, f"{location}.output")
                else:
                    output_path = None
                sha256 = asset.get("sha256", "")
                if not re.fullmatch(r"[0-9a-fA-F]{64}", str(sha256)):
                    self.error(location, "sha256 must be 64 hexadecimal characters")
                elif output_path and output_path.is_file():
                    actual_sha = hashlib.sha256(output_path.read_bytes()).hexdigest()
                    if actual_sha.lower() != str(sha256).lower():
                        self.error(location, "declared SHA-256 does not match the output file")
                if role in VISUAL_ROLES:
                    if role != "ordinary-photo":
                        alpha = asset.get("alpha")
                        if not isinstance(alpha, dict) or alpha.get("has_alpha") is not True:
                            self.error(location, "alpha.has_alpha must be true")
                        if output_path and output_path.is_file():
                            png_ok, ratio, clear_border, visible = self.inspect_png(output_path)
                            if not png_ok:
                                self.error(location, "output must be a valid PNG file")
                            elif ratio < 0.05 or not clear_border or not visible:
                                self.error(location, "PNG must have meaningful transparent padding and visible content")
                            declared_ratio = alpha.get("transparent_pixel_ratio") if isinstance(alpha, dict) else None
                            if not isinstance(declared_ratio, (int, float)):
                                self.error(location, "alpha.transparent_pixel_ratio must be numeric")
                            elif abs(float(declared_ratio) - ratio) > 0.02:
                                self.error(location, "declared transparent_pixel_ratio differs from decoded pixels")
                    self.validate_photo_asset(asset, location, output_path)
                    self.validate_asset_isolation(asset, location, output_path)
                    if asset.get("status") != "approved":
                        self.error(location, "visual asset status must be approved after isolation review")
                if self.reaches("release") and asset.get("status") != "approved":
                    self.error(location, "asset status must be 'approved' before release")
                if role in TRUTH_SENSITIVE_ROLES and output_path and source_path:
                    if asset.get("strategy") in {
                        "preserved-raster",
                        "preserved-vector",
                        "preserved-truth-sensitive",
                    } and hashlib.sha256(output_path.read_bytes()).digest() != hashlib.sha256(
                        source_path.read_bytes()
                    ).digest():
                        self.error(location, "preserved output must match its approved source exactly")
        if self.reaches("pilot"):
            call_limit = budget.get("generation_call_limit") if isinstance(budget, dict) else None
            if isinstance(call_limit, int) and actual_calls > call_limit:
                self.error(name, "actual generation calls exceed generation_call_limit")
            reserved_retries = estimate.get("reserved_retries", 0) if isinstance(estimate, dict) else 0
            actual_retries = actual_calls - len(generated_assets) - len(processed_photo_assets)
            if self.reaches("release") and actual_retries > reserved_retries:
                self.error(name, "actual retries exceed the approved reserved_retries")

    def validate_assembly(self) -> None:
        name = "assembly-map.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        slides = doc.get("slides")
        if not isinstance(slides, list) or not slides:
            self.error(name, "slides must be a non-empty list")
            return
        self.check_unique(slides, "slide_id", f"{name}.slides")
        output_indices = [slide.get("output_index") for slide in slides if isinstance(slide, dict)]
        expected_indices = list(range(1, len(slides) + 1))
        if any(not isinstance(value, int) for value in output_indices) or sorted(output_indices) != expected_indices:
            self.error(name, "output_index values must be unique and continuous from 1")
        deck = self.docs.get("deck-spec.json", {})
        deck_order = [slide.get("id") for slide in deck.get("slides", []) if isinstance(slide, dict)]
        assembly_order = []
        if all(isinstance(value, int) for value in output_indices):
            assembly_order = [
                slide.get("slide_id")
                for slide in sorted(
                    (slide for slide in slides if isinstance(slide, dict)),
                    key=lambda item: item.get("output_index", 0),
                )
            ]
        if assembly_order != deck_order:
            self.error(name, "output_index order must match deck-spec slide order")
        for index, slide in enumerate(slides):
            if not isinstance(slide, dict):
                self.error(f"{name}.slides[{index}]", "must be an object")
                continue
            if slide.get("flattened_concept_used") is not False:
                self.error(
                    f"{name}.slides[{index}]",
                    "flattened_concept_used must be false; concepts are references only",
                )

        brief = self.docs.get("brief.json", {})
        delivery = brief.get("delivery", {}) if isinstance(brief, dict) else {}
        if "pptx" in delivery.get("formats", []) and not doc.get("backend"):
            self.error(name, "backend is required when PPTX delivery is requested")
        if "pptx" in delivery.get("formats", []):
            if doc.get("status") not in {"structure-assembled-visual-review-pending", "visual-review-passed"}:
                self.error(name, "assembly status must declare pending or passed visual review")
            elif self.reaches("qa") and doc.get("status") != "visual-review-passed":
                self.error(name, "assembly visual review must pass before QA")
            if self.quality_profile() == "production" and doc.get("render_backend") != "powerpoint":
                self.error(name, "production assembly requires actual PowerPoint render evidence")
            input_hashes = doc.get("input_sha256")
            for source_name in ("deck-spec.json", "scene-graph.json", "asset-manifest.json", "style-contract.json"):
                source_path = self.root / source_name
                if not source_path.is_file() or not isinstance(input_hashes, dict) or input_hashes.get(source_name) != hashlib.sha256(source_path.read_bytes()).hexdigest():
                    self.error(name, f"assembly source is stale: {source_name}")
            batch_override_path = self.root / "assembly-style-overrides-batch.json"
            if batch_override_path.is_file() and (
                not isinstance(input_hashes, dict)
                or input_hashes.get(batch_override_path.name) != hashlib.sha256(batch_override_path.read_bytes()).hexdigest()
            ):
                self.error(name, "assembly source is stale: assembly-style-overrides-batch.json")
            pptx_parts = [slide.get("pptx_part") for slide in slides if isinstance(slide, dict)]
            if any(not value for value in pptx_parts) or len(pptx_parts) != len(set(pptx_parts)):
                self.error(name, "each PPTX slide requires a unique pptx_part")
            output_value = doc.get("output")
            if not output_value:
                self.error(name, "PPTX output path is required")
                return
            output_path = self.validate_relative_file(output_value, f"{name}.output")
            if not output_path:
                self.error(name, "PPTX output file is required")
                return
            if output_path.suffix.lower() != ".pptx":
                self.error(name, "PPTX output must use the .pptx extension")
            declared_sha = doc.get("sha256", "")
            actual_sha = hashlib.sha256(output_path.read_bytes()).hexdigest()
            if declared_sha != actual_sha:
                self.error(name, "PPTX SHA-256 does not match the output file")
            renders = doc.get("renders")
            if not isinstance(renders, list) or len(renders) != len(slides):
                self.error(name, "every assembled slide needs a full-size rendered image")
            else:
                for index, record in enumerate(renders):
                    render_location = f"{name}.renders[{index}]"
                    if not isinstance(record, dict) or record.get("pptx_sha256") != actual_sha:
                        self.error(render_location, "render must bind the current PPTX hash")
                        continue
                    render_path = self.validate_relative_file(record.get("path"), f"{render_location}.path")
                    if render_path:
                        if not render_path.is_relative_to((self.root / "reviews" / "assembly").resolve()) or record.get("sha256") != hashlib.sha256(render_path.read_bytes()).hexdigest():
                            self.error(render_location, "render path or hash is invalid")
                        else:
                            try:
                                from PIL import Image
                                with Image.open(render_path) as image:
                                    if image.format != "PNG" or image.size != (1920, 1080):
                                        self.error(render_location, "render must be a full-size 1920x1080 PNG")
                                    image.verify()
                            except (ImportError, OSError, ValueError) as exc:
                                self.error(render_location, f"render cannot be decoded: {exc}")
            self.validate_pptx_content(output_path)
            self.validate_object_level_assembly(output_path)

    def validate_object_level_assembly(self, path: Path) -> None:
        """Check that the PPTX itself still contains every independently editable instance."""
        location = "assembly-map.json.output"
        try:
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
            presentation = Presentation(path)
        except (ImportError, OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
            self.error(location, f"cannot inspect editable PPTX objects: {exc}")
            return
        deck = self.docs.get("deck-spec.json", {})
        scene = self.docs.get("scene-graph.json", {})
        manifest = self.docs.get("asset-manifest.json", {})
        assembly = self.docs.get("assembly-map.json", {})
        style = self.docs.get("style-contract.json", {})
        override_path = self.root / "assembly-style-overrides.json"
        try:
            style_overrides = json.loads(override_path.read_text(encoding="utf-8")) if override_path.is_file() else {}
            batch_override_path = self.root / "assembly-style-overrides-batch.json"
            if batch_override_path.is_file():
                batch_overrides = json.loads(batch_override_path.read_text(encoding="utf-8"))
                style_overrides.setdefault("slides", {}).update(batch_overrides.get("slides", {}))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            self.error(location, f"cannot read assembly text-format overrides: {exc}")
            style_overrides = {}
        if not all(isinstance(doc, dict) for doc in (deck, scene, manifest, assembly)):
            return
        semantic_system = style.get("semantic_system", {}) if isinstance(style, dict) else {}
        text_roles = semantic_system.get("text_roles", {}) if isinstance(semantic_system, dict) else {}
        frame_roles = semantic_system.get("frame_roles", {}) if isinstance(semantic_system, dict) else {}
        deck_slides = [slide for slide in deck.get("slides", []) if isinstance(slide, dict)]
        scene_slides = {slide.get("slide_id"): slide for slide in scene.get("slides", []) if isinstance(slide, dict)}
        assembly_slides = {slide.get("slide_id"): slide for slide in assembly.get("slides", []) if isinstance(slide, dict)}
        assets = {asset.get("id"): asset for asset in manifest.get("assets", []) if isinstance(asset, dict)}
        if len(presentation.slides) != len(deck_slides):
            self.error(location, "editable PPTX slide count differs from deck-spec")
            return
        for slide_index, (slide, deck_slide) in enumerate(zip(presentation.slides, deck_slides), start=1):
            slide_id = deck_slide.get("id")
            slide_location = f"{location}.slide[{slide_index}]"
            scene_slide = scene_slides.get(slide_id, {})
            assembled = assembly_slides.get(slide_id, {})
            try:
                expected_notes = speaker_notes_text(deck_slide)
            except ValueError as exc:
                self.error(slide_location, str(exc))
            else:
                notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
                actual_notes = notes_frame.text if notes_frame is not None else ""
                if actual_notes != expected_notes:
                    self.error(slide_location, "PPTX speaker notes differ from the current deck spec")
                expected_notes_sha = hashlib.sha256(expected_notes.encode("utf-8")).hexdigest() if expected_notes else None
                if expected_notes or "speaker_notes_sha256" in assembled:
                    if assembled.get("speaker_notes_sha256") != expected_notes_sha:
                        self.error(slide_location, "assembly speaker notes hash differs from the current deck spec")
            slide_overrides = style_overrides.get("slides", {}).get(slide_id, {}) if isinstance(style_overrides, dict) else {}
            if slide_overrides and slide_overrides.get("concept_sha256") != hashlib.sha256((self.root / deck_slide.get("concept_path", "")).read_bytes()).hexdigest():
                self.error(slide_location, "text-format overrides are stale against the approved concept")
            if not isinstance(scene_slide, dict) or not isinstance(assembled, dict):
                self.error(slide_location, "missing current scene or assembly slide")
                continue
            elements = scene_slide.get("elements", [])
            if not isinstance(elements, list):
                self.error(slide_location, "scene elements must be a list")
                continue
            if any(type(item.get("z")) is not int for item in elements if isinstance(item, dict)):
                self.error(slide_location, "scene z-order is invalid")
                continue
            ordered = sorted((item for item in elements if isinstance(item, dict)), key=lambda item: (item["z"], str(item.get("instance_id", ""))))
            slots = [item for item in deck_slide.get("fill_slots", []) if isinstance(item, dict)]
            expected_names = [f"APW:visual:{item.get('instance_id')}:{item.get('component_id')}" for item in ordered]
            expected_names.extend(f"APW:slot:{slot.get('id')}" for slot in slots)
            actual_names = [shape.name for shape in slide.shapes]
            if actual_names != expected_names:
                self.error(slide_location, "PPTX object count, identity, or z-order differs from scene graph and fill slots")
                continue
            object_rows = assembled.get("objects")
            slot_rows = assembled.get("slots")
            if not isinstance(object_rows, list) or len(object_rows) != len(ordered) or not isinstance(slot_rows, list) or len(slot_rows) != len(slots):
                self.error(slide_location, "assembly object and slot records are incomplete")
                continue
            for index, element in enumerate(ordered):
                shape = slide.shapes[index]
                asset = assets.get(element.get("component_id"))
                row = object_rows[index]
                item_location = f"{slide_location}.object[{index}]"
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or not isinstance(asset, dict) or not isinstance(row, dict):
                    self.error(item_location, "visual instance must be a separate PPTX picture object")
                    continue
                expected_row = {
                    "instance_id": element.get("instance_id"), "component_id": element.get("component_id"),
                    "name": shape.name, "bbox": element.get("bbox"), "z": element.get("z"),
                    "rotation": element.get("rotation", 0), "asset_sha256": asset.get("sha256"),
                }
                placement_crop = asset.get("placement_crop")
                if isinstance(placement_crop, dict):
                    expected_row["placement_crop"] = placement_crop
                    if any(
                        type(placement_crop.get(side)) not in {int, float}
                        or not 0 <= placement_crop[side] < 0.45
                        or abs(getattr(shape, f"crop_{side}") - placement_crop[side]) > 0.0001
                        for side in ("left", "top", "right", "bottom")
                    ) or placement_crop["left"] + placement_crop["right"] >= 0.95 or placement_crop["top"] + placement_crop["bottom"] >= 0.95:
                        self.error(item_location, "PPTX image crop differs from the approved asset placement")
                if row != expected_row:
                    self.error(item_location, "assembly object record is stale against scene graph or asset")
                if hashlib.sha256(shape.image.blob).hexdigest() != asset.get("sha256"):
                    self.error(item_location, "PPTX picture bytes do not match its approved component")
                self.validate_pptx_shape_geometry(shape, element.get("bbox"), presentation, item_location)
                if abs(float(shape.rotation) - float(element.get("rotation", 0))) > 0.01:
                    self.error(item_location, "PPTX object rotation differs from scene graph")
            regions = {
                region.get("fill_slot_id"): region for region in deck_slide.get("readability", {}).get("text_regions", [])
                if isinstance(region, dict)
            }
            for index, slot in enumerate(slots):
                shape = slide.shapes[len(ordered) + index]
                row = slot_rows[index]
                slot_location = f"{slide_location}.slot[{index}]"
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or not shape.has_text_frame or not isinstance(row, dict):
                    self.error(slot_location, "fill slot must be a native editable text box")
                    continue
                content = slot.get("content", "")
                if not content and regions.get(slot.get("id"), {}).get("role") == "action-title":
                    content = deck_slide.get("title", "")
                if slot.get("kind") == "ordinary-data" and content:
                    self.error(slot_location, "data fill slot must stay blank")
                slot_format = slide_overrides.get("slots", {}).get(slot.get("id"), slot.get("format"))
                display_content = slot_format.get("display_prefix", "") + content if isinstance(slot_format, dict) else content
                expected_row = {
                    "id": slot.get("id"), "kind": slot.get("kind"), "name": shape.name,
                    "bbox": slot.get("bbox"), "content": display_content,
                    "font_pt": regions.get(slot.get("id"), {}).get("planned_font_pt", 20),
                }
                font_role = regions.get(slot.get("id"), {}).get("font_role")
                font_token = text_roles.get(font_role) if isinstance(text_roles, dict) and isinstance(font_role, str) else None
                if isinstance(font_token, dict):
                    expected_row["font_role"] = font_role
                if isinstance(slot_format, dict):
                    expected_row["format"] = slot_format
                    expected_row["font_pt"] = slot_format.get("font_pt", expected_row["font_pt"])
                if row != expected_row or shape.text != display_content:
                    self.error(slot_location, "PPTX editable text differs from approved fill-slot content")
                self.validate_pptx_shape_geometry(shape, slot.get("bbox"), presentation, slot_location)
                paragraph = shape.text_frame.paragraphs[0]
                expected_font = expected_row["font_pt"]
                if paragraph.font.size is None or abs(paragraph.font.size.pt - expected_font) > 0.1:
                    self.error(slot_location, "PPTX text size differs from the readability plan")
                if isinstance(font_token, dict):
                    if paragraph.font.name != font_token.get("font_family"):
                        self.error(slot_location, "PPTX font family differs from semantic text role")
                    expected_bold = slot_format.get("bold", font_token.get("bold")) if isinstance(slot_format, dict) else font_token.get("bold")
                    if paragraph.font.bold is not expected_bold:
                        self.error(slot_location, "PPTX font weight differs from semantic text role")
                    expected_color = slot_format.get("color", font_token.get("color")) if isinstance(slot_format, dict) else font_token.get("color")
                    try:
                        actual_color = paragraph.font.color.rgb
                    except (AttributeError, TypeError, ValueError):
                        actual_color = None
                    if not isinstance(expected_color, str) or actual_color != RGBColor.from_string(expected_color[1:]):
                        self.error(slot_location, "PPTX font color differs from semantic text role")
                if isinstance(slot_format, dict) and "align" in slot_format:
                    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(slot_format["align"])
                    if alignment is None or paragraph.alignment != alignment:
                        self.error(slot_location, "PPTX text alignment differs from the slot format")
                if isinstance(slot_format, dict) and "vertical_align" in slot_format:
                    vertical = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                                "bottom": MSO_ANCHOR.BOTTOM}.get(slot_format["vertical_align"])
                    if vertical is None or shape.text_frame.vertical_anchor != vertical:
                        self.error(slot_location, "PPTX text vertical alignment differs from the slot format")
                rich_runs = slot_format.get("runs") if isinstance(slot_format, dict) else None
                if rich_runs is not None:
                    if not isinstance(rich_runs, list) or len(paragraph.runs) != len(rich_runs) or any(
                        not isinstance(part, dict) or run.text != part.get("text")
                        for run, part in zip(paragraph.runs, rich_runs)
                    ):
                        self.error(slot_location, "PPTX rich text runs differ from the approved text")
                    else:
                        for run, part in zip(paragraph.runs, rich_runs):
                            expected_run_color = part.get("color", slot_format.get("color", font_token.get("color") if isinstance(font_token, dict) else "#F8FAFC"))
                            if (run.font.size is None or abs(run.font.size.pt - part.get("font_pt", expected_font)) > 0.1
                                or run.font.bold is not part.get("bold", expected_bold if isinstance(font_token, dict) else None)
                                or run.font.color.rgb != RGBColor.from_string(expected_run_color[1:])):
                                self.error(slot_location, "PPTX rich text styling differs from approved assembly formatting")
                if slot.get("kind") == "ordinary-data" and isinstance(frame_roles, dict):
                    frame_role = slot.get("frame_role")
                    frame_token = frame_roles.get(frame_role) if isinstance(frame_role, str) else None
                    if isinstance(frame_token, dict) and isinstance(frame_token.get("color"), str):
                        try:
                            actual_frame_color = shape.line.color.rgb
                        except (AttributeError, TypeError, ValueError):
                            actual_frame_color = None
                        if actual_frame_color != RGBColor.from_string(frame_token["color"][1:]):
                            self.error(slot_location, "PPTX data frame color differs from semantic frame role")

    def validate_pptx_shape_geometry(self, shape: Any, bbox: Any, presentation: Any, location: str) -> None:
        if not isinstance(bbox, list) or len(bbox) != 4 or any(type(value) not in {float, int} for value in bbox):
            self.error(location, "approved bbox is invalid")
            return
        expected = [round(bbox[0] * presentation.slide_width), round(bbox[1] * presentation.slide_height),
                    round(bbox[2] * presentation.slide_width), round(bbox[3] * presentation.slide_height)]
        actual = [shape.left, shape.top, shape.width, shape.height]
        if any(abs(current - target) > 2 for current, target in zip(actual, expected)):
            self.error(location, "PPTX object geometry differs from current scene or fill-slot bbox")

    def pilot_bindings(self) -> dict[str, Any] | None:
        """Bind only the trial slide, so later work on other slides does not stale approval."""
        plan = self.docs.get("pilot-plan.json", {})
        deck = self.docs.get("deck-spec.json", {})
        scene = self.docs.get("scene-graph.json", {})
        decomposition = self.docs.get("decomposition-map.json", {})
        manifest = self.docs.get("asset-manifest.json", {})
        if not all(isinstance(doc, dict) for doc in (plan, deck, scene, decomposition, manifest)):
            return None
        slide_id = plan.get("slide_id")
        deck_slide = next((slide for slide in deck.get("slides", []) if isinstance(slide, dict) and slide.get("id") == slide_id), None)
        scene_slide = next((slide for slide in scene.get("slides", []) if isinstance(slide, dict) and slide.get("slide_id") == slide_id), None)
        map_slide = next((slide for slide in decomposition.get("slides", []) if isinstance(slide, dict) and slide.get("slide_id") == slide_id), None)
        if not all(isinstance(item, dict) for item in (deck_slide, scene_slide, map_slide)):
            return None
        raw_ids = [item.get("component_id") for item in scene_slide.get("elements", []) if isinstance(item, dict)]
        if any(not isinstance(item, str) for item in raw_ids):
            return None
        component_ids = sorted(set(raw_ids))
        if plan.get("component_ids") != component_ids:
            return None
        assets = {asset.get("id"): asset for asset in manifest.get("assets", []) if isinstance(asset, dict)}
        if any(component_id not in assets for component_id in component_ids):
            return None
        concept_path = self.validate_relative_file(deck_slide.get("concept_path"), "pilot-plan.json.concept_path")
        if concept_path is None:
            return None
        plan_path = self.root / "pilot-plan.json"
        style_path = self.root / "style-contract.json"
        override_path = self.root / "assembly-style-overrides.json"
        return {
            "pilot_plan_sha256": hashlib.sha256(plan_path.read_bytes()).hexdigest(),
            "slide_id": slide_id,
            "concept_sha256": hashlib.sha256(concept_path.read_bytes()).hexdigest(),
            "deck_slide_sha256": value_sha256(deck_slide),
            "scene_slide_sha256": value_sha256(scene_slide),
            "decomposition_slide_sha256": value_sha256(map_slide),
            "style_sha256": hashlib.sha256(style_path.read_bytes()).hexdigest() if style_path.is_file() else None,
            "assembly_style_overrides_sha256": hashlib.sha256(override_path.read_bytes()).hexdigest() if override_path.is_file() else None,
            "assets": {component_id: value_sha256(assets[component_id]) for component_id in component_ids},
        }

    def validate_pilot(self) -> None:
        plan = self.docs.get("pilot-plan.json", {})
        location = "pilot-plan.json"
        if not isinstance(plan, dict):
            return
        if not isinstance(plan.get("selection_rationale"), str) or not plan["selection_rationale"].strip():
            self.error(location, "selection_rationale must explain why this is the representative busy slide")
        bindings = self.pilot_bindings()
        if bindings is None:
            self.error(location, "slide, component IDs, or current concept are inconsistent")
            return
        slide_id = bindings["slide_id"]
        scene = self.docs.get("scene-graph.json", {})
        scene_slide = next((item for item in scene.get("slides", []) if isinstance(item, dict) and item.get("slide_id") == slide_id), {})
        deck = self.docs.get("deck-spec.json", {})
        deck_slide = next((item for item in deck.get("slides", []) if isinstance(item, dict) and item.get("id") == slide_id), {})
        if len(scene_slide.get("elements", [])) < 2 or not deck_slide.get("fill_slots"):
            self.error(location, "pilot needs multiple visual instances and actual editable text/data space")
        boxes = [item.get("bbox") for item in scene_slide.get("elements", []) if isinstance(item, dict)]
        has_overlap = any(
            isinstance(first, list) and isinstance(second, list) and len(first) == len(second) == 4
            and all(type(value) in {int, float} for value in first + second)
            and first[0] < second[0] + second[2] and second[0] < first[0] + first[2]
            and first[1] < second[1] + second[3] and second[1] < first[1] + first[3]
            for index, first in enumerate(boxes) for second in boxes[index + 1:]
        )
        if not has_overlap:
            self.error(location, "representative pilot must exercise visual overlap/occlusion")
        if not self.require_pilot_review:
            return
        review_path = self.root / "pilot-review.json"
        if not review_path.is_file():
            self.error("pilot-review.json", "current human trial-assembly approval is missing")
            return
        try:
            review = json.loads(review_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            self.error("pilot-review.json", f"cannot read valid UTF-8 JSON ({exc})")
            return
        if not isinstance(review, dict):
            self.error("pilot-review.json", "must be an object")
            return
        if review.get("bindings") != bindings:
            self.error("pilot-review.json", "trial approval is stale against the current concept, plan, decomposition, scene, style, or selected assets")
        if review.get("slide_id") != slide_id:
            self.error("pilot-review.json", "slide_id differs from the pilot plan")
        run_state = self.docs.get("run-state.json", {})
        local_qa = REVIEW_POLICY.local_agent_review(run_state, review, status="agent-verified-local")
        human_approval = review.get("status") == "approved" and review.get("approval_phrase") == "确认单页试拼"
        actor_error = REVIEW_POLICY.review_actor_error(run_state, review)
        if actor_error:
            self.error("pilot-review.json", actor_error)
        if not (local_qa or human_approval):
            self.error("pilot-review.json", "trial needs current human approval or user-authorized local agent QA")
        if review.get("status") == "rejected":
            defects = review.get("defects")
            if not isinstance(defects, list) or not defects:
                self.error("pilot-review.json", "rejection needs object/page, evidence, owner stage, and local repair entry")
            else:
                for index, defect in enumerate(defects):
                    if not isinstance(defect, dict) or not (defect.get("component_id") or defect.get("slide_id")) or not defect.get("evidence") or defect.get("severity") not in {"blocker", "major", "minor"} or defect.get("owner_stage") not in {"concept", "inventory", "assets", "assembly"} or not defect.get("repair_entry"):
                        self.error(f"pilot-review.json.defects[{index}]", "rejection defect needs object/page, evidence, owner stage, and local repair entry")
                    elif self.validate_relative_file(defect["evidence"], f"pilot-review.json.defects[{index}].evidence") is None:
                        self.error(f"pilot-review.json.defects[{index}]", "rejection evidence must be a current project file")
        if not isinstance(review.get("reviewer"), str) or not review["reviewer"].strip():
            self.error("pilot-review.json", "reviewer is required")
        try:
            reviewed_at = datetime.fromisoformat(review.get("reviewed_at", ""))
            if reviewed_at.utcoffset() is None:
                raise ValueError("timezone missing")
        except (TypeError, ValueError):
            self.error("pilot-review.json", "reviewed_at must be timezone-aware ISO-8601")
        required_checks = {"single_object", "no_missing_objects", "placement", "seams", "color_and_z", "text_and_image_space", "truth_labels"}
        checks = review.get("checks")
        if not isinstance(checks, dict) or any(checks.get(key) is not True for key in required_checks):
            self.error("pilot-review.json", f"all trial checks must pass: {sorted(required_checks)}")
        if review.get("defects") != []:
            self.error("pilot-review.json", "unresolved defects block batch production")
        assembly_path = self.validate_relative_file(review.get("assembly_map_path"), "pilot-review.json.assembly_map_path")
        if assembly_path is None:
            return
        try:
            pilot_assembly = json.loads(assembly_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            self.error("pilot-review.json", f"pilot assembly map is invalid ({exc})")
            return
        if not isinstance(pilot_assembly, dict) or pilot_assembly.get("pilot_bindings") != bindings:
            self.error("pilot-review.json", "pilot PPTX source bindings are stale")
            return
        if self.quality_profile() == "production" and pilot_assembly.get("render_backend") != "powerpoint":
            self.error("pilot-review.json", "production trial needs a render from the verified PowerPoint backend")
        output_path = self.validate_relative_file(pilot_assembly.get("output"), "pilot-review.json.pptx")
        if output_path is None:
            return
        output_sha = hashlib.sha256(output_path.read_bytes()).hexdigest()
        if output_sha != pilot_assembly.get("sha256") or output_sha != review.get("pptx_sha256"):
            self.error("pilot-review.json", "trial PPTX hash is stale")
        render_rows = pilot_assembly.get("renders")
        if not isinstance(render_rows, list) or len(render_rows) != 1:
            self.error("pilot-review.json", "actual PPTX must have one full-size rendered trial slide")
            return
        render_row = render_rows[0]
        if not isinstance(render_row, dict):
            self.error("pilot-review.json", "trial render record must be an object")
            return
        render_path = self.validate_relative_file(render_row.get("path"), "pilot-review.json.render")
        if render_path is None:
            return
        if not render_path.is_relative_to((self.root / "reviews" / "pilot").resolve()):
            self.error("pilot-review.json", "trial render must be under reviews/pilot")
        try:
            from PIL import Image
            with Image.open(render_path) as image:
                if image.format != "PNG" or image.size != (1920, 1080):
                    self.error("pilot-review.json", "actual PPTX render must be 1920x1080 PNG")
        except (ImportError, OSError) as exc:
            self.error("pilot-review.json", f"actual PPTX render cannot be decoded ({exc})")
        if render_row.get("pptx_sha256") != output_sha or render_row.get("sha256") != hashlib.sha256(render_path.read_bytes()).hexdigest():
            self.error("pilot-review.json", "render is stale against the trial PPTX")
        evidence = review.get("evidence")
        if not isinstance(evidence, dict):
            self.error("pilot-review.json", "full-size review evidence is missing")
        else:
            for key in ("concept", "render", "packet"):
                record = evidence.get(key)
                path = self.validate_relative_file(record.get("path") if isinstance(record, dict) else None, f"pilot-review.json.evidence.{key}")
                if path and record.get("sha256") != hashlib.sha256(path.read_bytes()).hexdigest():
                    self.error("pilot-review.json", f"{key} evidence hash is stale")
            concept_evidence = evidence.get("concept") if isinstance(evidence.get("concept"), dict) else {}
            render_evidence = evidence.get("render") if isinstance(evidence.get("render"), dict) else {}
            if concept_evidence.get("sha256") != bindings["concept_sha256"] or render_evidence.get("sha256") != review.get("render_sha256"):
                self.error("pilot-review.json", "concept or render evidence is not the reviewed subject")
            isolation = evidence.get("isolation")
            manifest = self.docs.get("asset-manifest.json", {})
            assets = {item.get("id"): item for item in manifest.get("assets", []) if isinstance(item, dict)}
            if not isinstance(isolation, list) or any(not isinstance(item, dict) for item in isolation) or {item.get("component_id") for item in isolation if isinstance(item, dict)} != set(bindings["assets"]):
                self.error("pilot-review.json", "one isolation sheet is required for every pilot asset")
            else:
                for item in isolation:
                    component_id = item.get("component_id")
                    current = assets[component_id].get("isolation_review", {})
                    for field in ("image", "packet"):
                        path_key, sha_key = f"{field}_path", f"{field}_sha256"
                        if item.get(path_key) != current.get(path_key) or item.get(sha_key) != current.get(sha_key):
                            self.error("pilot-review.json", f"{component_id} isolation evidence is stale")
                        path = self.validate_relative_file(item.get(path_key), f"pilot-review.json.isolation.{component_id}.{field}")
                        if path and item.get(sha_key) != hashlib.sha256(path.read_bytes()).hexdigest():
                            self.error("pilot-review.json", f"{component_id} isolation evidence hash is stale")
            for key, size in (("side_by_side", (3840, 1080)), ("overlay", (1920, 1080))):
                record = evidence.get(key)
                path = self.validate_relative_file(record.get("path") if isinstance(record, dict) else None, f"pilot-review.json.evidence.{key}")
                if path is None:
                    continue
                if record.get("sha256") != hashlib.sha256(path.read_bytes()).hexdigest():
                    self.error("pilot-review.json", f"{key} evidence hash is stale")
                try:
                    from PIL import Image
                    with Image.open(path) as image:
                        if image.format != "PNG" or image.size != size:
                            self.error("pilot-review.json", f"{key} must be a 100%-scale PNG")
                except (ImportError, OSError) as exc:
                    self.error("pilot-review.json", f"{key} cannot be decoded ({exc})")
        if review.get("render_sha256") != hashlib.sha256(render_path.read_bytes()).hexdigest():
            self.error("pilot-review.json", "reviewed render hash is stale")
        if isinstance(evidence, dict):
            side_record = evidence.get("side_by_side")
            overlay_record = evidence.get("overlay")
            side_path = self.root / side_record["path"] if isinstance(side_record, dict) and isinstance(side_record.get("path"), str) else None
            overlay_path = self.root / overlay_record["path"] if isinstance(overlay_record, dict) and isinstance(overlay_record.get("path"), str) else None
            concept_record = evidence.get("concept")
            concept_path = self.root / concept_record["path"] if isinstance(concept_record, dict) and isinstance(concept_record.get("path"), str) else None
            if all(path is not None and path.is_file() and path.resolve().is_relative_to(self.root.resolve()) for path in (side_path, overlay_path, concept_path)):
                try:
                    from PIL import Image
                    with Image.open(concept_path) as original, Image.open(render_path) as rendered_original, Image.open(side_path) as side_original, Image.open(overlay_path) as overlay_original:
                        concept = original.convert("RGB").resize((1920, 1080), Image.Resampling.LANCZOS)
                        rendered = rendered_original.convert("RGB")
                        side = side_original.convert("RGB")
                        overlay = overlay_original.convert("RGB")
                        if side.size != (3840, 1080) or side.crop((0, 0, 1920, 1080)).tobytes() != concept.tobytes() or side.crop((1920, 0, 3840, 1080)).tobytes() != rendered.tobytes():
                            self.error("pilot-review.json", "side-by-side pixels do not show the bound concept and actual render")
                        if overlay.size != (1920, 1080) or overlay.tobytes() != Image.blend(concept, rendered, 0.5).tobytes():
                            self.error("pilot-review.json", "overlay pixels do not match the bound concept and actual render")
                except (ImportError, OSError, ValueError) as exc:
                    self.error("pilot-review.json", f"cannot verify visual comparison pixels ({exc})")
        # Reuse the full-deck PPTX checks against a one-slide projection.
        original = {name: self.docs.get(name) for name in ("deck-spec.json", "scene-graph.json", "assembly-map.json")}
        try:
            self.docs["deck-spec.json"] = {**deck, "slides": [deck_slide]}
            self.docs["scene-graph.json"] = {**scene, "slides": [scene_slide]}
            self.docs["assembly-map.json"] = pilot_assembly
            self.validate_pptx_content(output_path)
            self.validate_object_level_assembly(output_path)
        finally:
            for name, value in original.items():
                if value is None:
                    self.docs.pop(name, None)
                else:
                    self.docs[name] = value

    def validate_qa(self) -> None:
        name = "qa-report.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        if doc.get("status") != "passed":
            self.error(name, "status must be 'passed'")
        checks = doc.get("checks_performed")
        if not isinstance(checks, list) or not checks:
            self.error(name, "checks_performed must be a non-empty list")
        elif self.reaches("release"):
            required_checks = {"contract", "alpha", "inventory", "overlay"}
            if self.quality_profile() == "production":
                required_checks.update({"reference-transfer", "composition-richness", "readability-full-size"})
            missing_checks = required_checks - {
                check for check in checks if isinstance(check, str)
            }
            if missing_checks:
                self.error(name, f"checks_performed missing: {sorted(missing_checks)}")
        evidence_files = doc.get("evidence_files")
        if not isinstance(evidence_files, list) or not evidence_files:
            self.error(name, "evidence_files must be a non-empty list")
        else:
            for index, evidence in enumerate(evidence_files):
                self.validate_relative_file(evidence, f"{name}.evidence_files[{index}]")
        if self.reaches("release"):
            for field in ("final_contact_sheet", "exception_register", "deliverable_index"):
                if not doc.get(field):
                    self.error(name, f"{field} is required for release")
                else:
                    self.validate_relative_file(doc.get(field), f"{name}.{field}")
            for field in ("differences", "exceptions"):
                if not isinstance(doc.get(field), list):
                    self.error(name, f"{field} must be a list")
        qa_slides = doc.get("slides")
        if not isinstance(qa_slides, list):
            self.error(name, "slides must be a list")
            qa_slides = []
        for index, slide in enumerate(qa_slides):
            if isinstance(slide, dict) and slide.get("status") != "passed":
                self.error(f"{name}.slides[{index}]", "slide QA has not passed")
            if isinstance(slide, dict) and not slide.get("evidence_files"):
                self.error(f"{name}.slides[{index}]", "evidence_files must not be empty")
        brief = self.docs.get("brief.json", {})
        formats = brief.get("delivery", {}).get("formats", []) if isinstance(brief, dict) else []
        if "pptx" in formats:
            assembly = self.docs.get("assembly-map.json", {})
            if doc.get("subject_sha256") != assembly.get("sha256"):
                self.error(name, "subject_sha256 must match the assembled PPTX SHA-256")
            self.validate_pptx_qa(doc)

    def qa_slide_bindings(self, deck_slide: dict, assembled: dict, render: dict) -> dict:
        scene = next((row for row in self.docs["scene-graph.json"]["slides"] if row.get("slide_id") == deck_slide["id"]), None)
        assets = {row["id"]: row for row in self.docs["asset-manifest.json"]["assets"]}
        return {
            "concept_sha256": hashlib.sha256((self.root / deck_slide["concept_path"]).read_bytes()).hexdigest(),
            "render_sha256": hashlib.sha256((self.root / render["path"]).read_bytes()).hexdigest(),
            "deck_slide_sha256": value_sha256(deck_slide), "scene_slide_sha256": value_sha256(scene),
            "assembly_slide_sha256": value_sha256(assembled),
            "style_sha256": hashlib.sha256((self.root / "style-contract.json").read_bytes()).hexdigest(),
            "assets": {component_id: value_sha256(assets.get(component_id)) for component_id in assembled["component_ids"]},
            "pptx_sha256": self.docs["assembly-map.json"]["sha256"],
        }

    def qa_deck_bindings(self, contact_sha256: str) -> dict:
        return {
            "reference_study_sha256": hashlib.sha256((self.root / "reference-study.json").read_bytes()).hexdigest(),
            "style_sha256": hashlib.sha256((self.root / "style-contract.json").read_bytes()).hexdigest(),
            "deck_sha256": hashlib.sha256((self.root / "deck-spec.json").read_bytes()).hexdigest(),
            "storyboard_sha256": hashlib.sha256((self.root / "storyboard.json").read_bytes()).hexdigest(),
            "pptx_sha256": self.docs["assembly-map.json"].get("sha256"),
            "contact_sheet_sha256": contact_sha256,
        }

    def validate_pptx_qa(self, report: dict[str, Any]) -> None:
        """Bind every human slide decision to current objects and verifiable pixels."""
        from PIL import Image, ImageChops

        name = "qa-report.json"
        deck = self.docs.get("deck-spec.json", {})
        scene = self.docs.get("scene-graph.json", {})
        assembly = self.docs.get("assembly-map.json", {})
        manifest = self.docs.get("asset-manifest.json", {})
        if not all(isinstance(doc, dict) for doc in (deck, scene, assembly, manifest)):
            return
        deck_slides = deck.get("slides", [])
        assembly_slides = assembly.get("slides", [])
        renders = assembly.get("renders", [])
        rows = report.get("slides")
        if not all(isinstance(items, list) for items in (deck_slides, assembly_slides, renders, rows)) or len(rows) != len(deck_slides) or len(renders) != len(deck_slides):
            self.error(name, "full-deck QA requires one current review and render per slide")
            return
        scenes = {row.get("slide_id"): row for row in scene.get("slides", []) if isinstance(row, dict)}
        assets = {row.get("id"): row for row in manifest.get("assets", []) if isinstance(row, dict)}
        version = report.get("version")
        if not isinstance(version, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,63}", version):
            self.error(name, "QA evidence version is required")
            return
        required_checks = {
            "independent_objects", "geometry", "hierarchy", "color_and_z", "content_truth",
            "visual_relationships", "readability_full_size", "main_image_space", "editability",
        }
        contact_inputs = []
        expected_evidence: set[str] = set()
        accepted_minor_issues = []
        if not isinstance(report.get("final_reviewer"), str) or not report["final_reviewer"].strip():
            self.error(name, "final_reviewer must identify the human whole-deck reviewer")
        actor_error = REVIEW_POLICY.review_actor_error(self.docs.get("run-state.json"), report, human_only=True)
        if actor_error:
            self.error(name, actor_error)
        deck_checks = report.get("deck_checks")
        if not isinstance(deck_checks, dict) or any(deck_checks.get(key) is not True for key in
                                                    ("reference_transfer", "composition_richness", "readability_full_size", "visual_coherence", "truth_consistency")):
            self.error(name, "all five human whole-deck checks must pass")
        for index, (deck_slide, assembled, render, row) in enumerate(zip(deck_slides, assembly_slides, renders, rows)):
            location = f"{name}.slides[{index}]"
            if not all(isinstance(item, dict) for item in (deck_slide, assembled, render, row)):
                self.error(location, "current slide evidence must be objects")
                continue
            slide_id = deck_slide.get("id")
            if row.get("slide_id") != slide_id or assembled.get("slide_id") != slide_id or assembled.get("output_index") != index + 1:
                self.error(location, "QA slide order differs from the assembled deck")
                continue
            concept_path = self.validate_relative_file(deck_slide.get("concept_path"), f"{location}.concept")
            render_path = self.validate_relative_file(render.get("path"), f"{location}.render")
            if not concept_path or not render_path:
                continue
            component_ids = assembled.get("component_ids", [])
            if not isinstance(component_ids, list) or any(not isinstance(value, str) for value in component_ids):
                self.error(location, "component_ids must be a list of IDs")
                continue
            expected_bindings = self.qa_slide_bindings(deck_slide, assembled, render)
            if row.get("bindings") != expected_bindings:
                self.error(location, "human slide approval is stale against current concept, render, style, scene, slide, asset, or PPTX")
            if row.get("status") != "passed" or not isinstance(row.get("checks"), dict) or any(row["checks"].get(key) is not True for key in required_checks):
                self.error(location, "all nine human visual checks must pass")
            if not isinstance(row.get("reviewer"), str) or not row["reviewer"].strip():
                self.error(location, "human reviewer is required")
            actor_error = REVIEW_POLICY.review_actor_error(self.docs.get("run-state.json"), row, human_only=True)
            if actor_error:
                self.error(location, actor_error)
            try:
                reviewed_at = datetime.fromisoformat(row.get("reviewed_at", ""))
                if reviewed_at.utcoffset() is None:
                    raise ValueError("timezone missing")
            except (TypeError, ValueError):
                self.error(location, "reviewed_at must be timezone-aware ISO-8601")
            issues = row.get("issues")
            if not isinstance(issues, list):
                self.error(location, "issues must be a list")
            else:
                for issue_index, issue in enumerate(issues):
                    issue_location = f"{location}.issues[{issue_index}]"
                    if not isinstance(issue, dict) or issue.get("severity") not in {"blocker", "major", "minor"} or issue.get("slide_id") != slide_id or not issue.get("evidence") or issue.get("owner_stage") not in {"concept", "inventory", "assets", "assembly", "qa"} or not issue.get("repair_entry"):
                        self.error(issue_location, "issue needs slide, severity, evidence, owner stage, and local repair entry")
                        continue
                    attempts = issue.get("repair_attempts")
                    if type(attempts) is not int or not 0 <= attempts <= 3:
                        self.error(issue_location, "repair_attempts must be 0..3; never flatten after the cap")
                    if issue.get("disposition") != "resolved" and (issue.get("severity") != "minor" or issue.get("disposition") != "accepted-at-G5"):
                        self.error(issue_location, "unresolved blocker or major issue prevents QA/release")
                    if issue.get("severity") == "minor" and issue.get("disposition") == "accepted-at-G5":
                        accepted_minor_issues.append(issue)
                    self.validate_relative_file(issue["evidence"], f"{issue_location}.evidence")
            base = f"reviews/qa/{version}/{slide_id}"
            expected_paths = {
                "side_by_side": f"{base}-side-by-side.png",
                "overlay": f"{base}-overlay.png",
                "packet": f"{base}-review.md",
            }
            evidence = row.get("evidence")
            if not isinstance(evidence, dict):
                self.error(location, "side-by-side, overlay, and review packet are required")
                continue
            paths: dict[str, Path] = {}
            for key, relative in expected_paths.items():
                record = evidence.get(key)
                if not isinstance(record, dict) or record.get("path") != relative:
                    self.error(location, f"{key} evidence path is missing or from another version")
                    continue
                path = self.validate_relative_file(relative, f"{location}.evidence.{key}")
                if path:
                    paths[key] = path
                    if record.get("sha256") != hashlib.sha256(path.read_bytes()).hexdigest():
                        self.error(location, f"{key} evidence hash is stale")
            expected_files = {deck_slide["concept_path"], render["path"], *expected_paths.values()}
            row_files = row.get("evidence_files")
            if not isinstance(row_files, list) or any(not isinstance(item, str) for item in row_files) or set(row_files) != expected_files:
                self.error(location, "slide evidence_files must contain exactly current concept, render, side-by-side, overlay, and packet")
            expected_evidence.update(expected_files)
            if "side_by_side" in paths and "overlay" in paths:
                try:
                    with Image.open(concept_path) as concept, Image.open(render_path) as actual, Image.open(paths["side_by_side"]) as side, Image.open(paths["overlay"]) as blend:
                        expected_side = DECK_QA.side_by_side(concept, actual)
                        expected_blend = DECK_QA.overlay(concept, actual)
                        if side.size != expected_side.size or ImageChops.difference(side.convert("RGB"), expected_side).getbbox():
                            self.error(location, "side-by-side pixels do not compare the current concept and render")
                        if blend.size != expected_blend.size or ImageChops.difference(blend.convert("RGB"), expected_blend).getbbox():
                            self.error(location, "overlay pixels do not compare the current concept and render")
                        contact_inputs.append((slide_id, actual.copy()))
                except (OSError, ValueError) as exc:
                    self.error(location, f"visual evidence cannot be decoded: {exc}")
        contact_relative = f"reviews/qa/{version}/final-contact-sheet.png"
        if report.get("final_contact_sheet") != contact_relative:
            self.error(name, "final contact sheet path must belong to the current QA version")
        contact_path = self.validate_relative_file(contact_relative, f"{name}.final_contact_sheet")
        if contact_path:
            expected_evidence.add(contact_relative)
            if report.get("contact_sheet_sha256") != hashlib.sha256(contact_path.read_bytes()).hexdigest():
                self.error(name, "contact sheet hash is stale")
            if len(contact_inputs) == len(deck_slides):
                try:
                    with Image.open(contact_path) as actual_contact:
                        expected_contact = DECK_QA.contact_sheet(contact_inputs)
                        if actual_contact.size != expected_contact.size or ImageChops.difference(actual_contact.convert("RGB"), expected_contact).getbbox():
                            self.error(name, "contact sheet pixels do not match all current PPTX renders in deck order")
                except (OSError, ValueError) as exc:
                    self.error(name, f"contact sheet cannot be decoded: {exc}")
        expected_deck_bindings = self.qa_deck_bindings(report.get("contact_sheet_sha256"))
        if report.get("deck_bindings") != expected_deck_bindings:
            self.error(name, "whole-deck human approval is stale against references, story, style, deck, PPTX, or contact sheet")
        exception_relative = f"reviews/qa/{version}/exception-register.json"
        index_relative = f"reviews/qa/{version}/deliverable-index.json"
        for field, relative in (("exception_register", exception_relative), ("deliverable_index", index_relative)):
            if report.get(field) != relative:
                self.error(name, f"{field} must belong to the current QA version")
            expected_evidence.add(relative)
        verified_renders = report.get("verified_renders")
        if self.quality_profile() == "production":
            if not isinstance(verified_renders, list) or len(verified_renders) != len(deck_slides):
                self.error(name, "production QA needs a fresh PowerPoint re-render of every current slide")
            else:
                for index, record in enumerate(verified_renders, start=1):
                    relative = f"reviews/qa/{version}/verified-renders/slide-{index:02d}.png"
                    if not isinstance(record, dict) or record.get("path") != relative or record.get("pptx_sha256") != assembly.get("sha256"):
                        self.error(name, f"verified render {index} must bind the current PPTX and QA version")
                        continue
                    if not isinstance(renders[index - 1], dict):
                        self.error(name, f"declared render {index} must be an object")
                        continue
                    verified_path = self.validate_relative_file(relative, f"{name}.verified_renders[{index - 1}]")
                    current_path = self.validate_relative_file(renders[index - 1].get("path"), f"{name}.renders[{index - 1}]")
                    if not verified_path or not current_path:
                        continue
                    expected_evidence.add(relative)
                    if record.get("sha256") != hashlib.sha256(verified_path.read_bytes()).hexdigest():
                        self.error(name, f"verified render {index} hash is stale")
                    try:
                        with Image.open(verified_path) as verified, Image.open(current_path) as current:
                            if verified.size != current.size or ImageChops.difference(verified.convert("RGB"), current.convert("RGB")).getbbox():
                                self.error(name, f"verified render {index} pixels differ from the declared PowerPoint render")
                    except OSError as exc:
                        self.error(name, f"verified render {index} cannot be decoded: {exc}")
        elif verified_renders not in (None, []):
            self.error(name, "fixture QA must not claim production PowerPoint re-render evidence")
        report_files = report.get("evidence_files")
        if not isinstance(report_files, list) or any(not isinstance(item, str) for item in report_files) or set(report_files) != expected_evidence:
            self.error(name, "evidence_files must cover every current full-size page comparison and whole-deck record")
        if report.get("exceptions") != accepted_minor_issues:
            self.error(name, "QA exceptions must exactly list minor issues awaiting G5 acceptance")
        repairs = report.get("repairs")
        if not isinstance(repairs, list):
            self.error(name, "repairs must be a list")
        else:
            attempts_by_defect: dict[tuple[str, str], set[int]] = {}
            for index, repair in enumerate(repairs):
                location = f"{name}.repairs[{index}]"
                if not isinstance(repair, dict) or not isinstance(repair.get("slide_id"), str) or not isinstance(repair.get("defect_class"), str) or repair.get("owner_stage") not in {"concept", "inventory", "assets", "assembly", "qa"}:
                    self.error(location, "repair needs slide_id, defect_class, and owner_stage")
                    continue
                attempt = repair.get("attempt")
                if type(attempt) is not int or not 1 <= attempt <= 3:
                    self.error(location, "repair attempt must be 1..3; no flattened fallback after the cap")
                    continue
                key = (repair["slide_id"], repair["defect_class"])
                seen = attempts_by_defect.setdefault(key, set())
                if attempt in seen:
                    self.error(location, "repair attempt number is duplicated for one defect class")
                seen.add(attempt)
                if len(seen) > 3:
                    self.error(location, "one defect class exceeds three focused repairs")
                self.validate_relative_file(repair.get("artifact_path"), f"{location}.artifact_path")
        try:
            exception_doc = json.loads((self.root / exception_relative).read_text(encoding="utf-8"))
            index_doc = json.loads((self.root / index_relative).read_text(encoding="utf-8"))
            if exception_doc != {"status": "reviewed", "issues": report.get("exceptions", [])}:
                self.error(name, "exception register differs from reviewed QA exceptions")
            if index_doc != {
                "status": "ready", "pptx": {"path": assembly.get("output"), "sha256": assembly.get("sha256")},
                "contact_sheet": {"path": contact_relative, "sha256": report.get("contact_sheet_sha256")},
                "slides": [{"slide_id": row.get("slide_id"), "status": row.get("status")} for row in rows if isinstance(row, dict)],
            }:
                self.error(name, "deliverable index does not bind the current PPTX, contact sheet, and page statuses")
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            self.error(name, f"full-deck QA indexes cannot be read: {exc}")

    def expected_gate_decisions(self, gate: str) -> dict[str, Any]:
        storyboard = self.docs.get("storyboard.json", {})
        style = self.docs.get("style-contract.json", {})
        reference_study = self.docs.get("reference-study.json", {})
        calibration = self.docs.get("concept-calibration.json", {})
        concept_log = self.docs.get("concept-generation-log.json", {})
        deck = self.docs.get("deck-spec.json", {})
        manifest = self.docs.get("asset-manifest.json", {})
        assembly = self.docs.get("assembly-map.json", {})
        qa = self.docs.get("qa-report.json", {})
        if gate == "story" and isinstance(storyboard, dict):
            slides = [slide for slide in storyboard.get("slides", []) if isinstance(slide, dict)]
            return {
                "slide_sequence": [slide.get("id") for slide in slides],
                "action_titles": {slide.get("id"): slide.get("title") for slide in slides},
                "slide_messages": {slide.get("id"): slide.get("message") for slide in slides},
                "evidence_mapping": {
                    slide.get("id"): slide.get("evidence_ids", []) for slide in slides
                },
                "appendix_boundary": {
                    "slide_ids": [
                        slide.get("id")
                        for slide in slides
                        if slide.get("narrative_role") == "appendix"
                    ],
                    "policy": "declared",
                },
            }
        if gate == "style_anchors" and isinstance(style, dict):
            anchors = {
                anchor.get("role"): anchor
                for anchor in style.get("anchor_concepts", [])
                if isinstance(anchor, dict)
            }
            decisions = {
                "reference_selection": {
                    "selected_reference_ids": reference_study.get(
                        "selected_reference_ids", []
                    )
                    if isinstance(reference_study, dict)
                    else [],
                    "reference_contact_sheet": reference_study.get(
                        "reference_contact_sheet"
                    )
                    if isinstance(reference_study, dict)
                    else None,
                },
                "concept_calibration": {
                    "selected_variant_id": style.get("selected_calibration_variant_id"),
                    "selected_direction_id": style.get("selected_direction_id"),
                    "style_strength": style.get("style_strength"),
                    "layout_density": style.get("layout_density"),
                    "calibration_version": calibration.get("version") if isinstance(calibration, dict) else None,
                } if isinstance(calibration, dict) and calibration.get("status") == "generated" else {
                    "status": "not-required",
                    "reason": calibration.get("reason") if isinstance(calibration, dict) else None,
                },
                "composition_profile": style.get("composition_profile"),
                "readability_profile": style.get("readability_profile"),
                "visual_direction": style.get("direction"),
                "style_rules": style.get("rules"),
                "anchor_opening": anchors.get("opening"),
                "anchor_typical": anchors.get("typical"),
                "anchor_highest_risk": anchors.get("highest-risk"),
            }
            if self.semantic_system_active():
                decisions["semantic_system"] = style.get("semantic_system")
            return decisions
        if gate == "concept_contact_sheet" and all(
            isinstance(doc, dict) for doc in (deck, manifest, concept_log)
        ):
            slides = [slide for slide in deck.get("slides", []) if isinstance(slide, dict)]
            assets = [asset for asset in manifest.get("assets", []) if isinstance(asset, dict)]
            concept_records = {
                record.get("slide_id"): record
                for record in concept_log.get("slides", [])
                if isinstance(record, dict) and record.get("slide_id")
            } if isinstance(concept_log.get("slides"), list) else {}
            generated_contentful_log = concept_log.get("status") == "generated"
            if generated_contentful_log:
                copy_source_review: dict[str, Any] = {
                    "slides": {
                        slide_id: {
                            "version": record.get("version"),
                            "copy_map_sha256": value_sha256(record.get("copy_map")),
                            "source_lineage": {
                                source.get("id"): source.get("sha256")
                                for source in record.get("source_lineage", [])
                                if isinstance(source, dict) and source.get("id")
                            },
                            "text_review": {
                                "transcript_path": record.get("text_review", {}).get(
                                    "transcript_path"
                                ),
                                "transcript_sha256": record.get("text_review", {}).get(
                                    "transcript_sha256"
                                ),
                            },
                        }
                        for slide_id, record in concept_records.items()
                    }
                }
                contact_sheet_record = concept_log.get("review_contact_sheet")
                full_deck_review = {
                    "concept_paths": [slide.get("concept_path") for slide in slides],
                    "contentful_image_sha256": {
                        slide_id: record.get("output_sha256")
                        for slide_id, record in concept_records.items()
                    },
                    "review_contact_sheet": contact_sheet_record,
                    "concept_generation_log_sha256": hashlib.sha256(
                        (self.root / "concept-generation-log.json").read_bytes()
                    ).hexdigest()
                    if (self.root / "concept-generation-log.json").is_file()
                    else None,
                }
            else:
                copy_source_review = {
                    "status": "not-required",
                    "reason": concept_log.get("reason"),
                }
                full_deck_review = {
                    "concept_paths": [slide.get("concept_path") for slide in slides]
                }
            expected_decisions = {
                "full_deck_contact_sheet": full_deck_review,
                "contentful_copy_source_review": copy_source_review,
                "readability_review": {
                    "policy": "full-size-human-review",
                    "profile": style.get("readability_profile") if isinstance(style, dict) else None,
                    "slides": {
                        slide.get("id"): slide.get("readability")
                        for slide in slides
                    },
                },
                "layout_richness_review": {
                    "slides": {
                        slide.get("id"): {
                            "layout_archetype": slide.get("composition", {}).get(
                                "layout_archetype"
                            ),
                            "complexity": slide.get("composition", {}).get("complexity"),
                            "information_unit_count": len(
                                slide.get("composition", {}).get("information_units", [])
                            ),
                            "visual_layer_count": len(
                                slide.get("composition", {}).get("visual_layers", [])
                            ),
                        }
                        for slide in slides
                        if isinstance(slide.get("composition"), dict)
                    },
                    "policy": "review-rendered-contact-sheet",
                },
                "blank_table_chart_regions": {
                    "fill_slot_ids": [
                        slot.get("id")
                        for slide in slides
                        for slot in slide.get("fill_slots", [])
                        if isinstance(slot, dict)
                        and slot.get("kind") in {"table", "chart", "data"}
                    ],
                    "policy": "blank-fill-slots",
                },
                "special_typography_exceptions": {
                    "asset_ids": [
                        asset.get("id")
                        for asset in assets
                        if asset.get("role") == "special-typography"
                    ],
                    "policy": "explicit-only",
                },
                "slide_specific_deviations": {
                    "items": [
                        {"slide_id": slide.get("id"), **exception}
                        for slide in slides
                        if isinstance(slide.get("semantic_design"), dict)
                        if isinstance(slide["semantic_design"].get("exceptions"), list)
                        for exception in slide["semantic_design"]["exceptions"]
                        if isinstance(exception, dict)
                    ] if self.semantic_system_active() else [],
                    "policy": "declared",
                },
            }
            if self.semantic_system_active():
                expected_decisions["cross_slide_semantics_review"] = {
                    "policy": "contact-sheet-and-full-size-human-review",
                    "semantic_system": style.get("semantic_system") if isinstance(style, dict) else None,
                    "slides": {slide.get("id"): slide.get("semantic_design") for slide in slides},
                }
            return expected_decisions
        if gate == "decomposition":
            decomposition = self.docs.get("decomposition-map.json", {})
            slides = decomposition.get("slides", []) if isinstance(decomposition, dict) else []
            slides = [slide for slide in slides if isinstance(slide, dict)]
            return {
                "object_inventory": {
                    slide.get("slide_id"): slide.get("items") for slide in slides
                },
                "coverage_review": {
                    slide.get("slide_id"): slide.get("coverage_review") for slide in slides
                },
                "review_evidence": {
                    slide.get("slide_id"): {
                        "review_image": slide.get("review_image"),
                        "review_packet": slide.get("review_packet"),
                    }
                    for slide in slides
                },
            }
        if gate == "release" and all(
            isinstance(doc, dict) for doc in (manifest, assembly, qa)
        ):
            assets = [asset for asset in manifest.get("assets", []) if isinstance(asset, dict)]
            assembly_slides = [
                slide for slide in assembly.get("slides", []) if isinstance(slide, dict)
            ]
            return {
                "final_render": {
                    "contact_sheet": qa.get("final_contact_sheet"),
                    "evidence_files": qa.get("evidence_files", []),
                },
                "qa_summary": {
                    "status": qa.get("status"),
                    "final_reviewer": qa.get("final_reviewer"),
                    "deck_checks": qa.get("deck_checks"),
                    "slides": {
                        slide.get("slide_id"): {"status": slide.get("status"), "checks": slide.get("checks")}
                        for slide in qa.get("slides", []) if isinstance(slide, dict)
                    } if isinstance(qa.get("slides"), list) else {},
                },
                "concept_differences": {"items": qa.get("differences", [])},
                "qa_exceptions": {
                    "items": qa.get("exceptions", []),
                    "register": qa.get("exception_register"),
                },
                "component_usability": {
                    slide.get("slide_id"): {
                        "component_ids": slide.get("component_ids", []),
                        "fill_slot_ids": slide.get("fill_slot_ids", []),
                    }
                    for slide in assembly_slides
                },
                "deliverable_index": {
                    "path": qa.get("deliverable_index"),
                    "asset_outputs": [asset.get("output") for asset in assets],
                },
            }
        return {}

    def validate_state(self) -> None:
        name = "run-state.json"
        doc = self.docs.get(name)
        if not isinstance(doc, dict):
            return
        declared_stage = doc.get("stage")
        if declared_stage not in STAGES and declared_stage != "released":
            self.error(name, f"unknown stage '{declared_stage}'")
        gates = doc.get("gates", {})
        if not isinstance(gates, dict):
            self.error(name, "gates must be an object")
            return
        if declared_stage == "released":
            release_gate = gates.get("release")
            if not isinstance(release_gate, dict) or release_gate.get("status") != "approved":
                self.error(
                    f"{name}.stage",
                    "stage 'released' requires an approved G5 release gate",
                )
        route = doc.get("production_route")
        if not isinstance(route, dict):
            self.error(name, "production_route must be an object")
        else:
            mode = route.get("mode")
            planning_artifacts = route.get("planning_artifacts")
            if mode == "full-tracked":
                if planning_artifacts != "required":
                    self.error(
                        name,
                        "full-tracked route requires planning_artifacts='required'",
                    )
            elif mode == "fast":
                if planning_artifacts != "skipped":
                    self.error(name, "fast route requires planning_artifacts='skipped'")
                if route.get("approval_phrase") != "确认快速路线":
                    self.error(name, "fast route requires exact approval phrase '确认快速路线'")
                eligibility = route.get("eligibility")
                if not isinstance(eligibility, dict):
                    self.error(name, "fast route eligibility must be an object")
                else:
                    if eligibility.get("consequential_decisions_approved") is not True:
                        self.error(
                            name,
                            "fast route requires all consequential decisions to be approved",
                        )
                    if eligibility.get("high_stakes") is not False:
                        self.error(name, "fast route is forbidden for high-stakes work")
                    if eligibility.get("multi_session") is not False:
                        self.error(name, "fast route is forbidden for multi-session work")
                brief_gate = gates.get("brief_authority")
                if (
                    not isinstance(brief_gate, dict)
                    or brief_gate.get("status") != "approved"
                ):
                    self.error(name, "fast route requires a current G1 approval")
            else:
                self.error(name, f"unknown production_route mode '{mode}'")
        gate_stage = {
            "brief_authority": "story",
            "story": "style",
            "style_anchors": "concept",
            "concept_contact_sheet": "inventory",
            "decomposition": "pilot" if self.pilot_required() else "assets",
            "release": "release",
        }
        deck = self.docs.get("deck-spec.json", {})
        concept_paths = {
            slide.get("concept_path")
            for slide in deck.get("slides", [])
            if isinstance(slide, dict) and isinstance(slide.get("concept_path"), str) and slide.get("concept_path")
        } if isinstance(deck, dict) else set()
        style = self.docs.get("style-contract.json", {})
        anchor_paths = {
            anchor.get("path")
            for anchor in style.get("anchor_concepts", [])
            if isinstance(anchor, dict) and isinstance(anchor.get("path"), str) and anchor.get("path")
        } if isinstance(style, dict) else set()
        required_gate_artifacts: dict[str, set[str]] = {
            "brief_authority": {"brief.json", "alignment.json", "research-bundle.json"},
            "story": {"storyboard.json"},
            "style_anchors": {
                "reference-study.json",
                "style-contract.json",
                "capability-report.json",
            }
            | anchor_paths,
            "concept_contact_sheet": concept_paths | {"concept-generation-log.json", "deck-spec.json"},
            "decomposition": {"decomposition-map.json"},
            "release": {"asset-manifest.json", "assembly-map.json", "qa-report.json"},
        }
        decomposition_map = self.docs.get("decomposition-map.json", {})
        if isinstance(decomposition_map, dict) and isinstance(decomposition_map.get("slides"), list):
            for slide in decomposition_map["slides"]:
                if not isinstance(slide, dict):
                    continue
                for key in ("concept_path", "master_path"):
                    path_value = slide.get(key)
                    if isinstance(path_value, str) and path_value:
                        required_gate_artifacts["decomposition"].add(path_value)
                for key in ("review_image", "review_packet"):
                    evidence = slide.get(key)
                    if isinstance(evidence, dict) and isinstance(evidence.get("path"), str):
                        required_gate_artifacts["decomposition"].add(evidence["path"])
        concept_log = self.docs.get("concept-generation-log.json", {})
        if isinstance(deck, dict):
            for slide in deck.get("slides", []):
                if not isinstance(slide, dict):
                    continue
                readability = slide.get("readability")
                review = readability.get("full_size_review") if isinstance(readability, dict) else None
                if isinstance(review, dict) and isinstance(review.get("path"), str):
                    required_gate_artifacts["concept_contact_sheet"].add(review["path"])
        if isinstance(concept_log, dict) and concept_log.get("status") == "generated":
            review_contact_sheet = concept_log.get("review_contact_sheet")
            if isinstance(review_contact_sheet, dict):
                contact_path = review_contact_sheet.get("path")
                if isinstance(contact_path, str) and contact_path:
                    required_gate_artifacts["concept_contact_sheet"].add(contact_path)
            for record in concept_log.get("slides", []):
                if not isinstance(record, dict):
                    continue
                for source in record.get("source_lineage", []):
                    if isinstance(source, dict) and isinstance(source.get("path"), str):
                        required_gate_artifacts["concept_contact_sheet"].add(source["path"])
                review = record.get("text_review", {})
                if isinstance(review, dict):
                    for path_field in ("transcript_path",):
                        path_value = review.get(path_field)
                        if isinstance(path_value, str) and path_value:
                            required_gate_artifacts["concept_contact_sheet"].add(path_value)
                    repair = review.get("repair")
                    if isinstance(repair, dict):
                        for path_field in ("input_path", "evidence_path"):
                            path_value = repair.get(path_field)
                            if isinstance(path_value, str) and path_value:
                                required_gate_artifacts["concept_contact_sheet"].add(path_value)
        reference_study = self.docs.get("reference-study.json", {})
        if isinstance(reference_study, dict):
            contact_sheet = reference_study.get("reference_contact_sheet")
            if isinstance(contact_sheet, dict):
                path = contact_sheet.get("path")
                if isinstance(path, str) and path:
                    required_gate_artifacts["style_anchors"].add(path)
        calibration = self.docs.get("concept-calibration.json", {})
        if isinstance(calibration, dict) and calibration.get("status") == "generated":
            required_gate_artifacts["style_anchors"].add("concept-calibration.json")
            board = calibration.get("contact_sheet", {})
            if isinstance(board, dict) and isinstance(board.get("path"), str):
                required_gate_artifacts["style_anchors"].add(board["path"])
            for variant in calibration.get("variants", []):
                if isinstance(variant, dict) and isinstance(variant.get("output_path"), str):
                    required_gate_artifacts["style_anchors"].add(variant["output_path"])
                if isinstance(variant, dict):
                    text_review = variant.get("text_review", {})
                    if isinstance(text_review, dict):
                        for field in ("transcript_path",):
                            path_value = text_review.get(field)
                            if isinstance(path_value, str) and path_value:
                                required_gate_artifacts["style_anchors"].add(path_value)
                        repair = text_review.get("repair")
                        if isinstance(repair, dict):
                            for field in ("input_path", "evidence_path"):
                                path_value = repair.get(field)
                                if isinstance(path_value, str) and path_value:
                                    required_gate_artifacts["style_anchors"].add(path_value)
            for source in calibration.get("source_lineage", []):
                if isinstance(source, dict) and isinstance(source.get("path"), str):
                    required_gate_artifacts["style_anchors"].add(source["path"])
        qa = self.docs.get("qa-report.json", {})
        if isinstance(qa, dict):
            release_paths: set[str] = set()
            for field in ("final_contact_sheet", "exception_register", "deliverable_index"):
                value = qa.get(field)
                if isinstance(value, str) and value:
                    release_paths.add(value)
            evidence_files = qa.get("evidence_files", [])
            if isinstance(evidence_files, list):
                release_paths.update(
                    path for path in evidence_files if isinstance(path, str) and path
                )
            for slide in (qa.get("slides") if isinstance(qa.get("slides"), list) else []):
                if not isinstance(slide, dict):
                    continue
                slide_evidence = slide.get("evidence_files", [])
                if isinstance(slide_evidence, list):
                    release_paths.update(
                        path for path in slide_evidence if isinstance(path, str) and path
                    )
            required_gate_artifacts["release"].update(release_paths)
        artifact_versions = doc.get("artifact_versions", {})
        if not isinstance(artifact_versions, dict):
            self.error(name, "artifact_versions must be an object")
            artifact_versions = {}
        for gate in GATES:
            if not self.reaches(gate_stage[gate]):
                continue
            gate_record = gates.get(gate)
            location = f"{name}.gates.{gate}"
            if not isinstance(gate_record, dict):
                self.error(location, "approval must be an object with evidence, not a bare status")
                self.error(location, "missing approver")
                self.error(location, "missing approved_at")
                continue
            if gate_record.get("status") != "approved":
                self.error(location, "status must be 'approved' by this stage")
            approver = gate_record.get("approver")
            if not isinstance(approver, str) or not approver.strip():
                self.error(location, "approver must be a non-empty string")
            approved_at = gate_record.get("approved_at")
            if not isinstance(approved_at, str) or not approved_at.strip():
                self.error(location, "approved_at must be a non-empty ISO-8601 string")
            else:
                try:
                    approval_time = datetime.fromisoformat(approved_at)
                    if approval_time.utcoffset() is None:
                        self.error(location, "approved_at must include a timezone offset")
                except ValueError:
                    self.error(location, "approved_at must be ISO-8601")
            artifact_version = gate_record.get("artifact_version")
            if not isinstance(artifact_version, str) or not artifact_version.strip():
                self.error(location, "artifact_version must be a non-empty string")
            if not isinstance(gate_record.get("approved_artifacts"), dict):
                self.error(location, "approved_artifacts must be an object")
            if gate_record.get("artifact_version") != artifact_versions.get(gate):
                self.error(location, "artifact_version does not match run-state artifact_versions")
            if gate == "decomposition":
                current_map = self.docs.get("decomposition-map.json", {})
                map_version = current_map.get("version") if isinstance(current_map, dict) else None
                if artifact_version != map_version:
                    self.error(location, "decomposition approval version must match map version")
            if gate == "brief_authority":
                approval_mode = gate_record.get("approval_mode")
                if approval_mode not in {"explicit", "local-affirmative", "batch-recommendations"}:
                    self.error(location, "approval_mode must be 'explicit', 'local-affirmative', or 'batch-recommendations'")
                approval_phrase = gate_record.get("approval_phrase")
                if not isinstance(approval_phrase, str) or not approval_phrase.strip():
                    self.error(location, "approval_phrase must be a non-empty string")
                alignment = self.docs.get("alignment.json", {})
                raw_decisions = alignment.get("decisions") if isinstance(alignment, dict) else None
                if not isinstance(raw_decisions, list):
                    self.error(location, "alignment decisions must be a list")
                    decisions: list[Any] = []
                else:
                    decisions = raw_decisions
                decision_map = {
                    decision.get("id"): decision
                    for decision in decisions
                    if isinstance(decision, dict)
                    and isinstance(decision.get("id"), str)
                }
                if approval_mode in {"explicit", "local-affirmative"} and isinstance(approval_phrase, str):
                    check_reply = REVIEW_POLICY.local_reply_error if approval_mode == "local-affirmative" else REVIEW_POLICY.explicit_reply_error
                    reply_error = check_reply(approval_phrase, self.docs.get("run-state.json"))
                    if reply_error:
                        self.error(location, reply_error)
                    if approval_mode == "local-affirmative":
                        current_packet = alignment.get("approval_packet", {}) if isinstance(alignment, dict) else {}
                        scoped_ids = current_packet.get("decision_ids") if isinstance(current_packet, dict) else None
                        if not isinstance(current_packet, dict) or current_packet.get("gate") != "brief_authority" or current_packet.get("version") != artifact_version or not isinstance(scoped_ids, list) or any(not isinstance(key, str) for key in scoped_ids) or len(scoped_ids) != len(CRITICAL_G1_DECISIONS) or set(scoped_ids) != CRITICAL_G1_DECISIONS:
                            self.error(location, "local G1 reply must cover the full current brief_authority packet")
                    decision_approvals = gate_record.get("decision_approvals")
                    if not isinstance(decision_approvals, dict):
                        self.error(location, "explicit approval requires per-decision approval evidence")
                    else:
                        if set(decision_approvals) != CRITICAL_G1_DECISIONS:
                            self.error(location, "explicit approval must cover every critical G1 decision")
                        for decision_id in sorted(CRITICAL_G1_DECISIONS):
                            approval = decision_approvals.get(decision_id)
                            decision = decision_map.get(decision_id, {})
                            if not isinstance(approval, dict) or approval.get("status") != "approved":
                                self.error(location, f"missing explicit approval for '{decision_id}'")
                            elif approval.get("value_sha256") != value_sha256(
                                decision.get("value") if isinstance(decision, dict) else None
                            ):
                                self.error(location, f"explicit approval value is stale for '{decision_id}'")
                if approval_mode == "batch-recommendations":
                    if gate_record.get("approval_phrase") != "全部按推荐":
                        self.error(location, "batch approval requires the exact phrase '全部按推荐'")
                    packet = alignment.get("approval_packet", {}) if isinstance(alignment, dict) else {}
                    if not isinstance(packet, dict):
                        self.error(location, "alignment approval_packet must be an object")
                    else:
                        if packet.get("gate") != "brief_authority":
                            self.error(location, "batch approval packet must target brief_authority")
                        if packet.get("version") != gate_record.get("artifact_version"):
                            self.error(location, "batch approval packet version is stale")
                        if packet.get("recommendations_visible") is not True:
                            self.error(location, "batch approval requires all recommendations to be visible")
                        packet_decisions = packet.get("decision_ids")
                        packet_ids_are_valid = (
                            isinstance(packet_decisions, list)
                            and all(
                                isinstance(decision_id, str) and decision_id
                                for decision_id in packet_decisions
                            )
                        )
                        if (
                            not packet_ids_are_valid
                            or set(packet_decisions) != CRITICAL_G1_DECISIONS
                        ):
                            self.error(location, "batch approval packet must cover every critical G1 decision")
                    for decision_id in sorted(CRITICAL_G1_DECISIONS):
                        decision = decision_map.get(decision_id, {})
                        recommendation = decision.get("recommendation", {}) if isinstance(decision, dict) else {}
                        if (
                            decision.get("basis") != "recommended"
                            or not isinstance(recommendation, dict)
                            or recommendation.get("value") != decision.get("value")
                            or not recommendation.get("reason")
                            or not recommendation.get("impact")
                        ):
                            self.error(
                                location,
                                f"batch approval lacks a complete accepted recommendation for '{decision_id}'",
                            )
            elif gate in GATE_REQUIRED_DECISIONS:
                packet = gate_record.get("approval_packet")
                packet_decisions: dict[str, Any] = {}
                if not isinstance(packet, dict):
                    self.error(location, f"{gate} approval_packet must be an object")
                else:
                    if packet.get("gate") != gate:
                        self.error(location, f"approval_packet must target '{gate}'")
                    if packet.get("version") != gate_record.get("artifact_version"):
                        self.error(location, "approval_packet version is stale")
                    raw_packet_decisions = packet.get("decisions")
                    if not isinstance(raw_packet_decisions, dict):
                        self.error(location, "approval_packet.decisions must be an object")
                    else:
                        packet_decisions = raw_packet_decisions
                        required_decisions = self.required_gate_decisions(gate)
                        if set(packet_decisions) != required_decisions:
                            self.error(
                                location,
                                f"approval_packet must cover: {sorted(required_decisions)}",
                            )
                        for decision_id in sorted(required_decisions):
                            decision = packet_decisions.get(decision_id)
                            if (
                                not isinstance(decision, dict)
                                or decision.get("status") != "approved"
                                or not has_meaningful_value(decision.get("value"))
                            ):
                                self.error(
                                    location,
                                    f"approval_packet decision '{decision_id}' is not approved",
                                )
                            elif decision.get("value_sha256") != value_sha256(
                                decision.get("value")
                            ):
                                self.error(
                                    location,
                                    f"approval value hash is stale for '{decision_id}'",
                                )
                        expected_decisions = self.expected_gate_decisions(gate)
                        for decision_id in sorted(required_decisions):
                            decision = packet_decisions.get(decision_id)
                            if (
                                isinstance(decision, dict)
                                and decision.get("value") != expected_decisions.get(decision_id)
                            ):
                                self.error(
                                    location,
                                    f"approval_packet value differs from current artifacts for '{decision_id}'",
                                )
                approval_mode = gate_record.get("approval_mode")
                approval_phrase = gate_record.get("approval_phrase")
                if not isinstance(approval_phrase, str) or not approval_phrase.strip():
                    self.error(location, "approval_phrase must be a non-empty string")
                if gate == "release":
                    if approval_mode != "release":
                        self.error(location, "release approval_mode must be 'release'")
                    if approval_phrase != "确认发布":
                        self.error(location, "release requires the exact phrase '确认发布'")
                else:
                    if approval_mode not in {"explicit", "local-affirmative", "batch-recommendations"}:
                        self.error(
                            location,
                            "approval_mode must be 'explicit', 'local-affirmative', or 'batch-recommendations'",
                        )
                    if gate == "decomposition" and approval_mode not in {"explicit", "local-affirmative"}:
                        self.error(location, "decomposition requires explicit human approval of the numbered review")
                    if approval_mode in {"explicit", "local-affirmative"} and isinstance(approval_phrase, str):
                        check_reply = REVIEW_POLICY.local_reply_error if approval_mode == "local-affirmative" else REVIEW_POLICY.explicit_reply_error
                        reply_error = check_reply(approval_phrase, self.docs.get("run-state.json"))
                        if reply_error:
                            self.error(location, reply_error)
                        for decision_id in sorted(self.required_gate_decisions(gate)):
                            decision = packet_decisions.get(decision_id, {})
                            if isinstance(decision, dict) and decision.get(
                                "value_sha256"
                            ) != value_sha256(decision.get("value")):
                                self.error(
                                    location,
                                    f"explicit approval value is stale for '{decision_id}'",
                                )
                    if approval_mode == "batch-recommendations":
                        if approval_phrase != "全部按推荐":
                            self.error(
                                location,
                                "batch approval requires the exact phrase '全部按推荐'",
                            )
                        if not isinstance(packet, dict) or packet.get(
                            "recommendations_visible"
                        ) is not True:
                            self.error(
                                location,
                                "batch approval requires all recommendations to be visible",
                            )
                        for decision_id in sorted(self.required_gate_decisions(gate)):
                            decision = packet_decisions.get(decision_id, {})
                            recommendation = (
                                decision.get("recommendation", {})
                                if isinstance(decision, dict)
                                else {}
                            )
                            if (
                                not isinstance(recommendation, dict)
                                or recommendation.get("value") != decision.get("value")
                                or not recommendation.get("reason")
                                or not recommendation.get("impact")
                            ):
                                self.error(
                                    location,
                                    f"batch approval lacks a complete accepted recommendation for '{decision_id}'",
                                )
            approved_artifacts = gate_record.get("approved_artifacts")
            if not isinstance(approved_artifacts, dict):
                continue
            actor_error = REVIEW_POLICY.review_actor_error(doc, gate_record, human_only=True)
            if actor_error:
                self.error(location, actor_error)
            missing = required_gate_artifacts[gate] - set(approved_artifacts)
            if missing:
                self.error(location, f"approved_artifacts missing: {sorted(missing)}")
            for path_value, declared_sha in approved_artifacts.items():
                approved_path = self.validate_relative_file(
                    path_value, f"{location}.approved_artifacts.{path_value}"
                )
                if approved_path:
                    actual_sha = hashlib.sha256(approved_path.read_bytes()).hexdigest()
                    if actual_sha != declared_sha:
                        self.error(location, f"approved artifact hash is stale: {path_value}")
        rejected_gates = {
            gate_name
            for gate_name, gate_record in gates.items()
            if gate_name in GATES
            and isinstance(gate_record, dict)
            and gate_record.get("status") == "rejected"
        }
        qa_state = self.docs.get("qa-report.json")
        if not isinstance(qa_state, dict):
            optional_qa_path = self.root / "qa-report.json"
            if optional_qa_path.is_file():
                try:
                    optional_qa = json.loads(optional_qa_path.read_text(encoding="utf-8"))
                    if isinstance(optional_qa, dict):
                        qa_state = optional_qa
                except (OSError, UnicodeError, json.JSONDecodeError):
                    qa_state = None
        qa_failed = isinstance(qa_state, dict) and (
            qa_state.get("status") == "failed"
            or any(
                isinstance(slide, dict) and slide.get("status") == "failed"
                for slide in (qa_state.get("slides") if isinstance(qa_state.get("slides"), list) else [])
            )
        )
        rollback = doc.get("rollback")
        if rejected_gates and not isinstance(rollback, dict):
            self.error(
                name,
                f"rejected gate requires a rollback record: {sorted(rejected_gates)}",
            )
        if qa_failed and not isinstance(rollback, dict):
            self.error(name, "failed QA requires a rollback record")
        if rollback is not None:
            rollback_location = f"{name}.rollback"
            if not isinstance(rollback, dict):
                self.error(rollback_location, "must be an object or null")
            else:
                trigger = rollback.get("trigger")
                if rejected_gates and trigger != "rejection":
                    self.error(rollback_location, "rejected gates require trigger 'rejection'")
                if qa_failed and not rejected_gates and trigger != "qa-defect":
                    self.error(rollback_location, "failed QA requires trigger 'qa-defect'")
                expected_owner: str | None = None
                if trigger == "qa-defect":
                    defect_class = rollback.get("defect_class")
                    if not isinstance(defect_class, str) or not defect_class:
                        self.error(rollback_location, "defect_class must be a non-empty string")
                    else:
                        expected_owner = DEFECT_STAGE_OWNERS.get(defect_class)
                    if isinstance(defect_class, str) and expected_owner is None:
                        self.error(rollback_location, "unknown defect_class")
                elif trigger == "rejection":
                    source_gate = rollback.get("source_gate")
                    if not isinstance(source_gate, str) or not source_gate:
                        self.error(rollback_location, "source_gate must be a non-empty string")
                    else:
                        expected_owner = GATE_ROLLBACK_OWNERS.get(source_gate)
                    if isinstance(source_gate, str) and expected_owner is None:
                        self.error(rollback_location, "unknown source_gate")
                    if isinstance(source_gate, str) and source_gate not in rejected_gates:
                        self.error(rollback_location, "source_gate status must be 'rejected'")
                    if rejected_gates and rejected_gates != {source_gate}:
                        self.error(
                            rollback_location,
                            "rollback source_gate must match the rejected gate",
                        )
                    source_record = gates.get(source_gate) if isinstance(source_gate, str) else None
                    if isinstance(source_record, dict):
                        retained_epoch = (
                            source_record.get("artifact_version") is not None
                            or bool(source_record.get("approved_artifacts"))
                            or source_record.get("approval_packet") is not None
                            or source_gate in artifact_versions
                        )
                        if retained_epoch:
                            self.error(
                                rollback_location,
                                "rejected gate must clear its prior approval epoch",
                            )
                else:
                    self.error(rollback_location, "trigger must be 'rejection' or 'qa-defect'")
                if not rollback.get("reason"):
                    self.error(rollback_location, "reason is required")
                if expected_owner is not None:
                    if rollback.get("owner_stage") != expected_owner:
                        self.error(
                            rollback_location,
                            f"rollback owner_stage must be '{expected_owner}'",
                        )
                    if declared_stage != expected_owner:
                        self.error(
                            rollback_location,
                            f"run-state stage must be rollback owner '{expected_owner}'",
                        )
        if self.reaches("release") and declared_stage != "released":
            self.error(name, "stage must be 'released' for release validation")

    def validate_cross_references(self) -> None:
        deck = self.docs.get("deck-spec.json", {})
        scene = self.docs.get("scene-graph.json", {})
        manifest = self.docs.get("asset-manifest.json", {})
        if not all(isinstance(doc, dict) for doc in (deck, scene, manifest)):
            return
        deck_records = {slide.get("id"): slide for slide in deck.get("slides", []) if isinstance(slide, dict)}
        scene_records = {slide.get("slide_id"): slide for slide in scene.get("slides", []) if isinstance(slide, dict)}
        deck_slides = set(deck_records)
        scene_slides = set(scene_records)
        asset_slides = {
            slide_id
            for asset in manifest.get("assets", [])
            if isinstance(asset, dict)
            for slide_id in (
                asset.get("slide_ids")
                if isinstance(asset.get("slide_ids"), list)
                else [asset.get("slide_id")]
            )
            if slide_id
        }
        for slide_id in (scene_slides | asset_slides) - deck_slides:
            self.error("cross-reference", f"unknown slide_id '{slide_id}'")
        scene_ids = {
            element.get("component_id")
            for slide in scene.get("slides", [])
            if isinstance(slide, dict)
            for element in slide.get("elements", [])
            if isinstance(element, dict)
        }
        manifest_ids = {
            asset.get("id") for asset in manifest.get("assets", []) if isinstance(asset, dict)
        }
        for component_id in scene_ids - manifest_ids:
            self.error("cross-reference", f"scene component '{component_id}' missing from asset-manifest")
        for component_id in manifest_ids - scene_ids:
            self.error("cross-reference", f"manifest asset '{component_id}' missing from scene-graph")

        manifest_by_slide: dict[Any, set[Any]] = {}
        for asset in manifest.get("assets", []):
            if isinstance(asset, dict):
                slide_ids = (
                    asset.get("slide_ids")
                    if isinstance(asset.get("slide_ids"), list)
                    else [asset.get("slide_id")]
                )
                for slide_id in slide_ids:
                    if slide_id:
                        manifest_by_slide.setdefault(slide_id, set()).add(asset.get("id"))
        for slide_id, deck_slide in deck_records.items():
            raw_components = deck_slide.get("components", [])
            deck_components = {
                component
                for component in raw_components
                if isinstance(component, str) and component
            } if isinstance(raw_components, list) else set()
            scene_components = {
                element.get("component_id")
                for element in scene_records.get(slide_id, {}).get("elements", [])
                if isinstance(element, dict)
            }
            if deck_components != scene_components:
                self.error("cross-reference", f"slide '{slide_id}' deck components differ from scene-graph")
            if deck_components != manifest_by_slide.get(slide_id, set()):
                self.error("cross-reference", f"slide '{slide_id}' deck components differ from asset-manifest")
            deck_slot_ids = {
                slot.get("id")
                for slot in deck_slide.get("fill_slots", [])
                if isinstance(slot, dict)
            }
            for asset in manifest.get("assets", []):
                if not isinstance(asset, dict):
                    continue
                asset_slide_ids = (
                    asset.get("slide_ids")
                    if isinstance(asset.get("slide_ids"), list)
                    else [asset.get("slide_id")]
                )
                if slide_id not in asset_slide_ids:
                    continue
                fill_slot_ids = asset.get("fill_slot_ids", [])
                if not isinstance(fill_slot_ids, list):
                    continue
                for slot_id in fill_slot_ids:
                    if not isinstance(slot_id, str):
                        continue
                    if slot_id not in deck_slot_ids:
                        self.error(
                            "cross-reference",
                            f"asset '{asset.get('id')}' references unknown fill_slot '{slot_id}'",
                        )

        if self.reaches("assembly"):
            assembly = self.docs.get("assembly-map.json", {})
            if isinstance(assembly, dict):
                assembly_records = {
                    slide.get("slide_id"): slide
                    for slide in assembly.get("slides", [])
                    if isinstance(slide, dict)
                }
                if set(assembly_records) != deck_slides:
                    self.error("cross-reference", "assembly-map slide coverage differs from deck-spec")
                for slide_id, deck_slide in deck_records.items():
                    assembled = assembly_records.get(slide_id, {})
                    assembled_components = {
                        component
                        for component in assembled.get("component_ids", [])
                        if isinstance(component, str) and component
                    }
                    deck_component_ids = {
                        component
                        for component in deck_slide.get("components", [])
                        if isinstance(component, str) and component
                    }
                    if assembled_components != deck_component_ids:
                        self.error("cross-reference", f"slide '{slide_id}' assembly component_ids are incomplete")
                    deck_slots = {
                        slot.get("id")
                        for slot in deck_slide.get("fill_slots", [])
                        if isinstance(slot, dict)
                    }
                    if set(assembled.get("fill_slot_ids", [])) != deck_slots:
                        self.error("cross-reference", f"slide '{slide_id}' assembly fill_slot_ids are incomplete")

        if self.reaches("qa"):
            qa = self.docs.get("qa-report.json", {})
            if isinstance(qa, dict):
                slide_rows = qa.get("slides") if isinstance(qa.get("slides"), list) else []
                qa_slides = {
                    slide.get("slide_id") for slide in slide_rows if isinstance(slide, dict)
                }
                if qa_slides != deck_slides:
                    self.error("cross-reference", "qa-report slide coverage differs from deck-spec")

    def validate_bbox(self, bbox: Any, location: str) -> None:
        if not isinstance(bbox, list) or len(bbox) != 4:
            self.error(location, "bbox must be [x, y, width, height]")
            return
        if any(not isinstance(value, (int, float)) for value in bbox):
            self.error(location, "bbox values must be numbers")
            return
        x, y, width, height = bbox
        if x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > 1 or y + height > 1:
            self.error(location, "bbox must fit within normalized 0..1 canvas bounds")

    def validate_relative_file(self, value: Any, location: str) -> Path | None:
        if not isinstance(value, str):
            self.error(location, "must be a relative path string")
            return None
        relative = Path(value)
        if relative.drive or relative.anchor or relative.is_absolute() or ".." in relative.parts:
            self.error(location, "must stay inside the project directory")
            return None
        project_root = self.root.resolve()
        resolved = (project_root / relative).resolve()
        if not resolved.is_relative_to(project_root):
            self.error(location, "must stay inside the project directory")
            return None
        if not resolved.is_file():
            self.error(location, f"referenced file does not exist: {value}")
            return None
        return resolved

    def inspect_png(self, path: Path) -> tuple[bool, float, bool, bool]:
        try:
            data = path.read_bytes()
        except OSError:
            return False, 0.0, False, False
        if not data.startswith(b"\x89PNG\r\n\x1a\n"):
            return False, 0.0, False, False
        offset = 8
        width: int | None = None
        height: int | None = None
        bit_depth: int | None = None
        color_type: int | None = None
        interlace: int | None = None
        compressed = bytearray()
        saw_iend = False
        while offset + 12 <= len(data):
            length = int.from_bytes(data[offset : offset + 4], "big")
            chunk_type = data[offset + 4 : offset + 8]
            end = offset + 12 + length
            if end > len(data):
                return False, 0.0, False, False
            payload = data[offset + 8 : offset + 8 + length]
            expected_crc = int.from_bytes(data[offset + 8 + length : end], "big")
            if (zlib.crc32(chunk_type + payload) & 0xFFFFFFFF) != expected_crc:
                return False, 0.0, False, False
            if chunk_type == b"IHDR":
                if length != 13:
                    return False, 0.0, False, False
                width = int.from_bytes(payload[0:4], "big")
                height = int.from_bytes(payload[4:8], "big")
                bit_depth = payload[8]
                color_type = payload[9]
                interlace = payload[12]
            elif chunk_type == b"IDAT":
                compressed.extend(payload)
            elif chunk_type == b"IEND":
                saw_iend = True
                break
            offset = end
        valid = all(value is not None for value in (width, height, bit_depth, color_type, interlace)) and saw_iend and bool(compressed)
        if not valid or bit_depth != 8 or color_type not in {4, 6} or interlace != 0:
            return bool(valid), 0.0, False, False
        channels = 2 if color_type == 4 else 4
        stride = int(width) * channels
        try:
            raw = zlib.decompress(bytes(compressed))
        except zlib.error:
            return False, 0.0, False, False
        expected_length = int(height) * (stride + 1)
        if len(raw) != expected_length:
            return False, 0.0, False, False
        if any(raw[row * (stride + 1)] > 4 for row in range(int(height))):
            return False, 0.0, False, False
        try:
            from PIL import Image
        except ImportError:
            return False, 0.0, False, False
        try:
            with self.pixel_cache.decode(data) as image:
                if image.mode not in {"LA", "RGBA"} or image.size != (width, height):
                    return False, 0.0, False, False
                alpha = image.getchannel("A")
                transparent_pixels = alpha.histogram()[0]
                # An alpha value of 1 is visually transparent edge noise.
                clear_border = (
                    alpha.crop((0, 0, width, 1)).getextrema()[1] <= 1
                    and alpha.crop((0, height - 1, width, height)).getextrema()[1] <= 1
                    and alpha.crop((0, 0, 1, height)).getextrema()[1] <= 1
                    and alpha.crop((width - 1, 0, width, height)).getextrema()[1] <= 1
                )
        except (OSError, ValueError, Image.DecompressionBombError, Image.DecompressionBombWarning):
            return False, 0.0, False, False
        pixel_count = int(width) * int(height)
        return True, transparent_pixels / pixel_count, clear_border, transparent_pixels < pixel_count

    def validate_pptx_content(self, path: Path) -> None:
        location = "assembly-map.json.output"
        try:
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
                required_parts = {
                    "[Content_Types].xml",
                    "_rels/.rels",
                    "ppt/presentation.xml",
                    "ppt/_rels/presentation.xml.rels",
                }
                missing_parts = required_parts - set(names)
                if missing_parts:
                    self.error(location, f"PPTX is missing required Office parts: {sorted(missing_parts)}")
                parsed_xml: dict[str, ElementTree.Element] = {}
                for part in required_parts & set(names):
                    try:
                        parsed_xml[part] = ElementTree.fromstring(archive.read(part))
                    except ElementTree.ParseError:
                        self.error(location, f"PPTX part is not valid XML: {part}")
                content_types = parsed_xml.get("[Content_Types].xml")
                if content_types is not None:
                    overrides = {
                        item.attrib.get("PartName"): item.attrib.get("ContentType")
                        for item in content_types
                        if item.tag.endswith("Override")
                    }
                    expected_presentation_type = (
                        "application/vnd.openxmlformats-officedocument."
                        "presentationml.presentation.main+xml"
                    )
                    if overrides.get("/ppt/presentation.xml") != expected_presentation_type:
                        self.error(location, "PPTX content types do not declare a presentation main part")
                root_relationships = parsed_xml.get("_rels/.rels")
                if root_relationships is not None:
                    office_targets = {
                        rel.attrib.get("Target")
                        for rel in root_relationships
                        if rel.attrib.get("Type", "").endswith("/officeDocument")
                    }
                    if "ppt/presentation.xml" not in office_targets:
                        self.error(location, "PPTX root relationships do not target ppt/presentation.xml")
                slide_names = [
                    name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                ]
                deck = self.docs.get("deck-spec.json", {})
                expected_pages = len(deck.get("slides", [])) if isinstance(deck, dict) else 0
                if len(slide_names) != expected_pages:
                    self.error(location, "PPTX slide count differs from deck-spec")
                if content_types is not None:
                    slide_type = (
                        "application/vnd.openxmlformats-officedocument."
                        "presentationml.slide+xml"
                    )
                    for slide_name in slide_names:
                        if overrides.get(f"/{slide_name}") != slide_type:
                            self.error(location, f"PPTX content types do not declare slide: {slide_name}")
                presentation = parsed_xml.get("ppt/presentation.xml")
                slide_size = None
                if presentation is not None:
                    size = presentation.find(
                        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}sldSz"
                    )
                    if size is not None:
                        try:
                            slide_size = (int(size.attrib["cx"]), int(size.attrib["cy"]))
                        except (KeyError, ValueError):
                            self.error(location, "PPTX slide size is invalid")
                    if slide_size is None:
                        self.error(location, "PPTX presentation is missing a valid slide size")
                    presentation_rels = parsed_xml.get("ppt/_rels/presentation.xml.rels")
                    rel_targets: dict[str, str] = {}
                    if presentation_rels is not None:
                        for rel in presentation_rels:
                            rel_id = rel.attrib.get("Id")
                            target = rel.attrib.get("Target")
                            if rel_id and target:
                                rel_targets[rel_id] = posixpath.normpath(posixpath.join("ppt", target))
                    relationship_attribute = (
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    )
                    slide_ids = presentation.findall(
                        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}sldId"
                    )
                    linked_slides = [
                        rel_targets.get(slide_id.attrib.get(relationship_attribute, ""))
                        for slide_id in slide_ids
                    ]
                    if None in linked_slides or set(linked_slides) != set(slide_names):
                        self.error(location, "PPTX presentation-to-slide relationship chain is incomplete")
                    assembly = self.docs.get("assembly-map.json", {})
                    expected_part_order = [
                        slide.get("pptx_part")
                        for slide in sorted(
                            (
                                slide
                                for slide in assembly.get("slides", [])
                                if isinstance(slide, dict)
                            ),
                            key=lambda item: item.get("output_index", 0),
                        )
                    ] if isinstance(assembly, dict) else []
                    if linked_slides != expected_part_order:
                        self.error(location, "PPTX slide relationship order differs from assembly-map")
                concept_hashes: set[str] = set()
                for slide in deck.get("slides", []):
                    if not isinstance(slide, dict) or not slide.get("concept_path"):
                        continue
                    concept = self.root / slide["concept_path"]
                    if concept.is_file():
                        concept_hashes.add(hashlib.sha256(concept.read_bytes()).hexdigest())
                manifest = self.docs.get("asset-manifest.json", {})
                approved_media_hashes = {
                    hashlib.sha256((self.root / asset["output"]).read_bytes()).hexdigest()
                    for asset in manifest.get("assets", [])
                    if isinstance(asset, dict)
                    and asset.get("status") == "approved"
                    and asset.get("output")
                    and (self.root / asset["output"]).is_file()
                } if isinstance(manifest, dict) else set()
                for name in names:
                    if name.startswith("ppt/media/"):
                        media_sha = hashlib.sha256(archive.read(name)).hexdigest()
                        if media_sha in concept_hashes:
                            self.error(location, "PPTX embeds an approved concept render as media")
                        if media_sha not in approved_media_hashes:
                            self.error(location, f"PPTX embeds unapproved raster media: {name}")
                visual_parts = slide_names + [
                    name
                    for name in names
                    if re.fullmatch(r"ppt/(slideLayouts|slideMasters)/[^/]+\.xml", name)
                ]
                for slide_name in visual_parts:
                    try:
                        slide_root = ElementTree.fromstring(archive.read(slide_name))
                    except ElementTree.ParseError:
                        self.error(location, f"PPTX slide is not valid XML: {slide_name}")
                        continue
                    if slide_size is None:
                        continue
                    self.validate_part_relationships(archive, set(names), slide_name, slide_root)
                    self.detect_full_slide_rasters(slide_root, slide_name, slide_size)
        except (OSError, zipfile.BadZipFile):
            self.error(location, "PPTX output is not a readable Office package")

    def detect_full_slide_rasters(
        self,
        xml_root: ElementTree.Element,
        part_name: str,
        slide_size: tuple[int, int],
    ) -> None:
        location = "assembly-map.json.output"
        p = "http://schemas.openxmlformats.org/presentationml/2006/main"
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        if xml_root.find(f".//{{{p}}}bg//{{{a}}}blip") is not None:
            self.error(location, f"PPTX uses a raster slide background in {part_name}")
        containers = xml_root.findall(f".//{{{p}}}pic") + xml_root.findall(f".//{{{p}}}sp")
        slide_cx, slide_cy = slide_size
        for container in containers:
            if container.find(f".//{{{a}}}blip") is None:
                continue
            transform = container.find(f".//{{{a}}}xfrm")
            if transform is None:
                continue
            offset = transform.find(f"{{{a}}}off")
            extent = transform.find(f"{{{a}}}ext")
            if offset is None or extent is None:
                continue
            try:
                x, y = int(offset.attrib["x"]), int(offset.attrib["y"])
                cx, cy = int(extent.attrib["cx"]), int(extent.attrib["cy"])
            except (KeyError, ValueError):
                continue
            if (
                x <= slide_cx * 0.05
                and y <= slide_cy * 0.05
                and cx >= slide_cx * 0.9
                and cy >= slide_cy * 0.9
            ):
                self.error(
                    location,
                    f"PPTX contains a near-full-slide raster in {part_name}; flattened concepts are prohibited",
                )

    def validate_part_relationships(
        self,
        archive: zipfile.ZipFile,
        names: set[str],
        part_name: str,
        xml_root: ElementTree.Element,
    ) -> None:
        location = "assembly-map.json.output"
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        relationship_attribute = (
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        blips = xml_root.findall(f".//{{{a}}}blip")
        if not blips:
            return
        directory = posixpath.dirname(part_name)
        rel_name = posixpath.join(directory, "_rels", posixpath.basename(part_name) + ".rels")
        if rel_name not in names:
            self.error(location, f"PPTX media relationships are missing for {part_name}")
            return
        try:
            rel_root = ElementTree.fromstring(archive.read(rel_name))
        except ElementTree.ParseError:
            self.error(location, f"PPTX relationships are invalid XML: {rel_name}")
            return
        targets: dict[str, str] = {}
        for rel in rel_root:
            rel_id = rel.attrib.get("Id")
            target = rel.attrib.get("Target")
            if rel_id and target:
                targets[rel_id] = posixpath.normpath(posixpath.join(directory, target))
        for blip in blips:
            rel_id = blip.attrib.get(relationship_attribute)
            if not rel_id or rel_id not in targets or targets[rel_id] not in names:
                self.error(location, f"PPTX media relationship is unresolved in {part_name}")

    def check_unique(self, records: list[Any], field: str, location: str) -> None:
        seen: set[Any] = set()
        for index, record in enumerate(records):
            if not isinstance(record, dict):
                continue
            value = record.get(field)
            if not isinstance(value, str) or not value:
                self.error(f"{location}[{index}]", f"{field} must be a non-empty string")
            elif value in seen:
                self.error(f"{location}[{index}]", f"duplicate {field} '{value}'")
            else:
                seen.add(value)

    def run(self) -> list[str]:
        try:
            return self._run()
        finally:
            self.pixel_cache.clear()

    def validate_inventory_candidate(self, documents: dict) -> list[str]:
        if self.stage != "inventory" or set(documents) != {"component-master-manifest.json", "scene-graph.json", "decomposition-map.json"}:
            raise ValueError("inventory candidate must provide exactly the three inventory documents")
        try:
            return self._run(inventory_candidate=documents)
        finally:
            self.pixel_cache.clear()

    def _run(self, *, inventory_candidate: dict | None = None) -> list[str]:
        if not self.root.is_dir():
            return [f"project: directory does not exist: {self.root}"]
        self.load_required_documents(inventory_candidate)
        if (self.root / ".inventory-merge").exists():
            return self.errors
        if self.reaches("intake"):
            self.validate_brief()
            self.validate_state()
        if self.reaches("research"):
            self.validate_research()
        if self.reaches("story"):
            self.validate_alignment()
            self.validate_storyboard()
        if self.reaches("style"):
            self.validate_reference_study()
            self.validate_concept_calibration()
            self.validate_style()
            self.validate_capability_report()
        if self.reaches("spec"):
            self.validate_deck()
        if self.reaches("concept"):
            self.validate_concept_generation_log()
        if self.reaches("inventory"):
            self.validate_scene_graph()
            self.validate_manifest()
            self.validate_cross_references()
            self.validate_component_master_manifest()
            self.validate_decomposition_map(require_review=inventory_candidate is None)
        if self.reaches("pilot"):
            state = self.docs.get("run-state.json", {})
            gates = state.get("gates", {}) if isinstance(state, dict) else {}
            decomposition = gates.get("decomposition", {}) if isinstance(gates, dict) else {}
            if not isinstance(decomposition, dict) or decomposition.get("status") != "approved":
                self.error("run-state.json.gates.decomposition", "approved decomposition is required before assets")
            if self.stage == "pilot" or self.pilot_required():
                self.validate_pilot()
        if self.reaches("assembly"):
            self.validate_assembly()
        if self.reaches("qa"):
            self.validate_qa()
        return self.errors


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project_dir", type=Path)
    parser.add_argument("--stage", choices=STAGES + ["export"], default="release")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    kind_path = args.project_dir / "project-kind.json"
    if kind_path.exists() or kind_path.is_symlink():
        if args.stage not in {"concept", "assets", "assembly", "qa", "export"}:
            print("FAIL: local showcase checks stop at the selected delivery mode; formal production release is a separate route")
            return 1
        from apw import check_project

        report = check_project(args.project_dir.resolve(), args.stage)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if report["ok"] else 1
    if args.stage == "export":
        print("FAIL: strict production export uses export_final_package.py with the current production contract")
        return 1
    validator = ProjectValidator(args.project_dir.resolve(), args.stage)
    errors = validator.run()
    if errors:
        print(f"FAIL: {len(errors)} contract issue(s)")
        for issue in errors:
            print(f"- {issue}")
        return 1
    print(f"PASS: project contract is valid through stage '{args.stage}'")
    return 0


if __name__ == "__main__":
    sys.exit(main())
