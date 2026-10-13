"""Native, editable PowerPoint object builders for the explicit showcase contract.

This module deliberately implements a small, inspectable subset. It never turns
a slide concept into a background image; picture objects reference separate
manifest assets and remain individually movable/replacable in PowerPoint.
"""

from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
)
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


class NativeObjectError(ValueError):
    """Invalid or unsupported showcase object contract."""


_COLOR = re.compile(r"#[0-9A-Fa-f]{6}\Z")
_SHAPES = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rect": MSO_SHAPE.RECTANGLE,
    "rounded-rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "oval": MSO_SHAPE.OVAL,
    "diamond": MSO_SHAPE.DIAMOND,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "hexagon": MSO_SHAPE.HEXAGON,
}
_CHARTS = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
}


def _color(value: Any, field: str) -> RGBColor:
    if not isinstance(value, str) or not _COLOR.fullmatch(value):
        raise NativeObjectError(f"{field} must be a #RRGGBB color")
    return RGBColor.from_string(value[1:])


def _normalized_point(value: Any, field: str) -> tuple[float, float]:
    if not isinstance(value, dict):
        raise NativeObjectError(f"{field} must be an object with x and y")
    x, y = value.get("x"), value.get("y")
    if (type(x) not in (int, float) or type(y) not in (int, float)
            or not math.isfinite(x) or not math.isfinite(y)
            or not 0 <= x <= 1 or not 0 <= y <= 1):
        raise NativeObjectError(f"{field} coordinates must be finite normalized numbers")
    return float(x), float(y)


def _validate_paint(value: Any, field: str, *, allow_none: bool = True) -> dict[str, Any] | None:
    if value is None and allow_none:
        return None
    if not isinstance(value, dict):
        raise NativeObjectError(f"{field} must be a paint object or null")
    _color(value.get("color"), f"{field}.color")
    transparency = value.get("transparency", 0)
    if (type(transparency) not in (int, float) or not math.isfinite(transparency)
            or not 0 <= transparency <= 100):
        raise NativeObjectError(f"{field}.transparency must be between 0 and 100")
    result = {"color": value["color"], "transparency": transparency}
    if "width_pt" in value:
        width_pt = value["width_pt"]
        if (type(width_pt) not in (int, float) or not math.isfinite(width_pt)
                or width_pt <= 0 or width_pt > 100):
            raise NativeObjectError(f"{field}.width_pt must be between 0 and 100")
        result["width_pt"] = width_pt
    return result


def _set_alpha(color_element: Any, transparency: float) -> None:
    if transparency <= 0:
        return
    for old in color_element.findall(qn("a:alpha")):
        color_element.remove(old)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(round((100 - transparency) * 1000)))
    color_element.append(alpha)


def _set_line_arrowhead(line_element: Any, arrowhead: str) -> None:
    """Write a DrawingML tailEnd marker; python-pptx currently exposes no public setter."""
    if arrowhead == "none":
        return
    if arrowhead != "triangle":
        raise NativeObjectError("connector.arrowhead must be 'none' or 'triangle'")
    line = line_element.get_or_add_ln()
    tail = line.find(qn("a:tailEnd"))
    if tail is None:
        tail = OxmlElement("a:tailEnd")
        line.append(tail)
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def _set_solid_shape_fill(shape: Any, paint: dict[str, Any] | None) -> None:
    if paint is None or paint["transparency"] == 100:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(paint["color"], "fill.color")
    solid = shape._element.spPr.find(qn("a:solidFill"))
    if solid is not None:
        color_element = solid.find(qn("a:srgbClr"))
        if color_element is not None:
            _set_alpha(color_element, paint["transparency"])


def _set_solid_line(shape: Any, paint: dict[str, Any] | None) -> None:
    if paint is None or paint["transparency"] == 100:
        shape.line.fill.background()
        return
    shape.line.color.rgb = _color(paint["color"], "line.color")
    shape.line.width = Pt(paint.get("width_pt", 1.0))
    line_element = shape._element.spPr
    line_element = line_element.find(qn("a:ln"))
    if line_element is not None:
        solid = line_element.find(qn("a:solidFill"))
        if solid is not None:
            color_element = solid.find(qn("a:srgbClr"))
            if color_element is not None:
                _set_alpha(color_element, paint["transparency"])


def _set_textbox_fill(shape: Any, background_color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(background_color, "text.style.background_color")
    shape.line.fill.background()


def _style_text(shape: Any, text: str, style: dict[str, Any], bbox: tuple[int, int, int, int]) -> None:
    if not isinstance(text, str):
        raise NativeObjectError("text.text must be a string")
    if "\r" in text:
        raise NativeObjectError("text.text line breaks must use canonical LF characters")
    font_family = style.get("font_family")
    font_size = style.get("font_size_pt")
    if not isinstance(font_family, str) or not font_family.strip():
        raise NativeObjectError("text.style.font_family must be a non-empty string")
    if type(font_size) not in (int, float) or not math.isfinite(font_size) or font_size <= 0 or font_size > 200:
        raise NativeObjectError("text.style.font_size_pt must be a positive number")
    if not isinstance(style.get("bold", False), bool):
        raise NativeObjectError("text.style.bold must be a boolean")
    align = style.get("align", "left")
    valign = style.get("valign", "middle")
    if align not in {"left", "center", "right"}:
        raise NativeObjectError("text.style.align must be left, center, or right")
    if valign not in {"top", "middle", "bottom"}:
        raise NativeObjectError("text.style.valign must be top, middle, or bottom")
    _color(style.get("color"), "text.style.color")
    background = style.get("background_color")
    _color(background, "text.style.background_color")

    textbox = shape
    textbox.left, textbox.top, textbox.width, textbox.height = bbox
    if style.get("opaque_background", False):
        _set_textbox_fill(textbox, background)
    else:
        textbox.fill.background()
        textbox.line.fill.background()
    frame = textbox.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.06)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }[valign]
    for line_index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        paragraph.font.name = font_family
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = style.get("bold", False)
        paragraph.font.color.rgb = _color(style["color"], "text.style.color")


def _table_text(value: Any, field: str) -> str:
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise NativeObjectError(f"{field} must be a string or finite number")
    if isinstance(value, float) and not math.isfinite(value):
        raise NativeObjectError(f"{field} must be finite")
    return str(value)


def _status_is_illustrative(status: Any) -> bool:
    if not isinstance(status, str):
        return False
    normalized = status.strip().lower()
    return normalized in {
        "illustrative", "illustration", "example", "demo", "synthetic",
        "hypothetical", "示意", "示例", "演示", "示意数据", "演示数据", "假设",
    }


def _object_name(slide_id: str, instance_id: str, kind: str) -> str:
    return f"APW:showcase:{slide_id}:{instance_id}:{kind}"


def _base_row(element: dict[str, Any], slide_id: str, shape: Any, object_type: str) -> dict[str, Any]:
    shape.name = _object_name(slide_id, element["instance_id"], object_type)
    return {
        "instance_id": element["instance_id"],
        "type": object_type,
        "object_name": shape.name,
        "bbox": list(element["bbox"]),
        "z": element["z"],
        "editable": object_type != "picture" and object_type != "source-evidence",
    }


def _add_text(slide: Any, element: dict[str, Any], geometry: tuple[int, int, int, int], slide_id: str) -> dict[str, Any]:
    style = element.get("style")
    if not isinstance(style, dict):
        raise NativeObjectError(f"{element.get('instance_id')}: text.style must be an object")
    textbox = slide.shapes.add_textbox(*geometry)
    _style_text(textbox, element.get("text"), style, geometry)
    row = _base_row(element, slide_id, textbox, "text")
    row.update({
        "copy_id": element.get("copy_id"),
        "role": element.get("role"),
        "text": element["text"],
        "truth_status": element.get("truth_status"),
        "source_ids": element.get("source_ids", []),
        "font_family": style["font_family"],
        "font_size_pt": style["font_size_pt"],
        "color": style["color"],
        "background_color": style["background_color"],
    })
    return row


def _add_shape(slide: Any, element: dict[str, Any], geometry: tuple[int, int, int, int], slide_id: str) -> dict[str, Any]:
    shape_kind = element.get("shape_kind")
    if not isinstance(shape_kind, str) or shape_kind not in _SHAPES:
        raise NativeObjectError(f"{element.get('instance_id')}: unsupported shape_kind {shape_kind!r}")
    fill = _validate_paint(element.get("fill"), "shape.fill")
    line = _validate_paint(element.get("line"), "shape.line")
    shape = slide.shapes.add_shape(_SHAPES[shape_kind], *geometry)
    _set_solid_shape_fill(shape, fill)
    _set_solid_line(shape, line)
    row = _base_row(element, slide_id, shape, "shape")
    row.update({"shape_kind": shape_kind, "fill": fill, "line": line})
    return row


def _add_connector(slide: Any, element: dict[str, Any], slide_id: str, canvas: tuple[int, int]) -> dict[str, Any]:
    start = _normalized_point(element.get("from"), "connector.from")
    end = _normalized_point(element.get("to"), "connector.to")
    if start == end:
        raise NativeObjectError(f"{element.get('instance_id')}: connector endpoints must differ")
    points = element.get("points")
    if points is not None:
        if not isinstance(points, list) or len(points) != 2:
            raise NativeObjectError("only a straight connector with two endpoints is supported")
        point_values = [_normalized_point(point, "connector.points[]") for point in points]
        if point_values != [start, end]:
            raise NativeObjectError("connector.points must match from/to in order")
    bbox = element.get("bbox")
    left, top, width, height = bbox
    if (min(start[0], end[0]) < left - 1e-6 or min(start[1], end[1]) < top - 1e-6
            or max(start[0], end[0]) > left + width + 1e-6
            or max(start[1], end[1]) > top + height + 1e-6):
        raise NativeObjectError(f"{element.get('instance_id')}: connector endpoints must fit its bbox")
    line = _validate_paint(element.get("line"), "connector.line", allow_none=False)
    if "width_pt" not in line:
        raise NativeObjectError("connector.line.width_pt is required")
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        round(start[0] * canvas[0]),
        round(start[1] * canvas[1]),
        round(end[0] * canvas[0]),
        round(end[1] * canvas[1]),
    )
    _set_solid_line(shape, line)
    _set_line_arrowhead(shape._element.spPr, element.get("arrowhead", "none"))
    row = _base_row(element, slide_id, shape, "connector")
    row.update({
        "from": {"x": start[0], "y": start[1]},
        "to": {"x": end[0], "y": end[1]},
        "line": line,
        "arrowhead": element.get("arrowhead", "none"),
    })
    return row


def _series_and_categories(element: dict[str, Any]) -> tuple[list[str], list[tuple[str, list[int | float]]]]:
    categories = element.get("categories")
    series = element.get("series")
    if (not isinstance(categories, list) or not categories
            or any(not isinstance(item, str) or not item for item in categories)):
        raise NativeObjectError("chart.categories must be a non-empty list of labels")
    if not isinstance(series, list) or not series:
        raise NativeObjectError("chart.series must be a non-empty list")
    clean_series = []
    names = set()
    for index, item in enumerate(series):
        if not isinstance(item, dict) or not isinstance(item.get("name"), str) or not item["name"].strip():
            raise NativeObjectError(f"chart.series[{index}] needs a non-empty name")
        if item["name"] in names:
            raise NativeObjectError("chart series names must be unique")
        names.add(item["name"])
        values = item.get("values")
        if (not isinstance(values, list) or len(values) != len(categories)
                or any(type(value) not in (int, float) or not math.isfinite(value) for value in values)):
            raise NativeObjectError(f"chart.series[{index}].values must match categories with finite numbers")
        clean_series.append((item["name"], list(values)))
    return categories, clean_series


_CHART_STYLE_KEYS = {
    "font_family", "font_size_pt", "color", "series_colors", "value_axis",
    "category_axis_title", "value_axis_title", "data_label_number_format",
    "line_markers",
}
_VALUE_AXIS_STYLE_KEYS = {"minimum_scale", "maximum_scale", "major_unit", "number_format"}


def _validate_chart_style(value: Any, chart_kind: str, series_count: int) -> dict[str, Any]:
    """Validate project-selected chart presentation settings without setting defaults."""
    if value is None:
        return {}
    if not isinstance(value, dict):
        raise NativeObjectError("chart.style must be an object")
    unknown = set(value) - _CHART_STYLE_KEYS
    if unknown:
        raise NativeObjectError(f"chart.style contains unsupported keys: {', '.join(sorted(unknown))}")

    normalized = dict(value)
    if "font_family" in normalized:
        family = normalized["font_family"]
        if not isinstance(family, str) or not family.strip():
            raise NativeObjectError("chart.style.font_family must be a non-empty string")
    if "font_size_pt" in normalized:
        size = normalized["font_size_pt"]
        if type(size) not in (int, float) or not math.isfinite(size) or not 1 <= size <= 200:
            raise NativeObjectError("chart.style.font_size_pt must be between 1 and 200")
    if "color" in normalized:
        _color(normalized["color"], "chart.style.color")
    if "series_colors" in normalized:
        colors = normalized["series_colors"]
        if not isinstance(colors, list) or len(colors) != series_count:
            raise NativeObjectError("chart.style.series_colors must contain one color per chart series")
        for index, color in enumerate(colors):
            _color(color, f"chart.style.series_colors[{index}]")
    if "value_axis" in normalized:
        axis = normalized["value_axis"]
        if not isinstance(axis, dict):
            raise NativeObjectError("chart.style.value_axis must be an object")
        unknown_axis = set(axis) - _VALUE_AXIS_STYLE_KEYS
        if unknown_axis:
            raise NativeObjectError(f"chart.style.value_axis contains unsupported keys: {', '.join(sorted(unknown_axis))}")
        for field in ("minimum_scale", "maximum_scale", "major_unit"):
            if field in axis:
                number = axis[field]
                if type(number) not in (int, float) or not math.isfinite(number):
                    raise NativeObjectError(f"chart.style.value_axis.{field} must be a finite number")
                if field == "major_unit" and number <= 0:
                    raise NativeObjectError("chart.style.value_axis.major_unit must be positive")
        if ("minimum_scale" in axis and "maximum_scale" in axis
                and axis["minimum_scale"] >= axis["maximum_scale"]):
            raise NativeObjectError("chart.style.value_axis.minimum_scale must be below maximum_scale")
        if "number_format" in axis and (not isinstance(axis["number_format"], str) or not axis["number_format"]):
            raise NativeObjectError("chart.style.value_axis.number_format must be non-empty text")
    for field in ("category_axis_title", "value_axis_title", "data_label_number_format"):
        if field in normalized and (not isinstance(normalized[field], str) or not normalized[field].strip()):
            raise NativeObjectError(f"chart.style.{field} must be non-empty text")
    if "line_markers" in normalized:
        if not isinstance(normalized["line_markers"], bool):
            raise NativeObjectError("chart.style.line_markers must be a boolean")
        if chart_kind != "line":
            raise NativeObjectError("chart.style.line_markers is only supported for line charts")
    return normalized


def _apply_chart_font(font: Any, style: dict[str, Any]) -> None:
    if "font_family" in style:
        font.name = style["font_family"]
        character_properties = font._element
        east_asian_font = character_properties.find(qn("a:ea"))
        if east_asian_font is None:
            east_asian_font = OxmlElement("a:ea")
            character_properties.insert_element_before(
                east_asian_font,
                "a:cs",
                "a:sym",
                "a:hlinkClick",
                "a:hlinkMouseOver",
                "a:rtl",
                "a:extLst",
            )
        east_asian_font.set("typeface", style["font_family"])
    if "font_size_pt" in style:
        font.size = Pt(style["font_size_pt"])
    if "color" in style:
        font.color.rgb = _color(style["color"], "chart.style.color")


def _apply_chart_text_style(chart: Any, style: dict[str, Any]) -> None:
    if not style:
        return
    _apply_chart_font(chart.font, style)
    text_frames = [chart.chart_title.text_frame]
    if chart.category_axis.has_title:
        text_frames.append(chart.category_axis.axis_title.text_frame)
    if chart.value_axis.has_title:
        text_frames.append(chart.value_axis.axis_title.text_frame)
    for frame in text_frames:
        for paragraph in frame.paragraphs:
            _apply_chart_font(paragraph.font, style)
    for axis in (chart.category_axis, chart.value_axis):
        _apply_chart_font(axis.tick_labels.font, style)
    if chart.has_legend:
        _apply_chart_font(chart.legend.font, style)
    for plot in chart.plots:
        if plot.has_data_labels:
            _apply_chart_font(plot.data_labels.font, style)


def _chart_style(chart: Any, body_font_pt: float, title: str, style: dict[str, Any] | None = None) -> None:
    style = style or {}
    chart.has_legend = len(chart.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(body_font_pt)
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    if chart.chart_title.has_text_frame:
        chart.chart_title.text_frame.paragraphs[0].font.size = Pt(body_font_pt)
    chart.category_axis.tick_labels.font.size = Pt(body_font_pt)
    chart.value_axis.tick_labels.font.size = Pt(body_font_pt)
    chart.value_axis.has_major_gridlines = True
    data_label_position = (
        XL_LABEL_POSITION.ABOVE
        if chart.chart_type == XL_CHART_TYPE.LINE
        else XL_LABEL_POSITION.OUTSIDE_END
    )
    for plot in chart.plots:
        plot.has_data_labels = True
        plot.data_labels.position = data_label_position
        plot.data_labels.font.size = Pt(body_font_pt)

    value_axis_style = style.get("value_axis", {})
    if "minimum_scale" in value_axis_style:
        chart.value_axis.minimum_scale = value_axis_style["minimum_scale"]
    if "maximum_scale" in value_axis_style:
        chart.value_axis.maximum_scale = value_axis_style["maximum_scale"]
    if "major_unit" in value_axis_style:
        chart.value_axis.major_unit = value_axis_style["major_unit"]
    if "number_format" in value_axis_style:
        chart.value_axis.tick_labels.number_format = value_axis_style["number_format"]

    for field, axis in (("category_axis_title", chart.category_axis), ("value_axis_title", chart.value_axis)):
        if field in style:
            axis.has_title = True
            axis.axis_title.text_frame.text = style[field]
    if "data_label_number_format" in style:
        for plot in chart.plots:
            plot.data_labels.number_format = style["data_label_number_format"]

    for index, color in enumerate(style.get("series_colors", [])):
        rgb = _color(color, f"chart.style.series_colors[{index}]")
        series = chart.series[index]
        if chart.chart_type == XL_CHART_TYPE.LINE:
            series.format.line.color.rgb = rgb
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = rgb
            series.marker.format.line.color.rgb = rgb
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = rgb
            series.format.line.color.rgb = rgb

    if "line_markers" in style:
        for series in chart.series:
            series.marker.style = XL_MARKER_STYLE.CIRCLE if style["line_markers"] else XL_MARKER_STYLE.NONE
    _apply_chart_text_style(chart, style)


def _add_chart(
    slide: Any,
    element: dict[str, Any],
    geometry: tuple[int, int, int, int],
    slide_id: str,
    body_font_pt: float,
) -> dict[str, Any]:
    chart_kind = element.get("chart_type")
    if not isinstance(chart_kind, str) or chart_kind not in _CHARTS:
        raise NativeObjectError(f"{element.get('instance_id')}: unsupported chart_type {chart_kind!r}")
    categories, series = _series_and_categories(element)
    title = element.get("title")
    if not isinstance(title, str) or not title.strip():
        raise NativeObjectError("chart.title must be a non-empty string")
    units = element.get("units")
    if not isinstance(units, str) or not units.strip():
        raise NativeObjectError("chart.units must be non-empty text")
    truth_status = element.get("truth_status")
    source_ids = element.get("source_ids")
    if not isinstance(truth_status, str) or not truth_status.strip():
        raise NativeObjectError("chart.truth_status must be a non-empty string")
    if not isinstance(source_ids, list) or any(not isinstance(item, str) or not item for item in source_ids):
        raise NativeObjectError("chart.source_ids must be a list of non-empty strings")
    style = _validate_chart_style(element.get("style"), chart_kind, len(series))
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    graphic = slide.shapes.add_chart(_CHARTS[chart_kind], *geometry, data)
    chart = graphic.chart
    _chart_style(chart, body_font_pt, title, style)
    row = _base_row(element, slide_id, graphic, "chart")
    row.update({
        "chart_type": chart_kind,
        "title": title,
        "categories": categories,
        "series": [{"name": name, "values": values} for name, values in series],
        "units": units,
        "truth_status": truth_status,
        "source_ids": source_ids,
        "visible_status_label": title if _status_is_illustrative(truth_status) else None,
    })
    if style:
        row["style"] = style
    return row


def _add_table(
    slide: Any,
    element: dict[str, Any],
    geometry: tuple[int, int, int, int],
    slide_id: str,
    body_font_pt: float,
) -> dict[str, Any]:
    columns = element.get("columns")
    rows = element.get("rows")
    units = element.get("units")
    visible_label = element.get("visible_label")
    truth_status = element.get("truth_status")
    source_ids = element.get("source_ids")
    if not isinstance(columns, list) or not columns:
        raise NativeObjectError("table.columns must be a non-empty list")
    if any(not isinstance(column, dict) or not isinstance(column.get("label"), str) or not column["label"].strip()
           for column in columns):
        raise NativeObjectError("table.columns entries need non-empty labels")
    if not isinstance(rows, list) or not rows:
        raise NativeObjectError("table.rows must be a non-empty list")
    if not isinstance(units, str) or not units.strip():
        raise NativeObjectError("table.units must be non-empty text")
    if visible_label is not None and (not isinstance(visible_label, str) or not visible_label.strip()):
        raise NativeObjectError("table.visible_label must be non-empty text or null")
    if not isinstance(truth_status, str) or not truth_status.strip():
        raise NativeObjectError("table.truth_status must be a non-empty string")
    if not isinstance(source_ids, list) or any(not isinstance(item, str) or not item for item in source_ids):
        raise NativeObjectError("table.source_ids must be a list of non-empty strings")
    width, height = geometry[2], geometry[3]
    status_label = visible_label if truth_status in {"illustrative", "hypothetical"} else None
    row_count = len(rows) + 1 + int(status_label is not None)
    graphic = slide.shapes.add_table(row_count, len(columns), *geometry)
    table = graphic.table
    table_style = element.get("style", {})
    if not isinstance(table_style, dict):
        raise NativeObjectError("table.style must be an object")
    font_size = table_style.get("font_size_pt", body_font_pt)
    if type(font_size) not in (int, float) or not math.isfinite(font_size) or font_size <= 0:
        raise NativeObjectError("table.style.font_size_pt must be positive")
    for col_index, column in enumerate(columns):
        if "width_fraction" in column:
            fraction = column["width_fraction"]
            if type(fraction) not in (int, float) or not 0 < fraction <= 1:
                raise NativeObjectError("table column width_fraction must be in 0..1")
            table.columns[col_index].width = round(width * fraction)
    header_labels = [column["label"] for column in columns]
    if status_label is not None:
        if len(columns) > 1:
            table.cell(0, 0).merge(table.cell(0, len(columns) - 1))
        _set_table_cell(table.cell(0, 0), status_label, font_size, header=True, style=table_style)
    for col_index, label in enumerate(header_labels):
        _set_table_cell(table.cell(int(status_label is not None), col_index), label, font_size, header=True, style=table_style)
    clean_rows = []
    for row_index, row in enumerate(rows, start=1):
        cells = row.get("cells") if isinstance(row, dict) else None
        if not isinstance(cells, list) or len(cells) != len(columns):
            raise NativeObjectError(f"table.rows[{row_index - 1}].cells must match columns")
        values = [_table_text(value, f"table.rows[{row_index - 1}].cells[]") for value in cells]
        clean_rows.append(values)
        for col_index, value in enumerate(values):
            _set_table_cell(table.cell(row_index + int(status_label is not None), col_index), value, font_size, header=False, style=table_style)
    row = _base_row(element, slide_id, graphic, "table")
    row.update({
        "columns": [{"label": column["label"]} for column in columns],
        "units": units,
        "rows": [{"cells": values} for values in clean_rows],
        "truth_status": truth_status,
        "source_ids": source_ids,
        "visible_label": status_label,
        "visible_status_label": status_label,
        "table_shape": {
            "row_count": row_count,
            "column_count": len(columns),
            "has_visible_label_row": status_label is not None,
        },
    })
    return row


def _set_table_cell(cell: Any, value: str, font_size_pt: float, *, header: bool, style: dict[str, Any]) -> None:
    cell.text = value
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    cell.fill.solid()
    cell.fill.fore_color.rgb = _color(style.get("header_fill" if header else "body_fill", "#1E293B" if header else "#F8FAFC"), "table.style.fill")
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.name = style.get("font_family", "Microsoft YaHei")
        paragraph.font.size = Pt(font_size_pt)
        paragraph.font.bold = header
        paragraph.font.color.rgb = _color(style.get("header_color" if header else "body_color", "#F8FAFC" if header else "#0F172A"), "table.style.color")


def _asset_for_picture(element: dict[str, Any], assets: dict[str, dict[str, Any]], root: Path) -> tuple[Path, dict[str, Any]]:
    asset_id = element.get("asset_id")
    asset = assets.get(asset_id) if isinstance(asset_id, str) else None
    if not isinstance(asset, dict):
        raise NativeObjectError(f"{element.get('instance_id')}: picture asset_id is missing from asset manifest")
    relative = asset.get("path")
    expected = asset.get("sha256")
    if not isinstance(relative, str) or not relative:
        raise NativeObjectError(f"{asset_id}: asset manifest path is required")
    path = (root / relative).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise NativeObjectError(f"{asset_id}: picture path is missing or unsafe")
    if not isinstance(expected, str) or len(expected) != 64:
        raise NativeObjectError(f"{asset_id}: picture sha256 is missing")
    import hashlib

    if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
        raise NativeObjectError(f"{asset_id}: picture source hash is stale")
    return path, asset


def _fit_picture(slide: Any, path: Path, geometry: tuple[int, int, int, int]) -> Any:
    from PIL import Image

    left, top, width, height = geometry
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = width / height
    if ratio > box_ratio:
        placed_width = width
        placed_height = round(width / ratio)
    else:
        placed_height = height
        placed_width = round(height * ratio)
    placed_left = left + (width - placed_width) // 2
    placed_top = top + (height - placed_height) // 2
    return slide.shapes.add_picture(
        str(path),
        placed_left,
        placed_top,
        width=placed_width,
        height=placed_height,
    )


def add_native_object(
    slide: Any,
    element: dict[str, Any],
    *,
    slide_id: str,
    slide_width: int,
    slide_height: int,
    root: Path,
    assets: dict[str, dict[str, Any]],
    body_font_pt: float,
) -> tuple[dict[str, Any], str | None]:
    """Add one validated scene element and return its source-map row and notes text."""
    if not isinstance(element, dict):
        raise NativeObjectError("scene elements must be objects")
    instance_id = element.get("instance_id")
    object_type = element.get("type")
    if not isinstance(instance_id, str) or not instance_id.strip():
        raise NativeObjectError("scene element instance_id must be a non-empty string")
    if object_type not in {"text", "shape", "connector", "table", "chart", "picture", "source-evidence"}:
        raise NativeObjectError(f"{instance_id}: unsupported scene element type {object_type!r}")
    bbox = element.get("bbox")
    if (not isinstance(bbox, list) or len(bbox) != 4
            or any(type(value) not in (int, float) or not math.isfinite(value) for value in bbox)
            or bbox[0] < 0 or bbox[1] < 0 or bbox[2] <= 0 or bbox[3] <= 0
            or bbox[0] + bbox[2] > 1 or bbox[1] + bbox[3] > 1):
        raise NativeObjectError(f"{instance_id}: invalid normalized bbox")
    z = element.get("z")
    if type(z) is not int or z < 0:
        raise NativeObjectError(f"{instance_id}: z must be a non-negative integer")
    geometry = (
        round(bbox[0] * slide_width),
        round(bbox[1] * slide_height),
        round(bbox[2] * slide_width),
        round(bbox[3] * slide_height),
    )
    if object_type == "text":
        row = _add_text(slide, element, geometry, slide_id)
        return row, _provenance_note(row)
    if object_type == "shape":
        return _add_shape(slide, element, geometry, slide_id), None
    if object_type == "connector":
        return _add_connector(slide, element, slide_id, (slide_width, slide_height)), None
    if object_type == "table":
        row = _add_table(slide, element, geometry, slide_id, body_font_pt)
        return row, _provenance_note(row)
    if object_type == "chart":
        row = _add_chart(slide, element, geometry, slide_id, body_font_pt)
        return row, _provenance_note(row)
    asset_path, asset = _asset_for_picture(element, assets, root)
    if bbox[0] <= 0.05 and bbox[1] <= 0.05 and bbox[2] >= 0.90 and bbox[3] >= 0.90:
        raise NativeObjectError(f"{instance_id}: a full-slide or near-full-slide raster is prohibited")
    shape = _fit_picture(slide, asset_path, geometry)
    row = _base_row(element, slide_id, shape, object_type)
    row.update({
        "asset_id": element["asset_id"],
        "asset_sha256": asset["sha256"],
        "asset_path": asset["path"],
        "source_locator": asset.get("source_locator"),
        "truth_status": asset.get("truth_status"),
    })
    return row, _provenance_note(row)


def _provenance_note(row: dict[str, Any]) -> str:
    source_ids = row.get("source_ids")
    asset_locator = row.get("source_locator")
    details = []
    if row.get("truth_status"):
        details.append(f"状态：{row['truth_status']}")
    if source_ids:
        details.append("来源 ID：" + "、".join(source_ids))
    if asset_locator:
        details.append("素材来源：" + str(asset_locator))
    if not details:
        return ""
    return f"[{row['instance_id']}] " + "；".join(details)

