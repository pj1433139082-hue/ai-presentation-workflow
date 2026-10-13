from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parent
ASSET = ROOT / "assets" / "community-library.png"
OUT = ROOT / "deck-草稿.pptx"
FONT = "Microsoft YaHei"
W, H = 13.333333, 7.5
C = {
    "paper": "F5FAF9", "white": "FFFFFF", "navy": "17344E", "teal": "287A78",
    "aqua": "D9ECE9", "aqua2": "BFDCD8", "coral": "D77858", "yellow": "F0C85B",
    "ink": "213B4E", "muted": "526A78", "line": "B8CCCE", "soft": "E9F0EF",
    "pale": "F7E7D4", "green": "4C8277",
}

def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)

def shape(slide, kind, x, y, w, h, fill, line=None, line_width=1.0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(line_width)
    else:
        s.line.fill.background()
    return s

def rect(slide, x, y, w, h, fill, line=None, line_width=1.0):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, line_width)

def oval(slide, x, y, w, h, fill, line=None, line_width=1.0):
    return shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill, line, line_width)

def line(slide, x1, y1, x2, y2, color, width=1.2):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = rgb(color); s.line.width = Pt(width)
    return s

def text(slide, x, y, w, h, value, size, color=None, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.015):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.005); tf.margin_bottom = Inches(0.005)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align; p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.0
    run = p.add_run(); run.text = value
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = rgb(color or C["ink"])
    return box

def contain_picture(slide, path, x, y, w, h):
    im = Image.open(path); iw, ih = im.size
    scale = min(w / iw, h / ih); pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w-pw)/2), Inches(y + (h-ph)/2), width=Inches(pw), height=Inches(ph))

def add_bg(slide, color):
    rect(slide, 0, 0, W, H, color)

def footer(slide, n, extra=""):
    line(slide, 0.80, 6.92, 12.45, 6.92, C["line"], 0.8)
    qualifier = "虚构城市 · 服务方案示例｜非真实上线信息"
    if extra:
        qualifier += "｜" + extra
    text(slide, 0.84, 6.98, 10.85, 0.34, qualifier, 16, C["muted"])
    text(slide, 11.90, 6.98, 0.74, 0.34, f"{n:02d}", 18, C["teal"], bold=True, align=PP_ALIGN.RIGHT, margin=0.0)

def title(slide, value, kicker=None):
    if kicker:
        text(slide, 0.84, 0.30, 6.1, 0.35, kicker, 18, C["teal"], bold=True)
    text(slide, 0.84, 0.69 if kicker else 0.45, 11.8, 0.72, value, 36, C["navy"], bold=True)

def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value

def new_slide(prs, bg):
    s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s, bg); return s

def build():
    if not ASSET.exists():
        raise FileNotFoundError(f"Required separate illustration not found: {ASSET}")
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)

    # 01 — clear public-information cover with a separate cutout illustration.
    s = new_slide(prs, C["paper"])
    rect(s, 0.84, 0.84, 0.10, 0.42, C["coral"])
    text(s, 1.10, 0.82, 4.9, 0.42, "公共服务说明 · 示例", 18, C["teal"], bold=True)
    text(s, 0.90, 1.72, 5.95, 1.24, "社区图书馆\n服务升级", 42, C["navy"], bold=True)
    text(s, 0.94, 3.28, 5.60, 0.72, "让附近的人更容易借阅", 28, C["teal"], bold=True)
    text(s, 0.96, 4.24, 5.56, 0.94, "从借阅流程、空间导向到活动通知，\n把服务信息说清楚、放在看得到的地方。", 20, C["ink"])
    rect(s, 0.96, 5.54, 4.65, 0.56, C["aqua"])
    text(s, 1.12, 5.59, 4.35, 0.42, "虚构城市 · 方案概念，不代表真实馆舍", 18, C["navy"], bold=True)
    # A pale backing keeps the provided transparent asset on a quiet light ground.
    oval(s, 7.50, 1.28, 4.70, 4.70, C["white"])
    contain_picture(s, ASSET, 7.38, 1.36, 4.92, 4.88)
    text(s, 8.11, 6.18, 3.42, 0.38, "馆舍插画为独立整体对象", 18, C["muted"], align=PP_ALIGN.CENTER)
    footer(s, 1)
    notes(s, "这是一个面向公众的虚构服务方案示例，城市、馆舍和升级计划均为虚构。封面馆舍图是一张独立透明 PNG 插画，内部建筑、人物与绿植作为整体图片，不是逐项可编辑的矢量形状。实际对外说明应替换成经属地确认的馆址、服务和运营信息。")

    # 02 — a three-step route, not a row of repeated cards.
    s = new_slide(prs, C["white"])
    title(s, "借阅流程拆成三步", "01 / 到馆前后 · 流程提案")
    text(s, 0.88, 1.34, 11.2, 0.48, "每一步都把办理条件放在旁边，减少来回查找。", 20, C["muted"])
    line(s, 1.78, 2.75, 11.45, 2.75, C["line"], 2.2)
    steps = [
        (2.35, "01", "查书", "查看馆藏状态；\n检索入口待属地确认。", C["teal"]),
        (6.65, "02", "预约", "选择可预约资源；\n范围与规则待确认。", C["coral"]),
        (10.95, "03", "取书", "按通知到指定地点；\n凭证与保留期待确认。", C["navy"]),
    ]
    for cx, no, head, body, color in steps:
        oval(s, cx-0.68, 2.07, 1.36, 1.36, color)
        text(s, cx-0.62, 2.13, 1.24, 1.20, no, 22, C["white"], bold=True, align=PP_ALIGN.CENTER)
        text(s, cx-1.35, 3.66, 2.70, 0.56, head, 28, C["navy"], bold=True, align=PP_ALIGN.CENTER)
        text(s, cx-1.50, 4.32, 3.00, 0.90, body, 20, C["muted"], align=PP_ALIGN.CENTER)
    rect(s, 0.88, 5.67, 11.55, 0.80, C["aqua"])
    text(s, 1.08, 5.75, 11.10, 0.58, "待属地确认：检索入口 · 预约范围 · 保留期限 · 取书凭证", 18, C["navy"], bold=True)
    footer(s, 2, "流程为拟议示意")
    notes(s, "本页用“查书—预约—取书”说明一个可能的借阅路径，属于服务流程提案，不表示虚构城市已有网上检索或预约系统。需要由馆方核实实际入口、资源范围、保留期限和取书凭证，再决定线上与线下各自怎么办理。")

    # 03 — native, editable wayfinding map and spatial principles.
    s = new_slide(prs, C["paper"])
    title(s, "让入口、服务台和安静区一眼可辨", "02 / 空间导向 · 概念平面")
    text(s, 0.88, 1.40, 7.20, 0.40, "动线分区示意 · 非真实馆舍平面", 18, C["coral"], bold=True)
    # Map boundary and three editable service zones.
    rect(s, 0.88, 1.94, 7.32, 4.45, C["white"], C["navy"], 1.3)
    rect(s, 1.15, 2.20, 2.73, 1.18, "F5E4BA", C["yellow"], 1.0)
    text(s, 1.30, 2.45, 2.42, 0.64, "儿童共读区", 22, C["navy"], bold=True, align=PP_ALIGN.CENTER)
    rect(s, 4.12, 2.20, 3.78, 1.18, C["aqua"], C["teal"], 1.0)
    text(s, 4.28, 2.45, 3.46, 0.64, "安静阅览区", 22, C["navy"], bold=True, align=PP_ALIGN.CENTER)
    rect(s, 1.15, 4.42, 2.73, 1.45, C["teal"])
    text(s, 1.30, 4.72, 2.42, 0.80, "借还与\n咨询台", 22, C["white"], bold=True, align=PP_ALIGN.CENTER)
    rect(s, 4.12, 4.42, 3.78, 1.45, C["pale"], C["coral"], 1.0)
    text(s, 4.28, 4.72, 3.46, 0.80, "活动 / 共用空间", 22, C["navy"], bold=True, align=PP_ALIGN.CENTER)
    # Continuous circulation band and an explicit entrance marker.
    rect(s, 1.15, 3.55, 6.75, 0.56, C["soft"])
    text(s, 1.34, 3.60, 6.36, 0.44, "入口 → 服务台 → 阅览区域：连续通行路线（提案）", 18, C["teal"], bold=True, align=PP_ALIGN.CENTER)
    rect(s, 0.60, 3.43, 0.28, 0.84, C["coral"])
    text(s, 0.93, 6.48, 7.20, 0.32, "分区名称、通道和面积均为示意", 16, C["muted"])
    # Principles column.
    text(s, 8.66, 2.00, 3.62, 0.56, "导向原则", 24, C["navy"], bold=True)
    principles = [
        (2.78, "先看见服务台", "入口先指向服务台"),
        (4.00, "区分动静空间", "儿童区与安静区分开"),
        (5.22, "保持路径连续", "转角和高差处补上指引"),
    ]
    for y, head, body in principles:
        oval(s, 8.68, y+0.08, 0.38, 0.38, C["coral"])
        text(s, 9.20, y, 3.05, 0.44, head, 20, C["navy"], bold=True)
        text(s, 9.20, y+0.46, 3.05, 0.62, body, 18, C["muted"])
    footer(s, 3, "空间信息待属地核实")
    notes(s, "平面图由 PowerPoint 原生形状和文本组成，可直接改分区名称、颜色与路径。它不是某座真实图书馆的测绘图，也不代表建筑条件已经具备。设计建议强调服务台可见、动静区域区分，以及从入口到服务点的连续通行；真实方案仍要现场走读并核对无障碍条件。")

    # 04 — a public notice mockup with explicit blanks and a separate process.
    s = new_slide(prs, C["white"])
    title(s, "活动通知回答三件事：何时、何地、如何参加", "03 / 活动信息 · 空白模板")
    # The left notice is a layout template, not a fake event.
    rect(s, 0.90, 1.72, 4.34, 4.84, C["aqua"])
    rect(s, 0.90, 1.72, 4.34, 0.16, C["teal"])
    text(s, 1.22, 2.03, 3.70, 0.42, "活动公告模板", 18, C["teal"], bold=True)
    text(s, 1.18, 2.54, 3.76, 0.78, "活动名称待填", 26, C["navy"], bold=True)
    fields = [
        (3.50, "时间", "待属地确认"),
        (4.20, "地点", "待属地确认"),
        (4.90, "参与方式", "待属地确认"),
        (5.60, "对象 / 名额", "待属地确认"),
    ]
    for y, label, value in fields:
        line(s, 1.22, y-0.08, 4.88, y-0.08, C["line"], 0.8)
        text(s, 1.26, y, 1.38, 0.38, label, 18, C["muted"], bold=True)
        text(s, 2.68, y, 2.12, 0.38, value, 20, C["navy"])
    # Right side: sign-up route and safeguards.
    text(s, 5.90, 1.96, 5.70, 0.48, "读者从看到信息到参加活动", 23, C["navy"], bold=True)
    line(s, 6.17, 2.78, 6.17, 5.85, C["line"], 2.0)
    items = [
        (2.68, "看活动内容", "对象、主题与时长"),
        (3.63, "查时间地点", "地址与到场提示"),
        (4.58, "确认报名方式", "线上 / 线下路径待确认"),
        (5.53, "查看提醒规则", "变更与取消方式待确认"),
    ]
    for idx, (y, head, body) in enumerate(items, start=1):
        oval(s, 5.93, y+0.05, 0.48, 0.48, C["teal"] if idx % 2 else C["coral"])
        text(s, 5.96, y+0.05, 0.42, 0.46, str(idx), 18, C["white"], bold=True, align=PP_ALIGN.CENTER)
        text(s, 6.66, y-0.02, 2.78, 0.42, head, 20, C["navy"], bold=True)
        text(s, 9.45, y-0.02, 2.84, 0.46, body, 18, C["muted"])
    footer(s, 4, "活动名称、时段、地点和规则均待确认")
    notes(s, "左侧是面向公众的活动公告字段模板，没有虚构具体活动名称、日期、地点或名额。右侧按阅读顺序呈现应当明确的主题对象、到场位置、报名路径和变更规则。真实页面应先填入经馆方确认的信息，并说明活动变化时读者从哪里获取更新。")

    # 05 — verify local facts before publishing; staged rollout is a proposal.
    s = new_slide(prs, C["paper"])
    title(s, "上线前先把四项本地信息核实清楚", "04 / 发布前核对 · 待属地确认")
    line(s, 1.08, 2.06, 1.08, 5.82, C["line"], 2.0)
    required = [
        (2.02, "场馆地址与服务入口"),
        (2.94, "开放时段与节假日安排"),
        (3.86, "联系电话与反馈渠道"),
        (4.78, "借阅、预约与续借规则"),
    ]
    for y, label in required:
        oval(s, 0.86, y+0.11, 0.44, 0.44, C["white"], C["coral"], 1.5)
        text(s, 1.54, y, 3.56, 0.58, label, 20, C["navy"], bold=True)
        text(s, 5.10, y, 1.56, 0.58, "待属地确认", 18, C["coral"], bold=True, align=PP_ALIGN.RIGHT)
    line(s, 7.10, 1.86, 7.10, 6.24, C["line"], 1.2)
    text(s, 7.62, 1.98, 4.40, 0.56, "建议采用小步试行", 24, C["teal"], bold=True)
    phases = [
        (2.85, "01", "核规则文本"),
        (3.69, "02", "现场走读"),
        (4.53, "03", "小范围试行"),
        (5.37, "04", "收集问题并复盘"),
    ]
    for i, (y, no, label) in enumerate(phases):
        oval(s, 7.66, y, 0.52, 0.52, C["teal"] if i % 2 == 0 else C["coral"])
        text(s, 7.69, y, 0.46, 0.50, no, 18, C["white"], bold=True, align=PP_ALIGN.CENTER)
        text(s, 8.42, y, 3.80, 0.52, label, 20, C["navy"], bold=True)
    text(s, 7.63, 6.08, 4.70, 0.50, "运营安排由属地馆方或主管部门确认", 18, C["muted"])
    footer(s, 5)
    notes(s, "发布前先确认地址、时间、联系方式和借阅规则，避免公众带着错误信息到场。右侧的核规则、现场走读、小范围试行和复盘是建议步骤，不代表当地已经安排了试运行，也没有设置虚构日期、预算或服务效果。馆方或主管部门应决定何时发布、谁负责更新，以及如何收集反馈。")

    prs.save(OUT)
    print(f"WROTE {OUT} ({len(prs.slides)} slides, {prs.slide_width / Inches(1):.3f} x {prs.slide_height / Inches(1):.3f} in)")

if __name__ == "__main__":
    build()
