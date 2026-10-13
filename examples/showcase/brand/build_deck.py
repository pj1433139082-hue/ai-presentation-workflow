from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

ROOT = Path(__file__).resolve().parent
ASSET = ROOT / "assets" / "tea-branch.png"
OUT = ROOT / "deck-草稿.pptx"
FONT = "Microsoft YaHei"
W, H = 13.333333, 7.5
C = {
    "paper": "F6F0E4", "paper2": "EFE2CE", "ink": "25362E", "green": "315547",
    "deep": "203E34", "rust": "B85D40", "clay": "D98E67", "rose": "E8C8B0",
    "sage": "C7D1B8", "line": "C9BBA7", "muted": "645E53", "white": "FFFDF8",
}

def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)

def shape(slide, kind, x, y, w, h, fill, line=None, line_width=1.0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(line_width)
    else:
        s.line.fill.background()
    return s

def rect(slide, x, y, w, h, fill, line=None, line_width=1.0):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, line_width)

def oval(slide, x, y, w, h, fill, line=None, line_width=1.0):
    return shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill, line, line_width)

def line(slide, x1, y1, x2, y2, color, width=1.2):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = rgb(color); s.line.width = Pt(width)
    return s

def text(slide, x, y, w, h, value, size, color=None, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.015):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.005); tf.margin_bottom = Inches(0.005)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align; p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.0
    run = p.add_run(); run.text = value
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = rgb(color or C["ink"])
    return box

def contain_picture(slide, path, x, y, w, h):
    im = Image.open(path); iw, ih = im.size
    scale = min(w / iw, h / ih); pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w-pw)/2), Inches(y + (h-ph)/2), width=Inches(pw), height=Inches(ph))

def add_bg(slide, color):
    rect(slide, 0, 0, W, H, color)

def footer(slide, n):
    line(slide, 0.78, 6.92, 12.45, 6.92, C["line"], 0.8)
    text(slide, 0.82, 6.98, 10.88, 0.34, "虚构品牌示例｜方案演示，不对应真实产品、销量或供应链", 16, C["muted"])
    text(slide, 11.90, 6.98, 0.74, 0.34, f"{n:02d}", 18, C["rust"], bold=True, align=PP_ALIGN.RIGHT, margin=0.0)

def title(slide, value, kicker=None):
    if kicker:
        text(slide, 0.82, 0.30, 5.5, 0.35, kicker, 18, C["rust"], bold=True)
    text(slide, 0.82, 0.69 if kicker else 0.45, 11.85, 0.72, value, 36, C["ink"], bold=True)

def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value

def new_slide(prs, bg):
    s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s, bg); return s

def build():
    if not ASSET.exists():
        raise FileNotFoundError(f"Required separate illustration not found: {ASSET}")
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)

    # 01 — editorial cover: asymmetric type and one separate illustration.
    s = new_slide(prs, C["paper"])
    rect(s, 0.82, 0.78, 0.09, 0.40, C["rust"])
    text(s, 1.06, 0.77, 4.3, 0.42, "消费品牌提案 · 示例", 18, C["rust"], bold=True)
    text(s, 0.88, 1.72, 5.6, 1.02, "山间茶事", 46, C["deep"], bold=True)
    text(s, 0.92, 2.82, 5.55, 0.70, "把山场装进一盒茶", 32, C["rust"], bold=True)
    text(s, 0.94, 3.77, 5.15, 1.05, "从新品包装到日常触点，探索一种轻装、好分享的茶礼表达。", 22, C["ink"])
    line(s, 0.96, 5.27, 5.72, 5.27, C["line"], 1.0)
    text(s, 0.96, 5.42, 5.30, 0.38, "概念 SKU：雾岭春芽", 18, C["muted"])
    text(s, 0.96, 5.82, 5.30, 0.38, "提案示例 / 尚未上市", 18, C["muted"])
    oval(s, 7.18, 1.25, 4.82, 4.82, C["paper2"])
    contain_picture(s, ASSET, 6.66, 1.76, 5.92, 3.78)
    text(s, 8.16, 5.77, 3.35, 0.42, "单幅茶枝与山雾插画组件", 18, C["green"], align=PP_ALIGN.CENTER)
    footer(s, 1)
    notes(s, "本提案以虚构品牌“山间茶事”和虚构 SKU“雾岭春芽”为讨论载体。封面中的茶枝与山雾是一张单独放置的透明 PNG 插画；插画内部不作为可逐叶编辑的矢量对象。后续页面提出的是设计假设，不代表真实品牌、产地或销售情况。")

    # 02 — editorial timeline, three assumed use moments.
    s = new_slide(prs, C["white"])
    title(s, "让一盒茶适配三种日常时刻", "01 / 使用情境 · 设计假设")
    line(s, 1.28, 2.00, 1.28, 5.62, C["line"], 2.0)
    rows = [
        (2.05, "01", "轻装出门", "小份、好携带，让取用动作简单。", C["green"]),
        (3.37, "02", "午后独处", "冲泡信息直观，给阅读留出呼吸感。", C["rust"]),
        (4.69, "03", "顺手赠礼", "开盒有记忆点，祝福卡与外盒一起设计。", C["deep"]),
    ]
    for y, num, head, body, accent in rows:
        oval(s, 0.88, y+0.10, 0.82, 0.82, accent)
        text(s, 0.90, y+0.10, 0.78, 0.82, num, 20, C["white"], bold=True, align=PP_ALIGN.CENTER)
        text(s, 2.06, y, 2.78, 0.94, head, 26, C["ink"], bold=True)
        line(s, 4.80, y+0.18, 4.80, y+0.78, C["rose"], 2.0)
        text(s, 5.14, y+0.08, 6.55, 0.78, body, 22, C["muted"])
    rect(s, 0.88, 5.98, 11.58, 0.62, C["paper2"])
    text(s, 1.08, 6.06, 11.16, 0.45, "这三种时刻是提案假设，需通过真实访谈与试用调整。", 18, C["deep"], bold=True)
    footer(s, 2)
    notes(s, "本页把设计讨论拆成三种使用时刻：轻装携带、午后自饮、登门分享。它们是为了组织包装与触点内容而设定的假设，并非调研结论。真实项目应补充目标人群访谈、场景观察和产品信息，再决定是否保留这三类需求。")

    # 03 — editable package silhouettes rather than a repeated card grid.
    s = new_slide(prs, C["paper"])
    title(s, "一盒茶，拆成三种拿取尺度", "02 / 包装系统 · 结构概念")
    text(s, 0.87, 1.40, 11.3, 0.45, "以“礼赠 / 自用 / 初尝”组织盒型角色，具体规格和材质待确认。", 20, C["muted"])
    # Gift carton: front, spine and top are independent native shapes.
    shape(s, MSO_SHAPE.PARALLELOGRAM, 1.18, 2.45, 2.05, 0.42, C["sage"])
    rect(s, 1.18, 2.83, 2.05, 2.55, "E9DCC6", C["green"], 1.2)
    rect(s, 3.23, 2.91, 0.40, 2.47, "D6C3A4", C["green"], 1.0)
    rect(s, 1.18, 3.40, 2.05, 0.52, C["green"])
    text(s, 1.25, 3.47, 1.92, 0.38, "雾岭春芽", 22, C["white"], bold=True, align=PP_ALIGN.CENTER)
    line(s, 1.48, 4.25, 2.92, 4.25, C["rust"], 1.4)
    text(s, 1.32, 4.44, 1.76, 0.48, "山形纹样位置", 18, C["green"], align=PP_ALIGN.CENTER)
    # Canister: oval lid + body + editable label.
    oval(s, 5.23, 2.70, 2.04, 0.42, C["rust"])
    rect(s, 5.23, 2.91, 2.04, 2.42, "D7DED0", C["green"], 1.2)
    oval(s, 5.23, 5.12, 2.04, 0.42, "C6D0BE", C["green"], 1.0)
    rect(s, 5.42, 3.57, 1.66, 0.96, C["white"])
    text(s, 5.48, 3.68, 1.54, 0.36, "雾岭春芽", 20, C["deep"], bold=True, align=PP_ALIGN.CENTER)
    text(s, 5.52, 4.05, 1.46, 0.30, "净含量待填", 18, C["muted"], align=PP_ALIGN.CENTER)
    # Trial pouch: editable trapezoid with folded top.
    shape(s, MSO_SHAPE.TRAPEZOID, 9.35, 2.95, 2.12, 2.52, "E8C9B1", C["rust"], 1.2)
    rect(s, 9.50, 2.82, 1.82, 0.26, C["rust"])
    line(s, 9.64, 3.75, 11.15, 3.75, C["green"], 1.2)
    text(s, 9.59, 4.00, 1.64, 0.46, "试饮装", 22, C["deep"], bold=True, align=PP_ALIGN.CENTER)
    text(s, 0.96, 5.73, 2.92, 0.42, "礼赠套盒", 22, C["ink"], bold=True, align=PP_ALIGN.CENTER)
    text(s, 4.81, 5.73, 2.91, 0.42, "日常小罐", 22, C["ink"], bold=True, align=PP_ALIGN.CENTER)
    text(s, 8.93, 5.73, 2.95, 0.42, "初次试饮", 22, C["ink"], bold=True, align=PP_ALIGN.CENTER)
    text(s, 0.94, 6.20, 11.3, 0.42, "盒型、标签文字与净含量为视觉占位；需由产品、合规和供应链共同确认。", 18, C["muted"])
    footer(s, 3)
    notes(s, "本页使用 PowerPoint 原生矩形、椭圆和梯形构成三种包装轮廓，因此这些外框和标签文字可以在 PowerPoint 中分别移动与改写。虚构 SKU 名称仅用于展示版式；净含量、材质、茶类、来源以及任何产品承诺都没有在此页设定，真实项目需先确认事实和合规文案。")

    # 04 — three different tactile touchpoints: phone, counter, gift note.
    s = new_slide(prs, C["white"])
    title(s, "从第一眼，到试一口，再到带走", "03 / 品牌触点 · 路径示意")
    # Phone/product-detail page.
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 1.78, 2.42, 3.86, C["deep"])
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.08, 1.98, 2.16, 3.46, C["paper"])
    text(s, 1.26, 2.16, 1.80, 0.38, "产品详情页", 20, C["deep"], bold=True, align=PP_ALIGN.CENTER)
    rect(s, 1.32, 2.68, 1.68, 0.76, C["sage"])
    line(s, 1.48, 3.20, 2.74, 2.88, C["green"], 1.5)
    text(s, 1.27, 3.57, 1.80, 0.38, "风味信息位置", 18, C["muted"], align=PP_ALIGN.CENTER)
    line(s, 1.38, 4.17, 2.90, 4.17, C["line"], 1.2)
    line(s, 1.38, 4.42, 2.76, 4.42, C["line"], 1.2)
    text(s, 1.26, 4.81, 1.80, 0.38, "页面版式示意", 18, C["rust"], align=PP_ALIGN.CENTER)
    # Flow arrows stay as editable text glyphs.
    text(s, 3.65, 3.16, 0.72, 0.70, "→", 32, C["rust"], bold=True, align=PP_ALIGN.CENTER)
    # Sampling counter with native shapes and a tasting cup.
    rect(s, 4.73, 3.35, 3.00, 0.32, C["rust"])
    rect(s, 4.97, 3.67, 0.20, 1.43, C["green"])
    rect(s, 7.27, 3.67, 0.20, 1.43, C["green"])
    oval(s, 5.50, 2.43, 0.76, 0.42, C["paper2"], C["green"], 1.1)
    shape(s, MSO_SHAPE.TRAPEZOID, 5.57, 2.67, 0.62, 0.58, C["white"], C["green"], 1.1)
    oval(s, 6.57, 2.55, 0.47, 0.18, C["sage"])
    text(s, 4.84, 5.20, 2.84, 0.45, "线下试饮点", 22, C["ink"], bold=True, align=PP_ALIGN.CENTER)
    text(s, 4.46, 5.65, 3.60, 0.42, "冲泡与讲解，帮助理解风味", 18, C["muted"], align=PP_ALIGN.CENTER)
    text(s, 8.11, 3.16, 0.72, 0.70, "→", 32, C["rust"], bold=True, align=PP_ALIGN.CENTER)
    # Take-away note / small package.
    shape(s, MSO_SHAPE.PARALLELOGRAM, 9.35, 2.20, 2.50, 2.78, C["paper2"], C["rust"], 1.2)
    rect(s, 9.62, 2.57, 1.96, 1.93, C["white"])
    text(s, 9.80, 2.82, 1.60, 0.36, "留一段", 20, C["deep"], bold=True, align=PP_ALIGN.CENTER)
    text(s, 9.72, 3.20, 1.76, 0.38, "慢下来的时间", 18, C["deep"], bold=True, align=PP_ALIGN.CENTER)
    line(s, 9.90, 3.92, 11.28, 3.92, C["clay"], 1.4)
    text(s, 9.70, 5.20, 1.86, 0.45, "赠礼小卡", 22, C["ink"], bold=True, align=PP_ALIGN.CENTER)
    text(s, 9.02, 5.65, 3.20, 0.42, "不含真实二维码与购买链接", 18, C["muted"], align=PP_ALIGN.CENTER)
    text(s, 0.96, 6.31, 11.25, 0.38, "产品页、试饮台与小卡都是概念版式，并非已上线的销售或预约入口。", 18, C["rust"], bold=True)
    footer(s, 4)
    notes(s, "按从左到右的顺序，页面展示三种提案触点：用于说明产品的详情页、帮助理解风味的线下试饮台，以及便于分享的赠礼卡。手机和卡片内部的线条、标题与插图占位均为可编辑的 PowerPoint 对象。此示例没有真实二维码、店铺、购买入口或活动，不应被理解为已投放的品牌活动。")

    # 05 — phased validation and decisions, no invented targets.
    s = new_slide(prs, C["paper"])
    title(s, "先打样验证，再决定是否进入下一步", "04 / 推进建议 · 待真实委托方确认")
    line(s, 1.25, 2.42, 8.35, 2.42, C["line"], 2.0)
    phases = [
        (1.02, "01", "定信息", "茶品事实与合规文案先确认", C["green"]),
        (3.70, "02", "试结构", "盒型、材料与印刷样需打样", C["rust"]),
        (6.40, "03", "看反馈", "按渠道小范围观察后调整", C["deep"]),
    ]
    for x, n, head, body, accent in phases:
        oval(s, x+0.53, 1.96, 0.90, 0.90, accent)
        text(s, x+0.55, 2.01, 0.86, 0.78, n, 20, C["white"], bold=True, align=PP_ALIGN.CENTER)
        text(s, x, 3.04, 1.98, 0.55, head, 24, C["ink"], bold=True, align=PP_ALIGN.CENTER)
        text(s, x-0.06, 3.72, 2.12, 1.04, body, 20, C["muted"], align=PP_ALIGN.CENTER)
    text(s, 0.98, 5.30, 7.55, 0.72, "执行周期、成本和渠道安排由实际供应商与委托方补齐。", 18, C["muted"])
    # A warm editorial decision panel rather than another card grid.
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.16, 1.67, 3.30, 4.78, C["deep"])
    text(s, 9.48, 1.94, 2.68, 0.62, "打样前四项确认", 23, C["white"], bold=True)
    for i, item in enumerate(["茶品与风味事实", "包装规格与材料", "成本与供货周期", "试用反馈方式"]):
        y = 2.95 + i * 0.72
        oval(s, 9.50, y+0.04, 0.34, 0.34, C["clay"])
        text(s, 9.96, y, 2.15, 0.46, item, 18, C["white"])
    footer(s, 5)
    notes(s, "这页把建议动作放在确认顺序上，而不是承诺销量或投放时间。首先核对商品事实和合规文案，再用实物样确认包装结构，最后根据真实渠道设计试用观察。右侧四项是启动打样前的讨论清单，成本、周期、供应商和评估办法需要由真实委托方提供。")

    prs.save(OUT)
    print(f"WROTE {OUT} ({len(prs.slides)} slides, {prs.slide_width / Inches(1):.3f} x {prs.slide_height / Inches(1):.3f} in)")

if __name__ == "__main__":
    build()
