from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
import csv, json, hashlib

ROOT=Path(__file__).resolve().parent; ASSETS=ROOT/'assets'; OUT=ROOT/'deck-草稿.pptx'
W,H=13.333333,7.5; FONT='Microsoft YaHei'
C={'paper':'F1F3F3','white':'FFFFFF','ink':'182F3B','muted':'50636D','blue':'365C70','blue2':'608397','orange':'F28B2D','pale':'E6ECEE','pale2':'DCE5E8','dark':'142B37','line':'C8D2D6','red':'B84935','green':'3F7868','gray':'7B8A90'}
def rgb(h): return RGBColor.from_string(h)
def set_bg(s,color=C['paper']): s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(color)
def set_ea(r):
    r.font.name=FONT
    try:r._r.get_or_add_rPr().set(qn('a:ea'),FONT)
    except Exception:pass
def text(s,value,x,y,w,h,size=20,color='182F3B',bold=False,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,margin=.02):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=1.08
    r=p.add_run(); r.text=value; set_ea(r); r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    return box
def clear_theme_effect(shape):
 sp=shape._element.spPr
 if sp.find(qn('a:effectLst')) is None: sp.append(OxmlElement('a:effectLst'))
def rect(s,x,y,w,h,fill,line=None,radius=True):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    if line: sh.line.color.rgb=rgb(line); sh.line.width=Pt(1)
    else: sh.line.fill.background()
    clear_theme_effect(sh)
    if radius:
        try:sh.adjustments[0]=.10
        except Exception:pass
    return sh
def ln(s,x1,y1,x2,y2,color='C8D2D6',width=1.2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); sh.line.color.rgb=rgb(color); sh.line.width=Pt(width); return sh
def circle(s,x,y,d,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill); sh.line.fill.background(); clear_theme_effect(sh); return sh
def image_fit(s,p,x,y,w,h):
    with Image.open(p) as im: iw,ih=im.size
    sc=min(w/iw,h/ih); ww,hh=iw*sc,ih*sc
    return s.shapes.add_picture(str(p),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
def notes(s,v):s.notes_slide.notes_text_frame.text=v
def new_slide(prs):s=prs.slides.add_slide(prs.slide_layouts[6]);set_bg(s);return s
def header(s,title,kicker,page):
    text(s,kicker,.62,.34,8.8,.36,18,C['blue'],True)
    text(s,title,.62,.78,11.95,.62,32,C['ink'],True)
    rect(s,10.55,.35,2.18,.43,C['orange'])
    text(s,'示例数据 · 虚构',10.62,.39,2.04,.31,18,C['white'],True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    # Preserve open space below the title; omit the decorative title rule.
    text(s,f'{page:02d} / 05',11.47,7.00,1.23,.34,16,C['muted'],True,PP_ALIGN.RIGHT)
def footer(s):text(s,'示例工厂 A｜虚构经营数据，不是真实经营记录或财务披露',.62,7.00,10.65,.34,16,C['muted'])
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
prs.core_properties.title='制造业季度运营复盘｜本地展示草稿'
prs.core_properties.subject='虚构制造企业运营数据；用于展示汇报结构，不是真实财报'
prs.core_properties.author='本地展示样稿（示例）';prs.core_properties.keywords='展示样稿; 虚构经营数据; 制造业运营复盘'

# 01 cover
s=new_slide(prs);rect(s,0,0,.18,H,C['orange'],radius=False)
text(s,'制造运营 / 季度复盘',.82,.61,6.5,.38,18,C['blue'],True)
text(s,'把运营差距\n拆成可验证的动作',.82,1.30,6.2,1.62,38,C['ink'],True)
text(s,'示例工厂 A · Q3 示例周期\n数据、原因与动作均为虚构样例',.84,3.25,5.9,.88,22,C['muted'])
rect(s,.84,4.68,4.65,.72,C['dark'])
text(s,'本地展示草稿 · 非真实财报',1.05,4.86,4.25,.34,18,C['white'],True)
text(s,'复盘顺序：核口径 → 看差距 → 验改善',.84,5.72,5.8,.60,20,C['ink'],True,margin=0)
image_fit(s,ASSETS/'factory-hall.png',6.85,.60,6.10,5.95)
text(s,'01 / 05',11.47,7.00,1.23,.34,16,C['muted'],True,PP_ALIGN.RIGHT)
notes(s,'本地展示草稿；示例工厂 A、Q3 示例周期及所有经营数据均为虚构，不对应任何真实公司或财务披露。工厂为图像生成的透明 PNG 插图。委托式子智能体制作并由主智能体统筹验收，不代表完整 human G1–G5 生产链认证或人类批准记录。')

# 02 scorecard
s=new_slide(prs);header(s,'五项指标先给出同一张“计划—实际”底表','01 · 经营概览 / 虚构示例数据',2)
rect(s,.68,1.78,11.98,.58,C['dark'])
text(s,'示意周期计划量 13,000 件｜实际量 12,400 件｜计划达成率 95.4%',.94,1.91,11.35,.31,20,C['white'],True)
metrics=[
 (.68,2.63,'准时交付率','94%','计划 96%','差距 −2 个百分点',C['orange']),
 (4.76,2.63,'一次合格率','96.8%','计划 97.5%','差距 −0.7 个百分点',C['blue']),
 (8.84,2.63,'非计划停机率','3.2%','计划 2.5%','差距 +0.7 个百分点',C['red']),
 (2.72,4.47,'单位能耗','8.1 kWh/件','计划 7.8','+3.8%',C['orange']),
 (6.80,4.47,'产量','12,400 件','计划 13,000','−4.6%',C['blue'])]
for x,y,label,value,target,gap,accent in metrics:
    rect(s,x,y,3.80,1.70,C['white'],C['line'])
    rect(s,x,y,.12,1.70,accent,radius=False)
    text(s,label,x+.25,y+.17,3.30,.32,18,C['muted'],True)
    text(s,value,x+.25,y+.52,3.30,.48,27,C['ink'],True)
    text(s,f'{target}\n{gap}',x+.25,y+1.04,3.30,.64,18,C['blue'],margin=0)
text(s,'指标口径、分母与采集系统均为演示占位；先与真实台账对齐再复盘。',.78,6.43,11.6,.38,20,C['ink'],True)
footer(s)
notes(s,'全部数值为编造的样例经营数据，见 data-demo.csv。计算：产量达成率=12400/13000=95.4%；产量差距=(12400-13000)/13000≈-4.6%；准时交付率差距=94%-96%=-2个百分点；一次合格率差距=96.8%-97.5%=-0.7个百分点；非计划停机率差距=3.2%-2.5%=+0.7个百分点；单位能耗相对计划增幅=(8.1-7.8)/7.8≈+3.8%。这些数值仅用于排版演示，不代表真实财务或运营表现。')

# 03 hypothesis, not causes
s=new_slide(prs);header(s,'把“损失构成”标成待核查假设','02 · 偏差拆解 / 虚构样例归因',3)
rect(s,.68,1.80,11.98,.66,'F8E8DF')
text(s,'下图是虚构的排查优先级假设，不是经验证的根因或真实损失数据。',.94,1.97,11.30,.32,20,C['red'],True)
text(s,'假设性损失占比（虚构）',.82,2.72,5.4,.35,20,C['ink'],True)
text(s,'条长按 0–50% 统一比例；数字为虚构占比',2.85,5.91,5.8,.30,16,C['muted'])
# simple editable horizontal bars
hyp=[('换型等待',42,C['orange']),('设备故障',35,C['blue']),('物料供应',23,C['gray'])]
for i,(label,val,color) in enumerate(hyp):
    yy=3.33+i*.82
    text(s,label,.84,yy,1.85,.34,20,C['ink'],True)
    rect(s,2.85,yy+.02,4.58,.37,C['pale2'],radius=False)
    rect(s,2.85,yy+.02,4.58*val/50,.37,color,radius=False)
    text(s,f'{val}%',7.62,yy-.02,.84,.36,20,C['ink'],True)
rect(s,8.70,2.80,3.95,3.20,C['dark'])
text(s,'先验证三件事',9.02,3.08,3.2,.39,22,C['orange'],True)
for i,(n,t) in enumerate([('01','统一停机与换型定义'),('02','核对时间戳与工单'),('03','核对排产、维修与缺料')]):
    yy=3.75+i*.65
    circle(s,9.04,yy,.38,C['blue'])
    text(s,n,9.04,yy,.38,.38,18,C['white'],True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    text(s,t,9.57,yy-.01,2.75,.52,18,C['white'],True,margin=0)
text(s,'若数据不能复现，撤销假设，不将占比写成结论。',.84,6.20,11.6,.39,20,C['ink'],True)
footer(s)
notes(s,'42% / 35% / 23% 为作者编造的样例比例，恰好合计100%，只用于展示可编辑条形图。它们不来自真实工厂、工单或设备数据，也不证明换型、故障、物料是实际根因。检查项是供复盘会议讨论的示例；实际归因必须回到定义清晰、可追溯的停机记录，并允许撤销假设。')

# 04 action plan
s=new_slide(prs);header(s,'按 30 / 60 / 90 天设验证门，不先许诺收益','03 · 行动设计 / 待决策草案',4)
text(s,'以下是可讨论的推进节奏；不是已批准计划、既定责任或效果承诺。',.72,1.78,9.85,.36,20,C['muted'])
image_fit(s,ASSETS/'pump-motor.png',10.83,1.66,1.46,.72)
phases=[
 (.68,'0–30 天','校准数据','对齐停机、交付与能耗口径；抽样核对班次台账。','通过门：口径可复算',C['blue']),
 (4.76,'31–60 天','小范围试点','选一条线测试换型记录与点检闭环；保留对照周。','通过门：数据可追溯',C['orange']),
 (8.84,'61–90 天','复核再决定','比较试点前后波动；检查副作用和样本覆盖。','通过门：复核后再扩围',C['green'])]
for x,period,title,body,gate,color in phases:
    rect(s,x,2.50,3.80,3.16,C['white'],C['line'])
    rect(s,x,2.50,3.80,.16,color,radius=False)
    text(s,period,x+.25,2.83,3.22,.36,20,color,True)
    text(s,title,x+.25,3.34,3.22,.44,24,C['ink'],True)
    text(s,body,x+.25,4.00,3.24,.73,18,C['muted'])
    ln(s,x+.25,4.94,x+3.52,4.94,C['line'])
    text(s,gate,x+.25,5.13,3.22,.34,18,C['ink'],True)
rect(s,.68,6.00,11.98,.61,C['dark'])
text(s,'管理约束：数据可追溯、结果可复核后，再讨论扩大范围。',.94,6.11,11.35,.42,20,C['white'],True,margin=0)
footer(s)
notes(s,'30/60/90天结构是本地样稿提出的规划模板。没有真实项目授权、责任人、预算或目标；不构成已批准经营计划。任何试点需明确基线、对照/比较口径、观察窗口、质量与安全约束；只有数据可复算、效果可重复且未出现不可接受副作用后，才由实际管理者决策是否推广。')

# 05 close / ask for a decision
s=new_slide(prs);header(s,'下次复盘先问：数据能否支持这项决定？','04 · 复盘闭环 / 决策清单模板',5)
rect(s,.68,1.84,7.04,4.74,C['dark'])
text(s,'建议会议带走三项证据',1.00,2.16,6.2,.44,24,C['white'],True)
items=[
 ('A','指标口径表','目标、分母、时间窗、数据责任系统'),
 ('B','异常记录样本','工单 / 班次时间戳可追溯到原始记录'),
 ('C','试点复核页','比较条件、样本覆盖、质量与安全护栏')]
for i,(letter,title,desc) in enumerate(items):
    yy=2.97+i*1.02
    circle(s,1.02,yy,.48,C['orange'])
    text(s,letter,1.02,yy,.48,.48,18,C['white'],True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    text(s,title,1.72,yy-.02,2.35,.36,20,C['white'],True)
    text(s,desc,4.10,yy-.02,3.12,.65,18,'D7E1E4')
    if i<2: ln(s,1.02,yy+.68,7.34,yy+.68,'39505B',1)
text(s,'讨论后决策',8.30,2.03,3.9,.42,22,C['blue'],True)
rect(s,8.28,2.73,4.30,1.20,C['white'],C['line'])
text(s,'选一个试点范围',8.56,2.98,3.75,.35,22,C['ink'],True)
rect(s,8.28,4.15,4.30,1.20,C['white'],C['line'])
text(s,'确认指标口径与责任人',8.56,4.40,3.75,.52,20,C['ink'],True)
text(s,'未核验前：保持“待验证”',8.32,5.75,4.2,.40,20,C['red'],True)
footer(s)
notes(s,'本页为虚构经营复盘的闭环模板，列出会议可以要求的证据与决策项。并非实际公司已召开会议、分派责任或批准试点。商业真实信息来源：无；本展示所有数值均为编造示意。')

prs.save(OUT)
rows=[
 ['产量','件',13000,12400,'计划/实际均为虚构样例','实际/计划=95.4%; 差距=-4.6%'],
 ['准时交付率','%',96,94,'虚构样例','差距=-2个百分点'],
 ['一次合格率','%',97.5,96.8,'虚构样例','差距=-0.7个百分点'],
 ['非计划停机率','%',2.5,3.2,'虚构样例','差距=+0.7个百分点'],
 ['单位能耗','kWh/件',7.8,8.1,'虚构样例','相对计划增幅=+3.8%']]
with (ROOT/'data-demo.csv').open('w',newline='',encoding='utf-8-sig') as f:
 w=csv.writer(f); w.writerow(['metric','unit','plan_demo_only','actual_demo_only','provenance','calculation_or_status']); w.writerows(rows)
with (ROOT/'hypothesis-demo.csv').open('w',newline='',encoding='utf-8-sig') as f:
 w=csv.writer(f); w.writerow(['candidate_category','share_percent_demo_only','provenance']);
 for name,val in [('换型等待',42),('设备故障',35),('物料供应',23)]: w.writerow([name,val,'虚构假设，仅作排版样例；非根因结论'])
manifest={'deck':OUT.name,'slide_count':len(prs.slides),'aspect_ratio':'16:9','status':'本地展示样稿 / 草稿','editability':'所有文本、KPI卡、条形图和流程卡均为独立 PowerPoint 原生形状；图片为透明 PNG 单体素材。','data_status':'所有计划值、实际值和归因比例均为编造样例，不是真实经营记录或财报。','assets':[]}
for p in sorted(ASSETS.glob('*.png')):manifest['assets'].append({'file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(ROOT/'build-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'created {OUT} ({OUT.stat().st_size} bytes, {len(prs.slides)} slides)')





