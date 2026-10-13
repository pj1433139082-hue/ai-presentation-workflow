from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "deck-草稿.pptx"
FONT = "Microsoft YaHei UI"
W, H = 13.3333, 7.5

C = {
    "bg": "0A111A", "panel": "111D29", "panel2": "152636", "line": "294255",
    "white": "EFF6F9", "muted": "A2B4C1", "cyan": "63D7E8", "mint": "A2E6D4",
    "blue": "3F9DB3", "soft": "D3E3EA", "black": "071018"
}

def rgb(hexstr):
    return RGBColor.from_string(hexstr)

def set_font(run, size, color, bold=False, family=FONT):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            rpr.append(node)
        node.set("typeface", family)

def no_shadow(shape):
    sp_pr = getattr(shape._element, "spPr", None)
    if sp_pr is not None and sp_pr.find(qn("a:effectLst")) is None:
        sp_pr.append(OxmlElement("a:effectLst"))


def add_text(slide, x, y, w, h, text, size=20, color=None, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    no_shadow(box)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.015)
    tf.margin_right = Inches(0.015)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    run = p.add_run()
    run.text = text
    set_font(run, size, color or C["white"], bold, font)
    return box

def rect(slide, x, y, w, h, fill, line=None, radius=False, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.fill.transparency = transparency
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    no_shadow(sh)
    return sh

def oval(slide, x, y, w, h, fill=None, line=None, line_width=1.2, transparency=0):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        sh.fill.transparency = transparency
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_width)
    else:
        sh.line.fill.background()
    no_shadow(sh)
    return sh

def line(slide, x1, y1, x2, y2, color=None, width=1.2):
    sh = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    sh.line.color.rgb = rgb(color or C["line"])
    sh.line.width = Pt(width)
    no_shadow(sh)
    return sh

def bg(slide, color=C["bg"]):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)

def add_image(slide, name, x, y, w, h):
    path = ROOT / "assets" / name
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    fit_w, fit_h = iw * scale, ih * scale
    left, top = x + (w - fit_w) / 2, y + (h - fit_h) / 2
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(fit_w), Inches(fit_h))
    no_shadow(pic)
    return pic

def footer(slide, index, note, color=None):
    col = color or C["muted"]
    line(slide, .72, 6.92, 12.62, 6.92, C["line"], .65)
    add_text(slide, .75, 7.02, 11.25, .34, note, 16, col)
    add_text(slide, 12.12, 7.00, .45, .32, f"{index:02d}", 16, C["cyan"], align=PP_ALIGN.RIGHT)

def header(slide, title, section, sub=None):
    add_text(slide, .76, .42, 4.8, .28, section.upper(), 16, C["cyan"], bold=True)
    add_text(slide, .76, .82, 11.7, .62, title, 32, C["white"], bold=True)
    if sub:
        add_text(slide, .78, 1.53, 11.1, .42, sub, 18, C["muted"])

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

# 01 — typographic cover and generated hero object
s = prs.slides.add_slide(blank); bg(s)
oval(s, 8.17, .67, 4.45, 4.45, None, "173042", 1.1)
oval(s, 8.65, 1.15, 3.5, 3.5, None, "1E4052", .8)
add_text(s, .82, .62, 4.6, .35, "产品概念展示   /   01", 18, C["cyan"], bold=True)
add_text(s, .82, 1.42, 5.5, .72, "讯桥", 32, C["white"], bold=True)
add_text(s, .82, 2.18, 5.85, 1.35, "让知识\n有迹可循", 32, C["white"], bold=True)
rect(s, .84, 3.86, .66, .055, C["cyan"])
add_text(s, .82, 4.18, 5.85, .84, "把团队知识，变成可追问的线索", 22, C["soft"])
add_text(s, .84, 5.29, 5.6, .54, "检索概念 · 引用可查 · 人来判断", 18, C["mint"])
add_image(s, "knowledge-bridge.png", 7.52, 1.0, 4.7, 4.7)
footer(s, 1, "虚构展示方案｜功能、流程与示例均为概念稿，不代表现有产品能力")
notes(s, "开场：讯桥是一个虚构的团队知识检索产品概念。本套展示稿只讨论可以如何设计一类工作流，不声称产品已经实现，也不含真实效果或投资回报数据。")

# 02 — scattered questions converge on the need for context
s = prs.slides.add_slide(blank); bg(s)
header(s, "团队难找的，常常不是答案", "01 / THE TENSION",
       "更难的是找回当时的来源、版本与决策背景。")
# large open field and an editorial hub, not repeated cards
oval(s, 5.22, 2.16, 2.62, 2.62, "102432", "2C5968", 1.2)
oval(s, 5.57, 2.51, 1.92, 1.92, None, "63D7E8", 1.4)
add_text(s, 5.54, 2.82, 1.98, .65, "上下文", 24, C["white"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
add_text(s, 5.58, 3.48, 1.9, .36, "从哪里来？", 18, C["cyan"], False, PP_ALIGN.CENTER)
# dispersed question fragments
add_text(s, .95, 2.17, 3.85, .58, "「规范现行版是哪份？」", 20, C["white"], True)
add_text(s, 8.76, 2.1, 3.72, .62, "「为何调整交付节奏？」", 20, C["white"], True)
add_text(s, .98, 4.68, 3.52, .56, "「客户约束记在哪？」", 20, C["white"], True)
add_text(s, 9.0, 4.55, 3.1, .88, "资料分散于\n多个项目空间", 20, C["soft"])
line(s, 4.55, 2.74, 5.28, 3.03, "3D6C7B", 1.2)
line(s, 8.0, 3.0, 8.82, 2.72, "3D6C7B", 1.2)
line(s, 4.55, 4.82, 5.45, 4.25, "3D6C7B", 1.2)
line(s, 7.64, 4.18, 9.0, 4.74, "3D6C7B", 1.2)
add_text(s, .98, 5.78, 10.9, .48, "示意情境：找答案之前，先找回“这句话从哪里来”。", 20, C["mint"])
footer(s, 2, "问题语句均为虚构情境，不代表用户访谈或产品实测结论")
notes(s, "本页的四个问题是为展示稿编写的虚构例句，不是访谈摘录或统计数据。重点是知识检索需要保留来源与上下文，而不是只给出一个孤立答案。")

# 03 — editable product concept UI
s = prs.slides.add_slide(blank); bg(s)
header(s, "从问题走到原文，回答才有落点", "02 / THE CONCEPT",
       "把检索结果做成可核对的线索，而不是封闭的结论。")
add_text(s, .84, 2.05, 3.55, .37, "概念流程", 18, C["cyan"], True)
# left sequence
steps = [
    ("01", "提出问题", "用自然语言描述要找的决定"),
    ("02", "定位资料", "限定项目与资料范围"),
    ("03", "展示出处", "保留文件名、段落与日期"),
    ("04", "回到原文", "由团队成员确认是否适用")
]
for i, (n, title, desc) in enumerate(steps):
    yy = 2.58 + i * .86
    oval(s, .88, yy, .46, .46, "123041", "3D899B", .9)
    add_text(s, .89, yy+.02, .44, .36, n, 16, C["cyan"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    titles = ["提出问题", "定位资料", "展示出处", "回到原文"]
    descriptions = ["说清想找的决定", "限定项目和资料范围", "保留文件、段落与日期", "由团队成员核对适用性"]
    add_text(s, 1.58, yy-.01, 3.55, .34, titles[i], 20, C["white"], True)
    add_text(s, 1.58, yy+.34, 3.55, .35, descriptions[i], 18, C["muted"])
    if i < 3:
        line(s, 1.11, yy+.47, 1.11, yy+.81, "315565", .9)
# right-side editable mock screen, marked clearly as illustrative
rect(s, 5.28, 2.05, 7.16, 4.22, "101C28", "2A4354", True)
add_text(s, 5.62, 2.28, 2.7, .3, "概念界面 / 搜索示意", 16, C["muted"], True)
rect(s, 5.62, 2.78, 6.45, .7, "172B3A", "31566B", True)
add_text(s, 5.92, 2.93, 5.7, .38, "我们为什么选这个发布窗口？", 20, C["white"])
oval(s, 11.34, 2.91, .4, .4, C["cyan"])
add_text(s, 5.68, 3.72, 1.62, .3, "相关线索  01", 18, C["cyan"], True)
add_text(s, 5.68, 4.14, 5.95, .52, "项目复盘 · 决策纪要 · 2025.04", 20, C["white"], True)
line(s, 5.7, 4.79, 11.95, 4.79, "345064", .8)
add_text(s, 5.68, 4.94, 5.85, .65, "“发布日期需避开客户活动；上线窗口由项目负责人复核。”", 18, C["soft"])
add_text(s, 5.68, 5.72, 3.1, .33, "查看原文  →", 18, C["mint"], True)
add_text(s, 9.48, 5.71, 2.48, .32, "资料范围：项目空间", 16, C["muted"])
add_image(s, "search-prism.png", 10.75, .96, 1.45, 1.45)
footer(s, 3, "概念界面与引用句均为虚构示例，不代表已有系统或能力")
notes(s, "右侧界面完全由PowerPoint原生形状与文字组成，便于编辑；画面内的项目名、日期与引文为虚构示例，不是实际检索结果。演示时可强调每条线索都应回到原始材料核对。")

# 04 — flow as a horizontal pathway with a visible human decision point
s = prs.slides.add_slide(blank); bg(s)
header(s, "检索给出线索，判断仍由团队完成", "03 / THE FLOW",
       "一条有用的路径，需要把范围、引用与人工核对连在一起。")
line(s, 1.28, 3.18, 12.06, 3.18, "315161", 2.0)
flow = [
    (1.35, "01", "限定范围", "项目 / 时间 / 权限"),
    (4.15, "02", "匹配片段", "保留命中上下文"),
    (6.98, "03", "标出来源", "文件 / 段落 / 日期"),
    (9.82, "04", "人工核对", "回到原文再判断")
]
for i, (x, no, title, detail) in enumerate(flow):
    oval(s, x, 2.82, .72, .72, "102432", C["cyan"] if i==3 else "3B7183", 1.2)
    add_text(s, x+.02, 2.98, .68, .34, no, 16, C["cyan"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(s, x-.36, 3.85, 1.55, .42, title, 20, C["white"], True, PP_ALIGN.CENTER)
    add_text(s, x-.62, 4.39, 2.05, .55, detail, 18, C["muted"], False, PP_ALIGN.CENTER)
    if i == 3:
        oval(s, x-.15, 2.67, 1.02, 1.02, None, "6FE4D0", .8)
# boundary statement, text-only, not another card
rect(s, .84, 5.46, .07, .76, C["cyan"])
add_text(s, 1.12, 5.45, 10.9, .54, "上线前需要验证：资料权限是否正确？旧版本是否可辨？引用能否回到原文？", 20, C["soft"])
footer(s, 4, "流程为产品设计假设，权限与来源能力尚未验证")
notes(s, "该流程是概念上的设计路径，不代表当前系统已经实现。权限、版本处理与引用可追溯性都是需要在真实试点中专门验证的事项；人仍需根据原文和具体工作判断。")

# 05 — pilot proposal and measured questions, no outcome claims
s = prs.slides.add_slide(blank); bg(s)
header(s, "先用一小组问题，设计可验证的试点", "04 / THE PILOT",
       "先看边界是否清楚，再讨论效率或规模。")
# three-stage open timeline with different visual weights
line(s, 1.12, 3.23, 9.26, 3.23, "315161", 1.7)
stages = [
    (1.08, "范围", "一个项目组", "限定资料目录，\n确认访问人群"),
    (3.92, "观察", "一批真实问题", "记录来源命中、\n引用复核与误差"),
    (6.78, "复盘", "共同决定下一步", "核对权限与版本，\n修正后再决定扩展")
]
for i, (x, label, title, desc) in enumerate(stages):
    oval(s, x, 2.92, .62, .62, C["cyan"] if i == 0 else "143041", "63D7E8", 1.0)
    add_text(s, x-.03, 3.06, .68, .32, str(i+1).zfill(2), 16, C["black"] if i==0 else C["cyan"], True, PP_ALIGN.CENTER)
    add_text(s, x-.05, 2.22, 1.4, .38, label, 18, C["mint"], True)
    add_text(s, x-.05, 3.83, 2.44, .39, title, 20, C["white"], True)
    add_text(s, x-.05, 4.37, 2.5, .86, desc, 18, C["muted"])
# close with bright concept object
add_image(s, "search-prism.png", 9.28, 2.1, 2.5, 2.5)
add_text(s, 9.18, 4.85, 2.9, .58, "建议验证项\n非效果承诺", 18, C["cyan"], True, PP_ALIGN.CENTER)
add_text(s, .95, 5.78, 11.0, .44, "讯桥｜把答案带回证据里", 22, C["white"], True)
footer(s, 5, "虚构方案｜试点范围、指标与成效均待实际设计和验证")
notes(s, "结尾强调这是试点建议而非承诺结果。建议先用一组事先确认有权访问的资料与问题，记录来源命中、引用复核、权限边界和版本差错；在真实评估之前不应对效率提升或投资回报作判断。")

prs.save(OUT)
print(f"Created {OUT}")
print(f"Slides: {len(prs.slides)}")

