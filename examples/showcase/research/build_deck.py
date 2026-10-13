from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from PIL import Image
import csv, json, hashlib

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
OUT = ROOT / 'deck-草稿.pptx'
W, H = 13.333333, 7.5
FONT = 'Microsoft YaHei'
C = {
 'paper':'F5F7F4','white':'FFFFFF','ink':'102C36','muted':'45616A','teal':'0C7C82',
 'teal2':'69B9B1','pale':'DFEFEB','pale2':'EAF3F0','gold':'D3A84B','red':'AD4D3D',
 'line':'CBD9D5','deep':'0B2630','blue':'31758B','grid':'DCE6E2'
}

def rgb(h): return RGBColor.from_string(h)
def set_bg(slide, color=C['paper']):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(color)
def set_ea(run):
    run.font.name = FONT
    try: run._r.get_or_add_rPr().set(qn('a:ea'), FONT)
    except Exception: pass

def text(slide, value, x, y, w, h, size=20, color='102C36', bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.02, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    p.space_before = Pt(0); p.space_after = Pt(0)
    p.line_spacing = 1.08
    r = p.add_run(); r.text = value; set_ea(r)
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
    return box

def clear_theme_effect(shape):
    sp = shape._element.spPr
    if sp.find(qn('a:effectLst')) is None:
        sp.append(OxmlElement('a:effectLst'))

def rect(slide, x,y,w,h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(1)
    clear_theme_effect(shape)
    if radius:
        try: shape.adjustments[0] = 0.12
        except Exception: pass
    return shape

def line(slide,x1,y1,x2,y2,color='CBD9D5',width=1.2):
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb=rgb(color); s.line.width=Pt(width); return s

def circle(slide,x,y,d,fill):
    s=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb=rgb(fill); s.line.fill.background(); clear_theme_effect(s); return s

def image_contain(slide,path,x,y,w,h):
    with Image.open(path) as im: iw,ih=im.size
    scale=min(w/iw,h/ih); ww,hh=iw*scale,ih*scale
    return slide.shapes.add_picture(str(path), Inches(x+(w-ww)/2), Inches(y+(h-hh)/2), width=Inches(ww), height=Inches(hh))

def add_notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body

def header(slide, title, kicker, page):
    text(slide, kicker, .62,.36,11.4,.34,18,C['teal'],True)
    text(slide, title, .62,.78,12.05,.62,32,C['ink'],True)
    # Leave open space below the title; no decorative title rule.
    text(slide, f'{page:02d} / 05',11.47,7.00,1.23,.34,16,C['muted'],True,PP_ALIGN.RIGHT)

def footer(slide, source):
    text(slide, source, .62,7.00,10.65,.34,16,C['muted'])

def make_slide(prs):
    s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); return s

prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
prs.core_properties.title='烧杯试验设计草案｜本地展示样稿'
prs.core_properties.subject='水处理混凝实验方案示例；不含实测研究结论'
prs.core_properties.author='本地展示样稿（示例）'
prs.core_properties.keywords='展示样稿; 虚构示意数据; 水处理实验设计'

# 01 — cover
s=make_slide(prs)
rect(s,0,0,.18,H,C['teal'],radius=False)
text(s,'研究设计 / 方法方案',.82,.63,5.8,.38,18,C['teal'],True)
text(s,'烧杯试验\n设计草案',.82,1.33,6.4,1.62,42,C['ink'],True)
text(s,'把剂量与 pH 拆成可比较的条件，\n让每一步都能复核。',.84,3.25,6.4,1.05,24,C['muted'])
rect(s,.84,4.82,4.85,.70,C['deep'],radius=True)
text(s,'方案演示 · 不代表已开展实验',1.05,4.98,4.45,.34,18,C['white'],True)
text(s,'水处理实验设计｜本地展示样稿',.84,6.77,6.7,.34,18,C['muted'])
image_contain(s,ASSETS/'jar-test-beaker.png',7.4,.42,5.4,6.45)
text(s,'01 / 05',11.47,7.00,1.23,.34,16,C['muted'],True,PP_ALIGN.RIGHT)
add_notes(s,'展示定位：独立本地示例，不是已完成研究或完整 human G1–G5 生产链认证。内容为实验方案草图，不给出真实实验结论。烧杯图为本地生成的透明 PNG 单体素材。')

# 02 — question and variables
s=make_slide(prs); header(s,'先把问题变成可比较的变量','01 · 研究问题与观测定义',2)
rect(s,.68,1.82,11.98,.96,C['deep'])
text(s,'问题',.95,2.02,.85,.34,18,C['teal2'],True)
text(s,'在同一批原水条件下，混凝剂剂量（及预设 pH）如何改变沉后水浊度？',1.86,1.99,10.35,.48,23,C['white'],True)
# three balanced cards
cards=[
    (.68,'操纵因素','剂量 D1–D5\n可选 pH P1–P3','先锁定水平，\n再随机安排条件'),
    (4.76,'主要响应','上清浊度为主指标\n记录沉前浊度与 pH','固定仪器、位置与时点，\n各组按同一规程操作'),
    (8.84,'控制条件','统一原水与容器\n固定混合/静置/取样','重复数预先确定，\n原始记录同步留存')]
for x,k,title,desc in cards:
    rect(s,x,3.12,3.80,2.50,C['white'],C['line'])
    circle(s,x+.22,3.38,.18,C['teal'])
    text(s,k,x+.50,3.28,2.9,.36,18,C['teal'],True)
    text(s,title,x+.24,3.82,3.34,.80,22,C['ink'],True)
    line(s,x+.24,4.72,x+3.52,4.72,C['line'],1)
    text(s,desc,x+.24,4.90,3.34,.64,18,C['muted'])
rect(s,.68,5.91,11.98,.70,C['pale'])
text(s,'实验参数须按原水特征、药剂与设备条件确定；本页不提供通用剂量或通用操作时间。',.94,6.08,11.35,.34,18,C['ink'],True)
footer(s,'方法背景：EPA Ireland（2002）；California SWRCB（2020）｜具体程序须本地化')
add_notes(s,'方法定位参考 EPA Ireland 的饮用水处理管理手册以及 California State Water Resources Control Board 的 Jar Testing Made Easy 指南。这里只借鉴 jar test 作为比较混凝条件的实验框架，不复制成通用 SOP。剂量区间、pH 水平、重复数、搅拌/沉降程序均须结合原水、药剂和设备，在实施前写入本地方案。')

# 03 — design matrix
s=make_slide(prs); header(s,'把条件表先画出来，再决定随机顺序','02 · 因素组合与记录结构',3)
text(s,'每格代表一个待执行条件；格内 D / P 编码仅用于说明结构。',.72,1.78,9.5,.38,18,C['muted'])
# table frame: 5 columns x3 rows
x0=1.48; y0=2.55; cw=1.55; ch=.74
text(s,'pH ↓  /  剂量 →',.72,2.10,2.9,.34,18,C['teal'],True)
for j in range(5):
    rect(s,x0+j*cw,y0-.47,cw-.08,.38,C['deep'],radius=True)
    text(s,f'D{j+1}',x0+j*cw,y0-.42,cw-.08,.27,18,C['white'],True,PP_ALIGN.CENTER)
for i in range(3):
    rect(s,.72,y0+i*ch,.61,ch-.08,C['teal'],radius=True)
    text(s,f'P{i+1}',.76,y0+i*ch+.14,.53,.30,18,C['white'],True,PP_ALIGN.CENTER)
    for j in range(5):
        fill=C['pale'] if (i+j)%2==0 else C['white']
        rect(s,x0+j*cw,y0+i*ch,cw-.08,ch-.08,fill,C['line'],True)
        text(s,f'P{i+1} · D{j+1}',x0+j*cw,y0+i*ch+.19,cw-.08,.27,18,C['ink'],True,PP_ALIGN.CENTER)
# control row
rect(s,.72,4.92,8.39,.62,'F3E9D1')
text(s,'空白/零剂量对照：按项目定义并纳入同一记录表',.94,5.09,7.95,.30,18,C['ink'],True)
# design side panel
rect(s,9.58,2.05,3.08,3.62,C['deep'])
text(s,'实施前锁定',9.88,2.35,2.5,.36,20,C['teal2'],True)
for i,(n,txt) in enumerate([('A','因素水平'),('B','重复数'),('C','随机顺序'),('D','排除规则')]):
    yy=2.94+i*.58
    circle(s,9.90,yy,.34,C['teal'])
    text(s,n,9.90,yy,.34,.34,18,C['white'],True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    text(s,txt,10.38,yy-.005,2.0,.33,18,C['white'],True)
text(s,'本矩阵是版式示意；条件数与重复数须由真实方案确定。',.78,5.94,11.8,.44,20,C['red'],True)
footer(s,'结构参考 jar-test 条件比较方法；不含真实剂量或已验证参数')
add_notes(s,'矩阵是虚构版式示例，P1–P3 与 D1–D5 并非推荐水平或实测值。需依据水源水质、混凝剂类型、设备容量、研究问题和统计计划确定剂量/pH 水平、对照设置、重复数、随机化和数据排除规则。条件运行顺序应留痕。')

# 04 — workflow
s=make_slide(prs); header(s,'用同一套记录链，把差异留给实验因素','03 · 执行流程与质控留痕',4)
steps=[
 ('01','取样与编码','记录原水批次、时间、初始浊度与 pH。'),
 ('02','配液与加药','核对药剂批号、浓度、剂量；逐杯记录。'),
 ('03','混合与沉降','按本地程序执行，保持组间条件一致。'),
 ('04','取样与读数','固定取样位置与时点，记录浊度、pH及异常。')]
for i,(n,t,d) in enumerate(steps):
    yy=1.90+i*1.12
    circle(s,.78,yy,.55,C['teal'] if i<3 else C['gold'])
    text(s,n,.78,yy,.55,.55,18,C['white'],True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    text(s,t,1.52,yy-.01,3.0,.38,22,C['ink'],True)
    text(s,d,1.52,yy+.43,5.25,.50,18,C['muted'])
    if i<3: line(s,1.05,yy+.56,1.05,yy+1.04,C['line'],2)
rect(s,7.27,1.88,5.28,4.55,C['pale2'])
image_contain(s,ASSETS/'floc-cluster.png',7.62,2.02,4.55,2.62)
text(s,'记录异常，而不是事后删点',7.64,4.92,4.54,.36,20,C['teal'],True)
text(s,'记录絮体、加药与读数异常。\n保留原始记录，按预设规则\n处理偏差。',7.64,5.42,4.40,.95,18,C['ink'])
footer(s,'示意流程；混合强度、持续时间、沉降时间等按本地方法确认')
add_notes(s,'图为生成的絮体透明 PNG 单体插图，仅作概念示意。步骤说明为方案级流程，不是经验证的操作程序。实际采样、快速/慢速混合、静置时长、测量方法、仪器校准和安全要求必须依适用标准、设备说明与机构 SOP 确定。任何偏差应记录并按预先规定处理。')

# 05 — fictional data visualization
s=make_slide(prs); header(s,'用示意图说明读数与局限','04 · 虚构数据呈现样例',5)
rect(s,.68,1.78,11.98,.62,'F8E8E4')
text(s,'示意数据（虚构）｜仅用于版式演示',.94,1.93,11.40,.32,20,C['red'],True)
text(s,'沉后浊度（NTU）',.82,2.60,4.2,.35,20,C['ink'],True)
# chart axes and bars
cx,cy,cw,ch=.98,3.05,7.15,2.90
for i,val in enumerate([20,15,10,5,0]):
    yy=cy + (20-val)/20*ch
    line(s,cx,yy,cx+cw,yy,C['grid'],1)
    text(s,str(val),.40,yy-.20,.52,.40,18,C['muted'],False,PP_ALIGN.RIGHT,MSO_ANCHOR.MIDDLE,0)
line(s,cx,cy,cx,cy+ch,C['muted'],1.3)
vals=[16,11,8,9,12]
bar_w=.72; gap=.58
for i,v in enumerate(vals):
    bx=cx+.55+i*(bar_w+gap); bh=ch*v/20; by=cy+ch-bh
    rect(s,bx,by,bar_w,bh,C['teal'] if i!=2 else C['gold'],radius=False)
    text(s,str(v),bx-.08,by-.36,bar_w+.16,.27,18,C['ink'],True,PP_ALIGN.CENTER)
    text(s,f'D{i+1}',bx-.08,cy+ch+.12,bar_w+.16,.30,18,C['ink'],True,PP_ALIGN.CENTER)
# right interpretation panel
rect(s,8.60,2.62,4.05,3.45,C['deep'])
text(s,'读图提醒',8.93,2.92,3.3,.40,22,C['teal2'],True)
text(s,'无重复与误差范围，\n也没有原水背景。',8.93,3.55,3.35,.70,20,C['white'],True)
line(s,8.93,4.45,12.30,4.45,'41616B',1)
text(s,'真实汇报需补齐：\n方法、重复数与离散度；\n原始记录及偏差说明。',8.93,4.70,3.30,1.05,18,C['white'])
text(s,'数据文件：data-demo.csv（仅供版式演示）',.82,6.46,8.8,.34,18,C['muted'])
footer(s,'数值由样稿作者编造，仅验证图表布局；无实验采集或外部数据库来源')
add_notes(s,'图中 D1–D5 的 NTU 数值 16、11、8、9、12 均为编造的版式演示数据，非实验结果，不支持剂量优选、因果或性能结论。真实研究须记录原水背景、药剂与方法、随机化、重复数、测量误差/离散度和异常处理；本样稿刻意不展示虚构结论。')

# speaker notes and save
prs.save(OUT)
# supporting demo data, all values explicitly fictional
with (ROOT/'data-demo.csv').open('w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f); w.writerow(['condition_code','settled_turbidity_NTU','provenance'])
    for i,v in enumerate([16,11,8,9,12],1): w.writerow([f'D{i}',v,'虚构示意数据；非实测；仅用于图表版式'])
# compact build manifest for reproducibility
manifest={
 'deck':OUT.name,'slide_count':len(prs.slides),'aspect_ratio':'16:9','status':'本地展示样稿 / 草稿',
 'editable_elements':'标题、段落、卡片、流程、矩阵、条形图均为独立 PowerPoint 原生对象；插图为透明 PNG 单体素材。',
 'data_status':'第5页及 data-demo.csv 全部为编造示意数据，不是研究结论。',
 'image_assets':[]
}
for p in sorted(ASSETS.glob('*.png')):
    manifest['image_assets'].append({'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size})
(ROOT/'build-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'created {OUT} ({OUT.stat().st_size} bytes, {len(prs.slides)} slides)')


