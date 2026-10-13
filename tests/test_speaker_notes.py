"""Native-note round trips and adversarial edits on synthetic, gated fixtures."""

import copy
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

from test_assembly import pptx_ready_fixture
from test_merge_page_inventory import inventory_fixture
from test_validate_project import ProjectValidator, gate_decision_values, value_sha256, write_json
from assemble_editable_pptx import assemble, speaker_notes_text


def set_fixture_notes(root: Path, text, *, required: bool = False) -> None:
    """Rebind synthetic human records only; never run on a real project."""
    deck = json.loads((root / "deck-spec.json").read_text(encoding="utf-8"))
    deck["slides"][0]["speaker_notes"] = text
    write_json(root, "deck-spec.json", deck)
    if required:
        brief = json.loads((root / "brief.json").read_text(encoding="utf-8"))
        brief["delivery"]["speaker_notes"] = "required"
        write_json(root, "brief.json", brief)
        alignment = json.loads((root / "alignment.json").read_text(encoding="utf-8"))
        for source in alignment["material_review"]["inspected_sources"]:
            if source["path"] == "brief.json":
                source["sha256"] = hashlib.sha256((root / "brief.json").read_bytes()).hexdigest()
        write_json(root, "alignment.json", alignment)
    state = json.loads((root / "run-state.json").read_text(encoding="utf-8"))
    for gate in state["gates"].values():
        for name in ("deck-spec.json", "brief.json", "alignment.json"):
            if name in gate.get("approved_artifacts", {}):
                gate["approved_artifacts"][name] = hashlib.sha256((root / name).read_bytes()).hexdigest()
    write_json(root, "run-state.json", state)


def inspect_object_contract(root: Path, output: Path) -> list[str]:
    validator = ProjectValidator(root, "assembly")
    validator.load_required_documents()
    validator.validate_object_level_assembly(output)
    return validator.errors


class SpeakerNotesTests(unittest.TestCase):
    def test_deck_only_reorder_cannot_bypass_approved_story(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inventory_fixture(root)
            self.assertEqual(ProjectValidator(root, "concept").run(), [])
            story_bytes = (root / "storyboard.json").read_bytes()
            state = json.loads((root / "run-state.json").read_text(encoding="utf-8"))
            g2 = copy.deepcopy(state["gates"]["story"])
            deck = json.loads((root / "deck-spec.json").read_text(encoding="utf-8"))
            deck["slides"].reverse()
            write_json(root, "deck-spec.json", deck)
            # Synthetic G4 rebinding reproduces the reviewed cross-gate gap;
            # G2 and the approved storyboard remain unchanged.
            g4 = state["gates"]["concept_contact_sheet"]
            g4["approved_artifacts"]["deck-spec.json"] = hashlib.sha256((root / "deck-spec.json").read_bytes()).hexdigest()
            for key, value in gate_decision_values(root, "concept_contact_sheet").items():
                g4["approval_packet"]["decisions"][key].update(value=value, value_sha256=value_sha256(value))
            write_json(root, "run-state.json", state)
            self.assertEqual((root / "storyboard.json").read_bytes(), story_bytes)
            self.assertEqual(state["gates"]["story"], g2)
            errors = ProjectValidator(root, "concept").run()
            self.assertTrue(any("slide IDs and order must match storyboard.json" in error for error in errors), errors)

    def test_matching_reordered_story_and_spec_satisfy_deck_contract(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inventory_fixture(root)
            validator = ProjectValidator(root, "concept")
            validator.load_required_documents()
            validator.docs["deck-spec.json"]["slides"].reverse()
            validator.docs["storyboard.json"]["slides"].reverse()
            # This checks content consistency only. Actual human G2/G4 records
            # must be reconciled separately for a real reordered project.
            validator.validate_deck()
            self.assertEqual(validator.errors, [])

    @unittest.skipUnless(sys.platform == "win32" and os.path.isfile("C:/Program Files/Microsoft Office/root/Office16/POWERPNT.EXE"), "PowerPoint renderer unavailable")
    def test_note_bearing_pptx_opens_and_renders_with_powerpoint(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_ready_fixture(root)
            text = "这是初步验证。\n\n请注意：B场景未测试，不能外推结论。"
            set_fixture_notes(root, text, required=True)
            script = Path(__file__).resolve().parents[1] / "scripts/assemble_editable_pptx.py"
            result = subprocess.run(
                [sys.executable, str(script), str(root), "--version", "powerpoint-notes"],
                capture_output=True, text=True, check=False, timeout=90,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(result.stderr, "")
            output = root / "assembly/deck-powerpoint-notes.pptx"
            self.assertEqual(ProjectValidator(root, "assembly").run(), [])
            self.assertEqual(Presentation(output).slides[0].notes_slide.notes_text_frame.text, text)

    def test_chinese_paragraphs_round_trip_into_native_notes(self):
        raw = '先看左侧的输入。范围：仅A数据集；B场景尚未测试。\r\n\r\n下一页解释 <验证计划>、"阈值" 与\t边界。'
        expected = raw.replace("\r\n", "\n")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_ready_fixture(root)
            set_fixture_notes(root, raw, required=True)
            self.assertEqual(ProjectValidator(root, "assets").run(), [])
            output = assemble(root, "notes-v1", render=False)
            actual = Presentation(output).slides[0].notes_slide.notes_text_frame.text
            self.assertEqual(actual, expected)
            assembly = json.loads((root / "assembly-map.json").read_text(encoding="utf-8"))
            self.assertEqual(assembly["slides"][0]["speaker_notes_sha256"], hashlib.sha256(expected.encode("utf-8")).hexdigest())
            self.assertEqual(inspect_object_contract(root, output), [])

    def test_legacy_deck_without_notes_still_assembles(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_ready_fixture(root)
            output = assemble(root, "legacy-notes", render=False)
            self.assertFalse(Presentation(output).slides[0].has_notes_slide)
            self.assertEqual(inspect_object_contract(root, output), [])

    def test_required_notes_cannot_be_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_ready_fixture(root)
            set_fixture_notes(root, " \n\t", required=True)
            errors = ProjectValidator(root, "assets").run()
            self.assertTrue(any("speaker_notes is required" in error for error in errors), errors)

    def test_bad_note_values_are_validation_errors(self):
        for value in (None, ["A", "B"], {"path": "speech.md"}, "invisible\x00text", "soft\x0bbreak"):
            with self.subTest(value=value), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                pptx_ready_fixture(root)
                set_fixture_notes(root, value)
                errors = ProjectValidator(root, "assets").run()
                self.assertTrue(any("speaker_notes must be plain text" in error or "XML-invalid" in error for error in errors), errors)

    def test_xml_invalid_unicode_is_rejected_before_serialization(self):
        for character in ("\ud800", "\ufffe", "\uffff"):
            with self.subTest(character=repr(character)), self.assertRaises(ValueError):
                speaker_notes_text({"speaker_notes": "text" + character})

    def test_tampered_or_deleted_notes_fail_even_with_updated_pptx_hash(self):
        for replacement in ("B场景已经取得确定效果。", ""):
            with self.subTest(replacement=replacement), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                pptx_ready_fixture(root)
                set_fixture_notes(root, "B场景未测试。")
                output = assemble(root, "tamper-notes", render=False)
                deck = Presentation(output)
                deck.slides[0].notes_slide.notes_text_frame.text = replacement
                deck.save(output)
                assembly = json.loads((root / "assembly-map.json").read_text(encoding="utf-8"))
                assembly["sha256"] = hashlib.sha256(output.read_bytes()).hexdigest()
                assembly["slides"][0]["speaker_notes_sha256"] = hashlib.sha256(replacement.encode("utf-8")).hexdigest()
                write_json(root, "assembly-map.json", assembly)
                errors = ProjectValidator(root, "assembly").run()
                self.assertTrue(any("PPTX speaker notes differ" in error for error in errors), errors)

    def test_injected_notes_fail_on_a_deck_with_no_notes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_ready_fixture(root)
            output = assemble(root, "injected-notes", render=False)
            deck = Presentation(output)
            deck.slides[0].notes_slide.notes_text_frame.text = "Unapproved claim."
            deck.save(output)
            errors = inspect_object_contract(root, output)
            self.assertTrue(any("PPTX speaker notes differ" in error for error in errors), errors)

    def test_missing_or_wrong_note_manifest_hash_fails(self):
        for bad_hash in (None, "0" * 64):
            with self.subTest(bad_hash=bad_hash), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                pptx_ready_fixture(root)
                set_fixture_notes(root, "Current approved explanation.")
                output = assemble(root, "wrong-note-hash", render=False)
                assembly = json.loads((root / "assembly-map.json").read_text(encoding="utf-8"))
                if bad_hash is None:
                    assembly["slides"][0].pop("speaker_notes_sha256")
                else:
                    assembly["slides"][0]["speaker_notes_sha256"] = bad_hash
                write_json(root, "assembly-map.json", assembly)
                errors = inspect_object_contract(root, output)
                self.assertTrue(any("assembly speaker notes hash differs" in error for error in errors), errors)

    def test_reordered_deck_and_selected_preview_keep_notes_with_slide_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_ready_fixture(root)
            deck = json.loads((root / "deck-spec.json").read_text(encoding="utf-8"))
            first = deck["slides"][0]
            first["speaker_notes"] = "S01 的说明。"
            seventh = copy.deepcopy(first)
            seventh["id"] = "S07"
            seventh["speaker_notes"] = "S07 的说明，与首页不同。"
            deck["slides"] = [seventh, first]
            write_json(root, "deck-spec.json", deck)
            scene = json.loads((root / "scene-graph.json").read_text(encoding="utf-8"))
            seventh_scene = copy.deepcopy(scene["slides"][0])
            seventh_scene["slide_id"] = "S07"
            seventh_scene["elements"][0]["instance_id"] = "I-S07"
            scene["slides"] = [seventh_scene, scene["slides"][0]]
            write_json(root, "scene-graph.json", scene)
            manifest = json.loads((root / "asset-manifest.json").read_text(encoding="utf-8"))
            manifest["assets"][0]["slide_ids"] = ["S01", "S07"]
            write_json(root, "asset-manifest.json", manifest)
            # Isolate assembly mapping; these synthetic reordered pages have no
            # real human approval and are never evidence for gate acceptance.
            with patch("assemble_editable_pptx.ProjectValidator") as validator:
                validator.return_value.run.return_value = []
                full = assemble(root, "reordered-notes", render=False)
                preview = assemble(root, "preview-S07-notes", render=False, preview_slide="S07")
                validator.return_value.pilot_bindings.return_value = {"slide_id": "S07"}
                pilot = assemble(root, "pilot-S07-notes", render=False, pilot_slide="S07")
            self.assertEqual([slide.notes_slide.notes_text_frame.text for slide in Presentation(full).slides],
                             [seventh["speaker_notes"], first["speaker_notes"]])
            for output in (preview, pilot):
                pages = Presentation(output).slides
                self.assertEqual(len(pages), 1)
                self.assertEqual(pages[0].notes_slide.notes_text_frame.text, seventh["speaker_notes"])


if __name__ == "__main__":
    unittest.main()
