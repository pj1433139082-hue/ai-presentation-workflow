import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from test_assembly import multiple_object_fixture
from test_validate_project import (
    ISOLATION_MODULE, ProjectValidator, gate_decision_values,
    refresh_fixture_decomposition_review, refresh_fixture_isolation_review,
    value_sha256, write_json,
)


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"


def make_contentful_pilot_fixture(root: Path) -> None:
    """Draw a real-size concept from the declared composition, not from a PPTX render."""
    multiple_object_fixture(root)
    canvas = Image.new("RGB", (1920, 1080), (15, 23, 42))
    draw = ImageDraw.Draw(canvas)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    def fixture_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
        return ImageFont.truetype(str(font_path), size) if font_path.is_file() else ImageFont.load_default(size=size)
    draw.text((154, 112), "A decision-ready workflow", font=fixture_font(68), fill=(56, 189, 248))
    draw.multiline_text((154, 285), "Pictures and copy remain\nindependently editable.", font=fixture_font(44), spacing=0, fill=(56, 189, 248))
    draw.text((154, 761), "Data to be entered", font=fixture_font(36), fill=(56, 189, 248))
    draw.rectangle((154, 853, 653, 940), outline=(56, 189, 248), width=2)
    scene = json.loads((root / "scene-graph.json").read_text(encoding="utf-8"))["slides"][0]
    manifest = json.loads((root / "asset-manifest.json").read_text(encoding="utf-8"))
    assets = {asset["id"]: asset for asset in manifest["assets"]}
    for item in sorted(scene["elements"], key=lambda element: element["z"]):
        x, y, width, height = item["bbox"]
        geometry = (round(x * 1920), round(y * 1080), round(width * 1920), round(height * 1080))
        with Image.open(root / assets[item["component_id"]]["output"]) as source:
            picture = source.convert("RGBA").resize(geometry[2:], Image.Resampling.LANCZOS)
            canvas.paste(picture, geometry[:2], picture)
    concept = root / "concepts" / "S01.png"
    canvas.save(concept)
    concept_sha = hashlib.sha256(concept.read_bytes()).hexdigest()
    deck = json.loads((root / "deck-spec.json").read_text(encoding="utf-8"))
    review = deck["slides"][0]["readability"]["full_size_review"]
    review["sha256"] = concept_sha
    review["reviewed_concept_sha256"] = concept_sha
    write_json(root, "deck-spec.json", deck)
    decomposition = json.loads((root / "decomposition-map.json").read_text(encoding="utf-8"))
    decomposition["slides"][0]["concept_sha256"] = concept_sha
    decomposition["slides"][0]["review_image"]["reviewed_concept_sha256"] = concept_sha
    write_json(root, "decomposition-map.json", decomposition)
    refresh_fixture_decomposition_review(root)
    for asset in manifest["assets"]:
        for reference in asset["references"]:
            if reference["path"] == "concepts/S01.png":
                reference["sha256"] = concept_sha
    write_json(root, "asset-manifest.json", manifest)
    refresh_fixture_isolation_review(root)
    manifest = json.loads((root / "asset-manifest.json").read_text(encoding="utf-8"))
    photo = next(asset for asset in manifest["assets"] if asset["id"] == "C02")
    photo_items = ISOLATION_MODULE.source_items_for_asset(decomposition, "C02")
    image_path = root / photo["isolation_review"]["image_path"]
    ISOLATION_MODULE.build_isolation_image(root, root / photo["output"], photo_items).save(image_path)
    packet_path = root / photo["isolation_review"]["packet_path"]
    packet_path.write_text("# C02 fixture isolation\n\n" + ISOLATION_MODULE.value_sha256(photo_items) + "\n\n"
                           + "\n".join(ISOLATION_MODULE.review_row(item) for item in photo_items) + "\n", encoding="utf-8")
    photo["isolation_review"]["image_sha256"] = hashlib.sha256(image_path.read_bytes()).hexdigest()
    photo["isolation_review"]["packet_sha256"] = hashlib.sha256(packet_path.read_bytes()).hexdigest()
    photo["isolation_review"]["reviewed_items_sha256"] = ISOLATION_MODULE.value_sha256(photo_items)
    write_json(root, "asset-manifest.json", manifest)
    state = json.loads((root / "run-state.json").read_text(encoding="utf-8"))
    for gate_name in ("style_anchors", "concept_contact_sheet", "decomposition"):
        gate = state["gates"][gate_name]
        for relative in gate["approved_artifacts"]:
            gate["approved_artifacts"][relative] = hashlib.sha256((root / relative).read_bytes()).hexdigest()
        for key, value in gate_decision_values(root, gate_name).items():
            gate["approval_packet"]["decisions"][key]["value"] = value
            gate["approval_packet"]["decisions"][key]["value_sha256"] = value_sha256(value)
    write_json(root, "run-state.json", state)


def make_pilot(root: Path) -> tuple[Path, Path]:
    multiple_object_fixture(root)
    write_json(root, "pilot-plan.json", {
        "slide_id": "S01", "component_ids": ["C01", "C02"],
        "selection_rationale": "Synthetic busy slide with independent arrows, a photo, native copy, and blank data.",
    })
    assert ProjectValidator(root, "pilot", require_pilot_review=False).run() == []
    result = subprocess.run([sys.executable, str(SCRIPTS / "assemble_editable_pptx.py"), str(root), "--version", "trial-v1", "--pilot-slide", "S01", "--no-render"], capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stdout + result.stderr
    map_path = root / "pilot" / "trial-v1" / "pilot-assembly-map.json"
    assembly = json.loads(map_path.read_text(encoding="utf-8"))
    pptx_path = root / assembly["output"]
    render = root / "reviews" / "pilot" / "trial-v1" / "slide-01.png"
    render.parent.mkdir(parents=True, exist_ok=True)
    # Fixture-only: production must render the actual PPTX with the verified backend.
    Image.new("RGB", (1920, 1080), (18, 37, 73)).save(render)
    assembly["renders"] = [{"path": render.relative_to(root).as_posix(), "sha256": hashlib.sha256(render.read_bytes()).hexdigest(), "pptx_sha256": assembly["sha256"]}]
    map_path.write_text(json.dumps(assembly, ensure_ascii=False, indent=2), encoding="utf-8")
    result = subprocess.run([sys.executable, str(SCRIPTS / "render_pilot_review.py"), str(root), "--version", "trial-v1"], capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stdout + result.stderr
    return pptx_path, map_path


def approve(root: Path) -> dict:
    review = json.loads((root / "pilot-review.json").read_text(encoding="utf-8"))
    review.update(status="approved", approval_phrase="确认单页试拼", reviewer="fixture-human", reviewed_at=datetime.now(timezone.utc).isoformat())
    review["checks"] = {key: True for key in review["checks"]}
    write_json(root, "pilot-review.json", review)
    return review


class PilotReviewTests(unittest.TestCase):
    @unittest.skipUnless(sys.platform == "win32" and os.path.isfile("C:/Program Files/Microsoft Office/root/Office16/POWERPNT.EXE"), "PowerPoint renderer unavailable")
    def test_actual_powerpoint_trial_render_and_packet(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_contentful_pilot_fixture(root)
            write_json(root, "pilot-plan.json", {
                "slide_id": "S01", "component_ids": ["C01", "C02"],
                "selection_rationale": "Synthetic representative page has arrows, photo, text, and data space.",
            })
            result = subprocess.run([sys.executable, str(SCRIPTS / "assemble_editable_pptx.py"), str(root), "--version", "real-v1", "--pilot-slide", "S01"], capture_output=True, text=True, check=False, timeout=90)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            result = subprocess.run([sys.executable, str(SCRIPTS / "render_pilot_review.py"), str(root), "--version", "real-v1"], capture_output=True, text=True, check=False)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            review = json.loads((root / "pilot-review.json").read_text(encoding="utf-8"))
            self.assertEqual(review["status"], "pending-human")
            self.assertEqual(review["evidence"]["side_by_side"]["path"], "reviews/pilot/real-v1/side-by-side.png")
            self.assertTrue((root / review["evidence"]["overlay"]["path"]).is_file())

    def test_pending_then_current_approval(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_pilot(root)
            errors = ProjectValidator(root, "pilot").run()
            self.assertTrue(any("trial needs current human approval or user-authorized local agent QA" in issue for issue in errors), errors)
            approve(root)
            self.assertEqual(ProjectValidator(root, "pilot").run(), [])

    def test_alpha_valid_but_fused_object_is_rejected_by_human_check(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_pilot(root)
            review = approve(root)
            review["checks"]["single_object"] = False
            review["status"] = "rejected"
            review["defects"] = [{"component_id": "C01", "severity": "blocker", "evidence": review["evidence"]["isolation"][0]["image_path"], "owner_stage": "assets", "repair_entry": "reconstruct only C01"}]
            write_json(root, "pilot-review.json", review)
            errors = ProjectValidator(root, "pilot").run()
            self.assertTrue(any("all trial checks must pass" in issue for issue in errors), errors)

    def test_missing_arrow_in_pptx_blocks_even_when_hashes_are_refreshed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_path, map_path = make_pilot(root)
            review = approve(root)
            deck = Presentation(pptx_path)
            arrow = deck.slides[0].shapes[0]
            arrow._element.getparent().remove(arrow._element)
            deck.save(pptx_path)
            assembly = json.loads(map_path.read_text(encoding="utf-8"))
            assembly["sha256"] = hashlib.sha256(pptx_path.read_bytes()).hexdigest()
            assembly["renders"][0]["pptx_sha256"] = assembly["sha256"]
            map_path.write_text(json.dumps(assembly, ensure_ascii=False, indent=2), encoding="utf-8")
            review["pptx_sha256"] = assembly["sha256"]
            write_json(root, "pilot-review.json", review)
            errors = ProjectValidator(root, "pilot").run()
            self.assertTrue(any("object count, identity, or z-order" in issue for issue in errors), errors)

    def test_large_manual_offset_blocks_even_when_hashes_are_refreshed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_path, map_path = make_pilot(root)
            review = approve(root)
            deck = Presentation(pptx_path)
            deck.slides[0].shapes[0].left += Inches(1.0)
            deck.save(pptx_path)
            assembly = json.loads(map_path.read_text(encoding="utf-8"))
            assembly["sha256"] = hashlib.sha256(pptx_path.read_bytes()).hexdigest()
            assembly["renders"][0]["pptx_sha256"] = assembly["sha256"]
            map_path.write_text(json.dumps(assembly, ensure_ascii=False, indent=2), encoding="utf-8")
            review["pptx_sha256"] = assembly["sha256"]
            write_json(root, "pilot-review.json", review)
            errors = ProjectValidator(root, "pilot").run()
            self.assertTrue(any("geometry differs from current scene" in issue for issue in errors), errors)

    def test_selected_asset_change_stales_approval(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_pilot(root)
            approve(root)
            manifest = json.loads((root / "asset-manifest.json").read_text(encoding="utf-8"))
            manifest["assets"][0]["prompt"] += " changed"
            write_json(root, "asset-manifest.json", manifest)
            errors = ProjectValidator(root, "pilot").run()
            self.assertTrue(any("trial approval is stale" in issue for issue in errors), errors)

    def test_other_manifest_metadata_does_not_stale_selected_page(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_pilot(root)
            approve(root)
            manifest = json.loads((root / "asset-manifest.json").read_text(encoding="utf-8"))
            manifest["later_page_note"] = "remaining pages may be produced after pilot approval"
            write_json(root, "asset-manifest.json", manifest)
            self.assertEqual(ProjectValidator(root, "pilot").run(), [])

    def test_fabricated_overlay_pixels_are_rejected_even_with_new_hash(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_pilot(root)
            review = approve(root)
            overlay_path = root / review["evidence"]["overlay"]["path"]
            Image.new("RGB", (1920, 1080), (0, 0, 0)).save(overlay_path)
            review["evidence"]["overlay"]["sha256"] = hashlib.sha256(overlay_path.read_bytes()).hexdigest()
            write_json(root, "pilot-review.json", review)
            errors = ProjectValidator(root, "pilot").run()
            self.assertTrue(any("overlay pixels do not match" in issue for issue in errors), errors)

    def test_production_assets_stage_requires_pilot_packet(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            multiple_object_fixture(root)
            brief = json.loads((root / "brief.json").read_text(encoding="utf-8"))
            brief["quality_profile"] = "production"
            write_json(root, "brief.json", brief)
            errors = ProjectValidator(root, "assets").run()
            self.assertTrue(any("pilot-plan.json" in issue for issue in errors), errors)


if __name__ == "__main__":
    unittest.main()
