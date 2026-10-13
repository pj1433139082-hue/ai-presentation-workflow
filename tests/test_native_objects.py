"""Round-trip tests for the explicit showcase native-object assembler."""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))

import showcase_contract  # noqa: E402
from assemble_editable_pptx import assemble_showcase  # noqa: E402
from native_objects import NativeObjectError, add_native_object  # noqa: E402


REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_IMAGE = REPO_ROOT / "examples" / "showcase" / "tech" / "assets" / "search-prism.png"


def _write_json(root: Path, relative: str, value: dict) -> None:
    path = root / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _text_element(instance_id: str, copy_id: str, role: str, text: str, bbox: list[float], z: int) -> dict:
    return {
        "instance_id": instance_id,
        "type": "text",
        "copy_id": copy_id,
        "role": role,
        "text": text,
        "bbox": bbox,
        "z": z,
        "style": {
            "font_family": "Arial",
            "font_size_pt": {"title": 28, "key-takeaway": 20, "body": 16}[role],
            "color": "#111111",
            "background_color": "#FFFFFF",
            "bold": role == "title",
            "align": "left",
            "valign": "middle",
        },
    }


def _make_project(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    _write_json(root, "project-kind.json", {"contract_version": "showcase-v1", "purpose": "showcase"})

    image_relative = "assets/prism.png"
    image_path = root / image_relative
    image_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(SOURCE_IMAGE, image_path)
    image_hash = hashlib.sha256(image_path.read_bytes()).hexdigest()
    shutil.copyfile(SOURCE_IMAGE, root / "assets" / "evidence.png")
    evidence_hash = hashlib.sha256((root / "assets" / "evidence.png").read_bytes()).hexdigest()

    concept_path = root / "concepts" / "S01.png"
    concept_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(SOURCE_IMAGE, concept_path)
    concept_hash = hashlib.sha256(concept_path.read_bytes()).hexdigest()

    image_assets = [
        {
            "id": "A-picture",
            "slide_id": "S01",
            "kind": "picture",
            "path": image_relative,
            "sha256": image_hash,
            "truth_status": "illustrative",
            "source_locator": "fixture source image already present in showcase examples",
        },
        {
            "id": "A-evidence",
            "slide_id": "S01",
            "kind": "source-evidence",
            "path": "assets/evidence.png",
            "sha256": evidence_hash,
            "truth_status": "illustrative",
            "source_locator": "fixture source image already present in showcase examples",
        },
    ]

    copy_rows = [
        {"id": "copy-title", "role": "title", "text": "示例页面", "truth_status": "illustrative", "source_ids": []},
        {"id": "copy-takeaway", "role": "key-takeaway", "text": "所有混合对象保持可编辑", "truth_status": "illustrative", "source_ids": []},
        {"id": "copy-units", "role": "body", "text": "单位：次", "truth_status": "illustrative", "source_ids": []},
    ]
    elements = [
        _text_element("title", "copy-title", "title", "示例页面", [0.05, 0.04, 0.42, 0.09], 0),
        _text_element("takeaway", "copy-takeaway", "key-takeaway", "所有混合对象保持可编辑", [0.05, 0.15, 0.55, 0.08], 1),
        _text_element("units-text", "copy-units", "body", "单位：次", [0.36, 0.69, 0.22, 0.07], 2),
        {
            "instance_id": "native-shape", "type": "shape", "shape_kind": "rounded-rectangle",
            "bbox": [0.05, 0.30, 0.20, 0.16], "z": 3,
            "fill": {"color": "#DDEEFF", "transparency": 0},
            "line": {"color": "#225588", "transparency": 0},
        },
        {
            "instance_id": "native-connector", "type": "connector",
            "bbox": [0.27, 0.30, 0.11, 0.11], "z": 4,
            "from": {"x": 0.27, "y": 0.30}, "to": {"x": 0.38, "y": 0.41},
            "line": {"color": "#225588", "transparency": 0, "width_pt": 1.5},
            "arrowhead": "triangle",
        },
        {
            "instance_id": "native-table", "type": "table",
            "bbox": [0.05, 0.51, 0.27, 0.19], "z": 5,
            "columns": [{"label": "阶段"}, {"label": "计数"}],
            "rows": [{"cells": ["A", 2]}, {"cells": ["B", 3]}],
            "units": "次", "truth_status": "illustrative", "source_ids": ["SOURCE-1"],
            "visible_label": "示例数据",
        },
        {
            "instance_id": "native-chart", "type": "chart",
            "bbox": [0.40, 0.31, 0.36, 0.34], "z": 6,
            "chart_type": "column", "title": "示例数据：阶段计数",
            "categories": ["A", "B"], "series": [{"name": "次数", "values": [2, 3]}],
            "units": "次", "truth_status": "illustrative", "source_ids": ["SOURCE-1"],
        },
        {
            "instance_id": "independent-picture", "type": "picture", "asset_id": "A-picture",
            "bbox": [0.77, 0.34, 0.16, 0.18], "z": 7,
        },
        {
            "instance_id": "independent-evidence", "type": "source-evidence", "asset_id": "A-evidence",
            "bbox": [0.79, 0.58, 0.12, 0.13], "z": 8,
        },
    ]

    deck_slide = {
        "id": "S01", "title": "示例页面", "message": "所有混合对象保持可编辑",
        "copy": copy_rows, "speaker_notes": "独立对象装配验证。",
    }
    documents = {
        "brief.json": {"project_id": "native-object-roundtrip-test", "source_map": [{"id": "SOURCE-1", "label": "示意用样例"}]},
        "deck-spec.json": {"slides": [deck_slide]},
        "reference-study.json": {"references": []},
        "style-contract.json": {
            "palette": {"background": "#FFFFFF"},
            "readability_profile": {
                "font_floors_pt": {"title": 24, "key-takeaway": 18, "body": 14},
                "minimum_contrast_ratio": 4.5,
            },
        },
        "scene-graph.json": {"slides": [{"slide_id": "S01", "canvas": {"width": 1920, "height": 1080}, "elements": elements}]},
        "asset-manifest.json": {"assets": image_assets},
        "concept-generation-log.json": {
            "slides": [{"slide_id": "S01", "output_path": "concepts/S01.png", "output_sha256": concept_hash}],
        },
        "run-state.json": {"mode": "showcase-local-draft"},
    }
    for name, document in documents.items():
        _write_json(root, name, document)
    return root


class NativeObjectAssemblyTests(unittest.TestCase):
    def test_assembly_round_trips_all_native_types_and_chart_data_remains_editable(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = _make_project(Path(temp_dir) / "showcase")

            def fake_validate(_root: Path, stage: str) -> dict:
                self.assertEqual(stage, "concept")
                return {"ok": True, "errors": [], "warnings": []}

            with patch.object(showcase_contract, "validate_showcase", side_effect=fake_validate):
                result = assemble_showcase(root, "native-test-v1", render=False)

            self.assertIsNone(result["render_backend"])
            self.assertEqual(result["renders"], [])
            self.assertIn("human review remain pending", result["review_boundary"])
            output = root / result["output"]
            self.assertTrue(output.is_file())
            reopened = Presentation(output)
            self.assertEqual(len(reopened.slides), 1)
            shapes = {shape.name: shape for shape in reopened.slides[0].shapes}
            mapped = {row["type"]: row for row in result["slides"][0]["objects"]}
            text_shape = shapes[mapped["text"]["object_name"]]
            self.assertTrue(text_shape.has_text_frame)
            text_sp_pr = text_shape._element.spPr
            self.assertIsNotNone(text_sp_pr.find(qn("a:noFill")))
            self.assertIsNotNone(text_sp_pr.find(qn("a:ln")).find(qn("a:noFill")))

            table = shapes[mapped["table"]["object_name"]].table
            self.assertEqual(table.cell(0, 0).text, "示例数据")
            self.assertEqual(table.cell(1, 0).text, "阶段")
            self.assertEqual(table.cell(2, 1).text, "2")

            chart_shape = shapes[mapped["chart"]["object_name"]]
            self.assertTrue(chart_shape.has_chart)
            chart = chart_shape.chart
            self.assertEqual(chart.chart_title.text_frame.text, "示例数据：阶段计数")
            self.assertEqual(list(chart.series[0].values), [2.0, 3.0])
            self.assertEqual([item.label for item in chart.plots[0].categories], ["A", "B"])
            self.assertFalse(chart.category_axis.has_title)
            self.assertFalse(chart.value_axis.has_title)
            self.assertEqual(chart.plots[0].data_labels.position, XL_LABEL_POSITION.OUTSIDE_END)
            self.assertNotIn("style", mapped["chart"])

            connector = shapes[mapped["connector"]["object_name"]]
            self.assertEqual(connector.shape_type, MSO_SHAPE_TYPE.LINE)
            self.assertIn("tailEnd", connector._element.xml)
            self.assertEqual(shapes[mapped["picture"]["object_name"]].shape_type, MSO_SHAPE_TYPE.PICTURE)
            self.assertEqual(shapes[mapped["source-evidence"]["object_name"]].shape_type, MSO_SHAPE_TYPE.PICTURE)
            self.assertTrue(all(row["editable"] for row in result["slides"][0]["objects"] if row["type"] not in {"picture", "source-evidence"}))

            replacement = CategoryChartData()
            replacement.categories = ["C", "D"]
            replacement.add_series("更新后的次数", [5, 8])
            chart.replace_data(replacement)
            edited_path = root / "edited-copy.pptx"
            reopened.save(edited_path)
            edited_chart = next(shape.chart for shape in Presentation(edited_path).slides[0].shapes if shape.has_chart)
            self.assertEqual([item.label for item in edited_chart.plots[0].categories], ["C", "D"])
            self.assertEqual(list(edited_chart.series[0].values), [5.0, 8.0])

    def test_project_chart_style_round_trips_axis_labels_series_color_and_markers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = _make_project(Path(temp_dir) / "showcase")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            style = {
                "font_family": "Microsoft YaHei",
                "font_size_pt": 11,
                "color": "#05213F",
                "series_colors": ["#08ACA8"],
                "value_axis": {
                    "minimum_scale": 0.6,
                    "maximum_scale": 1.1,
                    "major_unit": 0.1,
                    "number_format": "0.00",
                },
                "category_axis_title": "时间 (min)",
                "value_axis_title": "J/J0",
                "data_label_number_format": "0.00",
                "line_markers": True,
            }
            row, _ = add_native_object(
                slide,
                {
                    "instance_id": "configured-line",
                    "type": "chart",
                    "bbox": [0.08, 0.10, 0.82, 0.78],
                    "z": 0,
                    "chart_type": "line",
                    "title": "归一化通量（示例数据）",
                    "categories": ["0", "20", "40", "60"],
                    "series": [{"name": "J/J0", "values": [1.00, 0.93, 0.85, 0.78]}],
                    "units": "min",
                    "truth_status": "illustrative",
                    "source_ids": ["SOURCE-1"],
                    "style": style,
                },
                slide_id="S01",
                slide_width=presentation.slide_width,
                slide_height=presentation.slide_height,
                root=root,
                assets={},
                body_font_pt=10,
            )
            self.assertEqual(row["style"], style)
            first_path = root / "styled-chart.pptx"
            presentation.save(first_path)

            reopened = Presentation(first_path)
            chart = next(shape.chart for shape in reopened.slides[0].shapes if shape.has_chart)
            self.assertAlmostEqual(chart.value_axis.minimum_scale, 0.6)
            self.assertAlmostEqual(chart.value_axis.maximum_scale, 1.1)
            self.assertAlmostEqual(chart.value_axis.major_unit, 0.1)
            self.assertEqual(chart.value_axis.tick_labels.number_format, "0.00")
            self.assertEqual(chart.category_axis.axis_title.text_frame.text, "时间 (min)")
            self.assertEqual(chart.value_axis.axis_title.text_frame.text, "J/J0")
            self.assertEqual(chart.plots[0].data_labels.number_format, "0.00")
            self.assertEqual(chart.plots[0].data_labels.position, XL_LABEL_POSITION.ABOVE)
            self.assertEqual(chart.series[0].format.line.color.rgb, RGBColor(8, 172, 168))
            self.assertEqual(chart.series[0].marker.style, XL_MARKER_STYLE.CIRCLE)
            self.assertEqual(chart.series[0].marker.format.fill.fore_color.rgb, RGBColor(8, 172, 168))
            title_font = chart.chart_title.text_frame.paragraphs[0].font
            self.assertEqual(title_font.name, "Microsoft YaHei")
            self.assertEqual(title_font.size.pt, 11)
            self.assertEqual(title_font.color.rgb, RGBColor(5, 33, 63))
            self.assertEqual(chart.value_axis.tick_labels.font.name, "Microsoft YaHei")
            self.assertEqual(chart.category_axis.tick_labels.font.color.rgb, RGBColor(5, 33, 63))

            chart_ns = {
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
            }
            chart_text_properties = {}
            with zipfile.ZipFile(first_path) as archive:
                chart_part = next(
                    name for name in archive.namelist()
                    if name.startswith("ppt/charts/chart") and name.endswith(".xml")
                )
                chart_xml = ET.fromstring(archive.read(chart_part))
            for paragraph in chart_xml.findall(".//c:title/c:tx/c:rich/a:p", chart_ns):
                text = "".join(node.text or "" for node in paragraph.findall(".//a:t", chart_ns))
                if text not in {"归一化通量（示例数据）", "时间 (min)", "J/J0"}:
                    continue
                def_rpr = paragraph.find("a:pPr/a:defRPr", chart_ns)
                self.assertIsNotNone(def_rpr, f"missing default run properties for {text!r}")
                latin = def_rpr.find("a:latin", chart_ns)
                east_asian = def_rpr.find("a:ea", chart_ns)
                self.assertIsNotNone(latin, f"missing Latin font declaration for {text!r}")
                self.assertIsNotNone(east_asian, f"missing East Asian font declaration for {text!r}")
                self.assertEqual(latin.get("typeface"), "Microsoft YaHei")
                self.assertEqual(east_asian.get("typeface"), "Microsoft YaHei")
                child_order = [child.tag.rsplit("}", 1)[-1] for child in def_rpr]
                self.assertLess(child_order.index("latin"), child_order.index("ea"))
                if "cs" in child_order:
                    self.assertLess(child_order.index("ea"), child_order.index("cs"))
                chart_text_properties[text] = east_asian.get("typeface")
            self.assertEqual(
                chart_text_properties,
                {
                    "归一化通量（示例数据）": "Microsoft YaHei",
                    "时间 (min)": "Microsoft YaHei",
                    "J/J0": "Microsoft YaHei",
                },
            )

    def test_unstyled_line_chart_uses_above_labels_and_round_trips(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = _make_project(Path(temp_dir) / "showcase")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            element = {
                "instance_id": "unstyled-line",
                "type": "chart",
                "bbox": [0.08, 0.10, 0.82, 0.78],
                "z": 0,
                "chart_type": "line",
                "title": "归一化通量（示例数据）",
                "categories": ["0", "20", "40", "60"],
                "series": [{"name": "示例曲线", "values": [1.0, 0.93, 0.85, 0.78]}],
                "units": "min",
                "truth_status": "illustrative",
                "source_ids": ["SOURCE-1"],
            }
            add_native_object(
                slide,
                element,
                slide_id="S01",
                slide_width=presentation.slide_width,
                slide_height=presentation.slide_height,
                root=root,
                assets={},
                body_font_pt=16,
            )
            path = root / "unstyled-line.pptx"
            presentation.save(path)

            reopened = Presentation(path)
            chart = next(shape.chart for shape in reopened.slides[0].shapes if shape.has_chart)
            self.assertEqual(chart.plots[0].data_labels.position, XL_LABEL_POSITION.ABOVE)
            self.assertEqual(list(chart.series[0].values), [1.0, 0.93, 0.85, 0.78])

    def test_chart_style_rejects_mismatched_series_color_count_and_non_line_markers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = _make_project(Path(temp_dir) / "showcase")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            base = {
                "instance_id": "bad-style",
                "type": "chart",
                "bbox": [0.1, 0.1, 0.5, 0.5],
                "z": 0,
                "chart_type": "column",
                "title": "示意计数",
                "categories": ["A", "B"],
                "series": [{"name": "次数", "values": [2, 3]}],
                "units": "次",
                "truth_status": "illustrative",
                "source_ids": ["SOURCE-1"],
            }
            args = {
                "slide_id": "S01", "slide_width": presentation.slide_width,
                "slide_height": presentation.slide_height, "root": root,
                "assets": {}, "body_font_pt": 10,
            }
            with self.assertRaisesRegex(NativeObjectError, "one color per chart series"):
                add_native_object(slide, {**base, "style": {"series_colors": []}}, **args)
            with self.assertRaisesRegex(NativeObjectError, "only supported for line charts"):
                add_native_object(slide, {**base, "style": {"line_markers": True}}, **args)

    def test_near_full_slide_picture_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = _make_project(Path(temp_dir) / "showcase")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            with self.assertRaisesRegex(NativeObjectError, "full-slide or near-full-slide raster"):
                add_native_object(
                    slide,
                    {"instance_id": "flat-page", "type": "picture", "asset_id": "A-picture",
                     "bbox": [0.02, 0.02, 0.96, 0.96], "z": 0},
                    slide_id="S01", slide_width=presentation.slide_width, slide_height=presentation.slide_height,
                    root=root,
                    assets={"A-picture": {"id": "A-picture", "path": "assets/prism.png",
                                          "sha256": hashlib.sha256((root / "assets" / "prism.png").read_bytes()).hexdigest()}},
                    body_font_pt=16,
                )


if __name__ == "__main__":
    unittest.main()
