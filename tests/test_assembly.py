import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from test_validate_project import (
    ISOLATION_MODULE,
    ProjectValidator,
    gate_decision_values,
    make_valid_project,
    refresh_fixture_decomposition_review,
    refresh_fixture_isolation_review,
    value_sha256,
    write_json,
)


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "assemble_editable_pptx.py"
sys.path.insert(0, str(SCRIPT.parent))
from assemble_editable_pptx import apply_text_role, assemble


def pptx_ready_fixture(root: Path) -> None:
    make_valid_project(root)
    # The base validator fixture happens to reuse the same tiny PNG for the
    # concept and component; give the component independent pixels here.
    component_path = root / "assets" / "S01" / "C01.png"
    component = Image.new("RGBA", (160, 160), (0, 0, 0, 0))
    draw = ImageDraw.Draw(component)
    draw.polygon([(20, 65), (94, 65), (94, 35), (140, 80), (94, 125), (94, 95), (20, 95)], fill=(17, 194, 127, 255))
    component.save(component_path)
    component_sha = hashlib.sha256(component_path.read_bytes()).hexdigest()
    manifest = json.loads((root / "asset-manifest.json").read_text(encoding="utf-8"))
    manifest["assets"][0]["sha256"] = component_sha
    manifest["assets"][0]["generation_history"][0]["sha256"] = component_sha
    alpha = component.getchannel("A")
    manifest["assets"][0]["alpha"]["transparent_pixel_ratio"] = alpha.histogram()[0] / (component.width * component.height)
    write_json(root, "asset-manifest.json", manifest)
    refresh_fixture_isolation_review(root)
    brief = json.loads((root / "brief.json").read_text(encoding="utf-8"))
    brief["delivery"]["formats"] = ["component-library", "pptx"]
    write_json(root, "brief.json", brief)
    capability = json.loads((root / "capability-report.json").read_text(encoding="utf-8"))
    capability["assembly"] = {"status": "verified", "backend": "python-pptx"}
    write_json(root, "capability-report.json", capability)
    alignment = json.loads((root / "alignment.json").read_text(encoding="utf-8"))
    for decision in alignment["decisions"]:
        if decision["id"] == "delivery_form":
            decision["value"] = brief["delivery"]["formats"]
    for source in alignment["material_review"]["inspected_sources"]:
        if source["path"] == "brief.json":
            source["sha256"] = hashlib.sha256((root / "brief.json").read_bytes()).hexdigest()
    write_json(root, "alignment.json", alignment)
    state = json.loads((root / "run-state.json").read_text(encoding="utf-8"))
    gate = state["gates"]["brief_authority"]
    for name in ("brief.json", "alignment.json"):
        gate["approved_artifacts"][name] = hashlib.sha256((root / name).read_bytes()).hexdigest()
    gate["decision_approvals"]["delivery_form"]["value_sha256"] = value_sha256(brief["delivery"]["formats"])
    state["gates"]["style_anchors"]["approved_artifacts"]["capability-report.json"] = hashlib.sha256((root / "capability-report.json").read_bytes()).hexdigest()
    write_json(root, "run-state.json", state)


def multiple_object_fixture(root: Path) -> None:
    """A synthetic slide with two arrow placements, one photo, text and blank data."""
    pptx_ready_fixture(root)
    source = root / "sources" / "ordinary-photo.png"
    source.parent.mkdir(parents=True, exist_ok=True)
    photo = Image.new("RGB", (320, 180))
    photo.putdata([((x * 17 + y * 7) % 256, (x * 5 + y * 19) % 256, (x * 11 + y * 13) % 256)
                   for y in range(180) for x in range(320)])
    photo.save(source)
    source_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    rights = root / "rights" / "ordinary-photo.json"
    rights.parent.mkdir(parents=True, exist_ok=True)
    rights.write_text(json.dumps({
        "status": "confirmed", "approval_phrase": "确认素材使用权", "confirmed_by": "fixture-human",
        "source_path": "sources/ordinary-photo.png", "source_sha256": source_sha,
        "media_class": "ordinary-photo", "truth_sensitive": False,
        "classification_phrase": "确认普通照片", "classified_by": "fixture-human",
        "excluded_types": {"product_screenshot": True, "data_figure": True, "report": True},
    }), encoding="utf-8")
    rights_sha = hashlib.sha256(rights.read_bytes()).hexdigest()

    deck = json.loads((root / "deck-spec.json").read_text(encoding="utf-8"))
    deck["production_estimate"]["unique_components"] = 2
    slide = deck["slides"][0]
    slide["components"] = ["C01", "C02"]
    slide["fill_slots"][0]["bbox"] = [0.08, 0.09, 0.52, 0.12]
    slide["fill_slots"].extend([
        {"id": "F02", "kind": "ordinary-text", "bbox": [0.08, 0.23, 0.40, 0.16],
         "content": "Pictures and copy remain independently editable."},
        {"id": "F03", "kind": "ordinary-data", "bbox": [0.08, 0.79, 0.26, 0.08], "content": ""},
        {"id": "F04", "kind": "ordinary-text", "bbox": [0.08, 0.70, 0.28, 0.05],
         "content": "Data to be entered"},
    ])
    slide["readability"]["text_regions"].extend([
        {"fill_slot_id": "F02", "role": "body", "planned_font_pt": 22, "line_count": 2},
        {"fill_slot_id": "F03", "role": "key-caption", "planned_font_pt": 18, "line_count": 1},
        {"fill_slot_id": "F04", "role": "key-caption", "planned_font_pt": 18, "line_count": 1},
    ])
    write_json(root, "deck-spec.json", deck)
    scene = json.loads((root / "scene-graph.json").read_text(encoding="utf-8"))
    scene["slides"][0]["elements"] = [
        {"instance_id": "I01", "component_id": "C01", "bbox": [0.10, 0.48, 0.22, 0.20], "z": 1},
        {"instance_id": "I02", "component_id": "C02", "bbox": [0.55, 0.30, 0.32, 0.32], "z": 2},
        {"instance_id": "I03", "component_id": "C01", "bbox": [0.50, 0.51, 0.12, 0.12], "z": 3},
    ]
    write_json(root, "scene-graph.json", scene)

    decomposition = json.loads((root / "decomposition-map.json").read_text(encoding="utf-8"))
    inventory = decomposition["slides"][0]["items"]
    inventory[0]["bbox"] = [0.10, 0.48, 0.22, 0.20]
    inventory[1]["bbox"] = slide["fill_slots"][0]["bbox"]
    photo_item = dict(inventory[0], id="D03", label="one photo rectangle", role="ordinary-photo",
                      source_class="ordinary-photo", bbox=[0.55, 0.30, 0.32, 0.32], z=2,
                      instance_id="I02", component_id="C02", treatment="preserve-source",
                      occlusion="partially-occluded", overlaps=["I03"])
    second_arrow = dict(inventory[0], id="D04", label="reused arrow", bbox=[0.50, 0.51, 0.12, 0.12],
                        z=3, instance_id="I03", overlaps=["I02"])
    inventory.extend([
        photo_item, second_arrow,
        {"id": "D05", "kind": "fill-slot", "inventory_pass": "semantic", "label": "editable body",
         "role": "ordinary-text", "bbox": slide["fill_slots"][1]["bbox"], "fill_slot_id": "F02", "treatment": "editable"},
        {"id": "D06", "kind": "fill-slot", "inventory_pass": "semantic", "label": "blank user data",
         "role": "ordinary-data", "bbox": slide["fill_slots"][2]["bbox"], "fill_slot_id": "F03", "treatment": "editable"},
        {"id": "D07", "kind": "fill-slot", "inventory_pass": "semantic", "label": "data input label",
         "role": "ordinary-text", "bbox": slide["fill_slots"][3]["bbox"], "fill_slot_id": "F04", "treatment": "editable"},
    ])
    write_json(root, "decomposition-map.json", decomposition)
    refresh_fixture_decomposition_review(root)

    manifest = json.loads((root / "asset-manifest.json").read_text(encoding="utf-8"))
    arrow = manifest["assets"][0]
    photo_asset = {
        "id": "C02", "slide_id": "S01", "role": "ordinary-photo", "page_role": "ordinary-photo",
        "provider_id": "photo-postprocess-coach", "adapter_id": "photo-postprocess-v1",
        "style_contract_version": arrow["style_contract_version"],
        "style_contract_sha256": arrow["style_contract_sha256"],
        "source_kind": "ordinary-photo", "truth_sensitive": False, "deliverable_role": "slide-asset",
        "source_rights_confirmation": {"path": "rights/ordinary-photo.json", "sha256": rights_sha},
        "source_photo": {"path": "sources/ordinary-photo.png", "sha256": source_sha},
        "photo_output_mode": "rectangle", "strategy": "preserved-photo-rectangle",
        "editability": "component-editable", "content_policy": "preserve-source-photo",
        "background_requirement": "opaque-rectangle", "output_class": "opaque-photo-rectangle",
        "provider_transparency_behavior": "optional-subject-cutout",
        "provider_aspect_ratio_behavior": "preserve-subject-with-approved-crop",
        "contains_text": False, "contains_data": False, "contains_data_encoding_marks": False,
        "reference_bbox": [0.55, 0.30, 0.32, 0.32], "references": arrow["references"],
        "output": "sources/ordinary-photo.png", "sha256": source_sha,
        "alpha": None, "generation_attempts": 0, "generation_history": [],
        "extraction": {"route": "preserve-source"}, "status": "approved",
    }
    manifest["assets"].append(photo_asset)
    write_json(root, "asset-manifest.json", manifest)
    refresh_fixture_isolation_review(root)
    manifest = json.loads((root / "asset-manifest.json").read_text(encoding="utf-8"))
    photo_asset = manifest["assets"][1]
    photo_items = ISOLATION_MODULE.source_items_for_asset(decomposition, "C02")
    review_image = root / "reviews" / "assets" / "C02-isolation.png"
    ISOLATION_MODULE.build_isolation_image(root, source, photo_items).save(review_image)
    review_packet = root / "reviews" / "assets" / "C02-isolation.md"
    review_packet.write_text("# C02 fixture isolation\n\n" + ISOLATION_MODULE.value_sha256(photo_items) + "\n\n"
                             + "\n".join(ISOLATION_MODULE.review_row(item) for item in photo_items) + "\n", encoding="utf-8")
    photo_asset["isolation_review"] = {
        "status": "passed", "reviewer": "fixture-human",
        "image_path": "reviews/assets/C02-isolation.png", "image_sha256": hashlib.sha256(review_image.read_bytes()).hexdigest(),
        "packet_path": "reviews/assets/C02-isolation.md", "packet_sha256": hashlib.sha256(review_packet.read_bytes()).hexdigest(),
        "asset_sha256": source_sha, "reviewed_items_sha256": ISOLATION_MODULE.value_sha256(photo_items),
        "checks": {"single_object": True, "no_neighbor_residue": True, "edge_clean": True, "concept_match": True},
    }
    write_json(root, "asset-manifest.json", manifest)

    state = json.loads((root / "run-state.json").read_text(encoding="utf-8"))
    state["gates"]["brief_authority"]["approved_artifacts"]["rights/ordinary-photo.json"] = rights_sha
    for gate_name, names in {
        "concept_contact_sheet": ["deck-spec.json"],
        "decomposition": ["decomposition-map.json", "reviews/S01-decomposition.png", "reviews/S01-decomposition.md"],
    }.items():
        gate = state["gates"][gate_name]
        for name in names:
            gate["approved_artifacts"][name] = hashlib.sha256((root / name).read_bytes()).hexdigest()
        for key, value in gate_decision_values(root, gate_name).items():
            gate["approval_packet"]["decisions"][key]["value"] = value
            gate["approval_packet"]["decisions"][key]["value_sha256"] = value_sha256(value)
    write_json(root, "run-state.json", state)


class AssemblyTests(unittest.TestCase):
    def test_page_previews_use_distinct_paths_and_leave_final_map_untouched(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_ready_fixture(root)
            final_map = root / "assembly-map.json"
            before = final_map.read_bytes() if final_map.exists() else None
            with patch("assemble_editable_pptx.ProjectValidator") as validator:
                validator.return_value.run.return_value = []
                first = assemble(root, "batch-a-S01", render=False, preview_slide="S01")
                second = assemble(root, "batch-b-S01", render=False, preview_slide="S01")
            self.assertNotEqual(first, second)
            for version, output in (("batch-a-S01", first), ("batch-b-S01", second)):
                self.assertTrue(output.is_file())
                preview_map = json.loads((root / "previews" / version / "preview-assembly-map.json").read_text(encoding="utf-8"))
                self.assertEqual(preview_map["preview_slide"], "S01")
                self.assertEqual([slide["slide_id"] for slide in preview_map["slides"]], ["S01"])
            self.assertEqual(final_map.read_bytes() if final_map.exists() else None, before)

    def test_native_text_uses_fixed_semantic_role_token(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0]
        role = apply_text_role(
            paragraph,
            {"font_role": "action-title", "planned_font_pt": 34},
            {"action-title": {"font_family": "Microsoft YaHei", "bold": True, "color": "#F8FAFC"}},
            RGBColor(0, 0, 0),
        )
        self.assertEqual(role, "action-title")
        self.assertEqual(paragraph.font.name, "Microsoft YaHei")
        self.assertEqual(paragraph.font.size.pt, 34)
        self.assertIs(paragraph.font.bold, True)
        self.assertEqual(paragraph.font.color.rgb, RGBColor(248, 250, 252))

    def assemble_without_render(self, root: Path, version: str = "attack-v1") -> Path:
        result = subprocess.run([sys.executable, str(SCRIPT), str(root), "--version", version, "--no-render"], capture_output=True, text=True, check=False)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return root / "assembly" / f"deck-{version}.pptx"

    @staticmethod
    def refresh_output_hash(root: Path, output: Path) -> None:
        assembly = json.loads((root / "assembly-map.json").read_text(encoding="utf-8"))
        assembly["sha256"] = hashlib.sha256(output.read_bytes()).hexdigest()
        write_json(root, "assembly-map.json", assembly)

    def test_one_editable_picture_and_title_are_assembled_without_concept(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pptx_ready_fixture(root)
            issues = ProjectValidator(root, "assets").run()
            self.assertEqual(issues, [])
            result = subprocess.run([sys.executable, str(SCRIPT), str(root), "--version", "fixture-v1", "--no-render"], capture_output=True, text=True, check=False)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            output = root / "assembly" / "deck-fixture-v1.pptx"
            self.assertTrue(output.is_file())
            assembly = json.loads((root / "assembly-map.json").read_text(encoding="utf-8"))
            self.assertEqual(assembly["status"], "structure-assembled-visual-review-pending")
            self.assertEqual(len(assembly["slides"][0]["objects"]), 1)
            self.assertEqual(len(assembly["slides"][0]["slots"]), 1)

    def test_multi_component_photo_and_blank_data_stay_separate(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            multiple_object_fixture(root)
            issues = ProjectValidator(root, "assets").run()
            self.assertEqual(issues, [])
            result = subprocess.run([sys.executable, str(SCRIPT), str(root), "--version", "multi-v1", "--no-render"], capture_output=True, text=True, check=False)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            assembly = json.loads((root / "assembly-map.json").read_text(encoding="utf-8"))
            self.assertEqual([row["component_id"] for row in assembly["slides"][0]["objects"]], ["C01", "C02", "C01"])
            slots = {row["id"]: row["content"] for row in assembly["slides"][0]["slots"]}
            self.assertEqual(slots["F02"], "Pictures and copy remain independently editable.")
            self.assertEqual(slots["F03"], "")
            self.assertEqual(slots["F04"], "Data to be entered")

    def test_public_validator_rejects_missing_object_after_pptx_edit(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            multiple_object_fixture(root)
            output = self.assemble_without_render(root)
            deck = Presentation(output)
            victim = deck.slides[0].shapes[1]
            victim._element.getparent().remove(victim._element)
            deck.save(output)
            self.refresh_output_hash(root, output)
            errors = ProjectValidator(root, "assembly").run()
            self.assertTrue(any("object count, identity, or z-order" in issue for issue in errors), errors)

    def test_public_validator_rejects_manual_move_after_pptx_edit(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            multiple_object_fixture(root)
            output = self.assemble_without_render(root)
            deck = Presentation(output)
            deck.slides[0].shapes[0].left += Inches(0.25)
            deck.save(output)
            self.refresh_output_hash(root, output)
            errors = ProjectValidator(root, "assembly").run()
            self.assertTrue(any("geometry differs from current scene" in issue for issue in errors), errors)

    def test_public_validator_rejects_fabricated_data_in_blank_slot(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            multiple_object_fixture(root)
            output = self.assemble_without_render(root)
            deck = Presentation(output)
            data_box = next(shape for shape in deck.slides[0].shapes if shape.name == "APW:slot:F03")
            data_box.text = "9,999"
            deck.save(output)
            self.refresh_output_hash(root, output)
            errors = ProjectValidator(root, "assembly").run()
            self.assertTrue(any("PPTX editable text differs" in issue for issue in errors), errors)

    def test_invalid_scene_z_order_is_a_contract_error_not_a_crash(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            multiple_object_fixture(root)
            scene = json.loads((root / "scene-graph.json").read_text(encoding="utf-8"))
            scene["slides"][0]["elements"][0]["z"] = "front"
            write_json(root, "scene-graph.json", scene)
            errors = ProjectValidator(root, "assets").run()
            self.assertTrue(any("z must be a non-negative integer" in issue for issue in errors), errors)

    def test_public_validator_rejects_concept_image_substituted_for_component(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            multiple_object_fixture(root)
            output = self.assemble_without_render(root)
            deck = Presentation(output)
            deck.slides[0].shapes.add_picture(str(root / "concepts" / "S01.png"), Inches(1), Inches(1), Inches(4), Inches(2))
            deck.save(output)
            self.refresh_output_hash(root, output)
            errors = ProjectValidator(root, "assembly").run()
            self.assertTrue(any("concept render as media" in issue for issue in errors), errors)

    @unittest.skipUnless(sys.platform == "win32" and os.path.isfile("C:/Program Files/Microsoft Office/root/Office16/POWERPNT.EXE"), "PowerPoint renderer unavailable")
    def test_powerpoint_renders_current_pptx_at_full_size(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            multiple_object_fixture(root)
            result = subprocess.run([sys.executable, str(SCRIPT), str(root), "--version", "render-v1"], capture_output=True, text=True, check=False, timeout=90)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            assembly = json.loads((root / "assembly-map.json").read_text(encoding="utf-8"))
            self.assertEqual(len(assembly["renders"]), 1)
            self.assertEqual(len(assembly["slides"][0]["objects"]), 3)
            self.assertEqual(ProjectValidator(root, "assembly").run(), [])
            qa_errors = ProjectValidator(root, "qa").run()
            self.assertTrue(any("assembly visual review must pass before QA" in issue for issue in qa_errors), qa_errors)


if __name__ == "__main__":
    unittest.main()
