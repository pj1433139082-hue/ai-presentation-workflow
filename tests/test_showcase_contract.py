"""Contract regression tests use synthetic files only; these are not showcase QA."""
import hashlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from datetime import datetime
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from apw import check_project, init_showcase, status_project
from showcase_contract import validate_showcase, validate_concept_inputs, export_showcase
from test_native_objects import _make_project, _write_json
from assemble_editable_pptx import assemble_showcase
from showcase_contract import _validate_pptx_objects, _validate_review
from agent_qa import PAGE_CHECKS, DECK_CHECKS
from deck_qa_evidence import contact_sheet


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load(root, name):
    return json.loads((root / name).read_text(encoding="utf-8"))


def valid_concept(root):
    _make_project(root)
    Image.new("RGB", (1672, 941), "white").save(root / "concepts/S01.png")
    Image.new("RGB", (320, 200), "gray").save(root / "assets/prism.png")
    brief = load(root, "brief.json")
    brief.update(quality_profile="draft", audience="contract test", goal="check input lineage")
    brief["source_map"] = [{"id": "SOURCE-1", "kind": "manual-example", "locator": "synthetic contract test data", "truth_status": "illustrative"}]
    _write_json(root, "brief.json", brief)
    refs = {"selected_reference_ids": ["R1"], "references": [{"id": "R1", "status": "analyzed", "image_path": "assets/prism.png", "sha256": digest(root / "assets/prism.png"), "source_locator": "synthetic local reference", "usage_note": "contract test only", "transferable_rules": ["neutral field"], "project_translation": ["neutral test slide"], "do_not_copy": ["no source artwork claim"]}]}
    _write_json(root, "reference-study.json", refs)
    style = load(root, "style-contract.json")
    style["selected_reference_ids"] = ["R1"]
    _write_json(root, "style-contract.json", style)
    deck = load(root, "deck-spec.json")
    deck["slides"][0]["concept_path"] = "concepts/S01.png"
    _write_json(root, "deck-spec.json", deck)
    scene = load(root, "scene-graph.json")
    scene["contract_version"] = "showcase-v1"
    _write_json(root, "scene-graph.json", scene)
    assets = load(root, "asset-manifest.json")
    assets["assets"][0]["sha256"] = digest(root / "assets/prism.png")
    _write_json(root, "asset-manifest.json", assets)
    bind_inputs(root)
    return root


def bind_inputs(root):
    deck = load(root, "deck-spec.json")
    inputs = {key: digest(root / name) for key, name in {"brief": "brief.json", "deck_spec": "deck-spec.json", "reference_study": "reference-study.json", "style_contract": "style-contract.json"}.items()}
    slides = []
    for slide in deck["slides"]:
        slides.append({
            "slide_id": slide["id"],
            "status": "generated",
            "output_path": slide["concept_path"],
            "output_sha256": digest(root / slide["concept_path"]),
            "input_sha256": inputs,
            "copy_map": slide["copy"],
            "reference_images": [{"reference_id": "R1", "path": "assets/prism.png", "sha256": digest(root / "assets/prism.png")}],
        })
    _write_json(root, "concept-generation-log.json", {"status": "complete", "slides": slides})


def _assembled_test_project(root):
    valid_concept(root)
    mapping = assemble_showcase(root, "security-boundary-test", render=False)
    pptx = root / mapping["output"]
    render_path = root / "renders" / "S01.png"
    render_path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (1920, 1080), "white").save(render_path)
    assembly = load(root, "assembly-map.json")
    assembly["render_backend"] = "Synthetic test renderer (contract test only)"
    assembly["renders"] = [{"slide_id": "S01", "path": "renders/S01.png", "sha256": digest(render_path), "pptx_sha256": assembly["sha256"]}]
    _write_json(root, "assembly-map.json", assembly)
    return assembly, pptx


def _package_test_project(root):
    assembly, _ = _assembled_test_project(root)
    _write_json(root, "showcase-review.json", {
        "agent_review": {"actor": "agent", "reviewer": "synthetic-unit-test", "status": "complete", "outcome": "pass"},
        "user_feedback": {"actor": "user", "status": "rejected", "pptx_sha256": assembly["sha256"]},
    })
    return assembly


class ShowcaseContractTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = valid_concept(Path(self.temp.name) / "project")

    def tearDown(self):
        self.temp.cleanup()

    def errors(self, stage="concept"):
        return "\n".join(validate_showcase(self.root, stage)["errors"])

    def test_widescreen_generator_resolution_and_opaque_picture_are_valid(self):
        self.assertEqual(self.errors(), "")

    def test_concept_input_preflight_needs_no_scene_or_asset_manifest(self):
        style = load(self.root, "style-contract.json")
        style.pop("readability_profile", None)
        _write_json(self.root, "style-contract.json", style)
        bind_inputs(self.root)
        (self.root / "scene-graph.json").unlink()
        (self.root / "asset-manifest.json").unlink()
        report = validate_concept_inputs(self.root)
        self.assertTrue(report["ok"], report["errors"])
        self.assertEqual(report["facts"]["slide_ids"], ["S01"])

    def test_concept_input_preflight_keeps_concept_hash_binding(self):
        (self.root / "concepts/S01.png").write_bytes(b"changed content")
        report = validate_concept_inputs(self.root)
        self.assertFalse(report["ok"])
        self.assertTrue(any("SHA-256" in error for error in report["errors"]))

    def test_expected_slide_ids_is_optional_for_legacy_projects(self):
        self.assertNotIn("delivery", load(self.root, "brief.json"))
        report = validate_concept_inputs(self.root)
        self.assertTrue(report["ok"], report["errors"])

    def test_declared_expected_slide_ids_must_match_all_deck_pages_in_order(self):
        deck = load(self.root, "deck-spec.json")
        second = json.loads(json.dumps(deck["slides"][0]))
        second.update(id="S02", title="第二页", message="继续说明", concept_path="concepts/S02.png")
        second["copy"][0].update(id="S02-title", text="第二页")
        second["copy"][1].update(id="S02-takeaway", text="继续说明")
        deck["slides"].append(second)
        _write_json(self.root, "deck-spec.json", deck)
        Image.new("RGB", (1672, 941), "#eeeeee").save(self.root / "concepts/S02.png")

        brief = load(self.root, "brief.json")
        brief["delivery"] = {"expected_slide_ids": ["S01", "S02"]}
        _write_json(self.root, "brief.json", brief)
        bind_inputs(self.root)
        report = validate_concept_inputs(self.root)
        self.assertTrue(report["ok"], report["errors"])

        brief["delivery"]["expected_slide_ids"] = ["S02", "S01"]
        _write_json(self.root, "brief.json", brief)
        bind_inputs(self.root)
        report = validate_concept_inputs(self.root)
        self.assertFalse(report["ok"])
        self.assertTrue(any("must match deck-spec.json slide IDs in the same order" in error for error in report["errors"]), report["errors"])

    def test_declared_expected_slide_ids_must_be_nonempty_unique_strings(self):
        invalid_values = ([], ["S01", "S01"], [""], ["S01", 2], None)
        for expected_ids in invalid_values:
            with self.subTest(expected_ids=expected_ids):
                brief = load(self.root, "brief.json")
                brief["delivery"] = {"expected_slide_ids": expected_ids}
                _write_json(self.root, "brief.json", brief)
                bind_inputs(self.root)
                report = validate_concept_inputs(self.root)
                self.assertFalse(report["ok"], report["errors"])
                self.assertTrue(any("brief.json.delivery.expected_slide_ids" in error for error in report["errors"]), report["errors"])

    def test_illustrative_source_cannot_be_upgraded_to_verified_chart(self):
        scene = load(self.root, "scene-graph.json")
        chart = next(row for row in scene["slides"][0]["elements"] if row["type"] == "chart")
        chart["truth_status"] = "verified"
        _write_json(self.root, "scene-graph.json", scene)
        report = validate_showcase(self.root, "concept")
        self.assertFalse(report["ok"])
        self.assertTrue(any("truth_status" in error and "illustrative" in error for error in report["errors"]))

    def test_verified_copy_cannot_upgrade_nonverified_sources(self):
        for source_status in ("illustrative", "not-applicable"):
            with self.subTest(source_status=source_status):
                brief = load(self.root, "brief.json")
                brief["source_map"][0]["truth_status"] = source_status
                _write_json(self.root, "brief.json", brief)
                deck = load(self.root, "deck-spec.json")
                copy_row = deck["slides"][0]["copy"][0]
                copy_row["truth_status"] = "verified"
                copy_row["source_ids"] = ["SOURCE-1"]
                _write_json(self.root, "deck-spec.json", deck)
                bind_inputs(self.root)

                report = validate_concept_inputs(self.root)
                self.assertFalse(report["ok"], report["errors"])
                self.assertTrue(
                    any("verified copy cannot be supported by" in error and source_status in error for error in report["errors"]),
                    report["errors"],
                )

                # The same illustrative labels remain valid when they are not claimed as verified facts.
                copy_row["truth_status"] = "illustrative"
                copy_row["source_ids"] = []
                _write_json(self.root, "deck-spec.json", deck)
                bind_inputs(self.root)
                nonverified_report = validate_concept_inputs(self.root)
                self.assertTrue(nonverified_report["ok"], nonverified_report["errors"])

    def test_cutout_requires_real_alpha_only_when_declared(self):
        assets = load(self.root, "asset-manifest.json")
        assets["assets"][0]["requires_transparency"] = True
        _write_json(self.root, "asset-manifest.json", assets)
        self.assertIn("requires real transparent", self.errors())

    def test_source_reference_copy_or_concept_mutation_invalidates_lineage(self):
        cases = [("brief.json", lambda x: x.update(goal="changed")), ("reference-study.json", lambda x: x["references"][0].update(usage_note="changed")), ("deck-spec.json", lambda x: x["slides"][0].update(message="changed"))]
        for filename, mutation in cases:
            with self.subTest(filename=filename):
                original = (self.root / filename).read_bytes()
                doc = load(self.root, filename)
                mutation(doc)
                _write_json(self.root, filename, doc)
                self.assertIn("input hashes differ", self.errors())
                (self.root / filename).write_bytes(original)
        (self.root / "concepts/S01.png").write_bytes(b"invalid")
        self.assertIn("SHA-256", self.errors())

    def test_copy_and_unlabelled_data_cannot_pass(self):
        scene = load(self.root, "scene-graph.json")
        scene["slides"][0]["elements"][0]["text"] = "different"
        next(e for e in scene["slides"][0]["elements"] if e["type"] == "chart")["title"] = "真实数据"
        _write_json(self.root, "scene-graph.json", scene)
        self.assertIn("copy_id/text/role", self.errors())
        self.assertIn("visible", self.errors())

    def test_fixture_and_fabricated_human_gate_do_not_qualify(self):
        _write_json(self.root, "run-state.json", {"mode": "showcase-local-draft", "fixture_only": True, "human_gates": {"G5": "approved"}})
        self.assertIn("validation-fixture", self.errors())
        self.assertIn("must not claim", self.errors())

    def test_full_slide_raster_and_unsafe_asset_path_are_rejected(self):
        scene = load(self.root, "scene-graph.json")
        scene["slides"][0]["elements"][-2]["bbox"] = [0, 0, 1, 1]
        _write_json(self.root, "scene-graph.json", scene)
        assets = load(self.root, "asset-manifest.json")
        assets["assets"][0]["path"] = "../outside.png"
        _write_json(self.root, "asset-manifest.json", assets)
        self.assertIn("near-full-slide", self.errors())
        self.assertIn("path must stay inside", self.errors())

    def test_export_requires_actual_assembly_and_review(self):
        report = export_showcase(self.root, self.root / "package")
        self.assertFalse(report["ok"])
        self.assertFalse((self.root / "package").exists())

    def test_chart_axis_titles_are_not_concatenated_into_truth_title(self):
        scene = load(self.root, "scene-graph.json")
        chart = next(e for e in scene["slides"][0]["elements"] if e["type"] == "chart")
        chart["style"] = {"category_axis_title": "阶段", "value_axis_title": "次数"}
        _write_json(self.root, "scene-graph.json", scene)
        mapping = assemble_showcase(self.root, "test-axis-titles", render=False)
        pptx = self.root / mapping["output"]
        deck_rows = {r["id"]: r for r in load(self.root, "deck-spec.json")["slides"]}
        scenes = {r["slide_id"]: r for r in scene["slides"]}
        assembled = {r["slide_id"]: r for r in mapping["slides"]}
        errors = []
        _validate_pptx_objects(self.root, pptx, mapping, deck_rows, scenes, assembled, errors)
        self.assertEqual(errors, [])
        changed = Presentation(pptx)
        next(s.chart for s in changed.slides[0].shapes if s.has_chart).chart_title.text_frame.text = "伪造图表标题"
        changed.save(pptx)
        errors = []
        _validate_pptx_objects(self.root, pptx, mapping, deck_rows, scenes, assembled, errors)
        self.assertTrue(any("title/truth label" in error for error in errors))

    def test_unmapped_full_slide_overlay_is_rejected(self):
        _, pptx = _assembled_test_project(self.root)
        presentation = Presentation(pptx)
        slide = presentation.slides[0]
        overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height)
        overlay.name = "Unmapped full-slide overlay"
        presentation.save(pptx)
        deck_rows = {row["id"]: row for row in load(self.root, "deck-spec.json")["slides"]}
        scene_rows = {row["slide_id"]: row for row in load(self.root, "scene-graph.json")["slides"]}
        mapping = load(self.root, "assembly-map.json")
        assembled_rows = {row["slide_id"]: row for row in mapping["slides"]}
        errors = []
        _validate_pptx_objects(self.root, pptx, mapping, deck_rows, scene_rows, assembled_rows, errors)
        self.assertTrue(any("unmapped PPTX object" in error and "Unmapped full-slide overlay" in error for error in errors), errors)

    def test_chart_requires_a_readable_embedded_workbook_matching_source_data(self):
        _, pptx = _assembled_test_project(self.root)
        deck_rows = {row["id"]: row for row in load(self.root, "deck-spec.json")["slides"]}
        scene_rows = {row["slide_id"]: row for row in load(self.root, "scene-graph.json")["slides"]}
        mapping = load(self.root, "assembly-map.json")
        assembled_rows = {row["slide_id"]: row for row in mapping["slides"]}
        original = pptx.read_bytes()
        for case in ("missing", "unreadable", "source-mismatch"):
            with self.subTest(case=case):
                source = zipfile.ZipFile(io.BytesIO(original), "r")
                output = io.BytesIO()
                with source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as rewritten:
                    for name in source.namelist():
                        if case == "missing" and name.startswith("ppt/embeddings/"):
                            continue
                        data = source.read(name)
                        if case == "unreadable" and name.startswith("ppt/embeddings/"):
                            data = b"not a valid xlsx package"
                        elif case == "source-mismatch" and name.startswith("ppt/embeddings/"):
                            nested = io.BytesIO()
                            with zipfile.ZipFile(io.BytesIO(data), "r") as workbook, zipfile.ZipFile(nested, "w", zipfile.ZIP_DEFLATED) as changed:
                                for workbook_name in workbook.namelist():
                                    workbook_data = workbook.read(workbook_name)
                                    if workbook_name == "xl/worksheets/sheet1.xml":
                                        old_cell = b'<c r="B2" s="1"><v>2</v></c>'
                                        self.assertIn(old_cell, workbook_data)
                                        workbook_data = workbook_data.replace(old_cell, b'<c r="B2" s="1"><v>99</v></c>', 1)
                                    changed.writestr(workbook_name, workbook_data)
                            data = nested.getvalue()
                        rewritten.writestr(name, data)
                pptx.write_bytes(output.getvalue())
                errors = []
                _validate_pptx_objects(self.root, pptx, mapping, deck_rows, scene_rows, assembled_rows, errors)
                self.assertTrue(any("embedded workbook" in error for error in errors), errors)
        pptx.write_bytes(original)

    def test_package_verifier_rederives_gate_feedback_and_receipt_claims(self):
        _package_test_project(self.root)
        destination = self.root.parent / "security-boundary-package"
        report = {"ok": True, "errors": [], "warnings": [], "facts": {
            "agent_visual_review": {"status": "complete", "outcome": "pass"},
            "user_visual_feedback": "rejected",
        }}
        with patch("showcase_contract.validate_showcase", return_value=report):
            exported = export_showcase(self.root, destination)
            self.assertTrue(exported["ok"], exported["errors"])
            self.assertTrue(export_showcase(self.root, destination, verify_only=True)["ok"])

            manifest_path = destination / "PPT成稿" / "package-manifest.json"
            original_manifest = manifest_path.read_bytes()
            manifest = json.loads(original_manifest)
            manifest["human_gates"]["G4"] = "approved"
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            rejected = export_showcase(self.root, destination, verify_only=True)
            self.assertFalse(rejected["ok"])
            self.assertIn("human G1-G5", "\n".join(rejected["errors"]))

            manifest = json.loads(original_manifest)
            manifest["user_visual_feedback"] = "accepted"
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            rejected = export_showcase(self.root, destination, verify_only=True)
            self.assertFalse(rejected["ok"])
            self.assertIn("source receipt", "\n".join(rejected["errors"]))

            manifest = json.loads(original_manifest)
            manifest["formal_release"] = "approved"
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            rejected = export_showcase(self.root, destination, verify_only=True)
            self.assertFalse(rejected["ok"])
            self.assertIn("semantic claim fields", "\n".join(rejected["errors"]))

            manifest = json.loads(original_manifest)
            manifest_path.write_bytes(original_manifest)
            receipt_path = destination / "PPT成稿" / "execution-receipt.json"
            original_receipt = receipt_path.read_bytes()
            receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
            receipt["human_gates"]["G4"] = "approved"
            receipt_path.write_text(json.dumps(receipt, ensure_ascii=False), encoding="utf-8")
            for relative in ("PPT成稿/execution-receipt.json",):
                for row in manifest["files"]:
                    if row["path"] == relative:
                        row["sha256"] = digest(destination / relative)
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            rejected = export_showcase(self.root, destination, verify_only=True)
            self.assertFalse(rejected["ok"])
            self.assertIn("human G1-G5", "\n".join(rejected["errors"]))

            receipt = json.loads(original_receipt)
            receipt["agent_visual_review"]["outcome"] = "issues"
            receipt_path.write_text(json.dumps(receipt, ensure_ascii=False), encoding="utf-8")
            manifest = json.loads(original_manifest)
            for row in manifest["files"]:
                if row["path"] == "PPT成稿/execution-receipt.json":
                    row["sha256"] = digest(receipt_path)
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            rejected = export_showcase(self.root, destination, verify_only=True)
            self.assertFalse(rejected["ok"])
            self.assertIn("agent_visual_review differs", "\n".join(rejected["errors"]))

            receipt_path.write_bytes(original_receipt)
            readme_path = destination / "PPT成稿" / "README.md"
            readme_path.write_text("# 本地展示草稿\n\n已获正式批准。\n", encoding="utf-8")
            manifest = json.loads(original_manifest)
            for row in manifest["files"]:
                if row["path"] == "PPT成稿/README.md":
                    row["sha256"] = digest(readme_path)
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            rejected = export_showcase(self.root, destination, verify_only=True)
            self.assertFalse(rejected["ok"])
            self.assertIn("README semantic claims", "\n".join(rejected["errors"]))

    def test_unknown_discriminator_never_falls_back_to_production(self):
        _write_json(self.root, "project-kind.json", {"purpose": "showcase", "contract_version": "showcase-v999"})
        report = check_project(self.root, "concept")
        self.assertFalse(report["ok"])
        self.assertEqual(report["route"], "invalid")

    def test_status_points_at_first_unmet_stage(self):
        self.assertEqual(status_project(self.root)["next_stage"], "assembly")

    def test_complete_review_checks_contact_sheet_and_separate_user_feedback(self):
        # Synthetic contract-only unit evidence; this is never a real deck QA.
        image = Image.new("RGB", (1920, 1080), "white")
        image.save(self.root / "synthetic-render.png")
        contact_sheet([("S01", image)]).save(self.root / "synthetic-contact.png")
        stamp = datetime.now().astimezone().isoformat()
        render = {"slide_id":"S01", "path":"synthetic-render.png", "sha256":digest(self.root / "synthetic-render.png")}
        review = {"actor":"agent", "reviewer":"unit-test-fixture", "reviewed_at":stamp, "status":"complete", "pptx_sha256":"synthetic-binding", "renders":[dict(render, display_scale_percent=100)], "slides":[{"slide_id":"S01", "reviewer":"unit-test-fixture", "reviewed_at":stamp, "checks_performed":sorted(PAGE_CHECKS), "findings":[]}], "whole_deck_review":{"reviewer":"unit-test-fixture", "reviewed_at":stamp, "checks_performed":sorted(DECK_CHECKS), "findings":[], "contact_sheet":{"path":"synthetic-contact.png", "sha256":digest(self.root / "synthetic-contact.png")}}, "outcome":"pass"}
        docs = {"showcase-review.json":{"agent_review":review, "user_feedback":{"status":"not-recorded"}}}
        assembly = {"sha256":"synthetic-binding", "renders":[render]}
        errors, warnings, facts = [], [], {}
        _validate_review(self.root, docs, assembly, {"S01":{}}, errors, warnings, facts)
        self.assertEqual(errors, [])
        self.assertEqual(facts["user_visual_feedback"], "not-recorded")
        (self.root / "synthetic-contact.png").write_bytes(b"changed evidence")
        errors = []
        _validate_review(self.root, docs, assembly, {"S01":{}}, errors, [], {})
        self.assertTrue(any("contact" in error or "stale" in error for error in errors))

    def test_init_cannot_replace_existing_project(self):
        before = digest(self.root / "brief.json")
        self.assertFalse(init_showcase(self.root)["ok"])
        self.assertEqual(digest(self.root / "brief.json"), before)
        target = Path(self.temp.name) / "fresh"
        self.assertTrue(init_showcase(target)["ok"])
        self.assertFalse(check_project(target, "concept")["ok"])


if __name__ == "__main__":
    unittest.main()
