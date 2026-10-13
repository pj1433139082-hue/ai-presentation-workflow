import copy
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import Mock, patch

from PIL import Image, ImageChops
from pptx import Presentation

from test_pilot_review import make_contentful_pilot_fixture
from test_validate_project import ISOLATION_MODULE, ProjectValidator, REVIEW_MODULE, gate_decision_values, value_sha256, write_json


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"


def read(root: Path, name: str) -> dict:
    return json.loads((root / name).read_text(encoding="utf-8"))


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def two_page_fixture(root: Path) -> None:
    make_contentful_pilot_fixture(root)
    concept2 = root / "concepts" / "S02.png"
    with Image.open(root / "concepts" / "S01.png") as original:
        original.save(concept2)
    concept_sha = digest(concept2)
    deck = read(root, "deck-spec.json")
    second = copy.deepcopy(deck["slides"][0])
    second["id"] = "S02"
    second["concept_path"] = "concepts/S02.png"
    second["readability"]["full_size_review"].update(path="concepts/S02.png", sha256=concept_sha,
                                                        reviewed_concept_sha256=concept_sha)
    deck["slides"].append(second)
    write_json(root, "deck-spec.json", deck)
    scene = read(root, "scene-graph.json")
    second_scene = copy.deepcopy(scene["slides"][0])
    second_scene["slide_id"] = "S02"
    instance_map = {"I01": "I04", "I02": "I05", "I03": "I06"}
    for element in second_scene["elements"]:
        element["instance_id"] = instance_map[element["instance_id"]]
    scene["slides"].append(second_scene)
    write_json(root, "scene-graph.json", scene)
    manifest = read(root, "asset-manifest.json")
    for asset in manifest["assets"]:
        asset["slide_ids"] = ["S01", "S02"]
        second_ref = copy.deepcopy(asset["references"][0])
        second_ref.update(path="concepts/S02.png", sha256=concept_sha)
        asset["references"].append(second_ref)
    write_json(root, "asset-manifest.json", manifest)
    decomposition = read(root, "decomposition-map.json")
    second_map = copy.deepcopy(decomposition["slides"][0])
    second_map.update(slide_id="S02", concept_path="concepts/S02.png", concept_sha256=concept_sha)
    for item in second_map["items"]:
        if "instance_id" in item:
            item["instance_id"] = instance_map[item["instance_id"]]
        if "overlaps" in item:
            item["overlaps"] = [instance_map.get(value, value) for value in item["overlaps"]]
    second_map["review_image"].update(path="reviews/S02-decomposition.png", reviewed_concept_sha256=concept_sha)
    second_map["review_packet"]["path"] = "reviews/S02-decomposition.md"
    items_sha = value_sha256(second_map["items"])
    packet = root / second_map["review_packet"]["path"]
    packet.write_text("# S02 fixture decomposition\n\n" + items_sha + "\n\n" +
                      "\n".join(REVIEW_MODULE.review_row(item) for item in second_map["items"]) + "\n", encoding="utf-8")
    review = root / second_map["review_image"]["path"]
    REVIEW_MODULE.build_review_image(concept2, second_map["items"]).save(review)
    second_map["review_image"].update(sha256=digest(review), reviewed_items_sha256=items_sha)
    second_map["review_packet"].update(sha256=digest(packet), reviewed_items_sha256=items_sha)
    decomposition["slides"].append(second_map)
    write_json(root, "decomposition-map.json", decomposition)
    manifest = read(root, "asset-manifest.json")
    for asset in manifest["assets"]:
        source_items = ISOLATION_MODULE.source_items_for_asset(decomposition, asset["id"])
        image_path = root / asset["isolation_review"]["image_path"]
        packet_path = root / asset["isolation_review"]["packet_path"]
        ISOLATION_MODULE.build_isolation_image(root, root / asset["output"], source_items).save(image_path)
        packet_path.write_text(f"# {asset['id']} fixture isolation\n\n" + ISOLATION_MODULE.value_sha256(source_items) + "\n\n" +
                               "\n".join(ISOLATION_MODULE.review_row(item) for item in source_items) + "\n", encoding="utf-8")
        asset["isolation_review"].update(image_sha256=digest(image_path), packet_sha256=digest(packet_path),
                                         reviewed_items_sha256=ISOLATION_MODULE.value_sha256(source_items))
    write_json(root, "asset-manifest.json", manifest)
    storyboard = read(root, "storyboard.json")
    second_story = copy.deepcopy(storyboard["slides"][0])
    second_story["id"] = "S02"
    storyboard["slides"].append(second_story)
    write_json(root, "storyboard.json", storyboard)
    state = read(root, "run-state.json")
    for gate_name, names in {
        "story": ["storyboard.json"],
        "concept_contact_sheet": ["deck-spec.json", "concepts/S02.png"],
        "decomposition": ["decomposition-map.json", "concepts/S02.png", "reviews/S02-decomposition.png", "reviews/S02-decomposition.md"],
    }.items():
        gate = state["gates"][gate_name]
        for name in names:
            gate["approved_artifacts"][name] = digest(root / name)
        for key, value in gate_decision_values(root, gate_name).items():
            gate["approval_packet"]["decisions"][key]["value"] = value
            gate["approval_packet"]["decisions"][key]["value_sha256"] = value_sha256(value)
    state["gates"]["release"]["status"] = "pending"
    state["stage"] = "qa"
    write_json(root, "run-state.json", state)
    (root / "qa-report.json").unlink()


def make_rendered_deck(root: Path) -> None:
    two_page_fixture(root)
    errors = ProjectValidator(root, "assets").run()
    assert not errors, errors
    result = subprocess.run([sys.executable, str(SCRIPTS / "assemble_editable_pptx.py"), str(root), "--version", "two-v1"],
                            capture_output=True, text=True, check=False, timeout=120)
    assert result.returncode == 0, result.stdout + result.stderr
    result = subprocess.run([sys.executable, str(SCRIPTS / "prepare_deck_qa.py"), str(root), "--version", "two-v1"],
                            capture_output=True, text=True, check=False, timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr


def make_fixture_rendered_slide(root: Path) -> None:
    """Portable contract test; only the Windows test claims real PPTX rendering."""
    make_contentful_pilot_fixture(root)
    (root / "qa-report.json").unlink()
    reassemble_fixture_version(root, "fixture-v1")
    result = subprocess.run([sys.executable, str(SCRIPTS / "prepare_deck_qa.py"), str(root),
                             "--version", "fixture-v1"], capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stdout + result.stderr


def reassemble_fixture_version(root: Path, version: str) -> None:
    result = subprocess.run([sys.executable, str(SCRIPTS / "assemble_editable_pptx.py"), str(root),
                             "--version", version, "--no-render"], capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stdout + result.stderr
    assembly = read(root, "assembly-map.json")
    render = root / "reviews" / "assembly" / version / "slide-01.png"
    render.parent.mkdir(parents=True)
    with Image.open(root / "concepts" / "S01.png") as concept:
        concept.save(render)
    assembly["renders"] = [{"path": render.relative_to(root).as_posix(), "sha256": digest(render),
                            "pptx_sha256": assembly["sha256"]}]
    write_json(root, "assembly-map.json", assembly)


def review_all(root: Path) -> None:
    report = read(root, "qa-report.json")
    report["final_reviewer"] = "fixture-human"
    report["deck_checks"] = {key: True for key in report["deck_checks"]}
    for row in report["slides"]:
        row["status"] = "passed"
        row["reviewer"] = "fixture-human"
        row["reviewed_at"] = datetime.now(timezone.utc).isoformat()
        row["checks"] = {key: True for key in row["checks"]}
    write_json(root, "qa-report.json", report)


def finalize(root: Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run([sys.executable, str(SCRIPTS / "finalize_deck_qa.py"), str(root)],
                          capture_output=True, text=True, check=False, timeout=30)


class PortableDeckQATests(unittest.TestCase):
    def test_one_page_fixture_checks_current_evidence_without_powerpoint(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            review_all(root)
            self.assertEqual(finalize(root).returncode, 0)
            self.assertEqual(ProjectValidator(root, "qa").run(), [])
            report = read(root, "qa-report.json")
            side = root / report["slides"][0]["evidence"]["side_by_side"]["path"]
            with Image.open(side) as image:
                forged = Image.new("RGB", image.size, "#000000")
            forged.save(side)
            report["slides"][0]["evidence"]["side_by_side"]["sha256"] = digest(side)
            write_json(root, "qa-report.json", report)
            self.assertTrue(any("side-by-side pixels" in issue for issue in ProjectValidator(root, "qa").run()))

    def test_finalizer_rejects_path_traversal_before_writing(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "project"
            root.mkdir()
            make_fixture_rendered_slide(root)
            review_all(root)
            target = root.parent / "outside.json"
            target.write_text("keep", encoding="utf-8")
            report = read(root, "qa-report.json")
            report["exception_register"] = "../outside.json"
            write_json(root, "qa-report.json", report)
            result = finalize(root)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(target.read_text(encoding="utf-8"), "keep")

    def test_contact_sheet_has_both_rendered_pages_in_order(self):
        scripts = str(SCRIPTS)
        if scripts not in sys.path:
            sys.path.insert(0, scripts)
        from deck_qa_evidence import contact_sheet
        first = Image.new("RGB", (1920, 1080), "#ff0000")
        second = Image.new("RGB", (1920, 1080), "#0000ff")
        sheet = contact_sheet([("S01", first), ("S02", second)])
        self.assertEqual(sheet.size, (1340, 442))
        self.assertEqual(sheet.getpixel((30, 30)), (255, 0, 0))
        self.assertEqual(sheet.getpixel((690, 30)), (0, 0, 255))
        reversed_sheet = contact_sheet([("S02", second), ("S01", first)])
        self.assertIsNotNone(ImageChops.difference(sheet, reversed_sheet).getbbox())

    def test_slide_binding_rejects_stale_asset_and_style_even_with_current_pptx(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            review_all(root)
            self.assertEqual(finalize(root).returncode, 0)
            report = read(root, "qa-report.json")
            report["slides"][0]["bindings"]["assets"]["C01"] = "0" * 64
            report["slides"][0]["bindings"]["style_sha256"] = "1" * 64
            write_json(root, "qa-report.json", report)
            self.assertTrue(any("human slide approval is stale" in issue for issue in ProjectValidator(root, "qa").run()))

    def test_changed_asset_record_stales_human_review_after_assembly_manifest_hash_refresh(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            review_all(root)
            self.assertEqual(finalize(root).returncode, 0)
            manifest = read(root, "asset-manifest.json")
            manifest["assets"][0]["qa_probe_note"] = "record changed after slide review"
            write_json(root, "asset-manifest.json", manifest)
            assembly = read(root, "assembly-map.json")
            assembly["input_sha256"]["asset-manifest.json"] = digest(root / "asset-manifest.json")
            write_json(root, "assembly-map.json", assembly)
            errors = ProjectValidator(root, "qa").run()
            self.assertTrue(any("human slide approval is stale" in issue for issue in errors), errors)

    def test_subjective_deck_checks_are_not_preapproved(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            report = read(root, "qa-report.json")
            self.assertFalse(any(report["deck_checks"].values()))
            self.assertNotIn("reference-transfer", report["checks_performed"])
            review_all(root)
            report = read(root, "qa-report.json")
            report["deck_checks"]["composition_richness"] = False
            write_json(root, "qa-report.json", report)
            result = finalize(root)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("human whole-deck checks", result.stderr)
            self.assertEqual(read(root, "assembly-map.json")["status"], "structure-assembled-visual-review-pending")

    def test_rejected_review_is_archived_before_new_qa_version(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            report = read(root, "qa-report.json")
            report["slides"][0]["status"] = "rejected"
            report["slides"][0]["issues"] = [{"slide_id": "S01", "severity": "major",
                                                "evidence": report["slides"][0]["evidence"]["overlay"]["path"],
                                                "owner_stage": "assembly", "repair_entry": "fix placement",
                                                "repair_attempts": 1, "disposition": "open"}]
            write_json(root, "qa-report.json", report)
            old_bytes = (root / "qa-report.json").read_bytes()
            reassemble_fixture_version(root, "fixture-v2")
            result = subprocess.run([sys.executable, str(SCRIPTS / "prepare_deck_qa.py"), str(root),
                                     "--version", "fixture-v2"], capture_output=True, text=True, check=False)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual((root / "reviews" / "qa" / "fixture-v1" / "qa-report.snapshot.json").read_bytes(), old_bytes)
            self.assertEqual(read(root, "qa-report.json")["version"], "fixture-v2")
            self.assertEqual(read(root, "qa-report.json")["status"], "pending-human")

    def test_real_full_slide_picture_fallback_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            review_all(root)
            self.assertEqual(finalize(root).returncode, 0)
            assembly = read(root, "assembly-map.json")
            pptx_path = root / assembly["output"]
            presentation = Presentation(pptx_path)
            page = presentation.slides[0]
            for shape in list(page.shapes):
                shape._element.getparent().remove(shape._element)
            page.shapes.add_picture(str(root / "concepts" / "S01.png"), 0, 0,
                                    presentation.slide_width, presentation.slide_height)
            presentation.save(pptx_path)
            assembly["sha256"] = digest(pptx_path)
            assembly["renders"][0]["pptx_sha256"] = assembly["sha256"]
            write_json(root, "assembly-map.json", assembly)
            report = read(root, "qa-report.json")
            report["subject_sha256"] = assembly["sha256"]
            report["slides"][0]["bindings"]["pptx_sha256"] = assembly["sha256"]
            write_json(root, "qa-report.json", report)
            errors = ProjectValidator(root, "qa").run()
            self.assertTrue(any("full-slide" in issue or "concept image" in issue or "object count" in issue
                                for issue in errors), errors)

    def test_over_retry_cap_blocks_qa_even_if_issue_is_marked_resolved(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            review_all(root)
            report = read(root, "qa-report.json")
            report["slides"][0]["issues"] = [{
                "slide_id": "S01", "severity": "major", "disposition": "resolved",
                "evidence": report["slides"][0]["evidence"]["overlay"]["path"],
                "owner_stage": "assets", "repair_entry": "repair only the arrow",
                "repair_attempts": 4,
            }]
            write_json(root, "qa-report.json", report)
            result = finalize(root)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("repair_attempts must be 0..3", result.stderr)
            report["slides"][0]["issues"][0]["repair_attempts"] = 3
            report["repairs"] = [{"slide_id": "S01", "defect_class": "arrow-edge",
                                  "owner_stage": "assets", "artifact_path": "assets/S01/C01.png", "attempt": 4}]
            write_json(root, "qa-report.json", report)
            result = finalize(root)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("repair attempt must be 1..3", result.stderr)

    def test_malformed_slide_review_is_a_validation_error_not_a_crash(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            report = read(root, "qa-report.json")
            report["slides"] = [None]
            write_json(root, "qa-report.json", report)
            errors = ProjectValidator(root, "qa").run()
            self.assertTrue(any("slide evidence" in issue or "QA" in issue for issue in errors), errors)
            result = finalize(root)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(read(root, "assembly-map.json")["status"], "structure-assembled-visual-review-pending")

    def test_null_slides_and_nonobject_report_fail_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            report = read(root, "qa-report.json")
            report["slides"] = None
            write_json(root, "qa-report.json", report)
            self.assertTrue(ProjectValidator(root, "qa").run())
            self.assertNotEqual(finalize(root).returncode, 0)
            (root / "qa-report.json").write_text("[]", encoding="utf-8")
            self.assertTrue(ProjectValidator(root, "qa").run())
            self.assertNotEqual(finalize(root).returncode, 0)

    def test_production_reverification_rejects_declared_render_with_wrong_pixels(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_fixture_rendered_slide(root)
            (root / "qa-report.json").unlink()
            validator = ProjectValidator(root, "assembly")
            self.assertEqual(validator.run(), [])
            fake_validator = Mock()
            fake_validator.run.return_value = []
            fake_validator.docs = validator.docs
            fake_validator.quality_profile.return_value = "production"
            scripts = str(SCRIPTS)
            if scripts not in sys.path:
                sys.path.insert(0, scripts)
            import prepare_deck_qa

            def fake_fresh_render(_pptx: Path, out: Path, count: int) -> list[dict[str, str]]:
                self.assertEqual(count, 1)
                out.mkdir(parents=True)
                path = out / "slide-01.png"
                Image.new("RGB", (1920, 1080), "#000000").save(path)
                return [{"path": str(path)}]

            with patch.object(prepare_deck_qa, "ProjectValidator", return_value=fake_validator), \
                 patch.object(prepare_deck_qa, "render_powerpoint", side_effect=fake_fresh_render):
                with self.assertRaisesRegex(ValueError, "does not match a fresh PowerPoint render"):
                    prepare_deck_qa.prepare(root, "reverify-v1")

            def matching_render(_pptx: Path, out: Path, count: int) -> list[dict[str, str]]:
                self.assertEqual(count, 1)
                out.mkdir(parents=True)
                path = out / "slide-01.png"
                with Image.open(root / validator.docs["assembly-map.json"]["renders"][0]["path"]) as source:
                    source.save(path)
                return [{"path": str(path)}]

            with patch.object(prepare_deck_qa, "ProjectValidator", return_value=fake_validator), \
                 patch.object(prepare_deck_qa, "render_powerpoint", side_effect=matching_render):
                prepared = prepare_deck_qa.prepare(root, "reverify-v2")
            self.assertTrue(prepared.is_file())
            self.assertEqual(len(read(root, "qa-report.json")["verified_renders"]), 1)


@unittest.skipUnless(sys.platform == "win32" and os.path.isfile("C:/Program Files/Microsoft Office/root/Office16/POWERPNT.EXE"), "PowerPoint renderer unavailable")
class FullDeckQATests(unittest.TestCase):
    def test_two_real_powerpoint_slides_require_each_human_review_and_separate_g5(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_rendered_deck(root)
            presentation = Presentation(root / read(root, "assembly-map.json")["output"])
            self.assertEqual(len(presentation.slides), 2)
            report = read(root, "qa-report.json")
            self.assertEqual(report["status"], "pending-human")
            self.assertEqual(len(report["slides"]), 2)
            self.assertFalse(report["slides"][0]["checks"]["geometry"])
            self.assertIn("status must be 'passed'", " ".join(ProjectValidator(root, "qa").run()))
            review_all(root)
            report = read(root, "qa-report.json")
            report["slides"][1]["status"] = "rejected"
            write_json(root, "qa-report.json", report)
            failed = finalize(root)
            self.assertNotEqual(failed.returncode, 0)
            self.assertEqual(read(root, "qa-report.json")["status"], "pending-human")
            self.assertTrue(any("slide QA has not passed" in issue for issue in ProjectValidator(root, "release").run()))
            report["slides"][1]["status"] = "passed"
            write_json(root, "qa-report.json", report)
            passed = finalize(root)
            self.assertEqual(passed.returncode, 0, passed.stderr)
            self.assertEqual(ProjectValidator(root, "qa").run(), [])
            self.assertTrue(any("release" in item for item in ProjectValidator(root, "release").run()))

    def test_contact_sheet_and_page_evidence_reject_forged_pixels(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_rendered_deck(root)
            review_all(root)
            report = read(root, "qa-report.json")
            contact = root / report["final_contact_sheet"]
            with Image.open(contact) as image:
                size = image.size
            Image.new("RGB", size, "#000000").save(contact)
            report["contact_sheet_sha256"] = digest(contact)
            write_json(root, "qa-report.json", report)
            result = finalize(root)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("contact sheet pixels", result.stderr)

    def test_stale_asset_style_and_flattened_fallback_are_blocked(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_rendered_deck(root)
            review_all(root)
            self.assertEqual(finalize(root).returncode, 0)
            manifest = read(root, "asset-manifest.json")
            manifest["assets"][0]["status"] = "draft"
            write_json(root, "asset-manifest.json", manifest)
            self.assertTrue(any("stale" in item or "approved" in item for item in ProjectValidator(root, "qa").run()))
            manifest["assets"][0]["status"] = "approved"
            write_json(root, "asset-manifest.json", manifest)
            style = read(root, "style-contract.json")
            style["direction"] = "changed without new approval"
            write_json(root, "style-contract.json", style)
            self.assertTrue(any("stale" in item for item in ProjectValidator(root, "qa").run()))
            style["direction"] = read(root, "run-state.json")["gates"]["style_anchors"]["approval_packet"]["decisions"]["visual_direction"]["value"]
            write_json(root, "style-contract.json", style)
            assembly = read(root, "assembly-map.json")
            assembly["slides"][1]["flattened_concept_used"] = True
            write_json(root, "assembly-map.json", assembly)
            self.assertTrue(any("flattened_concept_used" in item for item in ProjectValidator(root, "qa").run()))


if __name__ == "__main__":
    unittest.main()
